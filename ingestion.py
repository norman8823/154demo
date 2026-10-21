"""Conventional horizontal P&L ingestion. Interpretation proposes; reviewed records calculate.

No accounting facts are supplied by a model. All amounts retain physical source refs.
"""
import calendar
import csv
import hashlib
import io
import json
import math
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from engine import Book, InputError, ROOT

CATEGORIES=['revenue','cogs','gross_profit','operating_expenses','ebitda','operating_profit','net_income','interest','tax','depreciation','other']
ROLES=['detail','total','subtotal','ignore']
SCENARIOS=['actual','budget','prior_year','ignore']
PROFILE_PATH=Path(os.environ.get('PORTFOLIO_MAPPING_PATH',str(ROOT/'.local'/'account_mappings.json')))
MONTHS={name.lower():i for i in range(1,13) for name in (calendar.month_name[i],calendar.month_abbr[i])}

def clean(text): return re.sub(r'[^a-z0-9]+',' ',str(text or '').lower()).strip()

def digest(text): return hashlib.sha256(text.encode()).hexdigest()[:18]

def period(text):
    if isinstance(text,(date,datetime)): return text.strftime('%Y-%m')
    text=str(text or '')
    # Explicit fiscal/quarter labels retain their grain; never map FY to December.
    q=re.search(r'\bQ([1-4])\s*(?:FY\s*)?(20\d{2}|\d{2})\b',text,re.I)
    if q:
        year=int(q[2]);year=year+2000 if year<100 else year
        return f'{year:04d}-'+('FQ' if 'fy' in q[0].lower() else 'Q')+q[1]
    q=re.search(r'\bFY\s*(20\d{2}|\d{2})\s*Q([1-4])\b',text,re.I)
    if q:return str(int(q[1])+(2000 if len(q[1])==2 else 0))+'-FQ'+q[2]
    annual=re.search(r'\b(?:FY|fiscal year)\s*(20\d{2}|\d{2})\b',text,re.I)
    if annual:return str(int(annual[1])+(2000 if len(annual[1])==2 else 0))+'-FY'
    canonical=re.fullmatch(r'20\d{2}-(?:FY|F?Q[1-4])',text)
    if canonical:return text
    m=re.search(r'\b(20\d{2})[-/](0?[1-9]|1[0-2])(?:[-/]\d{1,2})?\b',text)
    if m: return f'{int(m[1]):04d}-{int(m[2]):02d}'
    full_date=re.search(r'\b([A-Za-z]+)\s+\d{1,2},?\s+(20\d{2})\b',text)
    if full_date and full_date[1].lower() in MONTHS:return f'{int(full_date[2]):04d}-{MONTHS[full_date[1].lower()]:02d}'
    m=re.search(r"\b([A-Za-z]+)[\s\-/,.'’]+(\d{2}|20\d{2})\b",text)
    if m and m[1].lower() in MONTHS:
        year=int(m[2]);year=year+2000 if year<100 else year
        return f'{year:04d}-{MONTHS[m[1].lower()]:02d}'
    return ''

def column_period(value):
    if re.search(r'\bacq(?:uired|uisition)?\b|\bfounded\b|\binception\b',str(value or ''),re.I):return ''
    year=re.fullmatch(r'\s*(20\d{2})(?:\s+(?:actual|budget|plan|forecast|prior year))?\s*',str(value or ''),re.I)
    if year:return year[1]+'-FY'
    return period(value)


def period_grain(p):
    return 'annual' if str(p).endswith('-FY') else ('fiscal_quarter' if '-FQ' in str(p) else ('quarter' if re.fullmatch(r'20\d{2}-Q[1-4]',str(p)) else 'month'))


def period_label(p):
    if str(p).endswith('-FY'):return 'FY'+p[:4]
    if '-FQ' in str(p):return 'Q'+p[-1]+' FY'+p[:4]
    if re.fullmatch(r'20\d{2}-Q[1-4]',str(p)):return 'Q'+p[-1]+' '+p[:4]
    if re.fullmatch(r'20\d{2}-\d{2}',str(p)):return calendar.month_name[int(p[5:])]+' '+p[:4]
    return p


def amount(raw):
    if isinstance(raw,bool) or raw is None:return None
    if isinstance(raw,(float,int)):return float(raw) if math.isfinite(raw) else None
    text=str(raw).strip()
    if text in ('-','—'): return None  # A dash is not proof of zero.
    neg=text.startswith('(') and text.endswith(')')
    text=re.sub(r'[ $£€()]','',text)
    if ',' in text and not re.fullmatch(r'[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?',text):return None
    text=text.replace(',','')
    try: number=float(text)
    except ValueError:return None
    return (-number if neg else number) if math.isfinite(number) else None

# A finite vocabulary plus section structure, not one rule per client's account names.
ALIASES={
 'revenue':{'fee revenue','net fee revenue','total fee revenue','total net fee revenue','revenue','revenues','sales','net sales','income','turnover','total revenue','total revenues','total sales','total income','total turnover','net revenue'},
 'cogs':{'cost of sales','cost of services','cost of goods sold','cogs','direct costs','total cost of services','total cost of sales','total cost of goods sold','total direct costs','total cogs','cost of operations','total cost of operations'},
 'gross_profit':{'gross profit','gross income'},
 'operating_expenses':{'expenses','operating expenses','overhead','overheads','selling general and administrative','sga','total expenses','total operating expenses','total overhead','total overheads','selling general administrative','total sga','s g a','sg a','selling general administrative expenses','selling general and administrative'},
 'ebitda':{'ebitda','earnings before interest taxes depreciation and amortization'},
 'operating_profit':{'operating income','operating profit','net operating income','ebit'},
 'net_income':{'net income','net profit','net earnings','profit after tax','profit for the period','net income loss','net profit loss'},
 'interest':{'interest','interest expense','finance costs','interest paid'},
 'tax':{'income tax','income tax expense','corporation tax','tax expense','taxation','income taxes'},
 'depreciation':{'depreciation','amortization','depreciation and amortization','depreciation amortisation','depreciation and amortisation'},
}
AMBIGUOUS=re.compile(r'\b(reserve|deposit|advances?|suspense|clearing|ask my accountant|adjusted|exceptional|one off|prepayments?)\b')

@dataclass
class Account:
    id:str
    label:str
    row:int
    category:str
    role:str
    status:str
    context:str
    reference:str
    profile_key:str
    confirmed:bool=False

@dataclass
class FinancialTable:
    id:str
    file:str
    sheet:str
    header_row:int
    label_column:int
    company:str
    basis:str='unknown'
    currency:str='unknown'
    scale:float|None=None
    expense_sign:str='positive'
    columns:list=field(default_factory=list)
    accounts:list=field(default_factory=list)
    records:list=field(default_factory=list)
    issues:list=field(default_factory=list)
    metadata_refs:list=field(default_factory=list)
    label_header:str=''

@dataclass
class Intake:
    tables:list=field(default_factory=list)
    unresolved:list=field(default_factory=list)
    text_samples:list=field(default_factory=list)
    auxiliary:list=field(default_factory=list)
    warnings:list=field(default_factory=list)

class MappingRequired(InputError):
    def __init__(self,intake):
        self.intake=intake
        super().__init__('No usable dated financial statement was recognized in these inputs. Optionally select the financial table and period columns; approving every account is not required.')


def load_profiles(path=PROFILE_PATH):
    try:
        data=json.loads(Path(path).read_text())
        return data if isinstance(data,dict) else {}
    except (OSError,ValueError):return {}


def save_profiles(intake,path=PROFILE_PATH):
    profiles=load_profiles(path)
    for table in intake.tables:
        for a in table.accounts:
            if a.confirmed or a.status.startswith('Auto'): profiles[a.profile_key]={'category':a.category,'role':a.role}
        saved_layouts=profiles.setdefault('__layouts__',[])
        saved_layouts[:]=[l for l in saved_layouts if (l['company'],l['sheet'])!=(table.company,table.sheet)]
        saved_layouts.append({'company':table.company,'sheet':table.sheet,'header_row':table.header_row,'label_column':table.label_column,'label_header':table.label_header,'columns':[{'column':c['column'],'header':c['header'],'scenario':c['scenario']} for c in table.columns]})
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix('.tmp');temp.write_text(json.dumps(profiles,indent=2));temp.replace(path)


def propose(label,context,has_numbers):
    label=clean(label)
    if re.search(r'\bmargin\b',label) or label.endswith(' percent'):
        return 'other','ignore','Auto — reported ratio; not an additive account'
    if AMBIGUOUS.search(label):return 'other','detail','Review required — accounting meaning is ambiguous'
    # Presentation qualifiers do not change a reported metric's financial definition.
    label=re.sub(r'^(reported|consolidated) | (reported|consolidated)$','',label)
    for pattern,category in [(r'^depreciation\b|^amortization\b|^amortisation\b','depreciation'),(r'^interest expense\b','interest'),(r'^provision for income tax','tax')]:
        if re.search(pattern,label):return category,'detail','Auto — explicit financial account label'
    if label=='company ebitda':return 'ebitda','total','Auto — explicit company EBITDA'
    for category,aliases in ALIASES.items():
        if label in aliases:
            role='total' if has_numbers and (label.startswith('total ') or category not in {'revenue','cogs','operating_expenses'}) else 'detail'
            # A lone "Revenue" amount is a reported headline, not a detail hierarchy.
            if has_numbers and category in {'revenue','cogs','operating_expenses'}:role='total'
            return category,role,'Auto — explicit financial label'
    if label.startswith('total '):return context or 'other','subtotal','Auto — subtotal excluded from aggregation'
    if context in {'revenue','cogs','operating_expenses'}:
        return context,'detail','Auto — explicit statement section'
    return 'other','detail','Review required — unfamiliar account outside a financial section'


def unit_metadata(text):
    """Only stated denominations, never a company's industry or numerical magnitude."""
    text=re.sub(r'\b(?:U\.?S\.?|United States) dollars?\b','USD dollars',text,flags=re.I)
    currencies={c for c in ['USD','GBP','EUR','CAD','AUD'] if re.search(r'\b'+c+r'\b',text,re.I)}
    currency=next(iter(currencies)) if len(currencies)==1 else 'unknown'
    if not currencies and re.search(r'\$|\bdollars?\b',text,re.I):currency='$'
    scales=set()
    if re.search(r'\bmillions?\b|\$m\b',text,re.I):scales.add(1_000_000)
    if re.search(r'\bthousands?\b|\$000|\$k\b',text,re.I):scales.add(1000)
    if re.search(r'whole dollars|whole units|in dollars|amounts in (USD|GBP|EUR|CAD|AUD)',text,re.I):scales.add(1)
    if not scales and re.search(r'\bdollars?\b',text,re.I):scales.add(1)
    return currency,next(iter(scales)) if len(scales)==1 else None


def metadata(ws,header):
    cells=[c for row in ws.iter_rows(min_row=1,max_row=max(1,header-1)) for c in row if c.value is not None]
    text=' '.join(str(c.value) for c in cells)
    basis_flags=[value for value,pattern in [('accrual',r'\baccrual\b'),('cash',r'\bcash basis\b')] if re.search(pattern,text,re.I)]
    basis=basis_flags[0] if len(basis_flags)==1 else 'unknown'
    currency,scale=unit_metadata(text)
    refs=[(c.coordinate,str(c.value)) for c in cells]
    # Inherit only consistent explicit monetary declarations from workbook mastheads.
    # Local declarations take precedence; mixed workbook denominations never resolve ambiguity.
    def company_label(cells):
        return next((str(c.value).strip() for c in cells if len(str(c.value).strip())>3 and not period(c.value) and not re.search(r'income statement|profit.{0,5}loss|basis|USD|GBP|EUR|CAD|AUD|dollars?|report|thousand|million|synthetic|prepared|units',str(c.value),re.I)), 'Company not stated')
    company=company_label(cells)
    declarations=[]
    for sheet in ws.parent:
        prefix=[c for row in sheet.iter_rows(max_row=min(4,sheet.max_row),max_col=min(20,sheet.max_column)) for c in row if c.value is not None]
        if company=='Company not stated' or clean(company_label(prefix))!=clean(company):continue
        for row in sheet.iter_rows(max_row=min(4,sheet.max_row),max_col=min(20,sheet.max_column)):
            for cell in row:
                if isinstance(cell.value,str):
                    c,k=unit_metadata(cell.value)
                    if c!='unknown':declarations.append((c,k,sheet.title,cell.coordinate,cell.value))
    if currency=='unknown' and not re.search(r'\b(?:USD|GBP|EUR|CAD|AUD)\b',text,re.I):
        codes={d[0] for d in declarations}
        if len(codes)==1:
            currency=next(iter(codes));refs.extend(("'"+d[2]+"'!"+d[3],d[4]) for d in declarations if d[2]!=ws.title)
    if scale is None and not re.search(r'thousand|million|whole|\$[mk]|\$000',text,re.I):
        scales={d[1] for d in declarations if d[0]==currency and d[1] is not None}
        if currency!='unknown' and len(scales)==1:
            scale=next(iter(scales));refs.extend(("'"+d[2]+"'!"+d[3],d[4]) for d in declarations if d[2]!=ws.title and d[0]==currency)
    # Unscaled currency number formats identify source dollars, not the USD jurisdiction.
    formats={str(c.number_format) for row in ws.iter_rows(min_row=header+1,max_row=min(ws.max_row,header+50),max_col=min(ws.max_column,80)) for c in row if c.value is not None and '%' not in c.number_format and re.search(r'\$|\b(?:USD|GBP|EUR|CAD|AUD)\b',c.number_format)}
    format_codes={code for fmt in formats for code in ['USD','GBP','EUR','CAD','AUD'] if re.search(r'\b'+code+r'\b',fmt)}
    if currency=='unknown' and formats and not declarations and not re.search(r'\b(?:USD|GBP|EUR|CAD|AUD)\b|dollars?',text,re.I):
        currency=next(iter(format_codes)) if len(format_codes)==1 else ('$' if not format_codes else 'unknown')
    if format_codes and currency not in {'unknown','$'} and format_codes!={currency}:currency='unknown'
    if scale is None and formats and not any(d[1] is not None for d in declarations) and not re.search(r'thousand|million|\$[mk]|\$000',text,re.I):scale=1
    if formats:refs.extend((c.coordinate,'Number format: '+c.number_format) for row in ws.iter_rows(min_row=header+1,max_row=min(ws.max_row,header+50),max_col=min(ws.max_column,80)) for c in row if c.number_format in formats)
    return company or 'Company not stated',basis,currency,scale,list(dict.fromkeys(refs))


def discover_columns(ws,header,label_col):
    cols=[]
    top=' '.join(str(c.value or '') for row in ws.iter_rows(min_row=1,max_row=header-1) for c in row)
    context_periods=[period(str(c.value)) for row in ws.iter_rows(min_row=1,max_row=header-1) for c in row if period(str(c.value))]
    context_period=context_periods[0] if len(set(context_periods))==1 else ''
    for col in range(label_col+1,min(ws.max_column,80)+1):
        raw=ws.cell(header,col).value
        if raw is None or (amount(raw) is not None and not column_period(raw)):continue
        text=str(raw); norm=clean(text)
        if amount(raw) is not None and clean(ws.cell(header,label_col).value) not in {'','account','account name','description','line item','particulars','item','metric'}:continue
        # An entity's acquisition/inception date is not a reporting-period header.
        if not column_period(text) and period(text):continue
        # Two-level month / Actual-Budget headers, merged cells propagated only within the header.
        date_header=column_period(raw)
        upper=''
        if header>1:
            upper=ws.cell(header-1,col).value
            if not upper:
                for merged in ws.merged_cells.ranges:
                    if merged.min_row<=header-1<=merged.max_row and merged.min_col<=col<=merged.max_col:upper=ws.cell(merged.min_row,merged.min_col).value
        role_header=bool(re.search(r'actual|budget|plan|forecast|prior|\bpy\b',norm))
        p=date_header or (column_period(upper) if role_header else '')
        prior=bool(re.search(r'prior|previous|last year|\bpy\b',norm))
        banner=clean(ws.cell(header,label_col).value)
        scenario='budget' if re.search(r'budget|plan|forecast',norm) or re.search(r'\bbudget\b',banner) else ('prior_year' if prior or re.search(r'prior year|previous year',banner) else 'actual')
        if re.search(r'variance|\bvar\b|%|ytd|year to date|\btotal\b|fav|unfav',norm):continue
        is_period=bool(p) or bool(re.search(r'actual|budget|plan|prior|\bpy\b',norm))
        if not is_period:continue
        if not p and context_period:
            p=context_period
            if prior:p=f'{int(p[:4])-1:04d}'+p[4:]
        # Never guess a year from the machine clock.
        refs=[ws.cell(header,col).coordinate]
        if upper:refs.append(ws.cell(header-1,col).coordinate)
        dimension=str(upper or '') if upper and not period(upper) and clean(upper) not in {'month','year to date','actual','budget'} else ''
        cols.append({'column':col,'header':text,'period':p,'scenario':scenario,'reference':', '.join(refs),'confirmed':bool(p),'dimension':dimension,'grain':period_grain(p)})
    return cols


def inspect_files(files,profiles=None,layouts=None):
    profiles=profiles or {};layouts=layouts or {};intake=Intake()
    for name,data in files.items():
        if name.lower().endswith('.csv'):
            from openpyxl import Workbook
            wb=Workbook();ws=wb.active;ws.title='CSV'
            try:
                reader=csv.reader(io.StringIO(data.decode('utf-8-sig')))
                for row in reader:
                    # Worksheet row numbers retain physical CSV record-end rows, including comment lines.
                    for col,value in enumerate(row,1):ws.cell(reader.line_num,col,value)
            except (UnicodeError,csv.Error) as e:raise InputError(f'{name}: CSV must contain a conventional UTF-8 financial table.') from e
            book=None
        else:
            book=Book(name,data);book.strict_sum=False;wb=book.wb
        file_selection=layouts.get(name,{})
        table_count=0
        for ws in wb:
            sample=[{'ref':c.coordinate,'text':str(c.value)[:160]} for row in ws.iter_rows(max_row=min(ws.max_row,35),max_col=min(ws.max_column,20)) for c in row if isinstance(c.value,str) and amount(c.value) is None and not c.value.startswith('=')]
            intake.text_samples.append({'file':name,'sheet':ws.title,'cells':sample[:100]})
            if file_selection.get('sheet') and ws.title!=file_selection['sheet']:continue
            selected=layouts.get(name+'|'+ws.title,file_selection)
            choices=[]
            for row in range(1,min(ws.max_row,60)+1):
                for col in range(1,min(ws.max_column,8)+1):
                    label=clean(ws.cell(row,col).value)
                    if label in {'account','account name','account description','description','line item','ledger caption','particulars','financial statement','income statement','metric','measure','line','item'}:
                        columns=discover_columns(ws,row,col)
                        if columns:
                            if label in {'metric','measure','line','item'}:
                                financial_hints=sum(propose(ws.cell(r,col).value,'',True)[0]!='other' for r in range(row+1,min(ws.max_row,row+50)+1))
                                if financial_hints<2:continue
                            choices.append((row,col,columns))
            if not choices and not selected:
                # Conventional QB exports often leave the account-header cell blank.
                # Require both dated value columns and corroborating financial labels.
                for h in range(1,min(ws.max_row,60)+1):
                    candidates=[]
                    for c in range(1,min(ws.max_column,8)+1):
                        cols=discover_columns(ws,h,c)
                        if not cols:continue
                        if any(period(ws.cell(h,x['column']).value) and clean(ws.cell(h+1,x['column']).value) in {'actual','budget','plan'} for x in cols):continue
                        labels=[ws.cell(row,c).value for row in range(h+1,min(ws.max_row,h+35)+1)]
                        words=[clean(v) for v in labels if isinstance(v,str) and amount(v) is None]
                        hints=sum(propose(v,'',True)[0]!='other' for v in words)
                        numeric=0
                        for r in range(h+1,min(ws.max_row,h+50)+1):
                            for column in cols[:3]:
                                cell=ws.cell(r,column['column'])
                                try:value=amount(book.value(ws.title,cell.coordinate) if book else cell.value)
                                except (InputError,ArithmeticError,ValueError,TypeError,SyntaxError):value=None
                                if value is not None:numeric+=1;break
                        if len(words)>=3 and hints>=2 and len({propose(v,'',True)[0] for v in words}-{'other'})>=2 and numeric>=3:candidates.append((len(words),c,cols))
                    if candidates:
                        score=max(x[0] for x in candidates)
                        for _,c,cols in candidates:
                            if _==score:choices.append((h,c,cols))
            if not choices and not selected:
                for layout in profiles.get('__layouts__',[]):
                    if layout.get('sheet')!=ws.title:continue
                    h,c=layout['header_row'],layout['label_column']
                    if str(ws.cell(h,c).value)!=layout.get('label_header') or metadata(ws,h)[0]!=layout['company']:continue
                    cols=discover_columns(ws,h,c)
                    if not cols:
                        dates=[period(str(cell.value)) for row in ws.iter_rows(max_row=max(1,h-1)) for cell in row if period(str(cell.value))]
                        latest=dates[0] if len(set(dates))==1 else '';cols=[]
                        for original in layout['columns']:
                            raw=ws.cell(h,original['column']).value
                            if clean(raw)!=clean(original['header']):continue
                            p=period(raw) or latest
                            if p and original['scenario']=='prior_year':p=str(int(p[:4])-1)+p[4:]
                            cols.append(dict(original,period=p,reference=ws.cell(h,original['column']).coordinate))
                    if cols:choices.append((h,c,cols))
            if selected:
                header=int(selected['header_row']);label_col=int(selected['label_column'])
                if not (1<=header<=ws.max_row and 1<=label_col<=ws.max_column):raise InputError('Selected financial table coordinates are outside the worksheet.')
                selected_columns=selected.get('columns') or discover_columns(ws,header,label_col)
                if not selected_columns:
                    selected_columns=[{'column':c,'header':str(ws.cell(header,c).value),'period':'','scenario':'ignore','reference':ws.cell(header,c).coordinate} for c in range(label_col+1,ws.max_column+1) if ws.cell(header,c).value is not None]
                choices=[(header,label_col,selected_columns)]
            if not choices:continue
            # Multiple competing tables are not silently summed or selected.
            if len(choices)>1:
                intake.unresolved.append({'File':name,'Sheet':ws.title,'Reason':'Multiple possible financial tables; select one header row and account column.'});continue
            header,label_col,columns=choices[0]
            company,basis,currency,scale,refs=metadata(ws,header)
            keys=[(c['period'],c['scenario']) for c in columns if c['scenario']!='ignore']
            if len(keys)!=len(set(keys)) and any(c.get('dimension') for c in columns):
                # A branch/entity comparison is not another consolidated P&L.
                totals=[c for c in columns if re.search(r'\btotal\b',c.get('dimension',''),re.I)]
                entries=[]
                for row in range(header+1,ws.max_row+1):
                    label=ws.cell(row,label_col).value
                    category,role,_=propose(label,'',True)
                    if category not in {'revenue','cogs','gross_profit','ebitda'}:continue
                    for c in totals:
                        ref=ws.cell(row,c['column']).coordinate
                        try:value=amount(book.value(ws.title,ref)) if book else amount(ws[ref].value)
                        except (InputError,ArithmeticError,ValueError,TypeError):value=None
                        if value is not None:entries.append(dict(category=category,value=value*(scale or 1),period=c['period'],scenario=c['scenario'],source_ref=f"'{ws.title}'!{ref}",raw=ws[ref].value))
                intake.auxiliary.append(dict(file=name,sheet=ws.title,company=company,currency=currency,scale=scale,entries=entries))
                continue
            table=FinancialTable(digest(name+'|'+ws.title),name,ws.title,header,label_col,selected.get('company',company),selected.get('basis',basis),selected.get('currency',currency),selected.get('scale',scale),selected.get('expense_sign','positive'),columns=columns,metadata_refs=refs,label_header=str(ws.cell(header,label_col).value or ''))
            stated=' '.join(text for _,text in refs)
            if re.search(r'\baccrual\b',stated,re.I) and re.search(r'\bcash basis\b',stated,re.I):
                table.issues.append('Accounting basis conflicts in the source; comparisons are unavailable pending review.')
            if not table.company.strip() or table.company=='Company not stated':
                table.issues.append('Company identity is not stated; this source is kept separate.')
                table.company=f'Company not stated ({name} · {ws.title})'
            for column in columns:
                if column.get('scenario')=='ignore':continue
                header_text=str(column.get('header',''))
                stated_currencies={code for code in ['USD','GBP','EUR','CAD','AUD'] if re.search(r'\b'+code+r'\b',header_text,re.I)}
                if len(stated_currencies)>1 or (stated_currencies and table.currency!='unknown' and stated_currencies!={table.currency}):
                    table.issues.append(f'Column {column["column"]} currency conflicts with table metadata and is excluded. No FX conversion is performed.')
                    column['invalid']=True
                header_scale=1000000 if re.search('million',header_text,re.I) else (1000 if re.search('thousand',header_text,re.I) else (1 if re.search('whole units|whole dollars',header_text,re.I) else None))
                if header_scale is not None and table.scale is not None and header_scale!=table.scale:
                    table.issues.append(f'Column {column["column"]} scale conflicts with table metadata and is excluded.')
                    column['invalid']=True
            if selected.get('ai_pending'):
                observed={c['column']:c for c in discover_columns(ws,header,label_col)}
                for column in columns:
                    source=observed.get(column['column'],{})
                    if column.get('scenario')!='ignore' and (not source.get('period') or (column.get('period'),column.get('scenario'))!=(source.get('period'),source.get('scenario'))):
                        column['invalid']=True
                        table.issues.append(f'AI-proposed column {column["column"]} is not verified against the source header; excluded pending review.')
            if not columns or not any(c.get('scenario')!='ignore' for c in columns):table.issues.append('Select at least one period / scenario column.')
            for column in columns:
                if column.get('scenario')=='ignore':continue
                if not re.fullmatch(r'20\d{2}-(?:0[1-9]|1[0-2]|FY|F?Q[1-4])',str(column.get('period',''))) or column.get('scenario') not in SCENARIOS:
                    table.issues.append(f'Confirm period and actual/budget role for column {column.get("column")}.')
                    column['invalid']=True
            context=''
            for row in range(header+1,ws.max_row+1):
                if row>header+1 and len(discover_columns(ws,row,label_col))>=2:break
                label=ws.cell(row,label_col).value
                if not isinstance(label,str) or not label.strip():continue
                values=[];missing_values=[]
                for column in columns:
                    if column.get('scenario')=='ignore' or column.get('invalid'):continue
                    col=int(column['column']);raw=ws.cell(row,col).value
                    if raw is None or str(raw).strip() in ('-','—'):
                        missing_values.append((column,raw,None));continue
                    try:val=amount(book.value(ws.title,ws.cell(row,col).coordinate) if book and isinstance(raw,str) and raw.startswith('=') else raw)
                    except (InputError,ArithmeticError,ValueError,TypeError,SyntaxError) as e:val=None;table.issues.append(f'{ws.title}!{ws.cell(row,col).coordinate}: {e}')
                    if val is None:
                        if raw is not None:table.issues.append(f'{ws.title}!{ws.cell(row,col).coordinate}: nonnumeric or unsupported formula; provide a numeric financial amount.')
                        missing_values.append((column,raw,None))
                        continue
                    values.append((column,raw,val))
                category,role,status=propose(label,context,bool(values))
                if values or any(raw is not None and str(raw).strip() not in ('-','—') for _,raw,_ in missing_values):values.extend(missing_values)
                if not values:
                    if category in {'revenue','cogs','operating_expenses'} and clean(label) in ALIASES.get(category,set()):
                        context=category;continue
                    if clean(label) in {'other expenses','other income','other income and expenses'}:
                        context='other';continue
                    if clean(label)=='below ebitda':context='below_ebitda';continue
                    if context in {'revenue','cogs','operating_expenses'} and missing_values:
                        values=missing_values
                        status='Review required — empty account line or section heading; confirm its role'
                    else:continue
                if role=='total' and category in {'revenue','cogs','operating_expenses','gross_profit','operating_profit','ebitda','net_income'}:context=''
                key=digest(clean(table.company)+'|'+clean(ws.title)+'|'+context+'|'+clean(label))
                saved=profiles.get(key,{})
                confirmed=saved.get('category') in CATEGORIES and saved.get('role') in ROLES
                if confirmed:category,role,status=saved['category'],saved['role'],'Confirmed — saved local mapping'
                account=Account(table.id+':'+str(row),label.strip(),row,category,role,status,context,f"'{ws.title}'!{ws.cell(row,label_col).coordinate}",key,confirmed)
                table.accounts.append(account)
                for column,raw,val in values:
                    table.records.append({'account_id':account.id,'label':account.label,'period':column['period'],'grain':column.get('grain',period_grain(column['period'])),'scenario':column['scenario'],'raw':raw,'raw_value':val,'cell':ws.cell(row,int(column['column'])).coordinate,'header_ref':column.get('reference',ws.cell(header,int(column['column'])).coordinate)})
            if 'expense_sign' not in selected:
                cost_ids={a.id for a in table.accounts if a.category in {'cogs','operating_expenses'} and a.role=='total'}
                costs=[r['raw_value'] for r in table.records if r['account_id'] in cost_ids and r['raw_value'] is not None]
                if costs and all(value<=0 for value in costs) and any(value<0 for value in costs):table.expense_sign='negative'
                elif any(value<0 for value in costs) and any(value>0 for value in costs):table.issues.append('Mixed signs in financial cost totals. Confirm the expense sign convention in mapping review.')
            if table.accounts or selected:intake.tables.append(table);table_count+=1
        if not table_count and not any(x['File']==name for x in intake.unresolved):
            intake.unresolved.append({'File':name,'Sheet':'','Reason':'Mapping required: select a conventional financial table, its account-label column, and period headers.'})
    return intake


def apply_decisions(intake,decisions):
    for table in intake.tables:
        for a in table.accounts:
            d=decisions.get(a.id)
            if not d:continue
            if d.get('category') not in CATEGORIES or d.get('role') not in ROLES:raise InputError('Choose a supported account category and aggregation role.')
            a.category=d['category'];a.role=d['role'];a.confirmed=bool(d.get('confirmed'))
            a.status='Confirmed — user review' if a.confirmed else 'Review required — proposed change'
    return intake


def needs_review(intake):
    return bool(intake.unresolved or any(t.issues or any(c['scenario']!='ignore' and not c.get('period') for c in t.columns) or any('Review required' in a.status and not a.confirmed for a in t.accounts) for t in intake.tables))


def material_pending(table):
    """Review suggestions only; neither a generation gate nor permission to omit amounts."""
    revenue_ids={a.id for a in table.accounts if a.category=='revenue' and a.role=='total'}
    revenue=max([abs(r['raw_value']) for r in table.records if r['account_id'] in revenue_ids and r['raw_value'] is not None] or [0])
    threshold=max(1,revenue*.01)
    return [a for a in table.accounts if not (a.confirmed or a.status.startswith('Auto')) and
            any(abs(r['raw_value'])>=threshold for r in table.records if r['account_id']==a.id and r['raw_value'] is not None)]


def normalized_records(table):
    accounts={a.id:a for a in table.accounts}
    out=[]
    for raw in table.records:
        a=accounts[raw['account_id']]
        accepted=a.confirmed or a.status.startswith('Auto')
        if (a.role=='ignore' and accepted) or raw['scenario']=='ignore':continue
        value=raw['raw_value']*(table.scale if table.scale is not None else 1) if raw['raw_value'] is not None else None
        if value is not None and a.category in {'cogs','operating_expenses','interest','tax','depreciation'} and table.expense_sign=='negative':value=-value
        out.append(dict(raw,context=a.context,category=a.category,role=a.role,value=value,company=table.company,source_file=table.file,source_sheet=table.sheet,source_ref=f"'{table.sheet}'!{raw['cell']}",basis=table.basis,currency=table.currency,scale=table.scale))
        out[-1].update(accepted=a.confirmed or a.status.startswith('Auto'),mapping_status=a.status,
                      sign_uncertain=any('Mixed signs' in issue for issue in table.issues))
    return out
