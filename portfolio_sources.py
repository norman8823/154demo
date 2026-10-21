"""Offline source snapshots and truthful calculation-to-original-cell navigation.

Only manifest-approved workbooks are read. Original coordinates and styles stay
intact; formula cells hold retained observed values, with original formula text
in comments. No source formulas, macros, hyperlinks or external links execute.
"""
from copy import copy
import hashlib
import json
import re
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.utils import get_column_letter, coordinate_to_tuple
from engine import ROOT, InputError

PREFIXES={'summit':'SHS','tallis':'TBA','harrow':'HRE'}
CELL_REFERENCE=re.compile(r"^'((?:[^']|'')+)'!(\$?[A-Z]{1,3}\$?[1-9]\d*)( comment)?$")


def snapshot_name(company_id, ordinal, original):
    if company_id not in PREFIXES:raise InputError('Unknown source company namespace.')
    safe=re.sub(r'[\[\]:*?/\\]',' ',original).strip(" '")
    return (PREFIXES[company_id]+f' {ordinal:02d} '+safe)[:31]


def canonical_source(company):
    manifest=json.loads((ROOT/'demo_results/manifest.json').read_text())
    entry=next((e for e in manifest['runs'] if e['run']==company.run),None)
    fields=('workbook','workbook_sha256','snapshot','snapshot_sha256')
    if entry is None or any(company.provenance.get(k)!=entry[k] for k in fields):
        raise InputError('Source provenance does not match the canonical approved manifest: '+company.id)
    path=(ROOT/entry['workbook']).resolve()
    if not path.is_relative_to(ROOT.resolve()) or hashlib.sha256(path.read_bytes()).hexdigest()!=entry['workbook_sha256']:
        raise InputError('Canonical source workbook is missing or changed: '+company.id)
    return path


def _snapshot(wb,company,source,ordinal,catalog_row):
    from portfolio_excel import link
    name=snapshot_name(company.id,ordinal,source.title)
    if name.casefold() in {s.casefold() for s in wb.sheetnames}:raise InputError('Source snapshot worksheet name collision.')
    ws=wb.create_sheet(name);ws.sheet_state='visible';ws.sheet_view.showGridLines=False
    facts={(f.get('sheet'),f.get('cell')):f for f in company.review.package['facts']}
    for row in source:
        for original in row:
            if original.value is None:continue
            target=ws.cell(original.row,original.column)
            # Workbook style IDs are local; copy components so the destination
            # registers its own font/fill/alignment/format IDs correctly.
            target.font=copy(original.font);target.fill=copy(original.fill)
            target.border=copy(original.border);target.alignment=copy(original.alignment)
            target.number_format=original.number_format
            target.protection=Protection(locked=True,hidden=False)
            if original.data_type=='f':
                fact=facts.get((source.title,original.coordinate))
                if fact is None or 'source_value' not in fact or fact['raw']!=original.value:
                    raise InputError('A source formula has no retained observed value: '+company.id+'/'+source.title+'!'+original.coordinate)
                target.value=fact['source_value']
                note=('READ-ONLY SOURCE SNAPSHOT\nOriginal formula (provenance only; not executed):\n'+original.value+
                      '\nRetained observed value in original source units: '+str(fact['source_value'])+
                      '\nValidated base-unit value: '+str(fact['value'])+'\nSource: '+company.provenance['workbook']+
                      '\nOriginal coordinate: '+source.title+'!'+original.coordinate)
                if original.comment:note+='\nOriginal comment: '+original.comment.text
                target.comment=Comment(note,'Portfolio Intelligence')
            else:
                target.value=original.value
                if original.comment:target.comment=copy(original.comment)
            if isinstance(target.value,str):target.data_type='s'
    for merged in source.merged_cells.ranges:ws.merge_cells(str(merged))
    for key,dimension in source.column_dimensions.items():
        ws.column_dimensions[key].width=dimension.width;ws.column_dimensions[key].hidden=False
    for index,dimension in source.row_dimensions.items():
        ws.row_dimensions[index].height=dimension.height;ws.row_dimensions[index].hidden=False
    ws.freeze_panes=source.freeze_panes or 'B6';ws.sheet_properties.tabColor='90A7B8'
    ws.protection.sheet=True;ws.protection.selectLockedCells=False;ws.protection.selectUnlockedCells=False
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_area=f'A1:{get_column_letter(source.max_column)}{source.max_row}'
    # Navigation goes to the right of the original grid; source coordinates never move.
    gutter=source.max_column+2;letter=get_column_letter(gutter);ws.column_dimensions[letter].width=48
    for row,text in [(1,'READ-ONLY SOURCE SNAPSHOT'),(2,company.name),(3,'Original sheet: '+source.title),(4,'← Source snapshot catalog'),(5,'← Source References')]:
        c=ws.cell(row,gutter,text);c.data_type='s';c.font=Font(name='Arial',size=10,color='18364A',bold=row==1);c.alignment=Alignment(wrap_text=True)
        if row==1:c.fill=PatternFill('solid',fgColor='DFEAF0')
    link(ws.cell(4,gutter),'Source Snapshots',catalog_row)
    link(ws.cell(5,gutter),'Source References',1)
    note='Read-only snapshot: original coordinates and source units retained. Formula cells display retained observed values; original formulas are in cell comments. No active source formulas or external workbook links. Review-only sheet protection is not encryption.'
    ws.cell(7,gutter,'Formula text is retained in cell comments.').comment=Comment(note,'Portfolio Intelligence')
    return ws


def attach_source_snapshots(wb,companies,records,source_rows):
    from portfolio_excel import literal,headers,row_values,link
    catalog=wb['Source Snapshots'];index=wb['Source References']
    # Replace the provenance-only catalog with one clickable row per original sheet.
    if catalog.max_row>=5:catalog.delete_rows(5,catalog.max_row-4)
    widths=[29,32,39,56,68,68]
    for col,width in enumerate(widths,1):catalog.column_dimensions[get_column_letter(col)].width=width
    headers(catalog,5,['Company','Original worksheet','Read-only source snapshot','Canonical approved workbook','Workbook SHA-256','Snapshot SHA-256'])
    catalog['A3']='Original cell coordinates and source units are preserved. Formula cells show approved retained observed values; original formula text remains in cell comments. Snapshots are visible and protected for review. Links work offline.'
    catalog['A3'].data_type='s';catalog.row_dimensions[3].height=44
    mapping={};original_bounds={};catalog_row=6
    for company in companies:
        source=load_workbook(canonical_source(company),data_only=False,keep_links=False)
        for ordinal,sheet in enumerate(source,1):
            snapshot=_snapshot(wb,company,sheet,ordinal,catalog_row)
            mapping[(company.id,sheet.title)]=snapshot.title
            original_bounds[(company.id,sheet.title)]=(sheet.max_row,sheet.max_column)
            row_values(catalog,catalog_row,[company.name,sheet.title,snapshot.title,company.provenance['workbook'],company.provenance['workbook_sha256'],company.provenance['snapshot_sha256']],42)
            link(catalog.cell(catalog_row,3),snapshot.title,1);catalog_row+=1
    catalog.auto_filter.ref=f'A5:F{catalog_row-1}';catalog.print_area=f'A1:F{catalog_row-1}'

    lookup={(c.id,e['Fact ID']):(c,e) for c,e in records}
    calculations={(c.id,f['id']):f for c in companies for f in c.review.package['skill_engine']['calculations']}
    referenced_calcs=[key for key in source_rows if key in calculations]
    visiting=set();validated=set()
    def validate(key):
        if key in visiting:raise InputError('Source calculation contains a dependency cycle.')
        if key in validated:return
        if key not in source_rows:raise InputError('A calculation operand is absent from the source index: '+key[1])
        visiting.add(key)
        if key in calculations:
            for operand in calculations[key]['inputs']:validate((key[0],operand))
        visiting.remove(key);validated.add(key)
    for key in referenced_calcs:validate(key)

    starts={};cursor=index.max_row+5
    for key in referenced_calcs:
        starts[key]=cursor;cursor+=len(calculations[key]['inputs'])+4

    def destination(key):
        if key in starts:return ('Source References','A'+str(starts[key]))
        company,evidence=lookup[key]
        match=CELL_REFERENCE.fullmatch(evidence['Reference'])
        if match is None:raise InputError('Unsupported source reference: '+company.id+'/'+evidence['Fact ID']+' '+evidence['Reference'])
        original,coordinate,comment=match.groups();original=original.replace("''", "'");coordinate=coordinate.replace('$','')
        sheet=mapping.get((company.id,original))
        if sheet is None:raise InputError('Unsupported source worksheet: '+company.id+'/'+original)
        row,col=coordinate_to_tuple(coordinate);bounds=original_bounds[(company.id,original)]
        if row>bounds[0] or col>bounds[1] or wb[sheet][coordinate].value is None:raise InputError('Source target is unavailable: '+sheet+'!'+coordinate)
        if comment and wb[sheet][coordinate].comment is None:raise InputError('Source comment target is unavailable.')
        if evidence['File']!=canonical_source(company).name:raise InputError('Evidence source file differs from canonical approved workbook.')
        return sheet,coordinate

    def place_link(cell,key):
        sheet,coordinate=destination(key);row,col=coordinate_to_tuple(coordinate)
        link(cell,sheet,row,column=get_column_letter(col))

    # The index link explicitly leads to either an original cell or a derived definition.
    for key,row in source_rows.items():place_link(index.cell(row,4),key)
    facts={(c.id,f['id']):f for c in companies for f in c.review.package['facts']}
    for key,start in starts.items():
        company,_=lookup[key];calc=calculations[key]
        values=[company.name,calc['id'],'DERIVED CALCULATION',calc['expression'],calc['value'],calc['unit'],calc['period'],calc['assumptions'],
                ('Illustrative sensitivity. ' if calc['illustrative'] else 'Historical arithmetic. ')+calc['limitation'],'↑ Source index']
        row_values(index,start,values,100);link(index.cell(start,10),'Source References',source_rows[key])
        headers(index,start+1,['Company','Operand ID','Type','Original reference / calculation','Validated value','Unit','Operand scope','Open operand','Qualification',''])
        for row,operand in enumerate(calc['inputs'],start+2):
            operand_key=(company.id,operand);_,evidence=lookup[operand_key];fact=calculations.get(operand_key) or facts.get(operand_key,{})
            row_values(index,row,[company.name,operand,'Derived calculation' if operand_key in starts else 'Original source cell',evidence['Reference'],evidence['Validated value'],fact.get('unit','source text'),
                                 fact.get('period','See source context'),'Open calculation →' if operand_key in starts else 'Open original cell →',fact.get('limitation','Source observation'), ''],72)
            place_link(index.cell(row,8),operand_key)
    index.print_area=f'A1:J{index.max_row}'
    if wb._external_links:raise InputError('Source snapshots must not include external workbook links.')
    return mapping
