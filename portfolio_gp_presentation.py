"""GP-only investment-review presentation over unchanged approved calculations."""
import textwrap
from datetime import date
from copy import copy
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break, RowBreak
from portfolio_investments import investments_for, investment_summary
from portfolio_comparison import comparison
from portfolio_dashboard import workflow_summary
from portfolio_outlook import investment_outlook
from portfolio_financial_reporting import internal_link

NAVY, MUTED, PALE = '18364A', '60798A', 'EFF5F7'
MILLIONS = '"$"0.00,,"M";("$"0.00,,"M");"—"'
HOLDINGS_HEAD = 14


def review_layout(ws, *, landscape=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if landscape else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=.35, right=.35, top=.35, bottom=.4, header=.15, footer=.2)
    ws.oddFooter.center.text = 'For Investment Team Review · Product Demo · Synthetic Data · Page &P'
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = MUTED


def assessment(companies):
    """Editorial framing; amounts/ranking come from the existing sponsor ledger."""
    positions = [(c, investment_summary([i])) for c, i in zip(companies, investments_for(companies))]
    total = investment_summary(investments_for(companies))
    ranked = sorted(positions, key=lambda p: (-p[1]['gross_moic'], p[0].id))
    best, weakest = ranked[0], ranked[-1]
    proceeds = f'USD {total["realized_proceeds"]/1e6:,.2f} million'
    value = f'USD {total["current_equity_value"]/1e6:,.2f} million'
    returns = (f'The illustrative sponsor portfolio holds {value} of current equity plus {proceeds} of realized proceeds. '
               f'Gross MOIC is {total["gross_moic"]:.2f}x and dated gross IRR is '
               + (f'{total["gross_irr"]:.2%}.' if total['gross_irr'] is not None else 'unavailable.'))
    if len(positions) > 1:
        returns += f' {best[0].name} is the strongest holding by gross MOIC; {weakest[0].name} is the weakest.'
    stats = workflow_summary(companies)
    operating = (f'{stats["below_plan"]} of {stats["comparable_budget"]} holdings with comparable EBITDA budgets are below plan in their own reporting periods. '
                 'Prioritize earnings recovery and cash delivery alongside the illustrative valuation gains. '
                 'Monthly and annual operating results remain separate; no consolidated revenue or EBITDA is presented.')
    return total, positions, returns, operating


def write_investment_review(wb, companies):
    index = wb.sheetnames.index('Portfolio Overview')
    wb.remove(wb['Portfolio Overview'])
    ws = wb.create_sheet('Portfolio Overview', index)
    review_layout(ws)
    for col in range(1, 13): ws.column_dimensions[get_column_letter(col)].width = 9

    def merged(row, value, start=1, end=12, *, size=11, bold=False, fill=None, color=NAVY, height=24, wrap=True):
        ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        cell = ws.cell(row, start, value)
        if isinstance(value, str): cell.data_type = 's'
        cell.font = Font(name='Calibri', size=size, bold=bold, color=color)
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=wrap)
        if fill:
            for col in range(start, end+1): ws.cell(row, col).fill = PatternFill('solid', fgColor=fill)
        ws.row_dimensions[row].height = height
        return cell

    def body(row, text, *, color=NAVY):
        count = sum(max(1, len(textwrap.wrap(p, width=120))) for p in text.splitlines())
        return merged(row, text, color=color, height=count*14+12)

    def heading(row, text):
        return merged(row, text, bold=True, size=12, fill=PALE, wrap=False, height=27)

    total, positions, returns, operating = assessment(companies)
    merged(1, 'Portfolio Overview', size=22, bold=True, height=35, wrap=False)
    merged(2, 'Portfolio Investment Review · Valuation as of '+total['valuation_date'], color=MUTED, height=22)
    merged(3, 'NL Intelligence · For Investment Team Review · Product Demo · Synthetic Data', color=MUTED, size=10, height=22)
    cell = merged(4, 'Open Portfolio Financial Performance →', size=10, height=22)
    internal_link(cell, 'Portfolio Financial Performance', 'A1')
    for start, label, key, fmt in ((1,'Current Equity Value','current_equity_value',MILLIONS),
                                  (4,'Gross MOIC','gross_moic','0.00"x"'),
                                  (7,'Gross IRR','gross_irr','0.00%;(0.00%);"—"'),
                                  (10,'Invested Capital','invested_capital',MILLIONS)):
        merged(5, label, start, start+2, bold=True, size=10, fill=PALE, wrap=False, height=23)
        c = merged(6, total[key], start, start+2, bold=True, size=23, fill=PALE, wrap=False, height=36)
        c.number_format = fmt
        internal_link(c, 'Investment Ledger', 'A5')
        merged(7, 'Synthetic sponsor data', start, start+2, color=MUTED, size=9, fill=PALE, height=19, wrap=False)
    ws.row_dimensions[8].height = 9
    heading(9, 'Portfolio Investment Assessment')
    body(10, returns)
    body(11, operating)
    ws.row_dimensions[12].height = 9
    heading(13, 'Holdings Positioning · Illustrative Sponsor Returns')
    for start, end, label in ((1,4,'Company'),(5,6,'Current Equity'),(7,8,'Gross MOIC'),(9,10,'Gross IRR'),(11,12,'Position')):
        merged(HOLDINGS_HEAD, label, start, end, bold=True, size=10, fill=NAVY, color='FFFFFF', wrap=False, height=24)
    ranked = sorted(positions, key=lambda p: (-p[1]['gross_moic'], p[0].id))
    for row, (company, result) in enumerate(positions, HOLDINGS_HEAD+1):
        role = ('Single holding' if len(ranked)==1 else
                ('Strongest MOIC' if company.id==ranked[0][0].id else ('Weakest MOIC' if company.id==ranked[-1][0].id else 'Positive MOIC' if result['gross_moic']>1 else 'At / below cost')))
        c = merged(row, company.name, 1, 4, bold=True, size=10, height=34)
        internal_link(c, company.name if company.name in wb else {'summit':'Summit Home Services','tallis':'Tallis Brook Accounting','harrow':'Harrow Ridge Environmental'}[company.id], 'A1')
        for start, key, fmt in ((5,'current_equity_value',MILLIONS),(7,'gross_moic','0.00"x"'),(9,'gross_irr','0.00%;(0.00%);"—"')):
            cell = merged(row, result[key], start, start+1, size=11, height=34)
            cell.number_format = fmt
            internal_link(cell, 'Investment Ledger', 'A'+str(6+row-HOLDINGS_HEAD-1))
        merged(row, role, 11, 12, size=9, color=MUTED, height=34)
    row = HOLDINGS_HEAD+1+len(positions)
    body(row, 'Current equity excludes realized proceeds. Total value equals realized proceeds plus current equity; gross returns are before fees, carry and fund expenses.', color=MUTED)
    row += 1
    heading(row, 'Portfolio Priorities · Operating Risks & Management Focus')
    for company in companies:
        row += 1
        cell = merged(row, company.name+' · '+comparison(company)['period_label'], bold=True, size=10, height=24)
        first = next(r for r in range(6, wb['Investment Issues'].max_row+1) if wb['Investment Issues'].cell(r,1).value==company.name)
        internal_link(cell, 'Investment Issues', 'A'+str(first))
        outlook = investment_outlook(company)
        risk = '. '.join(outlook.paragraphs[2].split('. ')[:2])
        if not risk.endswith('.'): risk += '.'
        priority = outlook.paragraphs[3].split('. ')[0]
        if not priority.endswith('.'): priority += '.'
        row += 1
        body(row, risk+' Priority: '+priority)
    row += 1
    heading(row, 'Review & Evidence')
    row += 1
    body(row, f'{workflow_summary(companies)["issues"]} complete accepted findings, management questions and qualifications remain in Investment Issues. Source References links to original worksheet cells and calculation operands. Proposed operating initiatives are illustrative, not established causes.', color=MUTED)
    row += 1
    for start, end, title, target in ((1,4,'Complete Investment Findings →','Investment Issues'),(5,8,'Income Statements →','Portfolio Financial Performance'),(9,12,'Ledger & Assumptions →','Investment Ledger')):
        c = merged(row, title, start, end, bold=True, size=10, height=25)
        internal_link(c, target, 'A1')
    row += 1
    body(row, 'All companies, investments and financial figures are synthetic. Sponsor valuations are illustrative assumptions, separate from source-verified operating observations. No net fund or LP returns are calculated.', color=MUTED)
    ws.freeze_panes = 'A9'
    ws.print_area = 'A1:L'+str(row)
    ws.page_setup.fitToHeight = 1
    return ws


def polish_statement(ws, companies):
    review_layout(ws, landscape=True)
    ws.page_setup.fitToHeight = 1
    ws.print_area = 'A1:'+get_column_letter(1+3*len(companies))+'21'
    ws.print_title_rows = '1:6'
    for cells in ws:
        for cell in cells:
            if cell.value is None: continue
            font = copy(cell.font); font.name = 'Calibri'
            if 6 <= cell.row <= 17: font.sz = 11
            cell.font = font
    for index, company in enumerate(companies):
        col = 2+index*3
        for row in (4,5):
            ws.cell(row,col).alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
        for j in range(col,col+3):
            ws.cell(6,j).alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
    ws.row_dimensions[4].height = 30; ws.row_dimensions[5].height = 27; ws.row_dimensions[6].height = 27
    ws['A6'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A6'].font = Font(name='Calibri',size=11,bold=True,color='FFFFFF')


def polish_ledger(ws):
    review_layout(ws, landscape=True)
    for cells in ws:
        for cell in cells:
            if cell.value is None: continue
            font = copy(cell.font); font.name = 'Calibri'; font.sz = 11 if cell.row != 1 else 21
            cell.font = font
    widths = (28,18,18,21,19,13,14,23,33,30)
    # Preserve the original grid and every cash-flow/provenance coordinate.
    for col, width in enumerate(widths,1): ws.column_dimensions[get_column_letter(col)].width = width
    for row in (5,):
        for cell in ws[row]: cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    for row in range(6,10):
        for col in range(1,11):
            c = ws.cell(row,col)
            c.alignment = Alignment(horizontal='left' if col in (1,9,10) else 'right', vertical='center', wrap_text=True)
            if row == 9:
                c.fill = PatternFill('solid',fgColor=NAVY)
                c.font = Font(name='Calibri',size=11,bold=True,color='FFFFFF')
        ws.row_dimensions[row].height = 60
    # Refit wrapped supporting text to the narrower, readable print layout.
    # Merged narrative rows use their full span, not the first column's width.
    for cells in ws:
        if cells[0].row < 11: continue
        if not any(cell.value is not None for cell in cells):
            ws.row_dimensions[cells[0].row].height = 8
            continue
        lines = 1
        for cell in cells:
            if cell.value is None: continue
            span = next((area for area in ws.merged_cells.ranges if cell.coordinate in area), None)
            width = sum(widths[col-1] for col in range(span.min_col,span.max_col+1)) if span else widths[cell.column-1]
            lines = max(lines, sum(max(1,len(textwrap.wrap(part,width=max(10,int(width*.9))))) for part in str(cell.value).splitlines()))
        ws.row_dimensions[cells[0].row].height = max(ws.row_dimensions[cells[0].row].height or 0, lines*14+12)
    for cells in ws:
        if isinstance(cells[0].value, date): ws.row_dimensions[cells[0].row].height = 52
    for row in range(11,15):
        ws.cell(row,1).alignment = Alignment(vertical='center', wrap_text=False)
    starts = [cell.row for cells in ws for cell in cells if cell.column==1 and cell.value=='Company investment']
    ws.row_breaks = RowBreak()
    for start in (17,*starts): ws.row_breaks.append(Break(id=start-1))
    ws.print_title_rows = '1:2'
    ws.print_area = 'A1:J'+str(ws.max_row)


def polish_gp_workbook(wb, companies):
    write_investment_review(wb, companies)
    polish_statement(wb['Portfolio Financial Performance'], companies)
    polish_ledger(wb['Investment Ledger'])
    wb.active = 0
