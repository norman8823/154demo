"""One shared native source viewer for library, mappings and accepted findings."""
import streamlit as st
from openpyxl.utils import get_column_letter
from engine import InputError
from portfolio_documents import (original_document, source_books, checked_coordinate,
    source_fact, evidence_locations, cell_record, display_value, worksheet_html)

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def open_source(company, sheet=None, cell=None):
    try:
        document = original_document(company)
        book, _ = source_books(document)
        if sheet is not None:
            if sheet not in book.sheetnames: raise InputError('The cited worksheet is not in the original workbook.')
            checked_coordinate(book[sheet], cell)
        default_sheet=next((s.title for s in book if s.max_column>1),book.sheetnames[0])
        serial=st.session_state.get('source_open_serial',0)+1
        st.session_state.source_open_serial=serial
        source_dialog(company, sheet or default_sheet, cell or 'A1',serial)
    except (InputError, OSError, ValueError) as exc:
        st.error('Source preview unavailable: '+str(exc))


@st.dialog('Original source workbook', width='large', dismissible=True)
def source_dialog(company, requested_sheet, requested_cell, request_id):
    try:
        document = original_document(company)
        book, cached = source_books(document)
        st.subheader(company.name)
        st.caption(document.filename+' · Read-only original · SHA-256 verified')
        prefix = 'source_'+str(request_id)+'_'+company.id+'_'+requested_sheet+'_'+requested_cell
        name = st.selectbox('Worksheet', book.sheetnames, index=book.sheetnames.index(requested_sheet), key=prefix+'_sheet')
        sheet = book[name]
        focus = requested_cell if name == requested_sheet else 'A1'
        coordinate = st.text_input('Inspect cell', value=focus, key=prefix+'_'+name+'_cell', help='Enter an Excel coordinate to inspect its original value and nearby cells.').strip().upper()
        row, col = checked_coordinate(sheet, coordinate)
        anchor = prefix+'_'+name+'_'+coordinate
        left, right = st.columns(2)
        first_row = int(left.number_input('First row', min_value=1, max_value=sheet.max_row, value=max(1,row-5), step=12, key=anchor+'_row'))
        first_col = int(right.number_input('First column', min_value=1, max_value=sheet.max_column, value=max(1,col-2), step=4, key=anchor+'_col', help='1 = A, 2 = B, etc. Four columns are shown per view.'))
        last_row=min(sheet.max_row,first_row+11); last_col=min(sheet.max_column,first_col+3)
        st.caption(f'{name} · {get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row} · {sheet.max_row} rows × {sheet.max_column} columns in original')
        st.markdown(worksheet_html(company,sheet,cached[name],first_row,first_col,coordinate),unsafe_allow_html=True)
        if not (first_row<=row<=last_row and first_col<=col<=last_col): st.info('The inspected cell is outside this range. Change the first row / column to return to it.')
        record=cell_record(company,sheet,cached[name],coordinate)
        st.markdown('**'+name.replace('$','\\$')+' · '+coordinate+'**')
        st.text('Source label: '+str(record['label']))
        st.text('Original cell value: '+display_value(record['value'],record['number_format']))
        st.caption(record['origin']+' · Original number format: '+record['number_format'])
        if record['fact'] and record['fact'].get('value') is not None:
            fact=record['fact']
            st.caption('Source-unit multiplier: '+str(fact.get('source_scale',1))+' · Verified standardized value: '+str(fact.get('value'))+' '+str(fact.get('currency',''))+' · '+str(fact.get('period','')))
        if record['formula']:
            with st.expander('Original formula text'):
                st.code(record['formula'],language=None)
                st.caption('Retained for inspection only; formulas are never executed in this viewer.')
        if record['comment']:
            with st.expander('Original cell comment'):
                st.text(record['comment'])
        st.caption('Bounded worksheet reconstruction; original coordinates and values are retained. Merged layouts and Excel artwork are not reproduced.')
        download,close=st.columns(2)
        download.download_button('Download Original',document.data,file_name=document.filename,mime=XLSX,on_click='ignore',key=prefix+'_download',width='stretch')
        if close.button('Close',key=prefix+'_close',width='stretch'): st.rerun()
    except (InputError,OSError,ValueError) as exc:
        st.error('Source preview unavailable: '+str(exc))
        if st.button('Close',key='invalid_source_close'): st.rerun()


def source_action(company, fact_id, key):
    try:
        fact=source_fact(company,fact_id)
    except (InputError,OSError,ValueError) as exc:
        st.caption('Original source unavailable: '+str(exc)); return
    if st.button('View in Source',key=key): open_source(company,fact['sheet'],fact['cell'])


def evidence_actions(company, source_ids, key):
    try: originals,derived,unsupported=evidence_locations(company,source_ids)
    except (InputError,OSError,ValueError) as exc:
        st.warning('Original sources unavailable: '+str(exc)); return
    if derived:
        with st.expander('Derived calculations & original operands'):
            for calculation in derived:
                st.text(calculation['label'])
                st.code(calculation.get('expression',calculation.get('raw','')),language=None)
                st.caption('Inputs: '+', '.join(calculation['inputs'])+' · '+calculation.get('limitation',''))
            st.caption('Calculated values are not original worksheet cells. Inspect their recorded source operands below.')
    if originals:
        index={f['id']:f for f in originals}
        identity=st.selectbox('Original source cell',list(index),format_func=lambda i:index[i]['ref']+' · '+(index[i].get('label') or i),key=key+'_operand')
        source_action(company,identity,key+'_view')
    if unsupported: st.caption(str(len(unsupported))+' supporting records have no supported single-cell location; their original evidence remains available above.')


def render_documents(companies, workbook_data):
    from portfolio_ui import page_heading
    from portfolio_comparison import comparison
    from portfolio_company import company_export
    page_heading('Documents & Reports','Original company reporting packages · Source-backed reporting',None)
    st.subheader('Source Documents')
    st.caption('Approved original workbooks, verified against the prepared-analysis manifest before preview or download.')
    for company in companies:
        with st.container(border=True,key='source_document_'+company.id):
            try: document=original_document(company)
            except (InputError,OSError,ValueError) as exc:
                st.error(company.name+' · '+str(exc));continue
            title,preview,download=st.columns([3,1,1.5],vertical_alignment='center')
            title.markdown('**'+company.name+'**')
            title.caption(document.filename+' · '+comparison(company)['period_label'])
            if preview.button('Preview',key='preview_'+company.id,width='stretch'): open_source(company)
            download.download_button('Download Original',document.data,file_name=document.filename,mime=XLSX,on_click='ignore',key='original_'+company.id,width='stretch')
    st.subheader('Generated Reports')
    st.caption('Generated portfolio and company analyses. Original reporting packages are listed separately above.')
    st.download_button('↓ Download Portfolio Excel',workbook_data,file_name='GP_Portfolio_Review.xlsx',mime=XLSX,key='documents_gp_download')
    ids={c.id:c for c in companies}
    selected=st.selectbox('Company report',list(ids),format_func=lambda identity:ids[identity].name,key='documents_company')
    company=ids[selected]
    choice=st.selectbox('Company Excel format',['Standard Excel report','Use supplied firm template'],key='documents_format')
    custom=choice=='Use supplied firm template'
    template=st.file_uploader('Optional company Excel template',type=['xlsx'],key='documents_template') if custom else None
    if template:
        from reports import export_report
        try: data=export_report(company.review,template.getvalue())
        except Exception:
            st.error('This template could not be populated. Choose the standard report or supplied firm template.');return
    else: data=company_export(company.review,custom)
    st.download_button('Download Company Excel',data[0],file_name=data[1],mime=XLSX,key='documents_company_download')
