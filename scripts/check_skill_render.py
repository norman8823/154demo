"""Offline Streamlit + Excel check against a completed benchmark receipt."""
import argparse
import io
import json
import copy
import os
from pathlib import Path
import sys
import urllib.request
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from ai_runtime import load_environment
from engine import ROOT
from streamlit.testing.v1 import AppTest
from openpyxl import load_workbook


def main():
    parser=argparse.ArgumentParser();parser.add_argument('workbook',type=Path);parser.add_argument('run',type=Path)
    parser.add_argument('--offline-replay',action='store_true')
    parser.add_argument('--company-selector',action='store_true')
    args=parser.parse_args();load_environment()
    os.environ['COPILOT_ANALYST_ENGINE']='skill'
    if args.offline_replay:
        import engine
        from report_cache import unpack,reused
        fixture=engine.Review(**unpack(json.loads((args.run/'review-packed.json').read_text())))
        assert fixture.package['skill_engine'].get('offline_replay')
        engine.analyze=lambda *a,**k:reused(copy.deepcopy(fixture),'offline-replay-fixture')
    def no_network(*a,**k):raise AssertionError('Rendering/checking must never spend or call the model.')
    urllib.request.urlopen=no_network
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=45)
    if not args.company_selector:at.session_state['uploaded_inputs']={args.workbook.name:args.workbook.read_bytes()}
    at.run();assert not at.exception,at.exception
    if args.company_selector:
        next(x for x in at.selectbox if x.label=='Company workbook').set_value(args.workbook.name).run()
        next(b for b in at.button if b.label=='Load company workbook').click().run()
        assert not at.exception,at.exception
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception,at.exception
    r=at.session_state['review'];assert r.package['skill_engine']['status']=='completed',r.package['skill_engine']['reason']
    assert 'Skill analyst' in r.narrative_mode or args.offline_replay
    assert r.package['inference'].get('reused') and r.package['inference']['attempts']==0
    from executive import present
    titles=[c['title'] for c in present(r)['cards']]
    dashboard='\n'.join(str(m.value) for m in at.markdown)
    assert all(t in dashboard for t in titles)
    for key in ['standard','custom']:
        data=(args.run/(key+'.xlsx')).read_bytes();w=load_workbook(io.BytesIO(data))
        assert {'Findings','Source Evidence','Analyst Execution','Financial Performance'}<=set(w.sheetnames)
        assert all(any(t in str(row) for row in w['Findings'].values) for t in titles)
        assert not any(c.data_type=='e' for s in w for row in s for c in row)
        assert any(c.hyperlink for row in w['Findings'] for c in row)
    next(b for b in at.button if b.label.startswith('View Evidence')).click().run()
    assert not at.exception,at.exception
    # Change template via the actual setup controls, then Generate uses same cache.
    at.radio(key='view').set_value('01  Upload & Setup').run()
    at.radio(key='report_format').set_value('Use existing firm template').run()
    next(c for c in at.checkbox if c.label=='Use supplied sample template').check().run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception,at.exception
    assert at.session_state['review'].package['inference']['attempts']==0
    w=load_workbook(io.BytesIO(at.session_state['export'][0]))
    assert 'Partner Review' in w.sheetnames and 'Analyst Execution' in w.sheetnames
    assert w['Partner Review']['C16'].data_type=='f'
    result=dict(company_selector=args.company_selector,mode='offline_replay' if args.offline_replay else 'completed_provider_cache',streamlit_dashboard='passed',evidence_navigation='passed',custom_template_generation='passed',
                standard_excel='passed',custom_excel='passed',new_model_calls=0,priority_titles=titles,
                existing_template_formula_preserved=True)
    (args.run/'render-check.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))


if __name__=='__main__':main()
