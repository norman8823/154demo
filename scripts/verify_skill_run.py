"""Verify a saved fresh application run without model calls or comparator inputs."""
import argparse
import io
import json
from pathlib import Path
import sys
import urllib.request
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from engine import Review
from report_cache import unpack,successful
from ai_runtime import load_environment
from skill_tools import WorkbookTools
from openpyxl import load_workbook


def main():
    p=argparse.ArgumentParser();p.add_argument('workbook',type=Path);p.add_argument('run',type=Path);a=p.parse_args();load_environment()
    def no_network(*a,**k):raise AssertionError('Verification must not call the provider.')
    urllib.request.urlopen=no_network
    review=Review(**unpack(json.loads((a.run/'review-packed.json').read_text())))
    state=review.package['skill_engine'];inf=review.package['inference']
    assert not inf.get('reused') and not state.get('offline_replay')
    assert any(e.get('provider_completed') and e.get('transport')=='provider' for e in inf['events'])
    t=WorkbookTools({a.workbook.name:a.workbook.read_bytes()})
    for e in state['tool_events']:
        try:t.invoke(e['name'],e['arguments'])
        except Exception:
            if e['status']!='error':raise
    for f in t.facts:t.verify_fact(f)
    assert t.calculations==state['calculations']
    refs={k for f in review.package['ai_insights'] for k in f['fact_ids']}
    assert refs<=t.facts.keys()|t.notes.keys()
    for f in review.package['ai_insights']:
        for q in f['source_quotes']:assert q['quote'] in t.notes[q['source_id']]['text']
    sheets={}
    for name in ['standard','custom']:
        w=load_workbook(io.BytesIO((a.run/(name+'.xlsx')).read_bytes()))
        assert {'Findings','Source Evidence','Analyst Execution','Financial Performance'}<=set(w.sheetnames)
        assert not any(c.data_type=='e' for s in w for row in s for c in row)
        for f in review.package['ai_insights']:assert any(f['title'] in str(row) for row in w['Findings'].values)
        assert any(c.hyperlink for row in w['Findings'] for c in row)
        if name=='custom':assert w['Partner Review']['C16'].data_type=='f'
        sheets[name]=w.sheetnames
    result=dict(status=state['status'],fresh_provider_run=True,successful=successful(review),numeric_facts_recomputed=len(t.facts),calculations_recomputed=len(t.calculations),cited_numeric_ids=len(refs&t.facts.keys()),cited_text_ids=len(refs&t.notes.keys()),source_errors=t.errors,coverage=t.coverage(),excel_sheets=sheets,new_model_calls=0,narrative_assurance=False)
    (a.run/'verification.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))


if __name__=='__main__':main()
