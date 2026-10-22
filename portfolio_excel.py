"""Operational GP workbook for review, management follow-up and source handoff."""
import io
import math
import textwrap
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from portfolio_comparison import comparison, metric_operand, COMPARABILITY
from portfolio_dashboard import workflow_summary
from portfolio_watchlist import watchlist, source_closure, evidence_for

NAVY, MUTED, PALE = '18364A', '60798A', 'EFF5F7'
MONEY = '"$"#,##0;[Red]("$"#,##0);"$0"'
PERCENT = '0.0%;[Red](0.0%);"0.0%"'
PP = '0.0" pp";[Red](0.0" pp");"0.0 pp"'
SHEETS = ('Portfolio Overview','Company Comparison','Investment Issues','Summit Home Services',
          'Tallis Brook Accounting','Harrow Ridge Environmental','Source References','Source Snapshots','Reporting Output','Investment Ledger','Portfolio Financial Performance')
COMPANY_SHEETS = dict(zip(('summit','tallis','harrow'), SHEETS[3:6]))
OVERVIEW_HEAD = 28


def literal(ws, row, column, value):
    c = ws.cell(row,column,value if value is not None else 'Unavailable')
    if isinstance(c.value,str): c.data_type = 's'
    c.font = Font(name='Arial',size=10,color=NAVY)
    c.alignment = Alignment(vertical='top',wrap_text=True)
    return c


def link(cell, sheet, row, label=None, column='A'):
    location="'"+sheet.replace("'","''")+"'!"+column+str(row)
    cell.hyperlink=Hyperlink(ref=cell.coordinate,location=location)
    if label is not None:
        cell.value=label;cell.data_type='s'
    cell.font=Font(name='Arial',size=10,color='206482',underline='single')


def title_sheet(wb,name,widths,description):
    ws=wb.create_sheet(name)
    for col,width in enumerate(widths,1):ws.column_dimensions[ws.cell(1,col).column_letter].width=width
    for row,text in [(1,name),(2,'SYNTHETIC PREPARED REVIEW · For Investment Team Review'),(3,description)]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=len(widths))
        c=literal(ws,row,1,text);c.font=Font(name='Arial',size=21 if row==1 else 10,bold=row==1,color=NAVY if row==1 else MUTED)
    ws.row_dimensions[1].height=32;ws.row_dimensions[2].height=22;ws.row_dimensions[3].height=44
    ws.sheet_view.showGridLines=False;ws.freeze_panes='C6';ws.sheet_properties.tabColor=NAVY
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.page_margins.left=.25;ws.page_margins.right=.25
    ws.oddFooter.center.text='Synthetic prepared assessments · Page &P of &N'
    return ws


def headers(ws,row,labels):
    for col,text in enumerate(labels,1):
        c=literal(ws,row,col,text);c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
    ws.row_dimensions[row].height=30


def row_values(ws,row,values,height=36):
    for col,value in enumerate(values,1):
        c=literal(ws,row,col,value)
        if row%2==0:c.fill=PatternFill('solid',fgColor=PALE)
        c.border=Border(bottom=Side(style='hair',color='DCE6EB'))
    ws.row_dimensions[row].height=height


def narrative(ws,row,label,text,last=None):
    last=last or ws.max_column
    literal(ws,row,1,label).font=Font(name='Arial',size=10,bold=True,color=NAVY)
    ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=last)
    literal(ws,row,2,text)
    # Leave enough height for original prose; no truncation or financial rewriting.
    chars=max(60,int(sum(ws.column_dimensions[get_column_letter(c)].width for c in range(2,last+1))*.85))
    lines=sum(max(1,len(textwrap.wrap(p,width=chars))) for p in str(text).splitlines())
    ws.row_dimensions[row].height=min(409,max(28,lines*14+12))


def reference_records(companies):
    items=watchlist(companies);output=[]
    for c in companies:
        v=comparison(c)
        ids=[s for key in ('revenue','ebitda','margin') for s in v[key].source_ids]
        from portfolio_financial_reporting import statement
        ids.extend(s for line in statement(c.review) for s in line.actual_refs+line.budget_refs)
        ids.extend(s for item in items if item.company_id==c.id for s in item.source_ids)
        ids=source_closure(c.review,ids)
        evidence=evidence_for(c.review,ids)
        if set(ids)!={e['Fact ID'] for e in evidence}:raise ValueError('A requested source reference is unavailable: '+c.name)
        output.extend((c,e) for e in evidence)
    return output


def write_references(ws,records):
    headers(ws,5,['Company','Source ID','Type','Original reference / calculation','Validated value','Original value / formula','Check / period / qualification','Calculation inputs','Canonical source workbook','Prepared run'])
    rows={}
    for row,(c,e) in enumerate(records,6):
        calculations={f['id']:f for f in c.review.package['skill_engine']['calculations']}
        calc=calculations.get(e['Fact ID'])
        vals=[c.name,e['Fact ID'],'Derived calculation' if calc else 'Source cell',e['Reference'],e['Validated value'],str(e['Original value / formula']),
              e['Check / use'],', '.join(calc['inputs']) if calc else '',c.provenance['workbook'],c.run]
        row_values(ws,row,vals,min(300,max(48,math.ceil(max(len(str(vals[5]))/62,len(str(vals[6]))/88))*14+12)))
        rows[(c.id,e['Fact ID'])]=row
    ws.auto_filter.ref=f'A5:J{5+len(records)}'
    return rows


def source_link(cell,company_id,ids,rows):
    if ids:
        target=rows.get((company_id,ids[0]))
        if target is None:raise ValueError('Source index target is missing.')
        link(cell,'Source References',target)


def write_overview(ws,companies,source_rows):
    summary=workflow_summary(companies)
    for start,label,value in [(1,'Reporting coverage',f'{summary["received"]} / {summary["companies"]}'),(4,'Below EBITDA plan',f'{summary["below_plan"]} / {summary["comparable_budget"]}'),
                               (7,'Accepted investment issues',summary['issues']),(10,'Act-now follow-ups',summary['act_now'])]:
        ws.merge_cells(start_row=5,start_column=start,end_row=5,end_column=start+2)
        ws.merge_cells(start_row=6,start_column=start,end_row=7,end_column=start+2)
        literal(ws,5,start,label).font=Font(name='Arial',size=10,color=MUTED)
        c=literal(ws,6,start,value);c.font=Font(name='Arial',size=27,bold=True,color=NAVY);c.fill=PatternFill('solid',fgColor=PALE)
    ws.row_dimensions[6].height=30;ws.row_dimensions[7].height=15
    ws.merge_cells('A9:L10');literal(ws,9,1,COMPARABILITY);ws.row_dimensions[9].height=32
    ws.merge_cells('A12:L12');literal(ws,12,1,'Reported EBITDA — Actual versus Budget · independent company-period axes').font=Font(name='Arial',size=12,bold=True,color=NAVY)
    # Visible chart inputs are also the auditable data behind native Excel charts.
    headers(ws,42,['Company','Actual','Budget','Reporting period','Source IDs'])
    for row,c in enumerate(companies,43):
        v=comparison(c);m=v['ebitda'];row_values(ws,row,[c.name,m.actual,m.budget,v['period_label'],', '.join(m.source_ids)],44)
        for col in (2,3):ws.cell(row,col).number_format=MONEY
        source_link(ws.cell(row,5),c.id,m.source_ids,source_rows)
        chart=BarChart();chart.type='col';chart.style=10;chart.title=c.name+'\n'+v['period_label'];chart.y_axis.title='Reported EBITDA · source dollars'
        chart.y_axis.scaling.min=0;chart.legend=None;chart.width=11.2;chart.height=7.2
        chart.add_data(Reference(ws,min_col=2,max_col=3,min_row=row,max_row=row),from_rows=True,titles_from_data=False)
        chart.set_categories(Reference(ws,min_col=2,max_col=3,min_row=42,max_row=42))
        for idx,color in enumerate(('173E57','B8CBD9')):
            point=DataPoint(idx=idx);point.graphicalProperties.solidFill=color;point.graphicalProperties.line.noFill=True;chart.series[0].data_points.append(point)
        chart.dataLabels=DataLabelList();chart.dataLabels.showVal=True;chart.dataLabels.dLblPos='outEnd'
        chart.dataLabels.showLegendKey=False;chart.dataLabels.showCatName=False;chart.dataLabels.showSerName=False
        chart.dataLabels.showPercent=False;chart.dataLabels.showLeaderLines=False
        chart.dataLabels.numFmt='"$"#,##0'
        ws.add_chart(chart,('A14','E14','I14')[row-43])
    labels=['Company','Reporting period','Revenue','Reported EBITDA','EBITDA budget','EBITDA variance','EBITDA vs budget','EBITDA margin','Issues','Act now','Company review','Source IDs']
    headers(ws,OVERVIEW_HEAD,labels)
    for row,c in enumerate(companies,OVERVIEW_HEAD+1):
        v=comparison(c);m=v['ebitda'];s=workflow_summary([c])
        row_values(ws,row,[c.name,v['period_label'],v['revenue'].actual,m.actual,m.budget,m.variance,m.variance_pct,v['margin'].actual,s['issues'],s['act_now'],'Open company →',', '.join(m.source_ids)],64)
        for col in (3,4,5,6):ws.cell(row,col).number_format=MONEY
        for col in (7,8):ws.cell(row,col).number_format=PERCENT
        link(ws.cell(row,1),COMPANY_SHEETS[c.id],1);link(ws.cell(row,11),COMPANY_SHEETS[c.id],1)
        source_link(ws.cell(row,12),c.id,m.source_ids,source_rows)
    ws.auto_filter.ref=f'A{OVERVIEW_HEAD}:L{OVERVIEW_HEAD+len(companies)}';ws.freeze_panes='C29'
    narrative(ws,34,'Review workflow','1. Identify an operating exception. 2. Open the company analysis. 3. Use Investment Issues for the management question and proposed follow-up. 4. Follow Source References to verify the original source. 5. Copy Reporting Output into your existing reporting process.',12)
    narrative(ws,36,'Prepared results','0 new analyst model calls / $0 current inference. Source and arithmetic checks preserve traceability; interpretations, causal explanations and illustrative sensitivities still require investment-team review.',12)
    narrative(ws,38,'Reporting Output','Flat, numeric, source-referenced rows for copying or mapping into an existing template. No automatic platform integration or compatibility is claimed.',12)
    link(ws.cell(38,1),'Reporting Output',1)
    ws.print_area='A1:L39'


def write_comparison(ws,companies,source_rows):
    headers(ws,5,['Financial attribute','Unit / definition']+[c.name for c in companies])
    specs=[('Reporting period','Explicit company period',None,None),('Revenue','Source dollars','revenue','actual'),('Revenue budget','Same period / source dollars','revenue','budget'),
           ('Revenue vs budget','Actual / budget − 1','revenue','variance_pct'),('Reported EBITDA','Source dollars; excludes adjusted sensitivities','ebitda','actual'),
           ('EBITDA budget','Same period / source dollars','ebitda','budget'),('EBITDA variance','Actual − budget / source dollars','ebitda','variance'),('EBITDA vs budget','Actual / budget − 1','ebitda','variance_pct'),
           ('EBITDA margin','Reported EBITDA / revenue','margin','actual'),('Budget EBITDA margin','Budget EBITDA / budget revenue','margin','budget'),('Margin variance','Percentage points','margin','variance'),
           ('Accounting basis','Unknown remains unknown',None,'basis'),('Prepared run','Approved Review identity',None,'run'),('Company review','Source-backed findings and follow-up',None,'link')]
    for row,(label,definition,key,field) in enumerate(specs,6):
        row_values(ws,row,[label,definition],38)
        for col,c in enumerate(companies,3):
            v=comparison(c)
            value=getattr(v[key],field) if key else (v['period_label'] if field is None else (v['basis'] if field=='basis' else (c.run if field=='run' else 'Open company →')))
            cell=literal(ws,row,col,value)
            if key:
                cell.number_format=(PP if field=='variance' else PERCENT) if key=='margin' else (PERCENT if field=='variance_pct' else MONEY)
                source_link(cell,c.id,v[key].source_ids,source_rows)
            elif field=='link':link(cell,COMPANY_SHEETS[c.id],1)
    ws.auto_filter.ref='A5:E19';ws.freeze_panes='C7'
    for col,c in enumerate(companies,3):
        link(ws.cell(5,col),COMPANY_SHEETS[c.id],1)
        ws.cell(5,col).font=Font(name='Arial',size=10,bold=True,color='FFFFFF',underline='single')


def attach_metric_proofs(wb, companies, source_rows):
    """Numeric comparison links explain reporting arithmetic and exact scenarios."""
    ws=wb['Source References'];cursor=ws.max_row+5
    specs=[(7,'revenue','actual'),(8,'revenue','budget'),(9,'revenue','variance_pct'),
           (10,'ebitda','actual'),(11,'ebitda','budget'),(12,'ebitda','variance'),(13,'ebitda','variance_pct'),
           (14,'margin','actual'),(15,'margin','budget'),(16,'margin','variance')]
    for col,c in enumerate(companies,3):
        values=comparison(c)
        for matrix_row,key,field in specs:
            value=getattr(values[key],field)
            if value is None:
                wb['Company Comparison'].cell(matrix_row,col).hyperlink=None
                continue
            scenarios=('actual','budget') if field.startswith('variance') else (field,)
            keys=('ebitda','revenue') if key=='margin' else (key,)
            expression=('(Actual EBITDA / actual revenue − budget EBITDA / budget revenue) × 100; percentage points' if key=='margin' and field=='variance' else
                        ('EBITDA / revenue' if key=='margin' else ('Actual / budget − 1' if field=='variance_pct' else ('Actual − budget' if field=='variance' else 'Validated '+field+' statement metric'))))
            start=cursor
            row_values(ws,start,[c.name,key+' · '+field,'REPORTING METRIC',expression,value,values[key].unit,values['period_label'],
                                 'Presentation arithmetic over approved statement metrics; not a new analyst finding.','','↑ Comparison'],76)
            link(ws.cell(start,10),'Company Comparison',matrix_row,column=get_column_letter(col))
            ws.cell(start,5).number_format=PP if key=='margin' and field=='variance' else (PERCENT if key=='margin' or field=='variance_pct' else MONEY)
            headers(ws,start+1,['Company','Metric / scenario','Type','Statement definition','Operand value','Unit','Period','Source ID','Open source index',''])
            cursor=start+2
            for scenario in scenarios:
                for operand in keys:
                    support=metric_operand(c,operand,scenario)
                    if support is None:raise ValueError('A reporting calculation operand is unavailable.')
                    amount,unit,definition,ids=support
                    for sid in ids:
                        row_values(ws,cursor,[c.name,operand+' · '+scenario,'Approved statement operand',definition,amount,unit,values['period_label'],sid,'Open original evidence →',''],52)
                        source_link(ws.cell(cursor,9),c.id,[sid],source_rows)
                        ws.cell(cursor,5).number_format=MONEY;cursor+=1
            link(wb['Company Comparison'].cell(matrix_row,col),'Source References',start)
            cursor+=3


def write_issues(ws,companies,source_rows):
    headers(ws,5,['Company','Report period','Finding ID','Original priority','Disposition','Investment issue','Source-supported impact','Management question','Proposed follow-up','Caveat / qualification','Company review','Source IDs'])
    ordered=[item for c in companies for item in watchlist(companies) if item.company_id==c.id]
    details={};detail_row=10+len(ordered)
    for row,item in enumerate(ordered,6):
        f=item.finding
        vals=[item.company_name,item.report_period,f['id'],f['priority'],f['disposition'],f['title'],f.get('impact','Not quantified'),f['management_question'],f['action'],item.presentation['caveat'],'Open company →',', '.join(f['fact_ids'])]
        height=min(260,max(64,max(len(textwrap.wrap(str(vals[col]),width=width)) for col,width in [(5,43),(7,48),(8,54),(9,48)])*13+12))
        row_values(ws,row,vals,height)
        link(ws.cell(row,11),COMPANY_SHEETS[item.company_id],1);source_link(ws.cell(row,12),item.company_id,f['fact_ids'],source_rows)
        details[(item.company_id,f['id'])]=detail_row
        link(ws.cell(row,6),'Investment Issues',detail_row)
        fields=[('Finding',f['title']),('Company / reporting package',item.company_name+' · '+item.report_period),('Original classification',f['classification']),
                ('Evidence status',f['evidence_status']),('Original priority / disposition',f['priority']+' · '+f['disposition']),('Conclusion',f['conclusion']),
                ('Accepted finding narrative',f['text']),('Evidence',f['evidence']),('Financial implication',f['financial_implication']),('Investment assessment',f['assessment']),
                ('Management question',f['management_question']),('Proposed follow-up',f['action']),('Caveat',item.presentation['caveat'])]
        fields.extend(('Source quote · '+q['source_id'],q['quote']) for q in f.get('source_quotes',[]))
        for label,value in fields:
            narrative(ws,detail_row,label,value,12)
            if label=='Finding':
                link(ws.cell(detail_row,1),'Investment Issues',row,'↑ Issue index')
                ws.cell(detail_row,2).font=Font(name='Arial',size=13,bold=True,color=NAVY)
                ws.cell(detail_row,2).fill=PatternFill('solid',fgColor='DCEAF0')
            detail_row+=1
        for sid in f['fact_ids']:
            narrative(ws,detail_row,'Source reference',sid,12);source_link(ws.cell(detail_row,2),item.company_id,[sid],source_rows);detail_row+=1
        detail_row+=2
    ws.auto_filter.ref=f'A5:L{5+len(ordered)}'
    return details


def write_company(ws,c,items,source_rows,details):
    headers(ws,7,['Financial metric','Actual','Budget','Variance / pp','Units / definition','Source IDs'])
    v=comparison(c)
    for row,(key,label) in enumerate([('revenue','Revenue'),('ebitda','Reported EBITDA'),('margin','EBITDA margin')],8):
        m=v[key];row_values(ws,row,[label,m.actual,m.budget,m.variance,m.definition,', '.join(m.source_ids)],56)
        for col in (2,3):ws.cell(row,col).number_format=PERCENT if key=='margin' else MONEY
        ws.cell(row,4).number_format=PP if key=='margin' else MONEY
        source_link(ws.cell(row,6),c.id,m.source_ids,source_rows)
    narrative(ws,5,'Business context',c.review.package.get('company_profile',{}).get('description','See approved company package.'),6)
    narrative(ws,12,'All accepted findings',str(len(items))+' findings with original questions, actions, caveats and sources in Investment Issues.',6)
    link(ws.cell(12,2),'Investment Issues',next(r for (slug,_),r in details.items() if slug==c.id))
    row=14
    for item in items[:3]:
        f=item.finding
        for label,value in [('Priority finding',f['title']),('Classification / scope',item.presentation['severity']),('Accepted finding',f['text']),
                            ('Management question',f['management_question']),('Proposed follow-up',f['action']),('Caveat',item.presentation['caveat'])]:
            narrative(ws,row,label,value,6)
            if label=='Priority finding':link(ws.cell(row,2),'Investment Issues',details[(c.id,f['id'])])
            row+=1
        for sid in f['fact_ids'][:3]:
            narrative(ws,row,'Source reference',sid,6);source_link(ws.cell(row,2),c.id,[sid],source_rows);row+=1
        row+=1
    ws.freeze_panes='B8';ws.auto_filter.ref='A7:F10'


def write_output(ws,companies,source_rows):
    headers(ws,5,['Company','Reporting period','Metric','Actual','Budget','Variance','Variance unit','Variance %','Metric unit','Source IDs','Prepared run'])
    row=6
    for c in companies:
        v=comparison(c)
        for key in ('revenue','ebitda','margin'):
            m=v[key]
            row_values(ws,row,[c.name,v['period_label'],m.name,m.actual,m.budget,m.variance,'percentage points' if key=='margin' else m.unit,m.variance_pct,m.unit,', '.join(m.source_ids),c.run],46)
            for col in (4,5):ws.cell(row,col).number_format=PERCENT if key=='margin' else MONEY
            ws.cell(row,6).number_format=PP if key=='margin' else MONEY;ws.cell(row,8).number_format=PERCENT
            source_link(ws.cell(row,10),c.id,m.source_ids,source_rows);row+=1
    ws.auto_filter.ref=f'A5:K{row-1}'


def export_portfolio(companies):
    if not companies:raise ValueError('Select at least one verified company.')
    wb=Workbook();wb.remove(wb.active)
    overview=title_sheet(wb,SHEETS[0],[26,28,18,18,18,18,17,17,12,14,22,40],'Operating review → Management follow-up → Source-linked Excel handoff. Prepared September 2026 / FY2026 company assessments; no new inference.')
    compare=title_sheet(wb,SHEETS[1],[30,43]+[39]*len(companies),COMPARABILITY)
    issues=title_sheet(wb,SHEETS[2],[29,24,15,15,18,48,25,55,60,55,22,45],'Accepted findings grouped by company. Impacts retain source scope and are never summed. Full original narratives, questions, actions and qualifications follow the index.')
    company_sheets={c.id:title_sheet(wb,COMPANY_SHEETS[c.id],[25,24,24,25,46,58],comparison(c)['period_label']+' · '+c.review.company+' · '+c.run) for c in companies}
    sources=title_sheet(wb,SHEETS[6],[29,22,20,45,23,68,95,48,52,26],'Company-qualified source and calculation index. Validated monetary figures use base units; source snapshots retain original units and coordinates. All interpretation caveats remain visible.')
    catalog=title_sheet(wb,SHEETS[7],[29,54,40,70,70],'Canonical approved source workbooks from demo_results/manifest.json. Source snapshot navigation is catalogued here.')
    output=title_sheet(wb,SHEETS[8],[29,30,23,21,21,23,23,19,18,60,26],'Plain tabular output for copying or mapping into an existing reporting template. No automatic integration or compatibility is claimed. Company periods remain separate.')
    title_sheet(wb,'Investment Ledger',[29,24,24,26,23,24,22,23,54,52],'Separate synthetic illustrative sponsor investments and transparent return supporting schedules. Not source-verified operating financials.')
    records=reference_records(companies);source_rows=write_references(sources,records)
    write_overview(overview,companies,source_rows);write_comparison(compare,companies,source_rows)
    details=write_issues(issues,companies,source_rows)
    for c in companies:write_company(company_sheets[c.id],c,[i for i in watchlist(companies) if i.company_id==c.id],source_rows,details)
    write_output(output,companies,source_rows)
    headers(catalog,5,['Company','Canonical approved workbook','Prepared run','Workbook SHA-256','Snapshot SHA-256'])
    for row,c in enumerate(companies,6):row_values(catalog,row,[c.name,c.provenance['workbook'],c.run,c.provenance['workbook_sha256'],c.provenance['snapshot_sha256']],52)
    catalog.auto_filter.ref=f'A5:E{5+len(companies)}'
    from portfolio_sources import attach_source_snapshots
    attach_source_snapshots(wb,companies,records,source_rows)
    attach_metric_proofs(wb,companies,source_rows)
    from portfolio_investment_excel import attach_investment_excel
    attach_investment_excel(wb,companies)
    from portfolio_financial_reporting import attach_portfolio_statement
    attach_portfolio_statement(wb,companies,source_rows)
    wb.move_sheet('Portfolio Financial Performance',offset=10-wb.sheetnames.index('Portfolio Financial Performance'))
    # Source snapshots keep their original grid; never write navigation into it.
    for name in SHEETS:
        if name not in wb.sheetnames:continue
        ws=wb[name]
        if ws.title!='Portfolio Overview':
            ws.print_area=f'A1:{ws.cell(ws.max_row,ws.max_column).coordinate}';ws.print_title_rows='1:5'
        if ws.title!='Portfolio Overview':link(literal(ws,4,1,'← Portfolio Overview'),'Portfolio Overview',1)
    stream=io.BytesIO();wb.save(stream);data=stream.getvalue()
    check=load_workbook(io.BytesIO(data))
    if any(cell.data_type in ('e','f') for ws in check for row in ws for cell in row):raise ValueError('GP workbook contains an unsafe formula or error.')
    return data
