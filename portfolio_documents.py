"""Read-only access to the three approved original workbooks; no formula execution."""
import hashlib
import io
import re
from dataclasses import dataclass
from html import escape
from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from engine import InputError
from portfolio_data import REGISTRY
from portfolio_sources import canonical_source, CELL_REFERENCE
from portfolio_watchlist import source_closure


@dataclass(frozen=True)
class SourceDocument:
    filename: str
    sha256: str
    data: bytes


def original_document(company):
    """Hash the exact bytes returned for download/preview, on every request."""
    if (company.id, company.name, company.sector, company.run) not in REGISTRY:
        raise InputError('This company is not in the approved source-document allowlist.')
    path = canonical_source(company)
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    if digest != company.provenance['workbook_sha256'] or path.suffix.lower() != '.xlsx':
        raise InputError('Original source verification failed. No document was served.')
    return SourceDocument(path.name, digest, data)


def source_books(document):
    return (load_workbook(io.BytesIO(document.data), data_only=False, keep_links=False),
            load_workbook(io.BytesIO(document.data), data_only=True, keep_links=False))


def checked_coordinate(sheet, coordinate):
    if not isinstance(coordinate, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*', coordinate):
        raise InputError('Invalid source cell coordinate.')
    row, col = coordinate_to_tuple(coordinate)
    if row > sheet.max_row or col > sheet.max_column:
        raise InputError('This cell is outside the original worksheet.')
    return row, col


def source_fact(company, fact_id, document=None):
    document = document or original_document(company)
    fact = next((f for f in company.review.package['facts'] if f['id'] == fact_id), None)
    if fact is None:
        raise InputError('No original source cell is recorded for this source ID.')
    match = CELL_REFERENCE.fullmatch(fact.get('ref', ''))
    if not match or fact.get('file') != document.filename:
        raise InputError('Unsupported original source reference.')
    sheet, cell, _ = match.groups()
    sheet = sheet.replace("''", "'"); cell = cell.replace('$', '')
    if sheet != fact.get('sheet') or cell != fact.get('cell'):
        raise InputError('Source coordinate and recorded provenance disagree.')
    return fact


def evidence_locations(company, ids):
    """Derived records remain calculations; only recorded originals become targets."""
    document = original_document(company)
    calculations = {f['id']: f for f in company.review.package['skill_engine']['calculations']}
    originals, derived, unsupported = [], [], []
    for identity in source_closure(company.review, ids):
        if identity in calculations:
            derived.append(calculations[identity])
        else:
            try: originals.append(source_fact(company, identity, document))
            except InputError: unsupported.append(identity)
    return originals, derived, unsupported


def cell_record(company, sheet, cached, coordinate):
    checked_coordinate(sheet, coordinate)
    cell = sheet[coordinate]
    fact = next((f for f in company.review.package['facts']
                 if f.get('sheet') == sheet.title and f.get('cell') == coordinate), None)
    formula = cell.value if cell.data_type == 'f' else None
    value = cached[coordinate].value if formula else cell.value
    origin = 'Original workbook cached value' if formula else 'Original workbook value'
    if formula and value is None and fact and fact.get('raw') == formula:
        value = fact.get('source_value'); origin = 'Retained source-unit observation from approved Review'
    label = fact.get('label') if fact else None
    if not label:
        # An original row label, not an inferred metric definition.
        label = next((sheet.cell(cell.row, c).value for c in range(1, cell.column)
                      if sheet.cell(cell.row, c).data_type != 'f' and isinstance(sheet.cell(cell.row, c).value, str)), None)
    return dict(coordinate=coordinate, value=value, formula=formula, origin=origin,
                label=label or 'No row label recorded', fact=fact, number_format=cell.number_format,
                comment=cell.comment.text if cell.comment else None)


def display_value(value, fmt='General'):
    if value is None: return '—'
    if not isinstance(value, (float, int)) or isinstance(value, bool): return str(value)
    if '%' in fmt: return f'{value:.1%}'
    decimals = 0
    m = re.search(r'0\.([0#]+)', fmt)
    if m: decimals = min(len(m[1]), 4)
    text = f'{abs(value):,.{decimals}f}' if '#' in fmt or '0' in fmt else str(abs(value))
    if '$' in fmt: text = '$' + text
    return ('(' + text + ')' if value < 0 else text)


def worksheet_html(company, sheet, cached, first_row, first_col, target, rows=12, columns=4):
    """A bounded coordinate grid, not a scrolling iframe or editable spreadsheet."""
    if rows not in range(1, 17) or columns not in range(1, 7): raise InputError('Preview range is too large.')
    checked_coordinate(sheet, f'{get_column_letter(first_col)}{first_row}')
    last_row = min(sheet.max_row, first_row + rows - 1)
    last_col = min(sheet.max_column, first_col + columns - 1)
    head = '<thead><tr><th aria-label="Row number"></th>' + ''.join('<th scope="col">'+get_column_letter(c)+'</th>' for c in range(first_col, last_col+1)) + '</tr></thead>'
    body = []
    for r in range(first_row, last_row+1):
        cells = []
        for c in range(first_col, last_col+1):
            original = sheet.cell(r,c)
            record = cell_record(company, sheet, cached, original.coordinate)
            text = display_value(record['value'], original.number_format)
            if record['formula'] and record['value'] is None: text = record['formula']
            css = (' source-target' if original.coordinate == target else '') + (' source-bold' if original.font.bold else '')
            cells.append('<td class="'+css.strip()+'" data-coordinate="'+original.coordinate+'" title="'+escape(original.coordinate+' · '+str(record['formula'] or record['value']), quote=True)+'">'+escape(text)+'</td>')
        body.append('<tr><th scope="row">'+str(r)+'</th>'+''.join(cells)+'</tr>')
    return '<table class="source-grid" aria-label="Original worksheet cell grid">'+head+'<tbody>'+''.join(body)+'</tbody></table>'
