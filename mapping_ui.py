"""Small mapping-assisted onboarding panel inside the existing intake view."""
import copy
import os
import streamlit as st
import pandas as pd
from openpyxl.utils import get_column_letter
from engine import Book, InputError, partition_files
from ingestion import inspect_files,apply_decisions,needs_review,material_pending,load_profiles,save_profiles,CATEGORIES,ROLES,SCENARIOS
from mapping_ai import suggest


def prepare(files,summary=None,ai_authorized=False):
    _,files=partition_files(files)
    layouts=copy.deepcopy(st.session_state.get('financial_layouts',{}))
    decisions=copy.deepcopy(st.session_state.get('account_decisions',{}))
    profiles=load_profiles()
    try:intake=apply_decisions(inspect_files(files,profiles,layouts),decisions)
    except InputError as e:
        st.error(str(e));return decisions,layouts,profiles,True
    if summary is not None:
        import calendar
        from ingestion import normalized_records,period as parse_period
        with summary:
            companies=list(dict.fromkeys(t.company for t in intake.tables))
            periods=sorted({c['period'] for t in intake.tables for c in t.columns if c['period'] and parse_period(c['period'])==c['period'] and c['scenario']=='actual' and not c.get('invalid')})
            period=(__import__('ingestion').period_label(periods[-1])) if periods else ''
            if companies:st.markdown('**'+(' · '.join(companies)+' — '+period).strip(' —')+'**')
            if intake.tables:st.caption('✓ Income statement recognized')
            if intake.auxiliary:st.caption('✓ Branch reporting recognized · reconciliation only')
            usable=any(r['accepted'] and r['period'] and r['scenario']=='actual' and r['value'] is not None for t in intake.tables for r in normalized_records(t))
            if usable:st.caption('Financial analysis available · supported figures only')
            else:st.caption('Financial statement recognition is limited. Review mappings or try another supported statement.')
            pending=sum(len(material_pending(t)) for t in intake.tables)
            if pending:st.caption(f'{pending} account mapping{"s" if pending!=1 else ""} may need review. You can continue.')
            covered={(t.file,t.sheet) for t in intake.tables}|{(a['file'],a['sheet']) for a in intake.auxiliary}
            if any((s['file'],s['sheet']) not in covered for s in intake.text_samples):st.caption('Additional sheets present · supporting schedules are checked during generation.')
    with st.expander('Review Account Mappings',expanded=False):
        st.caption('Optional. Reported financial totals and recognized sections are used automatically. Unconfirmed ambiguous accounts are excluded from calculated metrics, not a reason to block the report. Review material ambiguities or correct source interpretation here. AI proposals require source validation, not approval of every line.')
        if st.button('Suggest account mappings and headers with AI',disabled=not (os.environ.get('OPENAI_API_KEY') and ai_authorized),help='Optional: set OPENAI_API_KEY and optionally OPENAI_MODEL. Sends text headers/account labels only, never amounts or raw workbooks.'):
            try:
                with st.spinner('Interpreting account labels and headers…'):
                    from ai_runtime import inference_session
                    with inference_session(authorized=ai_authorized):suggestions=suggest(intake)
                for a in suggestions['accounts']:decisions[a['id']]={'category':a['category'],'role':a['role'],'confirmed':False}
                for table in intake.tables:
                    changed=[c for c in suggestions['columns'] if c['table_id']==table.id]
                    if changed:
                        selected=layouts.setdefault(table.file+'|'+table.sheet,{'header_row':table.header_row,'label_column':table.label_column})
                        selected['columns']=[dict(c) for c in table.columns]
                        for c in selected['columns']:
                            proposal=next((p for p in changed if p['column']==c['column']),None)
                            if proposal:c.update(period=proposal['period'],scenario=proposal['scenario'])
                        selected['ai_pending']=True
                for proposal in suggestions['layouts']:
                    if not any(t.file==proposal['file'] for t in intake.tables):layouts[proposal['file']]=dict(proposal,ai_pending=True)
                st.session_state.account_decisions=decisions;st.session_state.financial_layouts=layouts;st.rerun()
            except InputError as e:st.warning(str(e))
        for issue in intake.unresolved:
            st.info(issue['File']+': '+issue['Reason'])
            name=issue['File']
            if not name.lower().endswith('.xlsx'):
                st.caption('For financial CSVs, provide a header row with an account/description column and dated period columns. A known scheduling CSV belongs with its supported rate card.');continue
            wb=Book(name,files[name]).wb
            identity=name+'|'+issue.get('Sheet','')
            worksheet=st.selectbox('Financial worksheet',wb.sheetnames,key='manual_sheet_'+identity)
            st.dataframe(pd.DataFrame(list(wb[worksheet].values)[:15]).astype(str),hide_index=False,width='stretch')
            h=st.number_input('Header row',min_value=1,max_value=max(1,wb[worksheet].max_row),value=1,key='manual_header_'+identity)
            c=st.number_input('Account-label column (A = 1)',min_value=1,max_value=max(1,wb[worksheet].max_column),value=1,key='manual_label_'+identity)
            if st.button('Use selected financial table',key='manual_apply_'+identity):
                columns=[]
                for col in range(int(c)+1,wb[worksheet].max_column+1):
                    if wb[worksheet].cell(int(h),col).value is not None:
                        columns.append({'column':col,'header':str(wb[worksheet].cell(int(h),col).value),'period':'','scenario':'ignore','reference':wb[worksheet].cell(int(h),col).coordinate})
                layouts[name]={'sheet':worksheet,'header_row':int(h),'label_column':int(c),'columns':columns}
                st.session_state.financial_layouts=layouts;st.rerun()
        for table in intake.tables:
            st.markdown('**'+table.file+' · '+table.sheet+'**')
            for issue in dict.fromkeys(table.issues):st.warning(issue)
            co,basis,currency,scale=st.columns(4)
            company=co.text_input('Company',table.company,key='company_'+table.id)
            bases=['unknown','accrual','cash'];basis_value=basis.selectbox('Accounting basis',bases,index=bases.index(table.basis),key='basis_'+table.id)
            currencies=['unknown','$','USD','GBP','EUR','CAD','AUD'];currency_value=currency.selectbox('Currency',currencies,index=currencies.index(table.currency),key='currency_'+table.id)
            scales={'Not stated':None,'Whole units':1,'Thousands':1000,'Millions':1000000};scale_name=next((k for k,v in scales.items() if v==table.scale),'Not stated')
            scale_value=scales[scale.selectbox('Source scale',list(scales),index=list(scales).index(scale_name),key='scale_'+table.id)]
            sign=st.selectbox('Expense presentation',['positive','negative'],index=0 if table.expense_sign=='positive' else 1,key='sign_'+table.id,help='Positive means costs are positive amounts; negative means costs are signed deductions. This changes cost normalization, not revenue or profit signs.')
            column_rows=[{'Column':c['column'],'Header':c['header'],'Period':c['period'],'Scenario':c['scenario']} for c in table.columns]
            columns=st.data_editor(pd.DataFrame(column_rows),key='columns_'+table.id,hide_index=True,width='stretch',disabled=['Column','Header'],column_config={'Period':st.column_config.TextColumn('Period (YYYY-MM)'),'Scenario':st.column_config.SelectboxColumn('Scenario',options=SCENARIOS,required=True)})
            account_rows=[{'ID':a.id,'Original account':a.label,'Standardized account':a.category,'Role':a.role,'Accepted':a.confirmed or a.status.startswith('Auto'),'Status':a.status,'Source reference':a.reference} for a in table.accounts]
            edited=st.data_editor(pd.DataFrame(account_rows,columns=['ID','Original account','Standardized account','Role','Accepted','Status','Source reference']),key='accounts_'+table.id,hide_index=True,width='stretch',disabled=['ID','Original account','Status','Source reference'],column_config={'ID':None,'Standardized account':st.column_config.SelectboxColumn('Standardized account',options=CATEGORIES,required=True),'Role':st.column_config.SelectboxColumn('Role',options=ROLES,required=True),'Accepted':st.column_config.CheckboxColumn('Accepted')})
            originals={a.id:a for a in table.accounts}
            for row in edited.to_dict('records'):
                a=originals[row['ID']]
                if (row['Standardized account'],row['Role'],bool(row['Accepted']))!=(a.category,a.role,a.confirmed or a.status.startswith('Auto')):
                    decisions[row['ID']]={'category':row['Standardized account'],'role':row['Role'],'confirmed':bool(row['Accepted'])}
            pending=material_pending(table)
            if pending:
                # Commit ambiguous choices together; avoid rerendering the editor mid-selection.
                with st.form('resolve_form_'+table.id,enter_to_submit=False):
                    updates={}
                    for account in pending:
                        decision=decisions.get(account.id,{'category':account.category,'role':account.role})
                        selected_category=st.selectbox('Classification for '+account.label,CATEGORIES,index=CATEGORIES.index(decision['category']),key='resolve_'+account.id)
                        accepted=st.checkbox('Confirm classification for '+account.label,key='accept_'+account.id)
                        updates[account.id]={'category':selected_category,'role':decision['role'],'confirmed':accepted}
                    applied=st.form_submit_button('Apply reviewed classifications')
                if applied:decisions.update(updates)
            selected={'header_row':table.header_row,'label_column':table.label_column,'company':company,'basis':basis_value,'currency':currency_value,'scale':scale_value,'expense_sign':sign,'columns':[{'column':int(c['Column']),'header':c['Header'],'period':c['Period'] or '','scenario':c['Scenario'],'reference':get_column_letter(int(c['Column']))+str(table.header_row)} for c in columns.to_dict('records')]}
            original=layouts.get(table.file+'|'+table.sheet,{})
            if original.get('ai_pending'):
                selected['ai_pending']=not st.checkbox('I reviewed the AI-proposed table and period interpretation',key='ai_confirm_'+table.id)
            layouts[table.file+'|'+table.sheet]=selected
        # Reinspect edited metadata/columns, then enforce ambiguity gates on the exact analysis inputs.
        reviewed=apply_decisions(inspect_files(files,profiles,layouts),decisions)
        blocked=needs_review(reviewed)
        st.caption('Role guide: detail = additive account; total = reported financial category total; subtotal = excluded from aggregation; ignore = intentionally excluded. All extracted amounts keep their original source cells.')
        if st.button('Confirm and save mappings locally',disabled=blocked):
            try:save_profiles(reviewed)
            except OSError:st.error('Could not save the mapping file. Check write access to .local/.');return decisions,layouts,profiles,True
            st.session_state.account_decisions=decisions;st.session_state.financial_layouts=layouts
            st.success('Confirmed account mappings saved locally for this company, worksheet and account context.')
        if blocked:st.info('You can generate now. Uncertain classifications or periods will be labeled unavailable where they affect calculations; supported metrics remain available.')
    st.session_state.account_decisions=decisions;st.session_state.financial_layouts=layouts
    return decisions,layouts,profiles,False
