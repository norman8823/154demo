"""Income-statement presentation over approved Reviews; no extraction or analysis.

The frozen exporter remains intact. This adapter replaces its generated financial
sheet, retaining all original evidence and custom-template output elsewhere.
"""
import io
import math
from dataclasses import dataclass
from types import SimpleNamespace
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from excel_layout import financial_table, period_label
from portfolio_comparison import comparison
from engine import InputError

PORTFOLIO_SHEET = 'Portfolio Financial Performance'
LINES = (
    ('revenue', 'Revenue'), ('cogs', 'Direct Costs / Cost of Services'),
    ('gross_profit', 'Gross Profit'), ('gm', 'Gross Margin'),
    ('operating_expenses', 'Operating Expenses'), ('ebitda', 'Reported EBITDA'),
    ('margin', 'EBITDA Margin'), ('depreciation', 'Depreciation / Amortization'),
    ('interest', 'Interest Expense'), ('tax', 'Income Taxes'), ('net_income', 'Net Income'),
)
NAVY = '18364A'
PALE = 'EFF5F7'
PERCENT = '0.0%;(0.0%);"—"'
# Margins retain fractional numeric values; their variances use percentage points.
POINTS = '0.0" pp";(0.0)" pp";"—"'
NOTE = ('Variance = Actual − Budget; expense increases are adverse. Margin variances are percentage points. '
        'N/A = no accepted comparable metric in the approved Review; not zero. '
        'Reported EBITDA excludes adjusted EBITDA sensitivities. Values are in original reported currency, '
        'base units; no FX conversion or annualization. Click a value for its supporting evidence.')

@dataclass(frozen=True)
class StatementLine:
    key: str
    label: str
    actual: float | None
    budget: float | None
    unit: str
    actual_refs: tuple
    budget_refs: tuple
    actual_definition: str
    budget_definition: str

    @property
    def variance(self):
        return (self.actual-self.budget)*(100 if self.unit == '%' else 1) if self.actual is not None and self.budget is not None else None


def statement(review):
    """Select existing approved totals once; never aggregate extracted observations."""
    rows = {r[13]: r for r in financial_table(review)}
    output = []
    for key, label in LINES:
        if key == 'margin':
            v = comparison(SimpleNamespace(review=review))['margin']
            a, b, unit = v.actual, v.budget, '%'
            period = v.period
        else:
            row = rows.get(key)
            a, b = (row[1], row[2]) if row else (None, None)
            unit = '%' if key == 'gm' else review.currency
            period = row[10] if row else ''
        support = {}
        for scenario, value in (('actual', a), ('budget', b)):
            if value is None:
                support[scenario] = ((), 'No accepted comparable metric in approved Review.')
                continue
            keys = ('ebitda', 'revenue') if key == 'margin' else (key,)
            refs, definitions = [], []
            for operand in keys:
                candidates = []
                for name, amount, source_unit, definition, ids in review.financial_rows:
                    parts = name.rsplit(' · ', 3)
                    if (len(parts) == 4 and parts[2:] == [scenario, operand]
                            and period_label(parts[1]) == period):
                        candidates.append((amount, definition, tuple(ids.split('; '))))
                if len(candidates) != 1:
                    raise InputError('Ambiguous or missing approved statement provenance: '+key+'/'+scenario)
                amount, definition, ids = candidates[0]
                if 'REQUIRES REVIEW' in definition:
                    raise InputError('Unresolved statement mapping: '+key+'/'+scenario)
                if key != 'margin' and not math.isclose(amount, value, rel_tol=1e-12):
                    raise InputError('Statement does not reconcile to approved metric: '+key)
                refs.extend(ids); definitions.append(definition)
            support[scenario] = (tuple(dict.fromkeys(refs)),
                                 ('Reported EBITDA / revenue; ' if key == 'margin' else '')+'; '.join(definitions))
        output.append(StatementLine(key, label, a, b, unit, support['actual'][0], support['budget'][0],
                                    support['actual'][1], support['budget'][1]))
    return tuple(output)


def internal_link(cell, sheet, coordinate):
    from openpyxl.worksheet.hyperlink import Hyperlink
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location="'"+sheet.replace("'", "''")+"'!"+coordinate)


def text(ws, row, col, value):
    cell = ws.cell(row, col, value)
    if isinstance(value, str): cell.data_type = 's'
    cell.font = Font(name='Arial', size=10, color=NAVY)
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    return cell


def banner(ws, row, start, end, value, dark=False, size=10):
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    c = text(ws, row, start, value)
    c.fill = PatternFill('solid', fgColor=NAVY if dark else PALE)
    c.font = Font(name='Arial', size=size, color='FFFFFF' if dark else NAVY, bold=True)
    return c


def setup(ws, title, columns):
    ws.sheet_view.showGridLines = False
    banner(ws, 1, 1, columns, title, True, 19); ws.row_dimensions[1].height = 38
    banner(ws, 2, 1, columns, '154 Partners · Portfolio Reporting Copilot · Product Demo · Synthetic Data')
    ws.row_dimensions[2].height = 26
    ws.column_dimensions['A'].width = 35
    for col in range(2, columns+1): ws.column_dimensions[get_column_letter(col)].width = 18
    ws.freeze_panes = 'B7'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3 if columns > 4 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.15, footer=.15)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = 'Product Demo · Synthetic Data | Page &P'


def number_format(line):
    if line.unit == '%': return PERCENT
    currency = 'USD' if line.unit == 'USD' else ('$' if line.unit == '$' else line.unit)
    prefix = '"'+currency.replace('"', '')+' "'
    return prefix+'#,##0;('+prefix+'#,##0);"—"'


def write_line(ws, row, line, start, proof=None):
    strong = line.key in {'revenue', 'gross_profit', 'ebitda', 'net_income'}
    label = text(ws, row, 1, line.label)
    label.alignment = Alignment(vertical='center', indent=0 if strong else 1)
    ws.row_dimensions[row].height = 29
    for col, value in enumerate((line.actual, line.budget, line.variance), start):
        c = text(ws, row, col, 'N/A' if value is None else value)
        c.alignment = Alignment(horizontal='right', vertical='center')
        c.number_format = POINTS if line.unit == '%' and col == start+2 else number_format(line)
        if proof and value is not None: internal_link(c, *proof)
        if value is None: c.font = Font(name='Arial', size=10, color='7A8993', italic=True)
    for col in [1, *range(start, start+3)]:
        c = ws.cell(row, col)
        if strong:
            c.fill = PatternFill('solid', fgColor='DFE9EF' if line.key == 'ebitda' else PALE)
            c.font = Font(name='Arial', size=10, color=NAVY, bold=True)
            c.border = Border(top=Side(style='thin', color='B7C7D0'), bottom=Side(style='double' if line.key in {'ebitda','net_income'} else 'hair', color='B7C7D0'))
    if line.variance is not None:
        adverse = line.variance > 0 if line.key in {'cogs','operating_expenses','depreciation','interest','tax'} else line.variance < 0
        ws.cell(row, start+2).font = Font(name='Arial', size=10, bold=strong, color='A44747' if adverse else '397A65')


def add_proof(ws, review, line, name, evidence_rows, company_id=None):
    """Exact current actual/budget operands, definitions and every original source ID."""
    first = ws.max_row+2
    banner(ws, first, 1, 6, name+' · '+line.label, True)
    for col, label in enumerate(('Scenario', 'Approved value', 'Unit', 'Approved definition', 'Source ID', 'Open source evidence'), 1): text(ws, first+1, col, label)
    row = first+2
    for scenario in ('actual','budget'):
        refs = getattr(line, scenario+'_refs')
        for sid in refs:
            key = (company_id, sid) if company_id else sid
            if key not in evidence_rows: raise InputError('Statement source evidence is unavailable: '+sid)
            for col, value in enumerate((scenario.title(), getattr(line,scenario), line.unit,
                                        getattr(line,scenario+'_definition'), sid, 'Open original evidence →'), 1): text(ws,row,col,value)
            internal_link(ws.cell(row,6), 'Source References' if company_id else 'Source Evidence', 'A'+str(evidence_rows[key]))
            ws.cell(row,2).number_format = number_format(line)
            ws.row_dimensions[row].height = 56
            row += 1
    ws.cell(first+1, 7, 'Variance = Actual − Budget; margins: (Actual − Budget) × 100 percentage points.').data_type = 's'
    return ws.title, 'A'+str(first)


def prepare_template(template):
    """Reserve the primary statement name without discarding a supplied layout.

    Move a colliding template sheet to a named appendix, preserving coordinates,
    styles and arithmetic. Rewrite explicit same-workbook references only; refuse
    ambiguous dynamic/3-D references rather than silently changing their meaning.
    """
    if template is None: return template, []
    wb=load_workbook(io.BytesIO(template))
    if 'Financial Performance' not in wb.sheetnames: return template, []
    from openpyxl.formula.tokenizer import Tokenizer
    from openpyxl.comments import Comment
    original='Financial Performance';archive='Template Financial Performance'
    if archive in wb.sheetnames: raise InputError('Template statement archive name is already used.')
    wb[original].title=archive
    prefix="'"+original+"'!";replacement="'"+archive+"'!"
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.data_type!='f': continue
                before=cell.value;tokens=Tokenizer(before).items
                for token in tokens:
                    if original.lower() not in token.value.lower(): continue
                    if token.type=='OPERAND' and token.subtype=='RANGE' and token.value.startswith(prefix):
                        token.value=replacement+token.value[len(prefix):]
                    else:
                        raise InputError('A template statement reference cannot be safely relocated: '+sheet.title+'!'+cell.coordinate)
                after='='+''.join(t.value for t in tokens)
                if after!=before:
                    cell.value=after
                    prior=cell.comment.text+'\n' if cell.comment else ''
                    cell.comment=Comment(prior+'Original template formula before statement-sheet relocation: '+before,'Portfolio Reporting Copilot')
        for name in sheet.defined_names.values():
            if original in name.attr_text:
                if not name.attr_text.startswith(prefix):raise InputError('Unsupported template named reference.')
                name.attr_text=replacement+name.attr_text[len(prefix):]
    for name in wb.defined_names.values():
        if original in name.attr_text:
            if not name.attr_text.startswith(prefix):raise InputError('Unsupported template named reference.')
            name.attr_text=replacement+name.attr_text[len(prefix):]
    # Chart/pivot references require a template-specific relocation contract.
    if any(sheet._charts or sheet._pivots for sheet in wb):
        raise InputError('Rename the Financial Performance template sheet before uploading a chart/pivot template.')
    stream=io.BytesIO();wb.save(stream)
    return stream.getvalue(), ['Supplied Financial Performance layout preserved as Template Financial Performance; explicit formula references relocated.']


def export_company_report(review, template=None):
    """Wrap the frozen exporter; all portfolio standard/custom/upload routes use this."""
    from reports import export_report
    template, presentation_warnings = prepare_template(template)
    data, filename, warnings = export_report(review, template)
    warnings += presentation_warnings
    wb = load_workbook(io.BytesIO(data))
    if 'Investment Summary' in wb:
        for row in wb['Investment Summary']:
            unit=row[2] if len(row)>2 else None
            if unit and isinstance(unit.value,str) and unit.value.startswith('ratio;'):
                unit.value='%'+unit.value[len('ratio'):]
                row[1].number_format=PERCENT
    index = wb.sheetnames.index('Financial Performance'); wb.remove(wb['Financial Performance'])
    ws = wb.create_sheet('Financial Performance', index)
    setup(ws, 'Financial Performance', 4)
    context = comparison(SimpleNamespace(review=review))
    banner(ws,3,1,4,review.company+' · '+context['period_label']);ws.row_dimensions[3].height=32
    banner(ws,4,1,4,'Income statement · '+review.currency+' base units · Accounting basis: '+context['basis']);ws.row_dimensions[4].height=26
    for col, label in enumerate(('Income statement','Actual','Budget','Variance / pp'),1):
        c=text(ws,6,col,label);c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
    ws.row_dimensions[6].height=26
    evidence=wb['Source Evidence']; sources={}
    for row in range(6,evidence.max_row+1):sources.setdefault(evidence.cell(row,1).value,row)
    proof=wb.create_sheet('Statement Evidence');setup(proof,'Financial statement evidence',7)
    for col,width in enumerate((15,20,13,65,23,30,40),1):proof.column_dimensions[get_column_letter(col)].width=width
    for row,line in enumerate(statement(review),7):
        target=add_proof(proof,review,line,review.company,sources) if line.actual_refs or line.budget_refs else None
        write_line(ws,row,line,2,target)
    banner(ws,19,1,4,NOTE);ws.row_dimensions[19].height=76
    banner(ws,20,1,4,'Statement Evidence explains all reported values and calculations. Extracted Detail retains the complete technical source inventory.');ws.row_dimensions[20].height=45
    internal_link(ws['A20'],'Statement Evidence','A1')
    banner(ws,21,1,4,review.basis_note);ws.row_dimensions[21].height=58
    if 'tallis' in review.company.lower():
        banner(ws,22,1,4,'Tallis: unresolved classifications prevent a complete expense reconciliation. Unsupported lines remain N/A; no balancing expense is inferred.');ws.row_dimensions[22].height=44
    ws.print_area='A1:D'+str(ws.max_row);ws.print_title_rows='1:6'
    ws.sheet_view.zoomScale=90
    for sheet in wb: sheet.sheet_view.tabSelected=False
    ws.sheet_view.tabSelected=True
    from openpyxl.workbook.views import BookView
    if not wb.views: wb.views=[BookView()]
    for view in wb.views: view.activeTab=index
    wb.active=index
    out=io.BytesIO();wb.save(out)
    return out.getvalue(),filename,warnings


def attach_portfolio_statement(wb, companies, evidence_rows):
    ws=wb.create_sheet(PORTFOLIO_SHEET)
    columns=1+3*len(companies);setup(ws,'Portfolio Financial Performance',columns)
    banner(ws,3,1,columns,'Company-level operating review · Monthly and annual periods shown separately · No consolidated total or GAAP consolidation')
    ws.row_dimensions[3].height=28
    for index,company in enumerate(companies):
        start=2+index*3;context=comparison(company)
        banner(ws,4,start,start+2,company.name);ws.row_dimensions[4].height=28
        banner(ws,5,start,start+2,context['period_label']+' · '+company.review.currency+' base units');ws.row_dimensions[5].height=28
        for col,label in enumerate(('Actual','Budget','Variance / pp'),start):
            c=text(ws,6,col,label);c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
        for row,line in enumerate(statement(company.review),7):
            proof=add_proof(wb['Source References'],company.review,line,company.name,evidence_rows,company.id) if line.actual_refs or line.budget_refs else None
            write_line(ws,row,line,start,proof)
    text(ws,6,1,'Income statement')
    banner(ws,19,1,columns,NOTE);ws.row_dimensions[19].height=48
    banner(ws,20,1,columns,'Tallis: unresolved classifications; no complete expense reconciliation or unsupported lower-statement totals inferred. Summit: cross-schedule difference retained in Investment Issues. Harrow: annual figures, never combined with monthly results.');ws.row_dimensions[20].height=44
    banner(ws,21,1,columns,'Source References contains exact statement definitions and operands. All original source snapshots, finding evidence and reporting limitations remain available.');ws.row_dimensions[21].height=30
    internal_link(ws['A21'],'Source References','A1')
    ws.print_area='A1:'+get_column_letter(columns)+'21';ws.print_title_rows='1:6'
    ws.sheet_view.zoomScale=80
    internal_link(wb['Portfolio Overview']['A4'],PORTFOLIO_SHEET,'A1')
    wb['Portfolio Overview']['A4']='Open Portfolio Financial Performance →'
    wb['Portfolio Overview']['A4'].font=Font(name='Arial',size=11,bold=True,color=NAVY)
