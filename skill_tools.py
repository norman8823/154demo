"""Workbook-only paged tools and reproducible arithmetic. No Python execution tool."""
import ast
import hashlib
import io
import json
import math
import re
import time
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter
from engine import Book, InputError
from ingestion import unit_metadata

VERSION = 'workbook-tools-2'
LIMITS = dict(upload_bytes=8_000_000, expanded_bytes=32_000_000, zip_entries=500,
              sheets=30, rows=2000, columns=80, cells=30000, grid_cells=120000, text_chars=8000,
              page_cells=500, page_bytes=36000, formula_steps=100000)


def preflight(files):
    if len(files) != 1:
        raise InputError('Skill engine supports exactly one uploaded .xlsx workbook; select legacy for packages.')
    name, data = next(iter(files.items()))
    if not name.lower().endswith('.xlsx') or len(data) > LIMITS['upload_bytes']:
        raise InputError('Skill engine requires .xlsx within the 8 MB upload limit.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            entries = z.infolist()
            if len(entries) > LIMITS['zip_entries'] or sum(e.file_size for e in entries) > LIMITS['expanded_bytes']:
                raise InputError('Workbook exceeds expanded archive/resource limits.')
            if len({e.filename for e in entries}) != len(entries):
                raise InputError('Duplicate workbook archive entries are unsupported.')
            cells = sheets = grid_cells = formula_work = 0
            for entry in entries:
                low = entry.filename.lower()
                if any(x in low for x in ('vbaproject', 'macrosheets/', 'embeddings/', 'activex/')):
                    raise InputError('Macros and embedded executable content are not supported.')
                if low.endswith(('.xml', '.rels')):
                    raw = z.read(entry)
                    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                        raise InputError('XML entity/DTD content is not supported.')
                    root = ET.fromstring(raw)
                    if low.startswith('xl/worksheets/') and low.endswith('.xml'):
                        sheets += 1
                        max_r=max_c=1
                        for node in root.iter():
                            tag = node.tag.rsplit('}', 1)[-1]
                            if tag == 'dimension' and node.get('ref'):
                                _, _, c, r = range_boundaries(node.get('ref'))
                                max_r=max(max_r,r);max_c=max(max_c,c)
                                if r > LIMITS['rows'] or c > LIMITS['columns']:
                                    raise InputError('Worksheet dimensions exceed skill limits.')
                            if tag == 'c':
                                cells += 1
                                _, _, c, r = range_boundaries(node.attrib['r'])
                                max_r=max(max_r,r);max_c=max(max_c,c)
                                if r > LIMITS['rows'] or c > LIMITS['columns']:
                                    raise InputError('Worksheet cell coordinates exceed skill limits.')
                            if tag == 'f':
                                expression=node.text or ''
                                # Existing local normalization runs before the analyst.
                                # Bound its formula expansion as well as our tool evaluator.
                                if len(expression)>4000 or ('"' in expression and '*' in expression):
                                    raise InputError('Formula exceeds safe local-normalization limits.')
                                for start,end in re.findall(r'(\$?[A-Z]{1,3}\$?[1-9]\d*)(?::(\$?[A-Z]{1,3}\$?[1-9]\d*))?',expression):
                                    c1,r1,c2,r2=range_boundaries(start+(':'+end if end else ''))
                                    if max(r1,r2)>LIMITS['rows'] or max(c1,c2)>LIMITS['columns']:
                                        raise InputError('Formula reference exceeds workbook limits.')
                                    formula_work+=max(1,(r2-r1+1)*(c2-c1+1))
                                if formula_work>LIMITS['formula_steps']:
                                    raise InputError('Aggregate formula expansion exceeds workbook limits.')
                            if tag in {'t', 'f'} and len(node.text or '') > LIMITS['text_chars']:
                                raise InputError('Workbook cell text/formula exceeds skill limits.')
                        grid_cells+=max_r*max_c
            if cells > LIMITS['cells'] or sheets > LIMITS['sheets'] or grid_cells>LIMITS['grid_cells']:
                raise InputError('Workbook cell/sheet count exceeds skill limits.')
    except InputError:
        raise
    except (ValueError, KeyError, zipfile.BadZipFile, ET.ParseError, OSError) as exc:
        raise InputError('Invalid or unsupported workbook archive.') from exc
    return name, data


class BoundedBook(Book):
    def __init__(self, name, data, deadline):
        self.name = name
        self.wb = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        self.strict_sum = False  # Excel SUM ignores blank referenced cells; direct blanks remain unavailable.
        self.cache = {}
        self.active = set()
        self.steps = 0
        self.deadline = deadline

    def value(self, sheet, coord):
        self.steps += 1
        if self.steps > LIMITS['formula_steps'] or len(self.active) > 80 or time.monotonic() >= self.deadline:
            raise InputError('Formula evaluation resource limit reached.')
        if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', coord) or sheet not in self.wb.sheetnames:
            raise InputError('Invalid workbook cell reference.')
        _, _, col, row = range_boundaries(coord)
        if row > LIMITS['rows'] or col > LIMITS['columns']:
            raise InputError('Formula reference outside workbook resource bounds.')
        return super().value(sheet, coord)

    def formula(self, sheet, expr):
        if '[' in expr or ']' in expr or len(expr) > 4000:
            raise InputError('External or oversized formula is unsupported.')
        for a, b in re.findall(r'(\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)', expr):
            c1, r1, c2, r2 = range_boundaries(a + ':' + b)
            if r2 > LIMITS['rows'] or c2 > LIMITS['columns'] or (r2-r1+1)*(c2-c1+1) > LIMITS['cells']:
                raise InputError('Formula range exceeds resource bounds.')
        return super().formula(sheet, expr)


def serial(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def finite(value):
    return type(value) in (int, float) and math.isfinite(value)


class WorkbookTools:
    def __init__(self, files, deadline=float('inf')):
        self.name, data = preflight(files)
        self.digest = hashlib.sha256(data).hexdigest()
        self.book = BoundedBook(self.name, data, deadline)
        self.cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
        self.facts = {}
        self.notes = {}
        self.calculations = []
        self.seen = set()
        self.errors = {}
        self.records = {}
        self.sheet_ids = {s.title: 's'+str(i) for i, s in enumerate(self.book.wb, 1)}
        for s in self.book.wb:
            self.records[s.title] = [c.coordinate for row in s for c in row if c.value is not None or c.comment]
        self.deadline = deadline

    def check(self):
        if time.monotonic() >= self.deadline:
            raise InputError('Skill engine elapsed-time limit reached.')

    def inventory(self):
        return dict(file=self.name, sha256=self.digest, sheets=[dict(sheet=s.title, id=self.sheet_ids[s.title],
                    rows=s.max_row, columns=s.max_column, populated=len(self.records[s.title]), state=s.sheet_state)
                    for s in self.book.wb], limits=LIMITS, coverage=self.coverage())

    def coverage(self):
        return [dict(sheet=s, total=len(coords), read=sum((s,c) in self.seen for c in coords)) for s,coords in self.records.items()]

    def _cell(self, sheet, coord):
        self.check()
        if sheet not in self.records or coord not in self.records[sheet]:
            raise InputError('Reference is not a populated cell in the uploaded workbook.')
        ws = self.book.wb[sheet]; c = ws[coord]; raw = serial(c.value)
        identity = self.sheet_ids[sheet]+'_'+coord
        error = None
        try:
            value = serial(self.book.value(sheet, coord))
            if type(value) is float and not math.isfinite(value):raise InputError('Nonfinite cell value.')
        except (InputError, ArithmeticError, ValueError, TypeError, SyntaxError, RecursionError) as exc:
            value = None; error = str(exc)
            self.errors[identity] = error
        label = str(ws.cell(c.row, 1).value or '')
        headings = [str(ws.cell(r, c.column).value) for r in range(max(1,c.row-50),c.row)
                    if isinstance(ws.cell(r,c.column).value,str) and not ws.cell(r,c.column).value.startswith('=')]
        context = ' '.join(str(ws.cell(r,1).value or '') for r in range(1,min(5,ws.max_row)+1))
        unit, currency, scale = 'unknown', 'unknown', None
        # Row unit columns override sheet denominations; no inferred USD jurisdiction.
        row_unit = str(ws.cell(c.row,2).value or '') if c.column > 2 else ''
        if re.match(r'^(?:DSO\b|days sales outstanding\b|WIP days\b|lock[- ]?up days\b)',label,re.I):unit='days';scale=1
        elif row_unit in {'%', 'ratio'} or '%' in c.number_format:unit='ratio';scale=1
        elif row_unit in {'#','tons','years'}:unit='count';scale=1
        elif not re.search(r'[$£€]',c.number_format) and not re.match(r'^(?:WIP|AR|revenue|fees|sales|costs?|expenses?|profit|EBITDA|receivables)\b',label,re.I) and any(re.search(r'^(?:clients (?:at|lost|added|retained|opening|closing|count)\b|(?:opening|closing|lost|new|active|number of) (?:clients|customers|employees|trucks|units|contracts)\b|(?:client|customer|employee|truck|unit|contract) count\b)',v,re.I) for v in [label,headings[-1] if headings else '']):
            unit='count';scale=1
        elif re.fullmatch(r'\$\s*/\s*[A-Za-z][A-Za-z /-]{0,40}',row_unit):
            unit='money';currency='$';scale=1
        elif row_unit in {'$', '$000'}:
            unit='money';currency='$';scale=1000 if row_unit=='$000' else 1
        else:
            currency, scale = unit_metadata(context)
            if currency!='unknown' and scale is not None:unit='money'
            elif 'Unit counts' in context:unit='count';scale=1
        basis='accrual' if re.search(r'accrual basis',context,re.I) else 'cash' if re.search(r'cash basis',context,re.I) else 'unknown'
        if finite(value):
            self.facts[identity] = dict(id=identity,file=self.name,sheet=sheet,ref=f"'{sheet}'!{coord}",cell=coord,
                label=label,value=value*(scale if unit=='money' else 1),source_value=value,raw=raw,
                unit=unit,currency=currency,scale=1 if unit=='money' else scale,source_scale=scale,source_unit_label=row_unit,
                source_period_labels=re.findall(r'\bFY\s?20\d{2}\b',context,re.I),period=' / '.join(headings[-2:]),column_label=' / '.join(headings[-2:]),scenario='source',
                entity=label if re.match(r'(?:(?:national|regional|key|major)\s+)?(?:account|customer|client|venue|branch|unit|region|division|entity)\s+\d',label,re.I) else sheet,company='',basis=basis,illustrative=False,inputs=[],source_refs=[],
                limitation='Units inferred from explicit sheet/row labels and formats; period, scope and narrative interpretation require review.')
        elif isinstance(value,str):
            self.notes[identity]=dict(id=identity,file=self.name,sheet=sheet,ref=f"'{sheet}'!{coord}",text=value)
        if c.comment:
            self.notes[identity+'_comment']=dict(id=identity+'_comment',file=self.name,sheet=sheet,ref=f"'{sheet}'!{coord} comment",text=c.comment.text)
        return dict(id=identity,cell=coord,raw=raw,value=value,unit=unit,scale=scale,
                    cached=serial(self.cached[sheet][coord].value) if c.data_type=='f' else None,
                    note=c.comment.text if c.comment else None,error=error)

    def read_range(self, sheet, cell_range, offset=0, limit=500):
        self.check()
        if sheet not in self.records:raise InputError('Unknown worksheet in uploaded workbook.')
        if type(offset) is not int or offset<0 or type(limit) is not int or not 1<=limit<=LIMITS['page_cells']:
            raise InputError('Invalid paging bounds.')
        if not re.fullmatch(r'[A-Z]+[1-9]\d*(?::[A-Z]+[1-9]\d*)?',cell_range):raise InputError('Use an A1 range without workbook/sheet paths.')
        a,b,c,d=range_boundaries(cell_range)
        if a>c or b>d or c>LIMITS['columns'] or d>LIMITS['rows']:raise InputError('Range exceeds bounds.')
        coords=[k for k in self.records[sheet] if a<=self.book.wb[sheet][k].column<=c and b<=self.book.wb[sheet][k].row<=d]
        rows=[];size=0
        for coord in coords[offset:offset+limit]:
            item=self._cell(sheet,coord)
            # Compact columns retain raw formula, computed value, cache and complete note.
            record=[item[k] for k in ('id','raw','value','unit','scale','cached','note','error')]
            if item['raw']==item['value']:record[2]=None  # means same as raw, not missing
            if item['cached'] is not None and item['cached']==item['value']:record[5]='=computed'
            if not finite(item['value']):record[3]=record[4]=None
            while record and record[-1] is None:record.pop()
            n=len(json.dumps(record,ensure_ascii=False).encode())
            if size+n>LIMITS['page_bytes']:
                if not size:raise InputError('A source cell exceeds the tool page byte limit; cannot return it without truncation.')
                break
            rows.append(record);size+=n;self.seen.add((sheet,coord))
        end=offset+len(rows)
        return dict(file=self.name,sheet=sheet,range=cell_range,columns=['id (suffix is cell)','raw','computed (null=same as raw unless error)','unit','source_scale','formula_cache (=computed means exact match)','comment','error'],omitted_trailing_fields='null',formula_semantics='SUM ignores blank references as Excel does. Direct blank cells remain unavailable; authored totals do not prove completeness of underlying inputs.',
                    cells=rows,total=len(coords),next_offset=end if end<len(coords) else None,
                    numeric_citation_rule='IDs identify values. Money placeholders use raw * source_scale in base currency units; calculations use normalized values.',coverage=self.coverage())

    def search(self, query, offset=0, limit=100):
        self.check()
        if not isinstance(query,str) or not 1<=len(query)<=150:raise InputError('Search requires literal text within 150 characters.')
        if type(offset) is not int or offset<0 or type(limit) is not int or not 1<=limit<=100:raise InputError('Invalid search paging bounds.')
        hits=[]
        for s,coords in self.records.items():
            for c in coords:
                cell=self.book.wb[s][c];text=str(cell.value or '')+' '+(cell.comment.text if cell.comment else '')
                if query.casefold() in text.casefold():hits.append((s,c))
        result=[];size=0
        for s,c in hits[offset:offset+limit]:
            item=self._cell(s,c);item.update(sheet=s,file=self.name)
            n=len(json.dumps(item).encode())
            if size+n>LIMITS['page_bytes']:
                if not size:raise InputError('A source cell exceeds the tool page byte limit; cannot return it without truncation.')
                break
            result.append(item);size+=n;self.seen.add((s,c))
        end=offset+len(result)
        return dict(matches=result,total=len(hits),next_offset=end if end<len(hits) else None)

    def calculate(self, expression, label, unit, scope, assumptions, illustrative=False):
        self.check()
        if len(expression)>1200 or len(label)>180 or len(scope)>500 or len(assumptions)>500:
            raise InputError('Calculation text exceeds bounds.')
        if unit not in {'money','ratio','count','multiple','pp','days','unknown'}:raise InputError('Unknown result unit.')
        value, refs, constants = self._arithmetic(expression)
        inferred=self._dimension(ast.parse(expression,mode='eval').body)
        if unit!=inferred and not (inferred=='ratio' and unit=='multiple'):
            raise InputError('Result unit does not match operand dimensions: expected '+inferred)
        if not refs:raise InputError('A calculation must reference workbook facts.')
        if not scope.strip():raise InputError('State period, entity, basis, units and comparison/bridge purpose.')
        if type(illustrative) is not bool:raise InputError('illustrative must be a boolean.')
        # Mechanics and explanatory scope text are not economic assumptions.
        economic_inputs=list(constants)
        # A typed day-count convention is a unit conversion, not a scenario.
        conventions=[]
        for node in ast.walk(ast.parse(expression,mode='eval')):
            if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id=='days':
                conventions.append(dict(function='days',day_count=ast.unparse(node.args[1])))
                if isinstance(node.args[1],ast.Constant):economic_inputs.remove(node.args[1].value)
        # A literal divisor equal to the number of distinct equally weighted
        # source terms is arithmetic, not an economic scenario assumption.
        def mean_terms(n):
            if isinstance(n,ast.Name):return [n.id]
            if isinstance(n,ast.BinOp) and isinstance(n.op,ast.Add):return mean_terms(n.left)+mean_terms(n.right)
            if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id=='sum':return [k for a in n.args for k in mean_terms(a)]
            return []
        for node in ast.walk(ast.parse(expression,mode='eval')):
            if isinstance(node,ast.BinOp) and isinstance(node.op,ast.Div) and isinstance(node.right,ast.Constant):
                terms=mean_terms(node.left)
                if len(terms)>1 and len(set(terms))==len(terms)==node.right.value and len({self.facts[k]['unit'] for k in terms})==1:
                    economic_inputs.remove(node.right.value);conventions.append(dict(function='unweighted_mean',terms=terms,divisor=node.right.value))
        economic_constants=[x for x in economic_inputs if x not in {0,1} and not (unit=='pp' and x==100)]
        illustrative=bool(illustrative or economic_constants or any(self.facts[k].get('illustrative') for k in refs))
        currencies={self.facts[k]['currency'] for k in refs if self.facts[k]['unit']=='money'}
        if len(currencies)>1:raise InputError('Cannot combine different or unknown currencies without a supported bridge.')
        identity='calc_'+str(len(self.calculations)+1)
        if len(self.calculations)>=100:raise InputError('Calculation count limit reached.')
        f=dict(id=identity,label=label,value=value,unit=unit,currency=next(iter(currencies),'unknown'),scale=1,
            file=self.name,sheet='Source-linked calculations',ref='Restricted arithmetic '+identity,raw=expression,
            inputs=refs,source_refs=[],expression=expression,operation='restricted_arithmetic',period=scope,
            scenario='illustrative' if illustrative else 'comparison',entity=scope,basis='see source inputs',
            illustrative=illustrative,assumptions=assumptions,constants=constants,conventions=conventions,
            limitation='Arithmetic verified; semantic compatibility, scope and interpretation require analyst review.')
        self.facts[identity]=f;self.calculations.append(f)
        return f

    def _arithmetic(self, expression, refresh=False):
        try:tree=ast.parse(expression,mode='eval')
        except (SyntaxError,ValueError) as exc:raise InputError('Invalid restricted expression.') from exc
        if len(list(ast.walk(tree)))>180:raise InputError('Expression operation limit exceeded.')
        refs=[];constants=[]
        def ev(n):
            if isinstance(n,ast.Expression):return ev(n.body)
            if isinstance(n,ast.Name):
                if n.id not in self.facts:raise InputError('Unknown or unread fact ID: '+n.id)
                source=self.facts[n.id]
                if not source.get('expression') and (source['sheet'],source['cell']) not in self.seen:
                    raise InputError('Calculation references an unread source cell.')
                if n.id not in refs:refs.append(n.id)
                return self.verify_fact(n.id) if refresh else self.facts[n.id]['value']
            if isinstance(n,ast.Constant) and finite(n.value):constants.append(n.value);return n.value
            if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)):return ev(n.operand)*(1 if isinstance(n.op,ast.UAdd) else -1)
            if isinstance(n,ast.BinOp) and isinstance(n.op,(ast.Add,ast.Sub,ast.Mult,ast.Div)):
                a,b=ev(n.left),ev(n.right)
                if isinstance(n.op,ast.Add):return a+b
                if isinstance(n.op,ast.Sub):return a-b
                if isinstance(n.op,ast.Mult):return a*b
                if b==0:raise InputError('Division by zero; blank is not zero.')
                return a/b
            if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id in {'sum','min','max','abs','days'} and not n.keywords and 1<=len(n.args)<=40:
                values=[ev(a) for a in n.args]
                if n.func.id=='days':
                    if len(values)!=2 or not 1<=values[1]<=366:raise InputError('days requires a ratio and a day count from 1 to 366.')
                    return values[0]*values[1]
                if n.func.id=='abs' and len(values)!=1:raise InputError('abs requires one operand.')
                return {'sum':sum,'min':min,'max':max,'abs':lambda a:abs(a[0])}[n.func.id](values)
            raise InputError('Only fact IDs, finite constants, + - * / and sum/min/max/abs/days are allowed.')
        try:result=ev(tree)
        except (OverflowError,ZeroDivisionError) as exc:raise InputError('Invalid calculation arithmetic.') from exc
        if not finite(result) or abs(result)>1e18:raise InputError('Calculation result outside finite bounds.')
        return result,refs,constants

    def _dimension(self,n):
        """Reject obvious unit errors; this does not certify temporal/entity scope."""
        if isinstance(n,ast.Name):return self.facts[n.id]['unit']
        if isinstance(n,ast.Constant):return 'scalar'
        if isinstance(n,ast.UnaryOp):return self._dimension(n.operand)
        if isinstance(n,ast.Call):
            if n.func.id=='days':
                if len(n.args)!=2 or self._dimension(n.args[0])!='ratio' or self._dimension(n.args[1]) not in {'scalar','count','days'}:raise InputError('days requires ratio and calendar-count operands.')
                return 'days'
            dims={self._dimension(a) for a in n.args}
            if len(dims)!=1:raise InputError('Cannot sum/compare incompatible units.')
            return next(iter(dims))
        a,b=self._dimension(n.left),self._dimension(n.right)
        if isinstance(n.op,(ast.Add,ast.Sub)):
            if a==b:return a
            if {a,b}<={'ratio','scalar'}:return 'ratio'
            raise InputError('Cannot add/subtract incompatible units.')
        if isinstance(n.op,ast.Mult):
            # Ratio values are fractions; pp values are hundredths. A label alone
            # cannot change scale. Require the conversion in the expression.
            if a=='ratio' and isinstance(n.right,ast.Constant) and n.right.value==100:return 'pp'
            if b=='ratio' and isinstance(n.left,ast.Constant) and n.left.value==100:return 'pp'
            if a in {'scalar','ratio'}:return b if b!='scalar' else a
            if b in {'scalar','ratio'}:return a
            if {a,b}=={'money','count'}:return 'money'
            raise InputError('Unsupported product dimensions.')
        if isinstance(n.op,ast.Div):
            if a=='pp' and isinstance(n.right,ast.Constant) and n.right.value==100:return 'ratio'
            if a=='ratio' and isinstance(n.right,ast.Constant) and n.right.value==.01:return 'pp'
            if a==b:return 'ratio'
            if b in {'scalar','ratio'}:return a
            if a=='money' and b=='count':return 'money'
            raise InputError('Unsupported quotient dimensions.')
        raise InputError('Unsupported dimensions.')

    def verify_fact(self, identity):
        self.check()
        f=self.facts.get(identity)
        if not f:raise InputError('Unknown numerical citation: '+identity)
        if f.get('expression'):
            value,refs,constants=self._arithmetic(f['expression'],refresh=True)
            inferred=self._dimension(ast.parse(f['expression'],mode='eval').body)
            if f['unit']!=inferred and not (inferred=='ratio' and f['unit']=='multiple'):
                raise InputError('Calculation unit or percentage-point conversion mismatch.')
            if refs!=f['inputs'] or constants!=f['constants']:raise InputError('Calculation provenance mismatch.')
        else:
            if (f['sheet'],f['cell']) not in self.seen:raise InputError('Cited source was not read.')
            value=self.book.number(f['sheet'],f['cell'])*(f['source_scale'] if f['unit']=='money' else 1)
        if not math.isclose(value,f['value'],rel_tol=1e-10,abs_tol=1e-8):raise InputError('Cited value differs from recalculated source: '+identity)
        return value

    def snapshot(self):
        """Lossless values/labels for all seen cells; formulas remain addressable.

        Remove repeated tool envelopes/history, not source evidence. Raw formulas
        and caches can be re-read; normalized numerical values retain exact IDs.
        """
        sheets=[]
        for sheet,coords in self.records.items():
            cells=[]
            from collections import Counter
            units=Counter(f['unit'] for f in self.facts.values() if f.get('sheet')==sheet)
            default_unit=units.most_common(1)[0][0] if units else 'unknown'
            for coord in coords:
                if (sheet,coord) not in self.seen:continue
                identity=self.sheet_ids[sheet]+'_'+coord
                if identity in self.facts:
                    f=self.facts[identity];cells.append([coord,f['value']]+([f['unit']] if f['unit']!=default_unit else []))
                elif identity in self.notes:cells.append([coord,self.notes[identity]['text']])
                else:cells.append([coord,serial(self.book.wb[sheet][coord].value),self.errors.get(identity)])
                if identity+'_comment' in self.notes:cells.append([coord+'_comment',self.notes[identity+'_comment']['text']])
            sheets.append(dict(sheet=sheet,id_prefix=self.sheet_ids[sheet]+'_',default_numeric_unit=default_unit,cells=cells))
        return dict(file=self.name,sha256=self.digest,coverage=self.coverage(),sheets=sheets,
            columns=['cell suffix: prepend sheet id_prefix to obtain exact fact ID','normalized value or verbatim source text','unit override if different from sheet default_numeric_unit'],
            interpretation='All numeric money values are BASE currency units, already scaled. Original formulas/caches/comments remain available via read_range. This snapshot retains every read cell value and label, not a relevance sample.',
            calculations=self.calculations,formula_errors=self.errors)

    def invoke(self, name, arguments):
        self.check()
        methods={'inventory':self.inventory,'read_range':self.read_range,'search':self.search,'calculate':self.calculate}
        if name not in methods:raise InputError('Tool is not allowed.')
        try:return methods[name](**arguments)
        except (TypeError,ValueError) as exc:raise InputError(str(exc)) from exc


def object_schema(props):
    return dict(type='object',properties=props,required=list(props),additionalProperties=False)


def tool_schema():
    text={'type':'string'};integer={'type':'integer'}
    specs=[('inventory','List only the uploaded workbook worksheets, dimensions and read coverage.',{}),
           ('read_range','Read all cells in a targeted A1 range, including labels, formulas, computed values, caches and comments. Page until next_offset is null; numeric null with no error means same as raw.',dict(sheet=text,cell_range=text,offset=integer,limit=integer)),
           ('search','Literal text search of source cells/comments only, paged.',dict(query=text,offset=integer,limit=integer)),
           ('calculate','Reproducible arithmetic using normalized numeric fact IDs. + - * / and sum/min/max/abs, plus days(ratio, calendar_day_count) for DSO or similar duration measures. Use unit days for days(), never ratio or money. Referenced blanks in authored Excel SUM are skipped as in Excel; direct blank cells are not zero. Mark economic scenarios illustrative; identity/growth constants and multiply-by-100 percentage-point conversions are mechanics. For percentage points use (ratioA-ratioB)*100 with unit pp. State period, scope, basis and purpose.',dict(expression=text,label=text,unit=dict(type='string',enum=['money','ratio','count','multiple','pp','days','unknown']),scope=text,assumptions=text,illustrative={'type':'boolean'}))]
    return [dict(type='function',name=n,description=d,parameters=object_schema(p),strict=True) for n,d,p in specs]
