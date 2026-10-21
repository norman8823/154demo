"""Bounded content-based supporting-schedule extraction; numbers always come from cells."""
import csv
import io
import re
from collections import defaultdict
from openpyxl import Workbook
from engine import Book,InputError
from ingestion import amount,clean,period,metadata,digest,discover_columns,column_period

MAX_ROWS=1000
MAX_COLS=80
MAX_SHEETS=30


def read_package(files):
    package=dict(facts=[],series=[],cumulative=[],aging=[],units=[],narratives=[],coverage=[],limitations=[])
    for name,data in files.items():
        if name.lower().endswith('.csv'):
            w=Workbook();ws=w.active;ws.title='CSV'
            try:
                for row in csv.reader(io.StringIO(data.decode('utf-8-sig'))):ws.append(row)
            except (UnicodeError,csv.Error):continue
            book=None
        else:
            try:book=Book(name,data);w=book.wb
            except InputError:continue
        for ws in list(w)[:MAX_SHEETS]:
            before=len(package['facts']);kind=set();used=set()
            if ws.max_row>MAX_ROWS or ws.max_column>MAX_COLS:
                package['limitations'].append(f'{name} · {ws.title}: exceeds bounded schedule size; not analyzed.');continue
            def fact(row,col,label,unit,p,scenario='actual',entity='',company='',basis='unknown',currency='unknown',scale=None):
                cell=ws.cell(row,col)
                try:value=amount(book.value(ws.title,cell.coordinate) if book else cell.value)
                except (InputError,ArithmeticError,ValueError,TypeError,SyntaxError,RecursionError):value=None
                if value is None:return None
                number=value*(scale if unit=='money' and scale is not None else 1)
                record=dict(grain='reported_month',id='pkg_'+digest(name+'|'+ws.title+'|'+cell.coordinate),file=name,sheet=ws.title,ref="'"+ws.title+"'!"+cell.coordinate,label=str(label),unit=unit,period=p,scenario=scenario,entity=entity,company=company,basis=basis,currency=currency,scale=scale,value=number,raw=cell.value)
                if record['id'] not in used:package['facts'].append(record);used.add(record['id'])
                return record
            # Time series: a row of explicitly dated columns, independent of worksheet name.
            for h in range(1,min(ws.max_row,60)+1):
                dates=[(c,column_period(ws.cell(h,c).value)) for c in range(1,ws.max_column+1) if re.fullmatch(r'20\d{2}-\d{2}',column_period(ws.cell(h,c).value)) and not re.search(r'ytd|year to date|variance',str(ws.cell(h,c).value),re.I)]
                if len(dates)<2:continue
                first=min(c for c,_ in dates)
                if first<=1:continue
                label_col=max(range(1,first),key=lambda c:sum(isinstance(ws.cell(r,c).value,str) and not str(ws.cell(r,c).value).startswith('=') for r in range(h+1,min(ws.max_row,h+80)+1)))
                co,basis,currency,scale,_=metadata(ws,h)
                unit_col=next((c for c in range(label_col+1,first) if clean(ws.cell(h,c).value) in {'unit','units','uom'}),None)
                for row in range(h+1,ws.max_row+1):
                    # Stop at a later dated header rather than borrowing its rows/units.
                    if row>h+1 and sum(bool(period(ws.cell(row,c).value)) for c in range(1,ws.max_column+1))>=2:break
                    label=ws.cell(row,label_col).value
                    if not isinstance(label,str) or label.startswith('=') or not label.strip():continue
                    explicit=str(ws.cell(row,unit_col).value or '') if unit_col else ''
                    fmt=ws.cell(row,dates[-1][0]).number_format
                    unit='ratio' if '%' in explicit+fmt or re.search(r'\bmargin\b.*%|%.*\bmargin\b',label,re.I) else ('count' if explicit in {'#','count','units'} else ('money' if '$' in explicit or explicit in {'USD','GBP','EUR','CAD','AUD'} or (not unit_col and currency!='unknown' and scale is not None) else 'unknown'))
                    if re.search(r'^dso\b|days sales outstanding|days inventory|days payable',label,re.I):unit='days'
                    row_currency=('$' if '$' in explicit else explicit) if unit=='money' and explicit else currency
                    row_scale=(1000 if re.search(r'000|thousand',explicit,re.I) else (1 if explicit in {'$','USD','GBP','EUR','CAD','AUD'} else scale))
                    percent_divisor=1
                    if unit=='ratio' and '%' in explicit and '%' not in fmt:
                        samples=[]
                        for c,_ in dates:
                            try:v=amount(book.value(ws.title,ws.cell(row,c).coordinate) if book else ws.cell(row,c).value)
                            except (InputError,ArithmeticError,ValueError,TypeError):v=None
                            if v is not None:samples.append(v)
                        if samples and any(abs(v)>1 for v in samples):percent_divisor=100
                        else:
                            unit='unknown';package['limitations'].append(f'{name} · {ws.title}: {label} has ambiguous unformatted percentage storage; direction analysis withheld.')
                    points=[]
                    for col,p in dates:
                        head=str(ws.cell(h,col).value)
                        if re.search(r'variance|prior year|\bpy\b',head,re.I):continue
                        scenario='budget' if re.search(r'budget|plan|forecast',head,re.I) else 'actual'
                        f=fact(row,col,label,unit,p,scenario,company=co,basis=basis,currency=row_currency,scale=row_scale)
                        if f:
                            if percent_divisor==100:f['value']/=100
                            points.append(f)
                    actual=[f for f in points if f['scenario']=='actual']
                    if len({f['period'] for f in actual})>=2:
                        identity='series_'+digest(name+'|'+ws.title+'|'+str(row))
                        package['series'].append(dict(id=identity,file=name,sheet=ws.title,company=co,label=label,unit=unit,currency=row_currency,scale=row_scale,basis=basis,explicit_unit=bool(unit_col),points=points));kind.add('dated schedules')
            # Explicit cumulative columns are a separate grain, never monthly observations.
            for h in range(2,min(ws.max_row,60)+1):
                cols={c:clean(ws.cell(h,c).value) for c in range(1,ws.max_column+1) if re.search(r'\bytd\b|year.to.date',str(ws.cell(h,c).value),re.I)}
                actual=next((c for c,v in cols.items() if 'actual' in v and 'prior' not in v),None)
                prior=next((c for c,v in cols.items() if 'prior' in v or 'last year' in v),None)
                budget=next((c for c,v in cols.items() if 'budget' in v or 'plan' in v),None)
                if not actual or not prior:continue
                ps={period(c.value) for row in ws.iter_rows(max_row=h-1) for c in row if period(c.value)}
                if len(ps)!=1:continue
                p=next(iter(ps));co,basis,currency,scale,_=metadata(ws,h)
                if scale is None or currency=='unknown':continue
                from ingestion import propose
                for row in range(h+1,ws.max_row+1):
                    for label_col in range(1,min(min(cols),5)):
                        label=ws.cell(row,label_col).value
                        category,role,status=propose(label,'',True)
                        if category not in {'revenue','ebitda'} or role!='total' or not status.startswith('Auto'):continue
                        points={}
                        for scenario,c in [('actual',actual),('prior_year',prior),('budget',budget)]:
                            if c:
                                source_period=(str(int(p[:4])-1)+p[4:]) if scenario=='prior_year' else p
                                f=fact(row,c,label,'money',source_period,scenario,company=co,basis=basis,currency=currency,scale=scale)
                                if f:f['grain']='ytd';points[scenario]=f
                        if {'actual','prior_year'}<=points.keys():package['cumulative'].append(dict(category=category,label=str(label),points=points));kind.add('cumulative financial comparisons')
                break

            # Dimensional Actual/Budget grids with entity names in a preceding header.
            for h in range(2,min(ws.max_row,60)+1):
                scenarios=[c for c in range(1,ws.max_column+1) if clean(ws.cell(h,c).value) in {'actual','budget','plan'}]
                if len(scenarios)<4:continue
                co,basis,currency,scale,_=metadata(ws,h-1)
                ps={period(c.value) for row in ws.iter_rows(max_row=h-1) for c in row if period(c.value)}
                if len(ps)!=1:
                    package['limitations'].append(f'{name} · {ws.title}: business-unit period is ambiguous.');continue
                p=next(iter(ps));entity='';cols=[]
                for c in range(min(scenarios),ws.max_column+1):
                    upper=ws.cell(h-1,c).value
                    if upper:entity=str(upper)
                    if c in scenarios and entity and not re.search(r'\btotal\b',entity,re.I):cols.append((c,entity,'budget' if clean(ws.cell(h,c).value) in {'budget','plan'} else 'actual'))
                label_col=max(range(1,min(scenarios)),key=lambda c:sum(isinstance(ws.cell(r,c).value,str) for r in range(h+1,ws.max_row+1)))
                for row in range(h+1,ws.max_row+1):
                    label=ws.cell(row,label_col).value
                    if not isinstance(label,str):continue
                    for c,entity,scenario in cols:
                        fmt=ws.cell(row,c).number_format
                        unit='ratio' if '%' in fmt or '%' in label else ('count' if re.search(r'headcount|calls completed|callbacks|units sold|employees',label,re.I) else ('money' if scale is not None and currency!='unknown' else 'unknown'))
                        f=fact(row,c,label,unit,p,scenario,entity,co,basis,currency,scale)
                        if f:package['units'].append(f);kind.add('business units')
                break
            # Accounts-receivable aging requires explicit aging buckets and a total.
            for h in range(1,min(ws.max_row,60)+1):
                headers={c:clean(ws.cell(h,c).value) for c in range(1,ws.max_column+1)}
                buckets={c:v for c,v in headers.items() if v in {'current','0 30','1 30','31 60','61 90','over 90','90','90 days','91','91 120','over 120'}}
                context=' '.join(str(c.value or '') for row in ws.iter_rows(max_row=h) for c in row)
                if not re.search(r'receivable|\bAR\b|trade debtors?',context,re.I):continue
                buckets={c:v for c,v in buckets.items() if v not in {'90','90 days'} or re.search(r'[+>]',str(ws.cell(h,c).value))}
                total=next((c for c,v in headers.items() if v in {'total','total ar','balance'}),None)
                if len(buckets)<3 or not total:continue
                co,basis,currency,scale,_=metadata(ws,h);ps={period(c.value) for row in ws.iter_rows(max_row=h-1) for c in row if period(c.value)}
                p=next(iter(ps)) if len(ps)==1 else ''
                labels=[c for c in range(1,min(buckets)) if headers[c] and headers[c] not in {'type','category'}]
                if not labels:continue
                label_col=labels[0];rows=[];reported=[];valid=True
                for row in range(h+1,ws.max_row+1):
                    label=ws.cell(row,label_col).value
                    if not label:continue
                    if re.search(r'^%|percent|reconciliation|trend',str(label),re.I):break
                    entries=[fact(row,c,label,'money',p,entity=str(label),company=co,basis=basis,currency=currency,scale=scale) for c in list(buckets)+[total]]
                    if not all(entries) or any(f['value']<0 for f in entries):valid=False;break
                    values={c:f for c,f in zip(list(buckets)+[total],entries)}
                    if abs(sum(values[c]['value'] for c in buckets)-values[total]['value'])>max(.01,abs(values[total]['value'])*1e-6):valid=False
                    item=dict(label=str(label),cells=values,total=values[total]['value'])
                    if re.search(r'^total\b',str(label),re.I):reported.append(item);break
                    rows.append(item)
                if not reported or not rows or any(abs(sum(r['cells'][c]['value'] for r in rows)-reported[0]['cells'][c]['value'])>.01 for c in list(buckets)+[total]):valid=False
                schedule=dict(file=name,sheet=ws.title,company=co,period=p,currency=currency,scale=scale,buckets=buckets,total_col=total,rows=rows,reported=reported,valid=valid and bool(p) and currency!='unknown' and scale is not None)
                ledger=[]
                for r in range(h+1,ws.max_row+1):
                    label=ws.cell(r,label_col).value
                    if not isinstance(label,str) or not re.search(r'per GL|general ledger|allowance|doubtful|net accounts receivable',label,re.I):continue
                    numeric=[]
                    for c in range(min(buckets),ws.max_column+1):
                        try:v=amount(book.value(ws.title,ws.cell(r,c).coordinate) if book else ws.cell(r,c).value)
                        except (InputError,ArithmeticError,ValueError,TypeError,SyntaxError):v=None
                        if v is not None:numeric.append(c)
                    if len(numeric)==1:
                        value=fact(r,numeric[0],label,'money',p,company=co,basis=basis,currency=currency,scale=scale)
                        if value:ledger.append(value)
                schedule['ledger']=ledger
                gl=next((f for f in ledger if re.search(r'per GL|general ledger',f['label'],re.I) and not re.search(r'allowance|doubtful',f['label'],re.I)),None)
                schedule['gl_tie']=abs(reported[0]['total']-gl['value'])<=.01 if gl and reported else None
                package['aging'].append(schedule);kind.add('receivables aging')
                if not schedule['valid']:package['limitations'].append(f'{name} · {ws.title}: aging totals, units or date could not be validated; exposure calculations withheld.')
                break
            # Narrative cells are quoted data, never system instructions.
            narrative_company=metadata(ws,min(5,ws.max_row+1))[0]
            narrative_periods={period(c.value) for row in ws.iter_rows(max_row=min(4,ws.max_row)) for c in row if period(c.value)}
            narrative_period=next(iter(narrative_periods)) if len(narrative_periods)==1 else ''
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value,str) and not c.value.startswith('=') and len(c.value.split())>=12:
                        package['narratives'].append(dict(id='quote_'+digest(name+'|'+ws.title+'|'+c.coordinate),file=name,sheet=ws.title,ref="'"+ws.title+"'!"+c.coordinate,text=c.value[:4000],company=narrative_company,period=narrative_period));kind.add('commentary')
            package['coverage'].append(dict(file=name,sheet=ws.title,recognized=sorted(kind),facts=len(package['facts'])-before))
        if len(w.worksheets)>MAX_SHEETS:package['limitations'].append(f'{name}: only the first {MAX_SHEETS} worksheets were inspected.')
    # A source cell has one immutable fact identity even when used in multiple schedules.
    package['facts']=list({f['id']:f for f in package['facts']}.values())
    return package
