"""Prepared-only GP dashboard. Run on 8502; the frozen analyst app is separate."""
import streamlit as st
import io
import pandas as pd
from openpyxl import load_workbook
from engine import InputError
from portfolio_data import load_companies, company_by_id
from portfolio_ui import shell, overview, footer, render_watchlist
from portfolio_excel import export_portfolio, SHEETS
from portfolio_watchlist import watchlist

st.set_page_config(page_title='Portfolio Intelligence · GP', page_icon='▣', layout='wide', initial_sidebar_state='expanded')


@st.cache_data(show_spinner=False)
def prepared_companies():
    return load_companies()


@st.cache_data(show_spinner=False)
def gp_workbook(companies):
    return export_portfolio(companies)


try:
    with st.spinner('Verifying prepared company reports…'):
        companies = prepared_companies()
except (InputError, OSError, ValueError, KeyError) as exc:
    st.error('Prepared reports could not be verified. Restore the checkpoint sources and run scripts/seed_demo_results.py --verify-only. '+str(exc))
    st.stop()

page = st.query_params.get('page', 'overview')
shell(page, companies)
workbook_data = gp_workbook(companies)
_, download_area = st.columns([3, 1])
download_area.download_button('↓ Download Portfolio Excel', workbook_data,
    file_name='GP_Portfolio_Review_Sep2026_FY2026.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    type='primary', width='stretch', key='gp_download')
if page == 'company':
    from portfolio_company import render_company
    company = company_by_id(companies, st.query_params.get('company', 'summit'))
    if company is None:
        st.warning('Unknown company. Select Portfolio Overview to open a prepared report.')
    else:
        render_company(company)
elif page == 'watchlist':
    st.title('Portfolio Watchlist')
    render_watchlist(companies)
elif page == 'excel':
    from portfolio_ui import stats
    st.title('One downloadable GP report')
    st.caption('Portfolio overview · Company comparison · Portfolio watchlist · Source references')
    stats(companies)
    st.info('Company figures remain separate. The downloadable workbook includes original finding narratives, caveats and source-linked calculations. Existing company reports remain available in each drilldown.')
    st.subheader('Workbook preview')
    book = load_workbook(io.BytesIO(workbook_data), data_only=True)
    for tab, name in zip(st.tabs(SHEETS), SHEETS):
        with tab:
            sheet = book[name]
            # Show the actual generated tables; watchlist detail narratives follow
            # its index in Excel and are also available in the portfolio watchlist.
            end = 5 + (len(companies) if name == SHEETS[0] else (len(companies)*3 if name == SHEETS[1] else (len(watchlist(companies)) if name == SHEETS[2] else sheet.max_row-5)))
            rows = list(sheet.iter_rows(min_row=5, max_row=end, values_only=True))
            frame = pd.DataFrame(rows[1:], columns=rows[0]).map(lambda x: str(x) if isinstance(x, (dict, list)) else x)
            st.dataframe(frame, hide_index=True, width='stretch')
            st.caption(f'{len(frame)} records · Preview reads the generated workbook')
else:
    overview(companies)
    render_watchlist(companies, compact=True)
footer()
