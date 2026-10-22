"""Company presentation contract: numbers, template ownership, and provenance."""
import io
import pytest
from openpyxl import load_workbook
from portfolio_data import load_companies
from portfolio_comparison import comparison
from portfolio_financial_reporting import export_company_report, statement
from portfolio_company_excel import company_filename
from portfolio_outlook import investment_outlook
from reports import export_report as frozen_export

@pytest.fixture(scope='module')
def companies():
    return load_companies()

def test_company_executive_numeric_kpis_context_and_complete_findings(companies):
    for company in companies:
        data, filename, _ = export_company_report(company.review)
        wb = load_workbook(io.BytesIO(data)); ws = wb['Executive Summary']
        expected = {'summit': 'Summit_Home_Services_Sep2026_MonthlyReview',
                    'tallis': 'Tallis_Brook_Accounting_Sep2026_MonthlyReview',
                    'harrow': 'Harrow_Ridge_Environmental_FY2026_AnnualReview'}[company.id]
        assert filename == expected + '.xlsx'
        assert ws['A1'].value == company.name
        assert ws['A2'].value.startswith(comparison(company)['period_label'])
        assert ws['A3'].value == 'Investment Review'
        assert ws['A4'].value == 'For Investment Team Review'
        assert len({ws.column_dimensions[chr(65+i)].width for i in range(12)}) == 1
        lines = {line.key: line for line in statement(company.review)}
        for col, key in (('A', 'revenue'), ('E', 'ebitda'), ('I', 'gm')):
            for row, attr in ((7, 'actual'), (9, 'budget'), (11, 'variance')):
                cell = ws[f'{col}{row}']
                assert cell.data_type == 'n'
                assert cell.value == pytest.approx(getattr(lines[key], attr))
                assert cell.hyperlink.location.startswith("'Financial Performance'!")
            assert ('%' in ws[f'{col}7'].number_format) if key == 'gm' else (company.review.currency in ws[f'{col}7'].number_format)
        assert 'pp' in ws['I11'].number_format
        texts = [cell.value for row in ws for cell in row if isinstance(cell.value, str)]
        for heading in ('Business Overview', 'Investment Outlook / Assessment', 'Priority Findings', 'Management Follow-Up'):
            cells = [cell for row in ws for cell in row if cell.value == heading]
            assert len(cells) == 1 and cells[0].font.bold and not cells[0].alignment.wrap_text
            assert any(cells[0].coordinate in area and area.size['columns'] == 12 for area in ws.merged_cells.ranges)
        assert company.review.package['company_profile']['description'] in texts
        outlook = investment_outlook(company)
        assert all(paragraph in texts for paragraph in outlook.paragraphs)
        for finding in company.review.package['ai_insights']:
            assert finding['conclusion'] in texts
            assert finding['action'] in texts
            assert finding['management_question'] in texts
        assert not any(text in ('money', 'ratio') for text in texts)
        assert all(dim.height <= 409 for dim in ws.row_dimensions.values() if dim.height)
        assert 'Investment Summary' not in wb
        assert wb.active.title == 'Executive Summary'
        # Complete analytical narratives and technical source inventories unchanged.
        original = load_workbook(io.BytesIO(frozen_export(company.review)[0]))
        for name in ('Findings', 'Source Evidence', 'Extracted Detail', 'Data Quality', 'Exceptions', 'Analyst Execution'):
            assert list(wb[name].values) == list(original[name].values), name


def test_filename_follows_approved_period_not_a_fixed_demo_date(companies):
    company = companies[0]
    context = dict(comparison(company), period='January 2027')
    assert company_filename(company.review, context) == 'Summit_Home_Services_Jan2027_MonthlyReview.xlsx'
    context.update(period='2028-FY', frequency='Annual')
    assert company_filename(company.review, context) == 'Summit_Home_Services_FY2028_AnnualReview.xlsx'
