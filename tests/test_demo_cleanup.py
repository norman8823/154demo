"""Presentation-only demo cleanup: contextual source actions and unchanged finance."""
import io
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from portfolio_data import load_companies
from portfolio_documents import evidence_locations
from portfolio_excel import export_portfolio
from portfolio_financial_reporting import export_company_report

@pytest.mark.parametrize('index',[0,1,2])
def test_summary_and_findings_only_keep_each_original_cell_action(index):
 at=AppTest.from_string(f'''
import streamlit as st
from portfolio_company import render_company_summary
from portfolio_data import load_companies
st.session_state.source_desktop=True
render_company_summary(load_companies()[{index}])
''',default_timeout=60).run()
 assert not at.exception
 assert [h.value for h in at.subheader]==['Investment Summary','Priority findings']
 company=load_companies()[index]
 findings=company.review.package['ai_insights']
 assert len([b for b in at.button if b.label=='View in Source'])==len(findings)
 assert not any(b.label in ('View source evidence','View Evidence  →') for b in at.button)
 assert not any(e.label=='Derived calculations & original operands' for e in at.expander)
 assert not any(s.label=='Original source cell' for s in at.selectbox)
 for finding in findings:
  assert any(b.key=='company_source_'+company.id+'_'+str(finding['id'])+'_view' for b in at.button)


def test_requested_id_columns_absent_other_columns_and_evidence_retained():
 companies=load_companies();book=load_workbook(io.BytesIO(export_portfolio(companies)))
 sheet=book['Portfolio Overview']
 assert [sheet.cell(14,c).value for c in (1,5,7,9,11)]==['Company','Current Equity','Gross MOIC','Gross IRR','Position']
 for row in sheet:
  assert not any(str(c.value).lower() in ('source ids','source id','fact ids','fact id') for c in row)
 assert len(sheet._charts)==0
 for c in companies:
  company=load_workbook(io.BytesIO(export_company_report(c.review)[0]))
  assert company['Financial Performance'].max_column==4
  assert [company['Financial Performance'].cell(6,col).value for col in range(1,5)]==['Income statement','Actual','Budget','Variance / pp']
  assert company['Source Evidence']['A5'].value=='Fact ID'
