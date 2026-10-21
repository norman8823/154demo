import csv
import io
from pathlib import Path
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from engine import analyze, demo_files, InputError, venue_key, ROOT
from reports import export_report, fields, TOKEN

@pytest.fixture
def source(): return demo_files()


def change_book(source,part,edit):
    files=dict(source)
    name=next(n for n in files if part in n)
    wb=load_workbook(io.BytesIO(files[name])); edit(wb)
    buf=io.BytesIO(); wb.save(buf); files[name]=buf.getvalue()
    return files


def test_calculations_and_lineage(source):
    r=analyze(source); m=r.metrics
    assert m['revenue']==5383752
    assert m['gm']==pytest.approx(1498370/5383752)
    assert m['gm_budget']==pytest.approx(1894600/5487900)
    assert m['subcontract_variance']==196803
    assert m['fill']==pytest.approx(16316/17525)
    assert m['leverage_reported']==pytest.approx(54580000/12730813)
    assert m['before_synergies']==10630813
    assert m['allowed_synergies']==pytest.approx(1594621.95)
    assert m['leverage_cap']==pytest.approx(54580000/12225434.95)
    assert m['unbilled']==20160
    assert m['underbilled']==7818.75
    assert m['billing_total']==27978.75
    assert m['duplicates_removed']==2 and m['schedule_rows']==64
    assert len(r.findings)==3
    assert any(e['Reference']=="'CEO Letter'!A6" for e in r.evidence)
    for finding in r.findings:
        assert any(e['Finding']==finding['id'] and e['Reference'] for e in r.evidence)
    assert set(x['CSV row'] for x in r.billing).isdisjoint({18,44})
    assert all(x['Invoice']=='(blank)' for x in r.billing if x['Exception']=='Uninvoiced worked hours')
    assert all(x['Invoice']!='(blank)' for x in r.billing if x['Exception']=='Below contract rate')


def test_uploaded_replacements_change_results_not_static(source):
    changed=change_book(source,'PlatformA',lambda w:setattr(w['P&L Trend']['M6'],'value',5000000))
    original=analyze(source); updated=analyze(changed)
    assert updated.metrics['revenue']-original.metrics['revenue']==5000000-4513673
    assert updated.metrics['gm']!=original.metrics['gm']
    assert str(updated.metrics['revenue']) in updated.memo.replace(',','')


def test_cap_is_bounded_by_reported_synergy(source):
    changed=change_book(source,'PlatformA',lambda w:setattr(w['Covenant']['B8'],'value',100000))
    r=analyze(changed)
    assert r.metrics['synergy_trim']==0
    assert r.metrics['leverage_cap']==r.metrics['leverage_reported']


def test_exact_duplicate_handling_and_multi_position_events(source):
    name=next(n for n in source if n.endswith('.csv'))
    original=analyze(source)
    files=dict(source)
    files[name]+=source[name].splitlines(keepends=True)[2]
    updated=analyze(files)
    assert updated.metrics['billing_total']==original.metrics['billing_total']
    assert updated.metrics['schedule_rows']==original.metrics['schedule_rows']
    assert updated.metrics['duplicates_removed']==3
    rows=list(csv.reader(io.StringIO(files[name].decode())))
    for row in rows[2:4]: row[-1]=''
    out=io.StringIO(); csv.writer(out).writerows(rows); files[name]=out.getvalue().encode()
    multi=analyze(files)
    assert len([x for x in multi.billing if x['Event ID']=='E9001'])==2

@pytest.mark.parametrize('venue',['Venue 4','Venue 04','Venue #4','venue 004','Venue 04 - Club Level'])
def test_venue_normalization(venue): assert venue_key(venue)=='Venue 4'


def test_unknown_venue_is_excluded_with_warning(source):
    files=dict(source); name=next(n for n in files if n.endswith('.csv'))
    files[name]=files[name].replace(b'Venue 21',b'Venue X')
    r=analyze(files)
    assert any('No exact supported' in w for w in r.warnings)


def test_standard_report_reopens_with_numeric_headlines(source):
    r=analyze(source); data,name,warnings=export_report(r)
    wb=load_workbook(io.BytesIO(data)); values=load_workbook(io.BytesIO(data),data_only=True)
    assert wb.sheetnames==['Executive Summary','Financial Performance','Findings','Source Evidence','Data Quality','Extracted Detail','Exceptions']
    assert name=='September_2026_Portfolio_Review.xlsx' and not warnings
    metrics={row[0]:row[1] for row in values['Financial Performance'].values}
    assert metrics['Platform A revenue']==r.metrics['revenue']
    assert metrics['Cap-only leverage']==pytest.approx(r.metrics['leverage_cap'])
    assert len([row for row in wb['Exceptions'].values if isinstance(row[0],int)])==3
    assert all(c.data_type!='f' for ws in wb for row in ws for c in row)


def test_custom_template_retains_layout_formula_style_and_guide(source):
    r=analyze(source); template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes()
    original=load_workbook(io.BytesIO(template))
    data,name,warnings=export_report(r,template)
    out=load_workbook(io.BytesIO(data))
    assert not warnings
    assert all(s in out.sheetnames for s in original.sheetnames)
    for ws in original:
        assert list(ws.merged_cells.ranges)==list(out[ws.title].merged_cells.ranges)
        for row in ws:
            for c in row:
                result=out[ws.title][c.coordinate]
                assert c._style==result._style
                if ws.title=='Template Guide' or c.data_type=='f' or not isinstance(c.value,str) or not TOKEN.search(c.value): assert c.value==result.value
    assert out['Partner Review']['C16'].data_type=='f'
    assert out['Partner Review']['C8'].value==r.metrics['revenue']
    assert out['Partner Review']['G8'].value==pytest.approx(r.metrics['gm'])
    assert out['Partner Review']['C12'].value==pytest.approx(r.metrics['leverage_cap'])
    assert out['Partner Review']['G16'].value==3
    assert 'Exceptions' in out.sheetnames and 'Source Evidence' in out.sheetnames
    assert not any(TOKEN.search(c.value) for row in out['Partner Review'] for c in row if isinstance(c.value,str))
    assert template==(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes()


def test_unmapped_fields_are_reported_and_retained(source):
    template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes()
    wb=load_workbook(io.BytesIO(template)); wb['Partner Review']['F6']='{{UNKNOWN_OWNER}}'
    buf=io.BytesIO(); wb.save(buf)
    data,_,warnings=export_report(analyze(source),buf.getvalue())
    assert warnings and 'UNKNOWN_OWNER' in warnings[0]
    assert load_workbook(io.BytesIO(data))['Partner Review']['F6'].value=='{{UNKNOWN_OWNER}}'


def test_export_does_not_execute_source_formula_text(source):
    r=analyze(source)
    r.memo='=HYPERLINK("https://example.invalid","test")'
    data,_,_=export_report(r)
    wb=load_workbook(io.BytesIO(data))
    assert all(c.data_type!='f' for ws in wb for row in ws for c in row)

@pytest.mark.parametrize('failure',['bad_workbook','bad_csv','bad_numeric','basis'])
def test_friendly_input_errors(source,failure):
    files=dict(source)
    if failure=='bad_workbook': files[next(n for n in files if n.endswith('.xlsx'))]=b'not excel'
    if failure=='bad_csv': files[next(n for n in files if n.endswith('.csv'))]=b'wrong,headers\n1,2\n'
    if failure=='bad_numeric': files=change_book(files,'PlatformA',lambda w:setattr(w['P&L Trend']['M6'],'value','bad number'))
    if failure=='basis': files=change_book(files,'PlatformA',lambda w:setattr(w['Cover']['A7'],'value','Cash'))
    if failure=='bad_csv':
        r=analyze(files)
        assert r.metrics['revenue']==5383752 and 'billing_total' not in r.metrics
        assert any('not analyzed' in warning for warning in r.warnings)
        export_report(r)
        return
    with pytest.raises(InputError): analyze(files)


def test_ui_demo_and_three_evidence_actions():
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30).run()
    assert not at.exception
    next(b for b in at.button if b.label=='Use Demo Files').click().run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception
    assert [m.value for m in at.metric]==['$5.38M','27.8%','$323K']
    for finding in range(1,4):
        if finding>1:
            at.radio(key='view').set_value('02  Portfolio Review').run()
        at.button(key=f'evidence_{finding}').click().run()
        assert not at.exception
        assert at.radio(key='selected').value==finding


def test_invalid_regeneration_cannot_download_previous_results():
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30).run()
    next(b for b in at.button if b.label=='Use Demo Files').click().run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    at.radio(key='view').set_value('01  Upload & Setup').run()
    next(r for r in at.radio if r.label=='Report format').set_value('Use existing firm template').run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert at.error and not at.exception
    at.radio(key='view').set_value('02  Portfolio Review').run()
    assert not at.metric
    assert any('Start with Upload' in info.value for info in at.info)
