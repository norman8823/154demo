"""Source-grounded acceptance: no answer-key data is used by runtime code."""
import io
from pathlib import Path
import pytest
from openpyxl import load_workbook,Workbook
from streamlit.testing.v1 import AppTest
from engine import analyze,ROOT,Book
from ingestion import inspect_files,apply_decisions,normalized_records
from financial_analysis import analyze_intake
from reports import export_report
from mapping_ai import automate

SUMMIT=ROOT/'tests/fixtures/Summit_Home_Services_Sep2026.xlsx'

def data(w):
    b=io.BytesIO();w.save(b);return b.getvalue()

def source():return {SUMMIT.name:SUMMIT.read_bytes()}

def standalone():
    """Keep the source's calculated facts before removing its supporting schedule."""
    w=load_workbook(SUMMIT);b=Book(SUMMIT.name,SUMMIT.read_bytes());b.strict_sum=False
    for row in w['Income Statement']:
        for c in row:
            if c.data_type=='f':c.value=b.value('Income Statement',c.coordinate)
    w.remove(w['Monthly Detail']);return w

def test_summit_one_click_matches_independent_source_expectations():
    r=analyze(source())
    expected={'revenue':3298000,'gross_profit':1433000,'operating_expenses':931000,
              'ebitda':502000,'net_income':183000,'revenue_budget':3207000,
              'ebitda_budget':659000,'revenue_prior_year':2971000,'ebitda_variance':-157000}
    for k,v in expected.items():assert r.metrics[k]==v
    assert r.metrics['gm']==pytest.approx(1433/3298)
    assert r.metrics['gm_variance_pp']==pytest.approx(100*(1433/3298-1577/3207))
    assert 'Overtime' in r.findings[0]['summary'] and '99,000' in r.findings[0]['summary']
    assert any('47,000' in f['summary'] and 'reconciliation' in f['title'] for f in r.findings)
    assert not any('Central' in f['summary'] for f in r.findings)
    assert all(any(e['Finding']==f['id'] and e['Reference'] for e in r.evidence) for f in r.findings)
    assert all('reserve looks inadequate' not in f['summary'].lower() for f in r.findings)
    assert r.package['aging'][0]['valid'] and r.package['aging'][0]['gl_tie'] is True
    assert any(s['theme']=='retention' for s in r.package['stats'])
    assert any('bridge' in row[3] for row in r.financial_rows if row[0].endswith('ebitda'))
    # Headline and monthly detail are corroborating sources, never additive packages.
    assert len([row for row in r.financial_rows if row[0].endswith('2026-09 · actual · revenue')])==1

def test_summit_changed_input_changes_analysis_not_static_answer_key():
    w=load_workbook(SUMMIT);w['Monthly Detail']['K18']=174
    r=analyze({'renamed_upload.xlsx':data(w)})
    assert r.metrics['ebitda']==492000 and r.metrics['gross_profit']==1423000
    assert r.metrics['ebitda_variance']==-167000
    assert any('109,000' in f['summary'] for f in r.findings)

@pytest.mark.parametrize('custom',[False,True])
def test_summit_excel_download_preserves_metrics_and_evidence(custom):
    r=analyze(source())
    template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    out,_,_=export_report(r,template)
    w=load_workbook(io.BytesIO(out))
    assert {'Financial Performance','Source Evidence','Exceptions'}<=set(w.sheetnames)
    assert any(row[1]==3298000 for row in w['Financial Performance'].values)
    assert any(row[1]==502000 for row in w['Financial Performance'].values)
    assert any('47,000' in str(v) for row in w['Exceptions'].values for v in row)
    if custom:assert w['Partner Review']['C16'].data_type=='f'

def test_summit_ui_generate_is_enabled_without_mapping_approval():
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30)
    at.session_state['uploaded_inputs']=source();at.run()
    assert not at.exception
    assert next(e for e in at.expander if e.label=='Review Account Mappings').proto.expanded is False
    button=next(b for b in at.button if b.label.startswith('Generate Portfolio Review'))
    assert not button.disabled
    button.click().run()
    assert not at.exception and at.session_state['review'].metrics['ebitda']==502000
    assert at.session_state['export']
    assert any('Reporting totals require reconciliation' in x.value for x in at.warning)
    assert any('$47K' in x.value for x in at.markdown)
    assert next(e for e in at.expander if e.label=='Data Quality & Coverage').proto.expanded is False

def test_unconfirmed_account_never_enters_derived_total_but_reported_totals_survive():
    w=standalone()
    ws=w['Income Statement'];ws['B9']='Acquisition reserve'
    i=inspect_files({'partial.xlsx':data(w)})
    r=analyze_intake(i)
    assert r.metrics['revenue']==3298000 and r.metrics['ebitda']==502000
    pending=next(x for x in r.normalized if x['label']=='Acquisition reserve')
    assert not pending['accepted']
    assert any('unconfirmed' in x for x in r.warnings)
    # Without the reported revenue total, partial accepted detail cannot become revenue.
    ws['B15']=None
    r=analyze({'partial.xlsx':data(w)})
    assert 'revenue' not in r.metrics and 'gm' not in r.metrics
    assert r.metrics['net_income']==183000

def test_isolated_unsupported_formula_does_not_block_other_metrics():
    w=standalone()
    w['Income Statement']['D48']='=UNSUPPORTED(1)'
    w['Income Statement']['D53']='=D44-SUM(D48:D51)'
    r=analyze({'formula.xlsx':data(w)})
    assert r.metrics['revenue']==3298000 and r.metrics['ebitda']==502000
    assert 'net_income' not in r.metrics
    assert any('unsupported' in x.lower() for x in r.warnings)

def test_conflicting_sources_withhold_only_conflicting_metrics_even_with_third_copy():
    w=load_workbook(SUMMIT)
    ws=w.copy_worksheet(w['Income Statement']);ws.title='Second statement';ws['D15']=3398
    third=w.copy_worksheet(w['Income Statement']);third.title='Third statement'
    r=analyze({'conflict.xlsx':data(w)})
    assert 'revenue' not in r.metrics and 'gm' not in r.metrics
    assert r.metrics['operating_expenses']==931000
    assert any('disagree' in f['summary'] for f in r.findings)

def test_unverified_ai_period_cannot_become_a_financial_fact():
    files=source();i=inspect_files(files);t=i.tables[0]
    cols=[dict(c) for c in t.columns];cols[1]['period']='2029-01'
    proposed={SUMMIT.name+'|'+t.sheet:dict(header_row=t.header_row,label_column=t.label_column,columns=cols,ai_pending=True)}
    r=analyze(files,layouts=proposed)
    assert r.report_period=='September 2026'
    assert not any('2029-01' in row[0] for row in r.financial_rows)
    assert any('not verified' in w for w in r.warnings)

def test_llm_labels_are_accepted_only_after_complete_source_total_validation(monkeypatch):
    import mapping_ai
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    w=Workbook();ws=w.active;ws.title='Statement';ws['A1']='Example Company';ws['A2']='USD whole dollars accrual basis'
    for row in [('Account','Mar-28 Actual'),('Alpha',600),('Beta',400),('Total revenue',1000),('Cost of sales',300),('Gross profit',700)]:ws.append(row)
    files={'unknown.xlsx':data(w)};i=inspect_files(files)
    pending=[a for t in i.tables for a in t.accounts if 'Review required' in a.status]
    assert {a.label for a in pending}=={'Alpha','Beta'}
    proposals={'accounts':[dict(id=a.id,category='revenue',role='detail',reason='Source revenue detail') for a in pending],'columns':[],'layouts':[]}
    monkeypatch.setattr(mapping_ai,'suggest',lambda *args:proposals)
    mapped=automate(i,files)
    assert all(a.status.startswith('Auto') for t in mapped.tables for a in t.accounts)
    assert analyze_intake(mapped).metrics['revenue']==1000
    # Same labels/AI answer cannot bypass a source total mismatch.
    ws['B6']=1100
    changed={'unknown.xlsx':data(w)}
    unchecked=automate(inspect_files(changed),changed)
    assert any('Review required' in a.status for t in unchecked.tables for a in t.accounts)

def test_llm_failure_is_nonblocking_and_never_approves_ambiguous_accounts(monkeypatch):
    import mapping_ai
    from engine import InputError
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    monkeypatch.setattr(mapping_ai,'suggest',lambda *args:(_ for _ in ()).throw(InputError('offline failure')))
    w=standalone();w['Income Statement']['B48']='Unfamiliar expense description'
    files={'unknown.xlsx':data(w)}
    r=analyze(files)
    assert r.metrics['revenue']==3298000
    assert any('AI interpretation was unavailable' in x for x in r.warnings)

def test_unrecognized_extra_file_does_not_block_the_original_demo():
    from engine import demo_files
    w=Workbook();w.active['A1']='Unsupported notes only'
    r=analyze({**demo_files(),'notes.xlsx':data(w)})
    assert r.metrics['revenue']==5383752 and r.metrics['billing_total']==27978.75
    assert any('not analyzed' in warning for warning in r.warnings)

def test_an_unconfirmed_ignore_proposal_cannot_hide_an_ambiguous_amount():
    p=ROOT/'tests/fixtures/B_Harbor.xlsx';i=inspect_files({p.name:p.read_bytes()})
    reserve=next(a for t in i.tables for a in t.accounts if a.label=='Mobilization reserve')
    apply_decisions(i,{reserve.id:dict(category='operating_expenses',role='ignore',confirmed=False)})
    r=analyze_intake(i)
    assert r.metrics['operating_expenses']==92000  # reported total survives, not an incomplete 80000
    assert any(x['label']==reserve.label and not x['accepted'] for x in r.normalized)
    assert any('unconfirmed' in warning for warning in r.warnings)
    for t in i.tables:
        for a in t.accounts:
            if a.category=='operating_expenses' and a.role=='total':a.role='ignore'
    assert 'operating_expenses' not in analyze_intake(i).metrics

def test_explicitly_conflicting_bases_are_not_treated_as_merely_unstated():
    w=load_workbook(ROOT/'tests/fixtures/A_Northstar.xlsx')
    w['Trading results']['B4']='USD in thousands; Accrual basis / Cash basis'
    r=analyze({'mixed.xlsx':data(w)})
    assert r.metrics['revenue']==1000000
    assert 'revenue_variance' not in r.metrics and 'ebitda_variance' not in r.metrics
    assert any('basis conflicts' in warning for warning in r.warnings)
