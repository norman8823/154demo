"""Reporting integrity, accepted-only selection, period labels and source navigation."""
import copy
import io
import pytest
from openpyxl import load_workbook
from engine import analyze, ROOT
from executive import present, briefing
from reports import export_report, validate_export
from excel_layout import financial_table, comparison_text
from test_ai_investment import config, files, fake_model


@pytest.mark.parametrize('custom',[False,True])
def test_authoritative_accepted_findings_and_clickable_source_records(monkeypatch,custom):
    def extend(result):
        original=result['insights'][0]
        for index in range(3):
            item=copy.deepcopy(original);item['title']=['Delivery recovery needs follow-up','Operational assurance needs review','Execution constraints warrant review'][index]
            result['insights'].append(item)
        rejected=copy.deepcopy(original);rejected['title']='Rejected unsupported amount';rejected['text']='The impact is $999,000.'
        result['insights'].append(rejected)
        return result
    fake_model(monkeypatch,extend)
    r=analyze(files(),ai_authorized=True);before=copy.deepcopy(r)
    template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    data,_,_=export_report(r,template);wb=load_workbook(io.BytesIO(data));summary=present(r)
    accepted={f['id'] for f in r.package['ai_insights']}
    assert len(accepted)==4 and len(summary['cards'])==3
    assert {c['id'] for c in summary['cards']}<=accepted
    assert {c['id'] for c in summary['secondary']}<=accepted
    finding_cells=[row for row in wb['Findings'] if isinstance(row[0].value,int)]
    assert {row[0].value for row in finding_cells}==accepted
    assert [row[0].value for row in finding_cells[:3]]==[c['id'] for c in summary['cards']]
    for row in finding_cells:
        target=row[6].hyperlink.target
        assert target.startswith("#'Source Evidence'!A")
        evidence_row=int(target.split('!A')[-1])
        assert wb['Source Evidence'].cell(evidence_row,2).value==row[0].value
        assert wb['Source Evidence'].cell(evidence_row,9).hyperlink.target==f"#'Findings'!A{row[0].row}"
    text=' '.join(str(c.value) for ws in wb for row in ws for c in row if c.value is not None)
    assert '999,000' not in text and 'Rejected unsupported amount' not in text
    assert 'input_tokens' not in text and 'Operational diagnostics' not in text
    assert 'Draft Partner Briefing' not in text and 'QUESTIONS FOR THE CFO' not in text
    labels={row[0].value:row[1].value for row in wb['Executive Summary']}
    assert labels['Investment Assessment']==briefing(summary)
    for i,q in enumerate(summary['questions'],1):assert labels[f'Question {i}']==q
    assert wb['Executive Summary'].page_setup.fitToHeight==1
    assert wb['Executive Summary']['A7'].font.color.rgb=='0018364A'
    assert r==before  # Exporting and projecting a cached review cannot mutate analysis.


def test_reporting_pivot_uses_explicit_months_and_keeps_unavailable_year_missing(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    r=analyze(files())
    row=next(row for row in financial_table(r) if row[0]=='Revenue')
    assert row[1:6]==[1000,1200,-200,pytest.approx(-1/6),None]
    assert row[7]==900 and row[8]==pytest.approx(1/9)
    assert row[10]=='Sep 2028'
    r.normalized[0]['currency']='unknown'
    row=next(row for row in financial_table(r) if row[0]=='Revenue')
    assert row[2] is None and row[7] is None


def test_ambiguous_legacy_comparison_is_replaced_from_cited_dates(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    r=analyze(files())
    refs=[e['Fact ID'] for e in r.evidence if e['Fact ID'].startswith('financial_') and 'revenue / total' in e['Check / use']]
    text=comparison_text('Revenue changed versus prior available period.',r,refs)
    assert 'prior available period' not in text
    assert 'Aug 2028' in text and 'month-over-month' in text


def test_original_source_values_are_literal_and_tampering_is_rejected(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    r=analyze(files());r.evidence[0]['Original value / formula']='=HYPERLINK("https://invalid.example","source")'
    data,_,_=export_report(r);wb=load_workbook(io.BytesIO(data))
    matching=[row for row in wb['Source Evidence'] if row[0].value==r.evidence[0]['Fact ID']]
    assert all(row[5].data_type=='s' for row in matching)
    for row in matching:row[5].value='changed'
    out=io.BytesIO();wb.save(out)
    with pytest.raises(ValueError,match='original source evidence changed'):validate_export(out.getvalue(),r)
