"""GP presentation, exact supporting-data preservation and print conventions."""
import io
from unittest.mock import patch
import pytest
from openpyxl import load_workbook
from portfolio_data import load_companies
from portfolio_excel import export_portfolio
from portfolio_investments import investments_for, investment_summary
from portfolio_gp_presentation import HOLDINGS_HEAD, assessment

@pytest.fixture(scope='module')
def exports():
    companies=load_companies()
    with patch('portfolio_gp_presentation.polish_gp_workbook',lambda *a:None):
        baseline=load_workbook(io.BytesIO(export_portfolio(companies)))
    current=load_workbook(io.BytesIO(export_portfolio(companies)))
    return companies,baseline,current


def signature(sheet, *, styles=False):
    return [(c.coordinate,c.value,c.data_type,c.number_format,
             (c.hyperlink.target,c.hyperlink.location) if c.hyperlink else None,
             (str(c.font),str(c.fill),str(c.border),str(c.alignment)) if styles else None)
            for row in sheet for c in row if c.value is not None]


def test_only_three_gp_sheets_change_and_every_supporting_amount_link_is_preserved(exports):
    companies,old,new=exports
    assert old.sheetnames==new.sheetnames
    for name in old.sheetnames:
        if name=='Portfolio Overview':continue
        presentation=name in ('Portfolio Financial Performance','Investment Ledger')
        assert signature(old[name],styles=not presentation)==signature(new[name],styles=not presentation),name
        assert str(old[name].merged_cells)==str(new[name].merged_cells),name
    assert len([i for c in companies for i in c.review.package['ai_insights']])==17


def test_executive_kpis_holdings_and_editorial_scope_reconcile(exports):
    companies,_,book=exports;ws=book['Portfolio Overview']
    total,positions,returns,operating=assessment(companies)
    assert ws['A10'].value==returns and ws['A11'].value==operating
    assert 'Summit Home Services is the strongest' in returns
    assert 'Tallis Brook Accounting is the weakest' in returns
    for col,key in ((1,'current_equity_value'),(4,'gross_moic'),(7,'gross_irr'),(10,'invested_capital')):
        c=ws.cell(6,col)
        assert c.data_type=='n' and c.value==pytest.approx(total[key])
        assert c.hyperlink.location=="'Investment Ledger'!A5"
    for index,(company,result) in enumerate(positions,HOLDINGS_HEAD+1):
        assert ws.cell(index,1).value==company.name
        assert ws.cell(index,5).value==result['current_equity_value']
        assert ws.cell(index,7).value==result['gross_moic']
        assert ws.cell(index,9).value==pytest.approx(result['gross_irr'])
    assert ws.page_setup.fitToWidth==1 and ws.page_setup.fitToHeight==1
    assert len({ws.column_dimensions[chr(65+i)].width for i in range(12)})==1
    for row in (9,13):
        assert ws.cell(row,1).font.bold and not ws.cell(row,1).alignment.wrap_text
    text=' '.join(str(c.value) for row in ws for c in row if c.value)
    for expected in ('September 2026 · monthly','FY2026 · annual','receivables','work-in-progress','capital requirements','not established causes'):
        assert expected in text
    for c in companies:
        for f in c.review.package['ai_insights']:
            assert f['text'] not in text  # Complete analysis stays on the findings sheet.


def test_statement_alignment_and_ledger_pagination(exports):
    companies,old,new=exports
    statement=new['Portfolio Financial Performance']
    assert statement.print_title_rows=='$1:$6'
    assert statement.page_setup.fitToHeight==1
    for col in (2,5,8):
        assert statement.cell(4,col).alignment.horizontal=='center'
        assert not statement.cell(4,col).alignment.wrap_text
        for c in range(col,col+3):assert statement.cell(6,c).alignment.horizontal=='center'
    assert '%' in statement['B10'].number_format and 'pp' in statement['D10'].number_format
    ledger=new['Investment Ledger']
    assert ledger.print_title_rows=='$1:$2'
    starts=[c.row for row in ledger for c in row if c.column==1 and c.value=='Company investment']
    assert [b.id for b in ledger.row_breaks.brk]==[16]+[r-1 for r in starts]
    assert ledger['B9'].value==58000000 and ledger['D9'].value==82000000 and ledger['F9'].value==1.5
    assert ledger['G9'].value==pytest.approx(investment_summary(investments_for(companies))['gross_irr'])
