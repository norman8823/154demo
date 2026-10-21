"""Independent fixture authoring. Expected values were recorded first in fixtures/expected_adaptive.json."""
from pathlib import Path
from openpyxl import Workbook
ROOT=Path(__file__).parent/'fixtures'

def create():
    wb=Workbook(); ws=wb.active; ws.title='Trading results'
    ws['B2']='Northstar Components';ws['B3']='Management income statement — April 2027';ws['B4']='USD thousands | Accrual basis'
    ws.append([])
    for col,val in zip('BCDEF',['Ledger caption','Apr 2027 Actual','Apr 2027 Plan','Mar 2027 Actual','Apr 2026 Actual']):ws[f'{col}7']=val
    rows=[('Revenue',),('Product shipments',720,750,650,600),('Service retainers',280,250,250,230),('Total revenue',1000,1000,900,830),('Cost of sales',),('Raw material consumption',320,300,270,240),('Shop-floor wages',180,160,160,150),('Total cost of sales',500,460,430,390),('Gross profit',500,540,470,440),('Operating expenses',),('Marketing programmes',60,50,40,35),('HQ payroll',90,90,85,80),('Facilities costs',30,30,30,25),('Total operating expenses',180,170,155,140),('EBITDA',320,370,315,300),('Net earnings',270,320,265,250)]
    for i,row in enumerate(rows,8):
        for j,val in enumerate(row,2):ws.cell(i,j,val)
    wb.save(ROOT/'A_Northstar.xlsx')
    wb=Workbook(); cover=wb.active;cover.title='Read me';cover['A1']='Monthly accounting export';cover['A2']='Figures in the Export Ledger tab'
    ws=wb.create_sheet('Export Ledger')
    for row,val in enumerate(['Harbor Field Services','QuickBooks profit and loss','Cash basis','USD (whole dollars)','Reporting month: February 2027'],1):ws.cell(row,2,val)
    for col,val in [(1,'GL code'),(2,'Description'),(4,'Feb 2027'),(5,'Jan 2027'),(6,'Total')]:ws.cell(9,col,val)
    rows=[('', 'Income'),('401','Field response receipts',270000,250000),('409','Maintenance subscriptions',130000,100000),('', 'Total income',400000,350000),('', 'Cost of goods sold'),('501','Subcontract crews',160000,145000),('509','Fleet materials',40000,30000),('', 'Total cost of goods sold',200000,175000),('', 'Gross profit',200000,175000),('', 'Expenses'),('601','Office payroll',70000,65000),('645','Mobilization reserve',12000,8000),('651','Depreciation',10000,9000),('', 'Total expenses',92000,82000),('', 'Net operating income',108000,93000),('', 'Other expenses'),('701','Interest expense',8000,7000),('702','Income tax expense',6000,5000),('', 'Net income',94000,81000)]
    for i,row in enumerate(rows,10):
        ws.cell(i,1,row[0]);ws.cell(i,2,row[1])
        if len(row)>2:ws.cell(i,4,row[2]);ws.cell(i,5,row[3]);ws.cell(i,6,row[2]+row[3])
    wb.save(ROOT/'B_Harbor.xlsx')
if __name__=='__main__':create()
