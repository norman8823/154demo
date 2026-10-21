"""Generic fixtures only: no holdout answers or paid inference in these tests."""
import copy
import io
import json
import time
import zipfile
from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from engine import InputError,Review
from skill_tools import WorkbookTools,preflight
from skill_output import validate,CHECKS


def files():
    w=Workbook();s=w.active;s.title='Accounts'
    s.append(['Example Services']);s.append(['$ in thousands']);s.append(['Accrual basis'])
    s.append(['Metric','FY2026','FY2025'])
    s.append(['Revenue',100,90]);s.append(['Operating cash flow',25,20]);s.append(['Cash capex',-15,-10]);s.append(['Cash after capex','=SUM(B6:B7)'])
    s.append(['Management note','Demand may soften.']);s.append(['Late row',1])
    b=io.BytesIO();w.save(b);return {'example.xlsx':b.getvalue()}


def read_all(t):
    for s in t.book.wb:
        off=0
        while True:
            page=t.read_range(s.title,f'A1:{s.cell(s.max_row,s.max_column).coordinate}',off,5)
            off=page['next_offset']
            if off is None:break


def finding():
    return dict(title='Cash availability requires attention',conclusion='Reinvestment reduces available cash.',evidence='Operating cash flow is {s1_B6}; cash capex is {s1_B7}.',
        financial_implication='Cash after capex is {calc_1}.',assessment='Further funding needs remain unverified.',action='Request a funded capital plan and lender availability confirmation.',
        management_question='Which capital commitments require additional funding?',category='cash_conversion',kind='risk',priority='high',disposition='investigate',
        evidence_status='evidence_backed_interpretation',fact_ids=['s1_B6','s1_B7','calc_1'],source_quotes=[dict(source_id='s1_B9',quote='Demand may soften.')])


def completion_checks():return [dict(area=a,status='tested',fact_ids=['s1_B6'],explanation='Relevant source reviewed; funding and verification uncertainty remains.') for a in CHECKS]


def proposal():return dict(findings=[finding()],rejected_candidates=[],coverage_limitations=['No lender terms supplied.'],completion_checks=completion_checks())


def calc(t):return t.calculate('s1_B6+s1_B7','Cash after capex','money','Company FY2026 cash flow; same denomination and period','')


def test_paging_preserves_later_cells_notes_formulas():
    t=WorkbookTools(files());read_all(t)
    assert all(c['read']==c['total'] for c in t.coverage())
    assert t.facts['s1_B10']['source_value']==1
    assert t.facts['s1_B8']['value']==10000
    assert not t.errors
    p=t.search('Demand',0,1);assert p['matches'][0]['id']=='s1_B9'
    assert p['next_offset'] is None


@pytest.mark.parametrize('method,args',[
 ('read_range',dict(sheet='../other.xlsx',cell_range='A1',offset=0,limit=5)),
 ('read_range',dict(sheet='Accounts',cell_range="'[other.xlsx]Sheet'!A1",offset=0,limit=5)),
 ('read_range',dict(sheet='Accounts',cell_range='A1:ZZ99999',offset=0,limit=5)),
 ('read_range',dict(sheet='Accounts',cell_range='A1:B5',offset=-1,limit=5)),
 ('search',dict(query='x',offset=0,limit=10000)),('shell',dict(command='cat .env'))])
def test_tools_cannot_access_other_files_or_exceed_bounds(method,args):
    with pytest.raises(InputError):WorkbookTools(files()).invoke(method,args)


@pytest.mark.parametrize('expression',['__import__("os").system("pwd")','s1_B6**100000','s1_B6/0','unknown+1','s1_B6.__class__','[s1_B6][0]'])
def test_restricted_calculation(expression):
    t=WorkbookTools(files());read_all(t)
    with pytest.raises(InputError):t.calculate(expression,'bad','money','same','')


def test_evidence_and_calculation_recomputed():
    t=WorkbookTools(files());read_all(t);calc(t)
    assert t.verify_fact('calc_1')==10000
    accepted,diag=validate(proposal(),t);assert len(accepted)==1,diag
    t.facts['calc_1']['value']=99999
    accepted,diag=validate(proposal(),t);assert not accepted and any(x['status']=='rejected' for x in diag)


def test_numeric_and_quote_rejection_preserves_valid_siblings():
    t=WorkbookTools(files());read_all(t);calc(t)
    p=proposal();bad=copy.deepcopy(finding());bad['evidence']='Cash is $900 million.';p['findings'].append(bad)
    a,d=validate(p,t);assert len(a)==1 and any(x['status']=='rejected' for x in d)
    p=proposal();p['findings'][0]['source_quotes'][0]['quote']='Guaranteed growth.'
    assert not validate(p,t)[0]


def test_assumptions_are_illustrative():
    t=WorkbookTools(files());read_all(t);calc(t)
    f=t.calculate('calc_1*4','Scenario','money','Quarter repeated; same scope','Assume constant quarterly amount')
    assert f['illustrative'] and f['constants']==[4]
    p=proposal();p['findings'][0]['financial_implication']='Cash is {calc_2}.';p['findings'][0]['fact_ids'].append('calc_2')
    accepted,diagnostics=validate(p,t)
    assert accepted and 'Includes illustrative calculations' in accepted[0]['financial_implication']
    assert any(d.get('code')=='illustrative_label_added' for d in diagnostics)
    assert t.facts['calc_2']['illustrative']
    p['findings'][0]['financial_implication']='Illustrative cash sensitivity is {calc_2}.'
    assert validate(p,t)[0]


def test_archive_bounds_and_macro_rejection():
    with pytest.raises(InputError):preflight(dict(files(),**{'second.xlsx':b'anything'}))
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:z.writestr('xl/vbaProject.bin',b'macro')
    with pytest.raises(InputError):preflight({'evil.xlsx':b.getvalue()})
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:z.writestr('xl/workbook.xml','<!DOCTYPE x [<!ENTITY x "bad">]><x/>')
    with pytest.raises(InputError):preflight({'evil.xlsx':b.getvalue()})


def response_calls(calls):
    return dict(status='completed',output=[dict(type='function_call',name=n,arguments=json.dumps(a),call_id='c'+str(i)) for i,(n,a) in enumerate(calls)],usage=dict(input_tokens=100,output_tokens=100,input_tokens_details=dict(cached_tokens=20)))


def transport():
    responses=[response_calls([('read_range',dict(sheet='Accounts',cell_range='A1:C10',offset=0,limit=500))]),
        response_calls([('calculate',dict(expression='s1_B6+s1_B7',label='Cash after capex',unit='money',scope='Company FY2026 cash flows',assumptions=''))]),
        dict(status='completed',output=[dict(type='message',content=[dict(type='output_text',text=json.dumps(proposal()))])],usage=dict(input_tokens=100,output_tokens=100))]
    def send(req):return responses.pop(0)
    return send


@pytest.fixture
def configured(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','fake')
    monkeypatch.setenv('OPENAI_MODEL','gpt-4o-mini')
    monkeypatch.setenv('COPILOT_AI_MAX_REVIEW_USD','.25')
    monkeypatch.delenv('OPENAI_INPUT_USD_PER_M',raising=False);monkeypatch.delenv('OPENAI_OUTPUT_USD_PER_M',raising=False)
    monkeypatch.setattr('engine._analyze_package',lambda *a,**k:Review(adaptive=True,company='Example Services',report_period='FY2026',basis_note='Accrual',currency='$',scale=1000,package=dict(facts=[],findings=[])))


def test_vertical_output_compatible_with_exports_and_cache(configured):
    from skill_engine import run
    from report_cache import successful,cache_key
    from reports import export_report
    from executive import present,report_findings
    r=run(files(),request_fn=transport());assert r.package['skill_engine']['status']=='completed',r.package['skill_engine']
    assert successful(r,require_real=False);assert not successful(r)
    assert r.package['inference']['events'][0]['cached_input_tokens']==20
    assert len(present(r)['cards'])==len(report_findings(r))==1
    for template in [None,Path('sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes()]:
        data,_,_=export_report(r,template);w=load_workbook(io.BytesIO(data))
        assert 'Analyst Execution' in w.sheetnames
        assert any('Cash availability' in str(row) for row in w['Findings'].values)
        assert any('10000' in str(row) for row in w['Source Evidence'].values)
    assert cache_key(files(),engine_name='skill')!=cache_key(files(),engine_name='legacy')


@pytest.mark.parametrize('limit,value,reason',[('ROUNDS','1','round'),('TOOL_CALLS','1','tool-call'),('SECONDS','1','elapsed')])
def test_execution_limits_fail_closed(configured,monkeypatch,limit,value,reason):
    from skill_engine import run
    monkeypatch.setenv('COPILOT_SKILL_'+limit,value)
    send=transport()
    if limit=='SECONDS':
        def send(req):time.sleep(1.02);return response_calls([])
    r=run(files(),request_fn=send)
    assert r.package['skill_engine']['status']=='failed'
    assert not r.package['ai_insights'] and not r.package['view']['cards']
    assert reason in r.package['skill_engine']['reason'] or limit in {'SECONDS','TOOL_CALLS'}


def test_spend_authorization_and_incomplete_response(configured,monkeypatch):
    from skill_engine import run
    r=run(files(),ai_authorized=False);assert not r.package['inference']['events']
    monkeypatch.setenv('COPILOT_AI_MAX_REVIEW_USD','.000001')
    r=run(files(),request_fn=transport());assert r.package['skill_engine']['status']=='failed' and not r.package['inference']['events']
    monkeypatch.setenv('COPILOT_AI_MAX_REVIEW_USD','.25')
    r=run(files(),request_fn=lambda req:dict(status='incomplete',output=[]))
    assert r.package['skill_engine']['status']=='failed'


def test_unread_sheets_not_published(configured):
    from skill_engine import run
    final=dict(status='completed',output=[dict(type='message',content=[dict(type='output_text',text=json.dumps(proposal()))])])
    r=run(files(),request_fn=lambda req:final)
    assert r.package['skill_engine']['status']=='failed' and 'coverage' in r.package['skill_engine']['reason']


def test_snapshot_retains_all_read_values_and_labels():
    t=WorkbookTools(files());read_all(t);calc(t)
    snapshot=t.snapshot();ids={s['id_prefix']+r[0] for s in snapshot['sheets'] for r in s['cells']}
    assert ids==set(t.facts)-{'calc_1'}|set(t.notes)
    assert 'Late row' in json.dumps(snapshot)
    assert snapshot['calculations'][0]['expression']=='s1_B6+s1_B7'


def test_units_cannot_be_relabelled_or_mixed():
    t=WorkbookTools(files());read_all(t)
    with pytest.raises(InputError):t.calculate('s1_B6+s1_B7','Cash','ratio','same','')
    t.calculate('s1_B6/s1_B5','Margin','ratio','same','')
    with pytest.raises(InputError):t.calculate('s1_B6+calc_1','Bad units','money','same','')


def test_cache_includes_methodology_tools_and_limits(monkeypatch):
    from report_cache import cache_key
    from skill_tools import LIMITS
    before=cache_key(files(),engine_name='skill')
    monkeypatch.setenv('COPILOT_SKILL_ROUNDS','7')
    assert cache_key(files(),engine_name='skill')!=before
    before=cache_key(files(),engine_name='skill')
    monkeypatch.setitem(LIMITS,'cells',29999)
    assert cache_key(files(),engine_name='skill')!=before


def test_sparse_large_grid_rejected():
    w=Workbook();w.active['CB2000']=1
    b=io.BytesIO();w.save(b)
    with pytest.raises(InputError):preflight({'sparse.xlsx':b.getvalue()})


def test_first_calculation_does_not_disable_further_tools(configured):
    from skill_engine import run
    send=transport();seen=[]
    def inspect(req):
        seen.append(json.loads(req.data));return send(req)
    r=run(files(),request_fn=inspect)
    assert r.package['skill_engine']['status']=='completed'
    assert seen[-1].get('tool_choice')!='none'


def test_receipt_settlement_keeps_unknown_reservations(monkeypatch):
    from ai_runtime import complete,inference_session
    import urllib.request
    monkeypatch.setenv('OPENAI_API_KEY','fake');monkeypatch.setenv('OPENAI_MODEL','gpt-4o-mini')
    class Reply:
        def __enter__(self):return io.StringIO(json.dumps(dict(status='completed',model='gpt-4o-mini',output=[],usage=dict(input_tokens=100,output_tokens=50))))
        def __exit__(self,*a):pass
    monkeypatch.setattr(urllib.request,'urlopen',lambda *a,**k:Reply())
    with inference_session(True) as s:
        complete(dict(input='test',max_output_tokens=100),'test',settle_usage=True)
        assert s.reserved==pytest.approx(s.events[0]['estimated_usd'])
        assert s.events[0]['settled_usd']<s.events[0]['reserved_upper_usd']


def test_historical_comparisons_are_not_scenarios():
    t=WorkbookTools(files());read_all(t)
    f=t.calculate('s1_B5/s1_C5-1','Growth','ratio','FY current versus prior','No assumptions; source comparison.')
    assert not f['illustrative']
    g=t.calculate('s1_B5*0.9','Scenario','money','Same company','',illustrative=False)
    assert g['illustrative']


def test_quote_case_rebinding_does_not_accept_changed_words():
    t=WorkbookTools(files());read_all(t);calc(t)
    p=proposal();p['findings'][0]['source_quotes'][0]['quote']='demand may soften.'
    a,d=validate(p,t);assert a and any(x.get('code')=='quote_case' for x in d)
    p['findings'][0]['source_quotes'][0]['quote']='Demand will soften.'
    assert not validate(p,t)[0]


def test_negated_causation_does_not_exempt_later_assertions():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    assert grounded('This is not proof that costs caused the decline.',['s1_B6'],t.facts,40)
    with pytest.raises(InputError):grounded('This is not proof that costs caused the decline; demand caused it.',['s1_B6'],t.facts,40)
    with pytest.raises(InputError):grounded('This is not proof that costs caused it, but demand caused it.',['s1_B6'],t.facts,40)


def test_fiscal_year_label_requires_own_source_header():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    assert grounded('FY2026 revenue was {s1_B5}.',['s1_B5'],t.facts,30)
    with pytest.raises(InputError):grounded('FY2027 revenue was {s1_B5}.',['s1_B5'],t.facts,30)
    with pytest.raises(InputError):grounded('FY2026 revenue was {s1_C5}.',['s1_C5'],t.facts,30)


def test_explicit_unreviewed_evidence_is_uncertainty():
    t=WorkbookTools(files());read_all(t);calc(t)
    p=proposal();p['findings'][0]['assessment']='No executed contract or customer confirmation was reviewed. The sensitivity should not be treated as expected loss.'
    assert validate(p,t)[0]
    p['findings'][0]['assessment']='The contract proves the expected loss.'
    assert not validate(p,t)[0]


@pytest.mark.parametrize('formula',['=SUM(A1:XFD1048576)','=XFD1048576','="text"*1000000000'])
def test_preflight_bounds_legacy_local_formula_expansion(formula):
    w=Workbook();w.active['A1']=formula;b=io.BytesIO();w.save(b)
    with pytest.raises(InputError):preflight({'bounded.xlsx':b.getvalue()})


def test_oversized_single_tool_cell_fails_explicitly(monkeypatch):
    from skill_tools import LIMITS
    t=WorkbookTools(files());monkeypatch.setitem(LIMITS,'page_bytes',10)
    with pytest.raises(InputError,match='without truncation'):t.read_range('Accounts','A1')
    with pytest.raises(InputError,match='without truncation'):t.search('Example')


def final_response(p):return dict(status='completed',output=[dict(type='message',content=[dict(type='output_text',text=json.dumps(p))])],usage=dict(input_tokens=100,output_tokens=100))


def test_partial_keeps_siblings_and_repairs_only_bad_findings(configured):
    from skill_engine import run
    from report_cache import successful
    from executive import present
    from reports import export_report
    send=transport();requests=[];bad=copy.deepcopy(finding());bad['title']='Unsupported special claim';bad['evidence']='Revenue is $999 billion.'
    p=proposal();p['findings'].append(bad)
    def mock(req):
        requests.append(json.loads(req.data))
        if len(requests)<3:return send(req)
        if len(requests)==3:return final_response(p)
        return final_response(dict(replacements=[],completion_checks=completion_checks()))
    r=run(files(),request_fn=mock)
    assert r.package['skill_engine']['status']=='partial'
    assert len(r.package['ai_insights'])==1 and not successful(r,require_real=False)
    assert 'PARTIAL' in r.memo and 'Unsupported special claim' not in r.memo
    assert len(present(r)['cards'])==1 and len(r.questions)==1
    repair=json.loads(requests[-1]['input'][1]['content'])
    assert [x['index'] for x in repair['rejected']]==[1]
    assert 'replacements' in requests[-1]['text']['format']['schema']['properties']
    for template in [None,Path('sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes()]:
        data,_,_=export_report(r,template);w=load_workbook(io.BytesIO(data))
        assert any('partial' in str(row).lower() for row in w['Analyst Execution'].values)
        assert not any('Unsupported special claim' in str(row) for row in w['Findings'].values)


def test_successful_targeted_patch_preserves_sibling(configured):
    from skill_engine import run
    send=transport();calls=[];p=proposal();bad=copy.deepcopy(finding());bad['evidence']='Cash is $999.';p['findings'].append(bad)
    def mock(req):
        calls.append(req)
        if len(calls)<3:return send(req)
        if len(calls)==3:return final_response(p)
        return final_response(dict(replacements=[dict(index=1,finding=finding())],completion_checks=completion_checks()))
    r=run(files(),request_fn=mock)
    assert r.package['skill_engine']['status']=='completed'
    assert len(r.package['ai_insights'])==2
    assert any(d.get('status')=='rejected' for v in r.package['skill_engine']['validation_history'] for d in v['diagnostics'])


def test_completion_checks_required_for_complete_review():
    t=WorkbookTools(files());read_all(t);calc(t);p=proposal();p['completion_checks']=[]
    accepted,d=validate(p,t)
    assert accepted and any(x.get('code')=='completion_check' for x in d)


def test_task_budget_reserves_settles_and_enforces_total(monkeypatch,tmp_path):
    from task_budget import update
    p=tmp_path/'budget.json';monkeypatch.setenv('COPILOT_TASK_BUDGET_FILE',str(p));monkeypatch.setenv('COPILOT_TASK_BUDGET_USD','2')
    a=update(amount=1.5);update(ticket=a,cost=.25)
    update(amount=1.5)
    with pytest.raises(InputError,match='ceiling'):update(amount=.3)
    assert json.loads(p.read_text())['committed_usd']==1.75
    monkeypatch.setenv('COPILOT_TASK_BUDGET_USD','3')
    with pytest.raises(InputError):update(amount=.1)


def test_accounting_idiom_is_not_an_unsupported_multiplier():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    assert grounded('These amounts should not be double counted.',['s1_B6'],t.facts,30)
    with pytest.raises(InputError):grounded('Cash will double.',['s1_B6'],t.facts,30)


@pytest.mark.parametrize('assessment',['Covenant definitions and liquidity headroom are not provided.','The package does not isolate price, volume or mix.','The evidence supports the observation, but not a forecast of its duration.'])
def test_explicit_missing_evidence_qualifies_assessment(assessment):
    t=WorkbookTools(files());read_all(t);calc(t);p=proposal();p['findings'][0]['assessment']=assessment
    assert validate(p,t)[0]


def test_derived_quarter_descriptor_uses_operand_provenance():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    t.facts['s1_B6']['column_label']='Q4';t.facts['s1_C6']['column_label']='Q1'
    t.calculate('s1_B6-s1_C6','Movement','money','same source reporting scope','')
    assert grounded('Q4 versus Q1 cash movement is {calc_1}.',['calc_1'],t.facts,30)
    with pytest.raises(InputError):grounded('Q3 cash movement is {calc_1}.',['calc_1'],t.facts,30)
    with pytest.raises(InputError):grounded('Q4 cash was 999.',['calc_1'],t.facts,30)


def test_derived_fiscal_descriptor_and_uncertainty_negation():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t);calc(t)
    assert grounded('FY2026 cash bridge is {calc_1}.',['calc_1'],t.facts,30)
    with pytest.raises(InputError):grounded('FY2027 cash bridge is {calc_1}.',['calc_1'],t.facts,30)
    p=proposal();p['findings'][0]['assessment']='The schedules do not verify the cause or contractual rights.'
    assert validate(p,t)[0]


def test_quoted_text_ids_rebound_only_with_exact_quotes():
    t=WorkbookTools(files());read_all(t);calc(t);p=proposal();p['findings'][0]['fact_ids'].append('s1_B9')
    accepted,d=validate(p,t);assert accepted and any(x.get('code')=='quoted_text_ids' for x in d)
    p=proposal();p['findings'][0]['fact_ids'].append('s1_A9')
    assert not validate(p,t)[0]
    p=proposal();p['findings'][0]['fact_ids'].append('s1_B9');p['findings'][0]['evidence']='The figure is {s1_B9}.'
    assert not validate(p,t)[0]


def test_per_unit_money_and_redundant_unit_tokens():
    from investment_output import grounded
    w=Workbook();s=w.active;s.title='Unit economics';s.append(['Example']);s.append(['Metric','Unit','Current']);s.append(['Service price','$/hour',50]);s.append(['Hours','tons',3])
    b=io.BytesIO();w.save(b);t=WorkbookTools({'example.xlsx':b.getvalue()});read_all(t)
    assert t.facts['s1_C3']['unit']=='money' and t.facts['s1_C3']['source_unit_label']=='$/hour'
    t.calculate('s1_C3*s1_C4','Revenue','money','Same service and reporting period','')
    assert grounded('Revenue is ${calc_1}.',['calc_1'],t.facts,30)=='Revenue is $150.'
    t.facts['pp']={'unit':'pp','value':2,'scale':1}
    assert grounded('Margin changed {pp} percentage points.',['pp'],t.facts,30)=='Margin changed 2.0 pp.'


def test_ui_cache_ignores_only_unchanged_automatic_layouts(monkeypatch):
    import ast,os
    from test_ai_investment import files as financialfiles
    from types import SimpleNamespace
    from ingestion import inspect_files
    import report_cache
    # Exercise the real compute function without executing Streamlit's page setup.
    node=next(n for n in ast.parse(Path('app.py').read_text()).body if isinstance(n,ast.FunctionDef) and n.name=='compute')
    captured=[]
    def analyze(*a,**kw):captured.append(a[2]);return Review()
    class Cache:
        def load(self,*a):return None
        def save(self,*a):return False
    monkeypatch.setenv('OPENAI_API_KEY','fake')
    monkeypatch.setattr(report_cache,'ReportCache',Cache)
    monkeypatch.setattr(report_cache,'cache_key',lambda *a,**k:'fixture')
    monkeypatch.setattr(report_cache,'successful',lambda *a,**k:False)
    ns=dict(os=os,copy=copy,st=SimpleNamespace(session_state={}),analyze=analyze,InputError=InputError)
    exec(compile(ast.Module(body=[node],type_ignores=[]),'app.py','exec'),ns)
    table=inspect_files(financialfiles()).tables[0]
    layout=dict(header_row=table.header_row,label_column=table.label_column,company=table.company,basis=table.basis,currency=table.currency,scale=table.scale,expense_sign=table.expense_sign,columns=[{k:c[k] for k in ('column','header','period','scenario','reference')} for c in table.columns])
    key=table.file+'|'+table.sheet
    ns['compute'](tuple(financialfiles().items()),layouts={key:layout},ai_authorized=True,engine_name='skill')
    assert captured[-1]=={}
    provenance_only=copy.deepcopy(layout)
    for column in provenance_only['columns']:column['reference']+=' additional generated header provenance'
    ns['compute'](tuple(financialfiles().items()),layouts={key:provenance_only},ai_authorized=True,engine_name='skill')
    assert captured[-1]=={}
    changed=dict(layout,company='Different entity')
    ns['compute'](tuple(financialfiles().items()),layouts={key:changed},ai_authorized=True,engine_name='skill')
    assert captured[-1]=={key:changed}
    ns['compute'](tuple(financialfiles().items()),layouts={key:layout},ai_authorized=True,engine_name='legacy')
    assert captured[-1]=={key:layout}


def test_percentage_points_require_explicit_conversion_and_reverification():
    t=WorkbookTools(files());read_all(t)
    t.calculate('s1_B6/s1_B5','Margin','ratio','Current company margin','')
    t.calculate('s1_C6/s1_C5','Margin','ratio','Prior company margin','')
    with pytest.raises(InputError,match='unit'):t.calculate('calc_1-calc_2','Margin change','pp','Comparable company periods','')
    f=t.calculate('(calc_1-calc_2)*100','Margin change','pp','Comparable company periods','')
    assert f['value']==pytest.approx(2.7777777778) and not f['illustrative']
    assert t.verify_fact(f['id'])==pytest.approx(f['value'])
    f['expression']='calc_1-calc_2';f['constants']=[];f['value']/=100
    with pytest.raises(InputError,match='unit'):t.verify_fact(f['id'])


def test_days_are_not_money_percent_or_sensitivity_and_sum_blanks_match_excel():
    from investment_output import display
    w=Workbook();s=w.active;s.title='Cash';s.append(['Amounts in $000']);s.append(['Receivables',200]);s.append(['Monthly revenue',400]);s.append(['DSO (days)','=B2/B3*30']);s.append(['Authored total','=SUM(B2:B3,B7)'])
    b=io.BytesIO();w.save(b);t=WorkbookTools({'sample.xlsx':b.getvalue()});read_all(t)
    assert t.facts['s1_B4']['unit']=='days' and t.facts['s1_B4']['value']==15
    assert display(t.facts['s1_B4'])=='15.0 days'
    f=t.calculate('days(s1_B2/s1_B3,30)','AR days','days','Reported monthly balance and revenue; 30-day convention','')
    assert f['value']==15 and not f['illustrative'] and f['conventions']
    assert t.verify_fact(f['id'])==15
    assert t.facts['s1_B5']['value']==600000
    with pytest.raises(InputError):t.verify_fact('s1_B7')
    with pytest.raises(InputError):t.calculate('days(s1_B2,30)','Bad dimension','days','same','')
    with pytest.raises(InputError):t.calculate('days(s1_B2/s1_B3,0)','Bad convention','days','same','')


def test_aging_descriptor_must_match_own_source_boundary():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t);t.facts['s1_B6']['label']='AR over 120 days'
    assert grounded('Over-120-day receivables were {s1_B6}.',['s1_B6'],t.facts,30)
    with pytest.raises(InputError):grounded('Over-90-day receivables were {s1_B6}.',['s1_B6'],t.facts,30)
    with pytest.raises(InputError):grounded('Over-120-day receivables were {s1_C6}.',['s1_C6'],t.facts,30)


def test_missing_support_bridge_is_explicit_uncertainty():
    t=WorkbookTools(files());read_all(t);calc(t)
    p=proposal();p['findings'][0]['assessment']='No complete operating variance bridge or acquisition perimeter reconciliation is supplied.'
    assert validate(p,t)[0]
    p['findings'][0]['assessment']='A complete operating variance bridge is supplied and proves the conclusion.'
    assert not validate(p,t)[0]


def test_numbered_source_entity_at_sentence_end_is_not_an_amount():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    t.facts['s1_B5']['entity']='Regional Client 07 (example industry)'
    assert grounded('Review Regional Client 07.',['s1_B5'],t.facts,30)
    with pytest.raises(InputError):grounded('Review Regional Client 08.',['s1_B5'],t.facts,30)
    with pytest.raises(InputError):grounded('Review Regional Client 07.5.',['s1_B5'],t.facts,30)
    with pytest.raises(InputError):grounded('Review Regional Client 07 for $100.',['s1_B5'],t.facts,30)


def test_negated_guarantee_preserves_later_assertions():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    assert grounded('Contract scale is not incremental fiscal-year earnings or guaranteed future revenue.',['s1_B5'],t.facts,40)
    assert grounded('Revenue is not guaranteed.',['s1_B5'],t.facts,40)
    for text in ['Revenue is guaranteed.','Revenue is not guaranteed, but earnings are guaranteed.','This is not weak evidence and guarantees revenue.']:
        with pytest.raises(InputError):grounded(text,['s1_B5'],t.facts,40)


def test_sheet_fiscal_heading_and_all_field_diagnostics():
    from investment_output import grounded
    w=Workbook();s=w.active;s.append(['Schedule FY2026']);s.append(['$ in thousands']);s.append(['Metric','Amount']);s.append(['Sales',100]);b=io.BytesIO();w.save(b)
    t=WorkbookTools({'generic.xlsx':b.getvalue()});read_all(t)
    assert grounded('FY2026 sales were {s1_B4}.',['s1_B4'],t.facts,30)
    with pytest.raises(InputError):grounded('FY2027 sales were {s1_B4}.',['s1_B4'],t.facts,30)
    t=WorkbookTools(files());read_all(t);calc(t);p=proposal()
    p['findings'][0]['evidence']='Sales were $100.';p['findings'][0]['action']='Collect $500.'
    accepted,diagnostics=validate(p,t);assert not accepted
    errors=next(d['field_errors'] for d in diagnostics if d.get('field_errors'))
    assert {e['field'] for e in errors}=={'evidence','action'}


def test_unclear_narrative_qualification_is_disclosed_not_numerical_rejection():
    t=WorkbookTools(files());read_all(t);calc(t);p=proposal()
    p['findings'][0]['assessment']='The reported cash bridge reconciles.'
    accepted,d=validate(p,t)
    assert accepted and 'investment-team review is required' in accepted[0]['assessment']
    assert any(x.get('code')=='narrative_review_required' for x in d)
    p['findings'][0]['evidence']='Cash is $999.'
    assert not validate(p,t)[0]


def test_reasoning_output_capacity_remains_bounded(monkeypatch):
    from skill_engine import settings
    monkeypatch.setenv('COPILOT_SKILL_OUTPUT_TOKENS','32000')
    assert settings()['output_tokens']==32000
    monkeypatch.setenv('COPILOT_SKILL_OUTPUT_TOKENS','32001')
    with pytest.raises(InputError):settings()


def test_negated_causal_quantification_does_not_assert_causation():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    assert grounded('These amounts do not quantify equipment-caused earnings loss.',['s1_B5'],t.facts,40)
    with pytest.raises(InputError):grounded('These amounts do not quantify equipment-caused loss, but demand caused the decline.',['s1_B5'],t.facts,40)
    with pytest.raises(InputError):grounded('Equipment caused the decline.',['s1_B5'],t.facts,40)


def test_mixed_currency_sheet_preserves_counts_days_and_aging_amounts():
    from investment_output import grounded
    w=Workbook();s=w.active;s.append(['Example professional services']);s.append(['US dollars']);s.append(['Practice','Clients at opening','Clients lost']);s.append(['Example office',610,98]);s.append(['WIP days',100]);s.append(['Lockup days (WIP days + DSO)',153]);s.append(['WIP over 120 days',40000]);b=io.BytesIO();w.save(b)
    t=WorkbookTools({'mixed.xlsx':b.getvalue()});read_all(t)
    assert t.facts['s1_B4']['unit']==t.facts['s1_C4']['unit']=='count'
    assert t.facts['s1_B5']['unit']==t.facts['s1_B6']['unit']=='days'
    assert t.facts['s1_B7']['unit']=='money'
    assert grounded('Lockup is {s1_B6} days.',['s1_B6'],t.facts,20)=='Lockup is 153.0 days.'
    with pytest.raises(InputError):grounded('Lockup is ${s1_B6}.',['s1_B6'],t.facts,20)


def test_unweighted_peer_average_is_not_an_economic_assumption():
    t=WorkbookTools(files());read_all(t)
    a=t.calculate('(s1_B5+s1_B6+s1_B7)/3','Unweighted historical peer average','money','Same currency and period; simple mean of three source observations','')
    assert not a['illustrative'] and a['conventions'][0]['function']=='unweighted_mean'
    b=t.calculate('(s1_B5+s1_B6+s1_B7)/4','Hypothetical allocation','money','Same currency and period; arbitrary divisor','')
    assert b['illustrative']


def test_skill_template_narrative_rows_fit_without_changing_styles_or_formulas():
    from reports import fit_template_narratives
    from copy import copy
    w=Workbook();s=w.active;s.title='Firm review';s.merge_cells('B4:H5');s['B4']='Evidence and specific decision. '*90
    from openpyxl.styles import Alignment
    s['B4'].alignment=Alignment(wrap_text=True,vertical='top');s.row_dimensions[4].height=24;s.row_dimensions[5].height=24;s['C8']='=SUM(C6:C7)'
    style=copy(s['B4']._style);merges=str(s.merged_cells);width=s.column_dimensions['B'].width
    fit_template_narratives(w,[('Firm review','B4')])
    assert s.row_dimensions[4].height>24 and s.row_dimensions[5].height>24
    assert s['B4']._style==style and str(s.merged_cells)==merges and s.column_dimensions['B'].width==width
    assert s['C8'].value=='=SUM(C6:C7)'


def test_publication_retains_bounded_prose_and_labels_recorded_sensitivities():
    t=WorkbookTools(files());read_all(t);calc(t)
    p=proposal();p['findings'][0]['conclusion']=' '.join(['Evidence remains unresolved.']*26)
    a,d=validate(p,t);assert a and any(x.get('code')=='prose_target_exceeded' for x in d)
    p['findings'][0]['conclusion']=' '.join(['Evidence']*161)
    assert not validate(p,t)[0]
    scenario=t.calculate('calc_1*4','Illustrative annualization','money','Quarterly source repeated','Assume unchanged quarterly cash generation')
    p=proposal();p['findings'][0]['fact_ids'].append(scenario['id']);p['findings'][0]['financial_implication']='Calculated exposure is {'+scenario['id']+'}.'
    a,d=validate(p,t);assert a and 'Includes illustrative calculations' in a[0]['financial_implication']
    assert 'illustrative calculations included' in a[0]['classification']
    t.facts[scenario['id']]['value']+=1
    assert not validate(p,t)[0]


def test_counted_source_groups_require_an_exact_raw_sum():
    from investment_output import grounded
    t=WorkbookTools(files());read_all(t)
    c=t.calculate('s1_B5+s1_B6','Selected buckets','money','Same unit and source period','')
    assert grounded('The two oldest buckets totaled {'+c['id']+'}.',[c['id']],t.facts,30)
    with pytest.raises(InputError):grounded('The three oldest buckets totaled {'+c['id']+'}.',[c['id']],t.facts,30)
    r=t.calculate('s1_B5/s1_B6','Ratio','ratio','Same unit and period','')
    with pytest.raises(InputError):grounded('The two oldest buckets totaled {'+r['id']+'}.',[r['id']],t.facts,30)
    assert grounded('This is one correlated financial issue.',['s1_B5'],t.facts,30)
    with pytest.raises(InputError):grounded('The exposure is one million.',['s1_B5'],t.facts,30)
