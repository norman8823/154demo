"""Packaged Excel files must be exact outputs for the approved release."""
import hashlib
import json
import shutil
from io import BytesIO
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from engine import InputError, ROOT
from portfolio_data import load_companies
from portfolio_excel import export_portfolio
from portfolio_financial_reporting import export_company_report
import portfolio_report_assets as assets


@pytest.fixture(scope='module')
def companies():
    return load_companies()


def _report_outputs(companies):
    by_id = {company.id: company for company in companies}
    yield 'gp_portfolio', export_portfolio(companies), assets.FILENAMES['gp_portfolio'], ()
    for company_id in ('summit', 'tallis', 'harrow'):
        data, filename, warnings = export_company_report(by_id[company_id].review)
        yield company_id, data, filename, tuple(warnings)


def _normalized_xlsx_parts(data):
    with ZipFile(BytesIO(data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    core = ET.fromstring(parts['docProps/core.xml'])
    for tag in ('{http://purl.org/dc/terms/}created',
                '{http://purl.org/dc/terms/}modified'):
        element = core.find(tag)
        if element is not None:
            element.text = '<timestamp>'
    parts['docProps/core.xml'] = ET.tostring(core, encoding='utf-8')
    return parts


def test_all_packaged_reports_are_byte_identical_to_approved_exports(companies):
    package = json.loads(assets.ASSET_MANIFEST.read_text())
    prepared_hash = hashlib.sha256((ROOT/'demo_results/manifest.json').read_bytes()).hexdigest()
    assert package['prepared_manifest_sha256'] == prepared_hash
    for report_id, expected, filename, warnings in _report_outputs(companies):
        payload, actual_filename, actual_warnings = assets.report_asset(report_id, companies)
        assert actual_filename == filename == assets.FILENAMES[report_id]
        assert actual_warnings == warnings
        # openpyxl stamps the current creation/modification time into the OOXML
        # core properties. All worksheet, formula, hyperlink, evidence and style
        # XML must otherwise match byte-for-byte.
        assert _normalized_xlsx_parts(payload) == _normalized_xlsx_parts(expected)
        record = package['reports'][report_id]
        assert hashlib.sha256(payload).hexdigest() == record['sha256']
        workbook = load_workbook(assets.ASSET_DIRECTORY/filename, data_only=False)
        assert workbook.sheetnames
        assert any(cell.hyperlink for sheet in workbook for row in sheet for cell in row)
        assert any(cell.has_style for sheet in workbook for row in sheet for cell in row)
        if report_id == 'gp_portfolio':
            assert {'Portfolio Overview', 'Investment Issues', 'Source References', 'Source Snapshots'} <= set(workbook.sheetnames)
            assert 'Portfolio Financial Performance' in workbook.sheetnames
        else:
            assert {'Financial Performance', 'Findings', 'Source Evidence', 'Statement Evidence'} <= set(workbook.sheetnames)


def test_required_asset_hash_mismatch_fails_closed(companies, tmp_path, monkeypatch):
    shutil.copytree(assets.ASSET_DIRECTORY, tmp_path/'reports')
    copied = tmp_path/'reports'
    target = copied/assets.FILENAMES['gp_portfolio']
    target.write_bytes(target.read_bytes()+b'changed')
    monkeypatch.setattr(assets, 'ASSET_DIRECTORY', copied)
    monkeypatch.setattr(assets, 'ASSET_MANIFEST', copied/'manifest.json')
    with pytest.raises(InputError, match='hash verification failed'):
        assets.report_asset('gp_portfolio', companies)


def test_checkpoint_identity_mismatch_fails_closed(companies, tmp_path, monkeypatch):
    shutil.copytree(assets.ASSET_DIRECTORY, tmp_path/'reports')
    copied = tmp_path/'reports'
    package_path = copied/'manifest.json'
    package = json.loads(package_path.read_text())
    package['prepared_manifest_sha256'] = '0'*64
    package_path.write_text(json.dumps(package))
    monkeypatch.setattr(assets, 'ASSET_DIRECTORY', copied)
    monkeypatch.setattr(assets, 'ASSET_MANIFEST', package_path)
    with pytest.raises(InputError, match='approved source manifest'):
        assets.report_asset('gp_portfolio', companies)
