"""Actual Chrome journeys/downloads and screenshots against localhost:8502 only."""
import io
import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(os.environ.get('COPILOT_GP_ARTIFACT_DIR', ROOT / 'artifacts/gp-v3')).resolve()
OUT.mkdir(parents=True, exist_ok=True)
URL = os.environ.get('COPILOT_GP_TEST_URL', 'http://127.0.0.1:8502')
assert urlparse(URL).hostname in ('127.0.0.1', 'localhost')
results = {'companies': {}, 'external_requests': [], 'browser_errors': []}


def check(page):
    expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state', 'notRunning', timeout=60000)
    expect(page.locator('[data-testid="stException"]')).to_have_count(0)
    assert not page.evaluate('document.documentElement.scrollWidth > innerWidth')


with sync_playwright() as p:
    browser = p.chromium.launch(channel='chrome', headless=True)
    page = browser.new_page(viewport={'width': 1510, 'height': 1600}, accept_downloads=True)
    page.on('pageerror', lambda error: results['browser_errors'].append(str(error)))
    def local_only(route):
        if urlparse(route.request.url).hostname not in ('127.0.0.1', 'localhost'):
            results['external_requests'].append(route.request.url)
            route.abort()
        else:
            route.continue_()
    page.route('**/*', local_only)
    page.goto(URL)
    expect(page.get_by_role('button', name='View Investment Detail →', exact=True)).to_have_count(3, timeout=60000)
    check(page)
    page.wait_for_function("[...document.querySelectorAll('.st-key-operating_cards [class*=\"st-key-gp_company_\"]')].filter(e=>e.getBoundingClientRect().height>450).length===3",timeout=60000)
    page.wait_for_timeout(250)
    boxes = page.locator('.st-key-operating_cards [class*=\"st-key-gp_company_\"]').evaluate_all('(es)=>es.map(e=>({width:e.getBoundingClientRect().width,height:e.getBoundingClientRect().height}))')
    assert all(boxes)
    assert max(b['width'] for b in boxes) - min(b['width'] for b in boxes) < 2
    assert max(b['height'] for b in boxes) - min(b['height'] for b in boxes) < 3
    page.screenshot(path=str(OUT/'overview-final.png'), full_page=True)
    expect(page.get_by_text('$82.00M',exact=True)).to_be_visible()
    expect(page.get_by_text('10.95%',exact=True)).to_be_visible()
    expect(page.get_by_text('Illustrative investment values and valuation context',exact=True)).to_have_count(0)
    with page.expect_download() as downloaded:
        page.get_by_role('button', name='↓ Download Portfolio Excel', exact=True).click()
    path = OUT/'portfolio-browser.xlsx'
    downloaded.value.save_as(path)
    gp = load_workbook(path)
    assert gp.sheetnames[:9] == ['Portfolio Overview', 'Company Comparison', 'Investment Issues', 'Summit Home Services', 'Tallis Brook Accounting', 'Harrow Ridge Environmental', 'Source References', 'Source Snapshots', 'Reporting Output']
    assert len(gp.sheetnames)==35 and 'Investment Ledger' in gp.sheetnames
    assert gp['Investment Ledger']['D9'].value==82000000
    assert gp['Investment Ledger']['B9'].value==58000000
    assert gp['Investment Ledger']['F9'].value==1.5
    def follow(cell):
        assert cell.hyperlink.target is None
        sheet,coordinate=re.fullmatch(r"'((?:[^']|'')+)'!([A-Z]+\d+)",cell.hyperlink.location).groups()
        return gp[sheet.replace("''", "'")][coordinate]
    # Inspect the actual downloaded artifact, including a derived multi-operand issue.
    chains=[]
    for row in range(6,23):
        issue=gp['Investment Issues'].cell(row,12);index=follow(issue)
        evidence=follow(gp['Source References'].cell(index.row,4))
        assert evidence.value is not None
        chains.append({'company':gp['Investment Issues'].cell(row,1).value,'source_id':gp['Source References'].cell(index.row,2).value,'destination':evidence.parent.title+'!'+evidence.coordinate})
    results['downloaded_issue_chains']=chains
    results['portfolio_download'] = {'bytes': path.stat().st_size, 'sheets': gp.sheetnames}
    for slug in ('summit', 'tallis', 'harrow'):
        page.locator('.st-key-open_'+slug+' button').click()
        expect(page.get_by_role('heading', name='Investment Summary', exact=True)).to_be_visible(timeout=60000)
        check(page)
        assert page.locator('.katex').count()==0
        page.screenshot(path=str(OUT/(slug+'-summary.png')), full_page=True)
        page.get_by_role('heading', name='Priority findings', exact=True).scroll_into_view_if_needed()
        page.screenshot(path=str(OUT/(slug+'-assessment-questions.png')), full_page=True)
        page.locator('[data-testid="stMain"]').evaluate('(e)=>e.scrollTo(0,0)')
        outputs = {}
        for custom in (False, True):
            page.get_by_role('combobox', name='Company Excel format', exact=True).click()
            page.get_by_role('option', name='Use supplied firm template' if custom else 'Standard Excel report', exact=True).click()
            if custom:
                expect(page.get_by_text('Optional company Excel template', exact=True)).to_be_visible(timeout=60000)
            else:
                expect(page.get_by_text('Optional company Excel template', exact=True)).to_have_count(0)
            check(page)
            with page.expect_download() as download:
                page.get_by_role('button', name='↓  Download Excel Report', exact=True).first.click()
            destination = OUT/(slug+('-custom' if custom else '-standard')+'.xlsx')
            download.value.save_as(destination)
            book = load_workbook(destination)
            assert {'Findings', 'Source Evidence', 'Analyst Execution'} <= set(book.sheetnames)
            assert any(c.hyperlink for row in book['Findings'] for c in row)
            if custom:
                assert book['Partner Review']['C16'].data_type == 'f'
            else:
                assert 'Partner Review' not in book.sheetnames
            outputs['custom' if custom else 'standard'] = destination.stat().st_size
        page.get_by_role('button',name='View in Source',exact=True).first.click()
        expect(page.get_by_role('dialog').locator('.source-target')).to_have_count(1,timeout=60000)
        page.screenshot(path=str(OUT/(slug+'-source-cell.png')),full_page=True)
        page.keyboard.press('Escape')
        expect(page.get_by_role('dialog')).to_have_count(0)
        page.get_by_role('button', name='← Portfolio Overview', exact=True).click()
        expect(page.get_by_role('heading', name='Portfolio Overview', exact=True)).to_be_visible(timeout=60000)
        results['companies'][slug] = {'summary': 'pass', 'evidence': 'pass', 'return': 'pass', 'downloads': outputs}
    page.get_by_role('button', name='Company Detail', exact=True).click()
    for slug,name in [('tallis','Tallis Brook Accounting'),('harrow','Harrow Ridge Environmental'),('summit','Summit Home Services')]:
        page.get_by_role('combobox',name='Company',exact=True).click()
        page.get_by_role('option',name=name,exact=True).click()
        expect(page.get_by_role('heading',name=name,exact=True)).to_be_visible(timeout=60000)
        check(page)
        assert 'company='+slug in page.url
        expect(page.get_by_role('combobox',name='Company Excel format',exact=True)).to_have_value('Standard Excel report')
    results['company_selector_switching']='pass'
    page.get_by_role('button', name='Investment Issues', exact=True).click()
    expect(page.get_by_role('heading', name='Investment Issues', exact=True).first).to_be_visible()
    expect(page.get_by_text('17 findings shown', exact=False)).to_be_visible()
    assert page.locator('.katex').count() == 0
    expect(page.get_by_role('heading',name='Tallis Brook Accounting',exact=True)).to_be_visible()
    page.screenshot(path=str(OUT/'investment-issues.png'), full_page=True)
    page.locator('.st-key-watch_tallis_10000 button').click()
    expect(page.get_by_role('heading',name='Investment Summary',exact=True)).to_be_visible(timeout=60000)
    assert 'finding=10000' in page.url
    page.locator('.st-key-company_source_tallis_10000_view button').click()
    expect(page.get_by_role('dialog').locator('.source-target')).to_have_attribute('data-coordinate','K40',timeout=60000)
    page.keyboard.press('Escape')
    expect(page.get_by_role('dialog')).to_have_count(0)
    results['deep_link_summary_evidence_regression'] = 'pass'
    page.get_by_role('button', name='Documents & Reports', exact=True).click()
    expect(page.get_by_role('heading', name='Source Documents', exact=True)).to_be_visible(timeout=60000)
    expect(page.get_by_role('button',name='Preview',exact=True)).to_have_count(3)
    expect(page.get_by_role('heading',name='Workbook preview',exact=True)).to_have_count(0)
    check(page)
    page.screenshot(path=str(OUT/'documents-library.png'),full_page=True)
    page.get_by_role('button',name='Data Foundation',exact=True).click()
    expect(page.get_by_role('heading',name='AI Workflow Automation Readiness',exact=True)).to_be_visible(timeout=60000)
    check(page)
    page.screenshot(path=str(OUT/'data-foundation.png'),full_page=True)
    with page.expect_download() as downloaded:
        page.get_by_role('button',name='↓ Download AI-ready data package',exact=True).click()
    package_path=OUT/'AI_Ready_Operating_Data_browser.zip';downloaded.value.save_as(package_path)
    consumer=OUT/'browser-consumer';consumer.mkdir(exist_ok=True)
    with zipfile.ZipFile(package_path) as archive:
        assert set(archive.namelist())=={'financial_facts.csv','metric_definitions.json','source_lineage.csv','quality_report.json','example_consumer.py','README.md'}
        archive.extractall(consumer)
    output=subprocess.run([sys.executable,'-I','-S',str(consumer/'example_consumer.py')],cwd=consumer,env={'PATH':os.defpath},check=True,capture_output=True,text=True,timeout=30)
    independent=json.loads(output.stdout)
    assert len(independent['comparison'])==3
    assert next(c for c in independent['comparison'] if c['company_id']=='harrow')['period']=='2026-FY'
    (OUT/'browser-independent-consumer.json').write_text(output.stdout)
    results['data_package_download']={'bytes':package_path.stat().st_size,'standalone_consumer':'pass','company_comparisons':3}
    page.get_by_role('button', name='Portfolio Overview', exact=True).click()
    expect(page.get_by_role('heading', name='Portfolio Overview', exact=True)).to_be_visible(timeout=60000)
    check(page)
    page.set_viewport_size({'width': 390, 'height': 1000})
    # Native Streamlit sidebar is collapsible on narrow screens.
    page.wait_for_timeout(500)  # Native sidebar animates/collapses during resize.
    sidebar = page.locator('[data-testid="stSidebar"]')
    if sidebar.get_attribute('data-collapsed') == 'false':
        page.get_by_role('button', name='Collapse sidebar', exact=True).click()
    expect(sidebar).to_have_attribute('data-collapsed', 'true')
    check(page)
    page.screenshot(path=str(OUT/'overview-mobile.png'), full_page=True)
    page.get_by_role('button',name='Expand sidebar',exact=True).click()
    page.get_by_role('button',name='Data Foundation',exact=True).click()
    expect(page.get_by_role('heading',name='AI Workflow Automation Readiness',exact=True)).to_be_visible(timeout=60000)
    if sidebar.get_attribute('data-collapsed')=='false':page.get_by_role('button',name='Collapse sidebar',exact=True).click()
    page.wait_for_timeout(350)
    check(page)
    page.screenshot(path=str(OUT/'data-foundation-mobile.png'),full_page=True)
    results['mobile_no_overflow'] = 'pass'
    assert not results['browser_errors'], results['browser_errors']
    assert not any('openai' in u for u in results['external_requests'])
    (OUT/'browser-results.json').write_text(json.dumps(results, indent=2))
    browser.close()
    print(json.dumps(results, indent=2))
