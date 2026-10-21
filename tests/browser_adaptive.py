"""Actual uploads, optional human mapping correction, map reuse and both adaptive Excel modes.
Run with a fresh app server and isolated PORTFOLIO_MAPPING_PATH for test profiles.
"""
from pathlib import Path
import json
import os
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'artifacts';OUT.mkdir(exist_ok=True)


def upload(page,name,title):
    page.goto(os.environ.get('COPILOT_TEST_URL','http://127.0.0.1:8501'))
    expect(page.get_by_role('heading',name='Generate a portfolio review',exact=True)).to_be_visible(timeout=30000)
    page.locator('input[type=file]').first.set_input_files(ROOT/'tests'/'fixtures'/name)
    expect(page.get_by_text('Review Account Mappings',exact=True)).to_be_visible()
    if not page.get_by_text(title,exact=True).is_visible():page.get_by_text('Review Account Mappings',exact=True).click()
    expect(page.get_by_text(title,exact=True)).to_be_visible()


def generate(page,title):
    page.get_by_role('button',name='Generate Portfolio Review').click()
    expect(page.get_by_role('heading',name=title+' Portfolio Review',exact=True)).to_be_visible(timeout=30000)
    expect(page.locator('[data-testid="stException"]')).to_have_count(0)


def download(page,name):
    with page.expect_download() as result:page.get_by_role('button',name='Download Excel Report').first.click()
    result.value.save_as(OUT/name)
    return load_workbook(OUT/name)

with sync_playwright() as p:
    browser=p.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',headless=True)
    page=browser.new_page(viewport={'width':1440,'height':1100})
    errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    upload(page,'A_Northstar.xlsx','A_Northstar.xlsx · Trading results')
    generate(page,'April 2027')
    expect(page.get_by_text('$1.00M',exact=True)).to_be_visible()
    expect(page.get_by_text('50.0%',exact=True)).to_be_visible()
    standard=download(page,'adaptive-A-standard.xlsx')
    assert standard.sheetnames==['Executive Summary','Financial Performance','Exceptions','Source Evidence']
    assert any(row[1]==1000000 for row in standard['Financial Performance'].values)
    page.get_by_role('button',name='View Evidence').first.click()
    expect(page.get_by_role('heading',name='See where every conclusion came from',exact=True)).to_be_visible()
    expect(page.get_by_role('heading',name='Northstar Components · EBITDA below budget',exact=True)).to_be_visible()
    page.get_by_text('01  Upload & Setup',exact=True).click()
    page.get_by_text('Use existing firm template',exact=True).click()
    page.locator('input[type=file]').nth(1).set_input_files(ROOT/'sample_templates'/'Sample_Firm_Review_Template.xlsx')
    expect(page.get_by_role('button',name='Remove Sample_Firm_Review_Template.xlsx',exact=True)).to_be_visible()
    generate(page,'April 2027')
    custom=download(page,'adaptive-A-custom.xlsx')
    assert custom['Partner Review']['C8'].value==1000000
    assert custom['Partner Review']['F5'].value=='April 2027'
    assert custom['Partner Review']['C16'].data_type=='f'
    page.screenshot(path=str(OUT/'adaptive-A-dashboard.png'),full_page=True)
    second=browser.new_page(viewport={'width':1440,'height':1100})
    upload(second,'B_Harbor.xlsx','B_Harbor.xlsx · Export Ledger')
    expect(second.get_by_role('button',name='Generate Portfolio Review')).to_be_enabled()
    # Resolve the single ambiguous account through actual visible onboarding controls.
    second.get_by_text('Review Account Mappings',exact=True).click() if not second.get_by_label('Confirm classification for Mobilization reserve',exact=True).is_visible() else None
    classification=second.get_by_test_id('stSelectbox').filter(has=second.get_by_text('Classification for Mobilization reserve',exact=True))
    combo=second.get_by_role('combobox',name='Classification for Mobilization reserve',exact=True)
    expect(second.get_by_test_id('stApp')).to_have_attribute('data-test-script-state','notRunning')
    # The editable grids finish measuring after the script completes; let their layout settle.
    second.wait_for_timeout(1500)
    combo.press('ArrowDown')
    second.get_by_role('option',name='operating_expenses',exact=True).click()
    expect(combo).to_have_value('operating_expenses')
    second.get_by_role('checkbox',name='Confirm classification for Mobilization reserve',exact=True).press('Space')
    second.get_by_role('button',name='Apply reviewed classifications',exact=True).click()
    expect(second.get_by_role('button',name='Generate Portfolio Review')).to_be_enabled()
    second.get_by_role('button',name='Confirm and save mappings locally',exact=True).click()
    expect(second.get_by_text('Confirmed account mappings saved locally for this company, worksheet and account context.',exact=True)).to_be_visible()
    generate(second,'February 2027')
    expect(second.get_by_text('$118K',exact=True)).to_be_visible()
    b=download(second,'adaptive-B-standard.xlsx')
    assert any(row[1]==118000 and row[0].endswith('ebitda') for row in b['Financial Performance'].values)
    third=browser.new_page()
    upload(third,'B_Harbor.xlsx','B_Harbor.xlsx · Export Ledger')
    expect(third.get_by_role('button',name='Generate Portfolio Review')).to_be_enabled()
    expect(third.get_by_label('Confirm classification for Mobilization reserve',exact=True)).to_have_count(0)
    generate(third,'February 2027')
    expect(third.locator('[data-testid="stException"]')).to_have_count(0)
    assert not errors
    (OUT/'adaptive-browser-results.json').write_text(json.dumps({'unfamiliar_actual_upload_A':'passed','A_standard_and_template_download':'passed','A_source_evidence':'passed','B_ambiguity_nonblocking':'passed','B_human_correction':'passed','B_local_mapping_reuse':'passed','B_financial_export':'passed','browser_errors':errors},indent=2))
    print('PASS: adaptive A/B uploads, human mapping correction, persisted reuse, financial evidence, standard/custom Excel.')
    browser.close()
