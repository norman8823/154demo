import io
from copy import deepcopy
import pytest
from openpyxl import load_workbook
from portfolio_data import load_companies
from portfolio_financial_reporting import statement, export_company_report, PORTFOLIO_SHEET, LINES
from portfolio_excel import export_portfolio
from reports import export_report as frozen_export

@pytest.fixture(scope='module')
def companies(): return load_companies()

# Independent checked fixtures: approved statement totals, not summed fact ledgers.
EXPECTED = {
 'summit': {'revenue':(3298000,3207000), 'cogs':(1865000,1630000), 'gross_profit':(1433000,1577000), 'operating_expenses':(931000,918000), 'ebitda':(502000,659000), 'depreciation':(93000,93000), 'interest':(165000,165000), 'tax':(61000,100000), 'net_income':(183000,301000)},
 'tallis': {'revenue':(3475418,3651832), 'cogs':(1875360,1860351), 'gross_profit':(1600058,1791481), 'ebitda':(684306,946225)},
 'harrow': {'revenue':(46650000,47006000), 'cogs':(31616000,30790000), 'gross_profit':(15034000,16216000), 'operating_expenses':(5322000,5321000), 'ebitda':(9712000,10895000), 'depreciation':(7104000,7104000), 'interest':(3996000,3996000), 'tax':(96000,96000), 'net_income':(-1484000,-301000)}
}


def test_approved_totals_margins_and_no_duplicate_detail_aggregation(companies):
 for c in companies:
  lines={l.key:l for l in statement(c.review)}
  for key,(a,b) in EXPECTED[c.id].items():
   assert (lines[key].actual,lines[key].budget)==(a,b)
   assert lines[key].variance==a-b
  for scenario in ('actual','budget'):
   assert getattr(lines['gm'],scenario)==pytest.approx(getattr(lines['gross_profit'],scenario)/getattr(lines['revenue'],scenario))
   assert getattr(lines['margin'],scenario)==pytest.approx(getattr(lines['ebitda'],scenario)/getattr(lines['revenue'],scenario))
  for key in ('gm','margin'):assert lines[key].variance==pytest.approx((lines[key].actual-lines[key].budget)*100)
  copy=deepcopy(c.review);copy.package['facts']*=2
  assert statement(copy)==statement(c.review)
 tallis={l.key:l for l in statement(companies[1].review)}
 for key in ('operating_expenses','depreciation','interest','tax','net_income'):
  assert (tallis[key].actual,tallis[key].budget,tallis[key].variance)==(None,None,None)


def signature(sheet):
 return [(c.coordinate,c.value,c.number_format,c.hyperlink.target if c.hyperlink else None,c.hyperlink.location if c.hyperlink else None) for row in sheet for c in row if c.value is not None]


def test_three_standard_company_exports_preserve_evidence_and_statement(companies):
 for c in companies:
  original=load_workbook(io.BytesIO(frozen_export(c.review)[0]))
  result=load_workbook(io.BytesIO(export_company_report(c.review)[0]))
  for name in original.sheetnames:
   if name not in ('Financial Performance','Investment Summary','Executive Summary'):assert signature(original[name])==signature(result[name]),name
  assert 'Investment Summary' not in result
  assert 'Partner Review' not in result
  sheet=result['Financial Performance'];assert sheet.max_column==4 and sheet.max_row<=22
  assert result.active.title=='Executive Summary'
  for row,line in enumerate(statement(c.review),7):
   assert sheet.cell(row,1).value==line.label
   for col,value in zip((2,3,4),(line.actual,line.budget,line.variance)):
    assert sheet.cell(row,col).value==pytest.approx(value) if value is not None else sheet.cell(row,col).value=='N/A'
   if line.actual is not None:assert sheet.cell(row,2).hyperlink.location.startswith("'Statement Evidence'!")
  assert not any(str(cell.value) in {'money','ratio'} or str(cell.value).startswith('financial_') for row in sheet for cell in row)
  assert '%' in sheet['B10'].number_format and 'pp' in sheet['D10'].number_format
  assert 'monthly' in sheet['A3'].value if c.id!='harrow' else 'annual' in sheet['A3'].value
  proof=result['Statement Evidence']
  links=[cell for row in proof for cell in row if cell.hyperlink]
  assert links
  for cell in links:
   address=cell.hyperlink.location.split('!')[1];assert result['Source Evidence'][address].value is not None


def test_gp_statement_reconciles_and_supports_every_operand(companies):
 book=load_workbook(io.BytesIO(export_portfolio(companies)));ws=book[PORTFOLIO_SHEET]
 assert len(book.sheetnames)==35
 assert [ws.cell(4,i).value for i in (2,5,8)]==[c.name for c in companies]
 assert [ws.cell(5,i).value for i in (2,5,8)]==['September 2026 · monthly · $ base units','September 2026 · monthly · USD base units','FY2026 · annual · $ base units']
 assert ws.max_column==10 and 'No consolidated total' in ws['A3'].value
 for index,c in enumerate(companies):
  for row,line in enumerate(statement(c.review),7):
   start=2+3*index
   for col,value in zip(range(start,start+3),(line.actual,line.budget,line.variance)):
    assert ws.cell(row,col).value==pytest.approx(value) if value is not None else ws.cell(row,col).value=='N/A'
   if line.actual is not None:
    coordinate=ws.cell(row,start).hyperlink.location.split('!')[1]
    proof=book['Source References'][coordinate];assert line.label in proof.value
 for row in ws:
  for cell in row:
   assert cell.value not in ('money','ratio') and not str(cell.value).startswith('financial_')
 assert book['Portfolio Overview']['A4'].hyperlink.location=="'Portfolio Financial Performance'!A1"
