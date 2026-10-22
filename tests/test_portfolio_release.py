"""Public release boundaries for prepared-only portfolio presentation."""
from copy import deepcopy
import io
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from engine import ROOT
from scripts.seed_demo_results import calculations_match


COMPANIES = {
    'summit': ('Summit Home Services', 'Summit_Home_Services_Sep2026_MonthlyReview.xlsx'),
    'tallis': ('Tallis Brook Accounting', 'Tallis_Brook_Accounting_Sep2026_MonthlyReview.xlsx'),
    'harrow': ('Harrow Ridge Environmental', 'Harrow_Ridge_Environmental_FY2026_AnnualReview.xlsx'),
}


@pytest.mark.parametrize(('slug', 'details'), COMPANIES.items())
def test_company_pages_only_offer_standard_download(slug, details):
    name, filename = details
    at = AppTest.from_file(str(ROOT/'portfolio_app.py'), default_timeout=90)
    at.query_params.update(page='company', company=slug)
    at.run()
    assert not at.exception
    assert any(name in str(block.value) for block in at.markdown)
    assert not any(item.label == 'Company Excel format' for item in at.selectbox)
    assert not at.get('file_uploader')
    assert [item.proto.label for item in at.get('download_button')] == ['Download Excel']
    payload, actual_filename, _warnings = at.session_state.export
    assert actual_filename == filename
    assert payload.startswith(b'PK')
    report = load_workbook(io.BytesIO(payload))
    assert report.active.title == 'Executive Summary'
    assert 'Partner Review' not in report.sheetnames
    assert {'Financial Performance', 'Findings', 'Source Evidence'} <= set(report.sheetnames)


def test_unknown_company_route_cannot_access_arbitrary_files():
    at = AppTest.from_file(str(ROOT/'portfolio_app.py'), default_timeout=90)
    at.query_params.update(page='company', company='../../.env')
    at.run()
    assert not at.exception
    assert any('Unknown company' in warning.value for warning in at.warning)
    assert not at.get('download_button') and not at.get('file_uploader')


def sample_calculation():
    return [dict(id='calc_34', label='Mean commodity price', value=92.4975,
                 unit='money', currency='$', scale=1, file='harrow.xlsx',
                 sheet='Source-linked calculations', ref='Restricted arithmetic calc_34',
                 raw='(s1_J12+s1_J13+s1_J14)/3', inputs=['s1_J12','s1_J13','s1_J14'],
                 source_refs=[], expression='(s1_J12+s1_J13+s1_J14)/3',
                 operation='restricted_arithmetic', period='FY2026', scenario='comparison',
                 entity='FY2026', basis='see source inputs', illustrative=False,
                 assumptions=[], constants=[3], conventions=[], limitation='Arithmetic verified.')]


def test_calculation_value_allows_only_tiny_float_representation_difference():
    expected = sample_calculation()
    observed = deepcopy(expected)
    observed[0]['value'] = 92.49749999999999
    assert calculations_match(observed, expected)
    observed[0]['value'] = 92.4975 + 1e-6
    assert not calculations_match(observed, expected)


@pytest.mark.parametrize('field', ('id', 'expression', 'raw', 'inputs', 'source_refs', 'unit', 'scale', 'period', 'basis', 'constants'))
def test_calculation_provenance_and_dimensions_remain_exact(field):
    expected = sample_calculation()
    observed = deepcopy(expected)
    value = observed[0][field]
    if isinstance(value, str):
        observed[0][field] = value + '_changed'
    elif type(value) in (int, float):
        observed[0][field] = value + 0.01
    else:
        observed[0][field] = value + ['tampered']
    assert not calculations_match(observed, expected)


@pytest.mark.parametrize('observed', (None, [], [dict(id='calc_34', value=92.4975)]))
def test_calculation_comparison_rejects_missing_or_malformed_records(observed):
    assert not calculations_match(observed, sample_calculation())
