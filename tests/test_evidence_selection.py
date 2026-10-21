"""Bounded source selection must preserve economic scope, notes and independence."""
import io
from openpyxl import Workbook
from ai_evidence import grid_facts,balanced_select
from executive import present,rank_insights
from engine import Review


def encoded(w):
    out=io.BytesIO();w.save(out);return out.getvalue()


def test_notes_to_right_do_not_hide_amounts_or_context():
    w=Workbook();s=w.active;s.title='Earnings bridge'
    for row in [['Example Group'],['September 2028; USD whole units; accrual basis'],['Item','Amount ($)','Source / basis'],['Reported EBITDA',1000,'Reported accounts'],['Compensation normalization',200,'Contract not yet effective']]:s.append(row)
    for cell in (s['B4'],s['B5']):cell.number_format='$#,##0'
    facts,_=grid_facts({'source.xlsx':encoded(w)})
    adjustment=next(f for f in facts if f['ref']=="'Earnings bridge'!B5")
    assert adjustment['value']==200 and adjustment['unit']=='money'
    assert adjustment['row_context']['C5']=='Contract not yet effective'
    assert "'Earnings bridge'!C5" in adjustment['context_refs']


def test_acquisition_dates_are_entities_and_date_bearing_count_headers_are_metrics():
    w=Workbook();s=w.active;s.title='Business units'
    for row in [['Example Group'],['September 2028; USD whole units'],['Measure','Legacy','East (acq. Mar-28)'],['EBITDA',100,20],['Margin',.3,.08]]:s.append(row)
    s['B5'].number_format=s['C5'].number_format='0.0%'
    s=w.create_sheet('Customer cohorts')
    for row in [['Example Group'],['USD whole units'],['Practice','Clients at October 1, 2027','Clients at September 30, 2028'],['East',100,90]]:s.append(row)
    facts,_=grid_facts({'source.xlsx':encoded(w)})
    refs={f['ref'] for f in facts}
    assert "'Business units'!C4" in refs and "'Business units'!C5" in refs
    assert "'Customer cohorts'!B4" in refs and "'Customer cohorts'!C4" in refs


def fact(identity,sheet,label,value,entity=''):
    return dict(id=identity,file='source.xlsx',sheet=sheet,ref=f"'{sheet}'!B{identity[-2:]}",label=label,value=value,unit='money',currency='USD',scale=1,entity=entity)


def test_large_early_sheet_cannot_starve_later_schedules_and_bounds_hold():
    entries=[fact(f'large_{i:02}', 'Large ledger','Detail',i) for i in range(90)]
    entries+=[fact('later_01','Cash cycle','Total WIP',1000),fact('bridge_01','Earnings bridge','Adjusted EBITDA',2000)]
    selected=balanced_select(entries,8)
    assert len(selected)==8
    assert {'Cash cycle','Earnings bridge'}<={f['sheet'] for f in selected}
    assert [f['id'] for f in selected]==[f['id'] for f in balanced_select(list(reversed(entries)),8)]


def test_largest_same_unit_exposure_survives_row_sampling():
    entries=[dict(fact(str(i),'Cash cycle','WIP over 120 days',v,entity),ref=f"'Cash cycle'!{chr(66+i)}10") for i,(entity,v) in enumerate([('North',20),('South',4000),('Firm total',4100)])]
    selected=balanced_select(entries,2)
    assert {f['entity'] for f in selected}=={'Firm total','South'}


def test_ranking_prefers_economic_relationship_and_penalizes_duplicate_claims():
    observations=[fact('margin_01','Profit','EBITDA margin',.08),fact('wip_01','Cash','WIP over 120 days',4000),fact('claim_01','Operations','Completion rate',.9)]
    def insight(id,title,refs,kind='risk'):
        return dict(id=id,title=title,text='Reported evidence warrants management review.',why_it_matters='This may affect earnings durability and cash conversion.',management_question='Which assumptions need investigation?',fact_ids=refs,kind=kind,classification='AI interpretation / hypothesis')
    r=Review(package={'facts':observations,'ai_insights':[insight(8000,'Simple threshold',['claim_01']),insight(8001,'Economic relationship',['margin_01','wip_01'],'relationship'),insight(8002,'Repeated relationship',['margin_01','wip_01'],'relationship'),insight(8003,'Other operating observation',['claim_01'])]})
    ranked=rank_insights(r)
    assert ranked[0]['id']==8001
    assert ranked[1]['id']==8000  # An independent finding precedes a duplicate high-scoring one.
    assert present(r)['cards'][0]['id']==8001
    assert len(present(r)['secondary'])==1


def test_exact_aging_bucket_is_a_source_label_only_when_bound_to_its_own_value():
    import pytest
    from investment_output import grounded
    from engine import InputError
    facts={'old':dict(label='WIP Over 120 days',unit='money',currency='USD',scale=1,value=4000),'recent':dict(label='WIP 0-30 days',unit='money',currency='USD',scale=1,value=1000)}
    assert grounded('WIP over 120 days of {old}.',['old'],facts,45)=='WIP over 120 days of $4,000.'
    for text,refs in [
        ('WIP over 150 days of {old}.',['old']),
        ('WIP over 120 days of {recent}.',['old','recent']),
        ('WIP over 120 days of {old}; losses are $999,000.',['old']),
        ('WIP over 120 days is growing.',['old']),
        ('WIP over 120 days of {old}, growing by two times.',['old']),
    ]:
        with pytest.raises(InputError):grounded(text,refs,facts,45)


def test_single_sheet_earnings_bridge_is_not_outranked_by_more_worksheets_alone():
    inputs=[fact('reported_01','Bridge','LTM reported EBITDA',100),fact('adjusted_01','Bridge','Adjusted EBITDA',160),fact('normalization_01','Bridge','Compensation normalization',60),fact('profit_01','Profit','EBITDA',30),dict(fact('retention_01','Customers','Retention',.99),unit='ratio'),dict(fact('realization_01','Delivery','Realization',.92),unit='ratio')]
    def insight(id,title,refs,kind):
        return dict(id=id,title=title,text='The cited evidence requires assessment.',why_it_matters='This may affect underwriting assumptions.',management_question='Which assumptions need investigation?',fact_ids=refs,kind=kind,classification='AI interpretation / hypothesis')
    r=Review(package={'facts':inputs,'ai_insights':[insight(8000,'Favorable current operations',['profit_01','retention_01','realization_01'],'opportunity'),insight(8001,'Adjusted earnings assumptions',['reported_01','adjusted_01','normalization_01'],'risk')]})
    ranked=rank_insights(r)
    assert ranked[0]['id']==8001
    assert ranked[0]['components']['underwriting_assumptions']>ranked[1]['components']['cross_sheet_context']
