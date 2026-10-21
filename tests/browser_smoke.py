"""Real-browser acceptance check; run with the app already on port 8501."""
import base64
import json
import os
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts'; OUT.mkdir(exist_ok=True)

def ready(page):
    expect(page.get_by_role('heading',name='Generate a portfolio review',exact=True)).to_be_visible(timeout=30000)

def generate(page):
    page.get_by_role('button',name='Generate Portfolio Review').click()
    expect(page.get_by_role('heading',name='September 2026 Portfolio Review',exact=True)).to_be_visible(timeout=30000)
    expect(page.locator('[data-testid="stException"]')).to_have_count(0)
    expect(page.get_by_test_id('stApp')).to_have_attribute('data-test-script-state','notRunning')

def get_report(page,destination):
    with page.expect_download() as dl:
        page.get_by_role('button',name='Download Excel Report').first.click()
    dl.value.save_as(OUT/destination)
    return load_workbook(OUT/destination)

with sync_playwright() as p:
    browser=p.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',headless=True)
    page=browser.new_page(viewport={'width':1440,'height':1050},device_scale_factor=1)
    errors=[];page.on('pageerror',lambda err:errors.append(str(err)))
    page.goto(os.environ.get('COPILOT_TEST_URL','http://127.0.0.1:8501'));ready(page)
    page.screenshot(path=str(OUT/'01-intake.png'),full_page=True)
    page.get_by_role('button',name='Use Demo Files',exact=True).click()
    generate(page)
    expect(page.get_by_text('$5.38M',exact=True)).to_be_visible()
    page.screenshot(path=str(OUT/'02-review.png'),full_page=True)
    standard=get_report(page,'standard.xlsx')
    assert standard.sheetnames==['Executive Summary','Financial Performance','Exceptions','Source Evidence']
    for index,title in enumerate(['Profitability is deteriorating','Covenant headroom needs review','Potential revenue leakage']):
        if index:
            page.get_by_text('02  Portfolio Review',exact=True).click()
            expect(page.get_by_role('heading',name='September 2026 Portfolio Review',exact=True)).to_be_visible()
        page.get_by_role('button',name='View Evidence').nth(index).click()
        expect(page.get_by_role('heading',name='See where every conclusion came from',exact=True)).to_be_visible()
        expect(page.locator('.st-key-evidence_panel').get_by_role('heading',name=title,exact=True)).to_be_visible()
        expect(page.get_by_role('radio',name=f'0{index+1} '+title,exact=True)).to_be_checked()
        expect(page.locator('[data-testid="stException"]')).to_have_count(0)
    page.screenshot(path=str(OUT/'03-evidence.png'),full_page=True)
    manual=browser.new_page(viewport={'width':1440,'height':1050})
    manual.goto(os.environ.get('COPILOT_TEST_URL','http://127.0.0.1:8501'));ready(manual)
    payload=[{'name':f.name,'data':base64.b64encode(f.read_bytes()).decode(),'type':'text/csv' if f.suffix=='.csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'} for f in sorted((ROOT/'inputs').iterdir())]
    transfer=manual.evaluate_handle('''items => { const dt=new DataTransfer(); for(const f of items){const b=Uint8Array.from(atob(f.data),c=>c.charCodeAt(0));dt.items.add(new File([b],f.name,{type:f.type}));} return dt;}''',payload)
    manual.locator('[data-testid="stFileUploaderDropzone"]').first.dispatch_event('drop',{'dataTransfer':transfer})
    for source in sorted((ROOT/'inputs').iterdir()):
        expect(manual.get_by_role('button',name='Remove '+source.name,exact=True)).to_be_visible()
    expect(manual.get_by_text('Demonstration financial package recognized',exact=True)).to_be_visible()
    generate(manual)
    uploaded=get_report(manual,'uploaded-standard.xlsx')
    for name in standard.sheetnames:
        # Financial content must agree; independent sessions have distinct run receipts.
        rows=lambda w:[row for row in w[name].values if row[0]!='Inference usage']
        assert rows(standard)==rows(uploaded),name
    manual.get_by_text('01  Upload & Setup',exact=True).click();ready(manual)
    manual.get_by_text('Use existing firm template',exact=True).click()
    manual.locator('input[type=file]').nth(1).set_input_files(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx')
    expect(manual.get_by_role('button',name='Remove Sample_Firm_Review_Template.xlsx',exact=True)).to_be_visible()
    generate(manual)
    custom=get_report(manual,'custom.xlsx')
    assert custom['Partner Review']['C8'].value==5383752
    assert custom['Partner Review']['C16'].data_type=='f'
    assert abs(custom['Partner Review']['C12'].value-4.464462836964341)<1e-10
    assert 'Source Evidence' in custom.sheetnames
    manual.get_by_role('button',name='View source evidence',exact=True).click()
    expect(manual.get_by_role('heading',name='Downloadable Excel report',exact=True)).to_be_visible()
    expect(manual.locator('[data-testid="stException"]')).to_have_count(0)
    # Targeted follow-up: actual single/partial uploads and sparse finding IDs.
    subsets=[('platform',['PlatformA'],{1,2}),('quickbooks',['QuickBooks'],set()),('schedule',['.csv'],set()),('billing',['QuickBooks','.csv'],{3})]
    partial_results={}
    for label,patterns,expected in subsets:
        partial=browser.new_page(viewport={'width':1440,'height':1050})
        partial.goto(os.environ.get('COPILOT_TEST_URL','http://127.0.0.1:8501'));ready(partial)
        paths=[path for path in sorted((ROOT/'inputs').iterdir()) if any(pattern in path.name for pattern in patterns)]
        partial.locator('input[type=file]').first.set_input_files(paths)
        for source in paths:
            expect(partial.get_by_role('button',name='Remove '+source.name,exact=True)).to_be_visible()
        expect(partial.get_by_text('Demonstration financial package recognized',exact=True)).to_be_visible()
        generate(partial)
        expect(partial.get_by_text('Data Quality & Coverage',exact=True)).to_be_visible()
        # The source-denominated platform also supports an existing AR concentration check.
        displayed=min(3,len(expected)+(1 if label=='platform' else 0))
        expect(partial.get_by_role('button',name='View Evidence')).to_have_count(displayed)
        report=get_report(partial,f'partial-{label}.xlsx')
        rows=[row for row in report['Exceptions'].values if isinstance(row[0],int)]
        assert {row[0] for row in rows}==expected
        partial.get_by_role('button',name='View source evidence',exact=True).click()
        expect(partial.get_by_role('heading',name='Downloadable Excel report',exact=True)).to_be_visible()
        expect(partial.locator('[data-testid="stException"]')).to_have_count(0)
        if expected=={3}: expect(partial.get_by_role('radio',name='03 Potential revenue leakage',exact=True)).to_be_checked()
        if label=='quickbooks':
            partial.screenshot(path=str(OUT/'04-quickbooks-only.png'),full_page=True)
            partial.get_by_text('01  Upload & Setup',exact=True).click();ready(partial)
            partial.get_by_text('Use existing firm template',exact=True).click()
            partial.locator('input[type=file]').nth(1).set_input_files(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx')
            expect(partial.get_by_role('button',name='Remove Sample_Firm_Review_Template.xlsx',exact=True)).to_be_visible()
            generate(partial)
            custom_partial=get_report(partial,'partial-quickbooks-custom.xlsx')
            assert custom_partial['Partner Review']['C8'].value=='Not available — additional data required.'
            assert custom_partial['Partner Review']['G16'].value==0
            assert custom_partial['Partner Review']['C16'].data_type=='f'
        partial_results[label]='passed'
        partial.close()
    assert not errors,errors
    (OUT/'browser-results.json').write_text(json.dumps({'demo':'passed','real_drag_drop':'passed','manual_matches_demo':'passed','three_evidence_actions':'passed','standard_download':'passed','custom_upload_and_download':'passed','partial_uploads':partial_results,'browser_errors':errors},indent=2))
    print('PASS: demo, drag/drop, equivalent uploaded results, three evidence routes, standard/custom Excel downloads, and single/partial uploads.')
    browser.close()
