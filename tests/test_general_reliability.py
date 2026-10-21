"""Conventional periods, exact source descriptors and question-only repair."""
import io,json,copy
import pytest
from openpyxl import Workbook,load_workbook
from engine import analyze,ROOT
from ingestion import period,period_grain
from excel_layout import financial_table
from investment_output import grounded,validate,InputError
from company_context import validate as validate_context,text as context_text
from reports import export_report,fields


@pytest.mark.parametrize('headers',[
 ['Account','FY2028 Actual','FY2028 Budget','FY2027 Actual'],
 ['Account','2028 Actual','2028 Budget','2027 Actual'],
 ['Account',2028,2027],
 ['Account','Q4 FY28 Actual','Q4 FY28 Budget','Q4 FY27 Actual'],
 ['Account','September 2028 Actual','September 2028 Budget','September 2027 Actual'],
])
def test_conventional_statement_formats(headers):
 w=Workbook();s=w.active;s.title='Statement'
 s.append(['Independent Example Services']);s.append(['USD whole units; accrual basis']);s.append(headers)
 vals=[['Total revenue',1000,1100,900],['Total direct costs',600,650,500],['Gross profit',400,450,400],['Total SG&A',200,230,210],['EBITDA',200,220,190]]
 for row in vals:s.append(row[:len(headers)])
 b=io.BytesIO();w.save(b);r=analyze({'example.xlsx':b.getvalue()})
 assert r.metrics['revenue']==1000 and r.metrics['ebitda']==200
 assert r.currency=='USD' and all(x['basis']=='accrual' for x in r.normalized)
 table=financial_table(r);assert table and any(row[13]=='revenue' and row[1]==1000 for row in table)
 if 'Budget' in str(headers):assert r.metrics['revenue_budget']==1100
 for t in [None,(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes()]:assert export_report(r,t)[0]


def test_period_grains_are_not_interchangeable():
 assert period('FY2026')=='2026-FY' and period('Q2 FY26')=='2026-FQ2'
 assert period('Q2 2026')=='2026-Q2' and period('September 2026')=='2026-09'
 assert len({period_grain(period(s)) for s in ['FY2026','Q2 FY26','Q2 2026','September 2026']})==4


def test_exact_source_descriptors_do_not_authorize_new_figures():
 facts={'q':dict(id='q',value=.12,unit='ratio',label='Downtime',column_label='South Q2',entity='South'),
        'age':dict(id='age',value=9,unit='count',label='13+ years',entity='South'),
        'account':dict(id='account',value=500,unit='money',currency='USD',scale=1,label='Revenue',entity='National Account 01 (many sites)')}
 assert '12.0%' in grounded('Q2 downtime is {q}.',['q'],facts,45)
 assert '9 units' in grounded('There are {age} units in the 13+ years age band.',['age'],facts,45)
 assert '$500' in grounded('National Account 01 revenue is {account}.',['account'],facts,45)
 for prose in ['Q4 downtime is {q}.','Q2 downtime is 14%.','Q2 losses total $200.','There are {age} units in the 15+ years age band.','National Account 02 revenue is {account}.']:
  with pytest.raises(InputError):grounded(prose,list(facts),facts,45)


def test_supported_finding_survives_question_only_errors():
 f=dict(id='cash',file='example.xlsx',sheet='Cash',ref='B2',label='Operating cash',value=200,unit='money',currency='USD',scale=1)
 i=dict(title='Cash conversion needs investigation',theme='financial',kind='risk',text='Operating cash is {cash}.',why_it_matters='This may affect cash conversion.',management_question='Invented $999 loss?',fact_ids=['cash'],claim=dict(quote_id='',quote='',fact_id='',operator='none',threshold=''))
 d=[];items,brief=validate(dict(insights=[i],briefing=[]),dict(facts=[f],commentary=[]),d)
 assert len(items)==1 and '$999' not in items[0]['management_question'] and items[0]['management_question'].endswith('?')
 assert any(x['section']=='questions' and x['status']=='rewritten' for x in d)
 i['text']='Operating cash is $999.';assert not validate(dict(insights=[i],briefing=[]),dict(facts=[f],commentary=[]))[0]


def test_company_context_omits_unproven_relationships_and_numbers():
 sources=[dict(id='s',file='example.xlsx',sheet='Business',ref='A1',text='North is a geographic operating territory providing collection services.')]
 p=dict(description='The company provides collection services.',citations=[dict(source_id='s',quote='collection services')],segments=[dict(name='North',kind='geographic territory',description='Collection operating territory.',citations=[dict(source_id='s',quote='North is a geographic operating territory')])])
 d=[];assert validate_context(p,sources,d)['segments'][0]['kind']=='geographic territory'
 bad=copy.deepcopy(p);bad['segments'][0]['kind']='subsidiary';assert not validate_context(bad,sources,[])['segments']
 bad=copy.deepcopy(p);bad['description']='The company earns $999.';assert not validate_context(bad,sources,[])
 bad=copy.deepcopy(p);bad['citations'][0]['quote']='Invented industry';assert not validate_context(bad,sources,[])




def test_company_citations_resolve_exact_text_and_unit_names_only():
 sources=[dict(id='wrong',file='example.xlsx',sheet='Services',ref='A1',text='Regional collection services.'),dict(id='heading',file='example.xlsx',sheet='Structure',ref='A1',text='Geographic territories'),dict(id='name',file='example.xlsx',sheet='Structure',ref='B1',text='Northern')]
 p=dict(description='The company provides collection services.',citations=[dict(source_id='heading',quote='Regional collection services.')],segments=[dict(name='Northern',kind='geographic territory',description='Collection operating territory.',citations=[dict(source_id='heading',quote='Geographic territories')])])
 d=[];v=validate_context(p,sources,d);assert v['citations'][0]['source_id']=='wrong' and v['segments'][0]['name']=='Northern'
 assert v['segments'][0]['citations'][-1]['source_id']=='name'
 bad=copy.deepcopy(p);bad['segments'][0]['name']='North';assert not validate_context(bad,sources,[])['segments']
 bad=copy.deepcopy(p);bad['citations'][0]['quote']='Unstated collection industry';assert not validate_context(bad,sources,[])


def test_source_only_summary_preserves_services_territories_and_financials():
 w=Workbook();s=w.active;s.title='Financial review'
 for row in [['Example Operations'],['USD whole units; accrual basis'],['Account','FY2028 Actual','FY2028 Budget','FY2027 Actual'],['Revenue'],['Service contracts',1000,1100,900],['Total revenue',1000,1100,900],['Total direct costs',500,550,450],['Gross profit',500,550,450],['EBITDA',200,220,180]]:s.append(row)
 s=w.create_sheet('Operations')
 for row in [['Example Operations'],['Operating KPIs by Territory, FY2028'],['Units as stated in each row'],[],[],['KPI','Unit','Upland Q1','Upland Q2','Lowland Q1','Lowland Q2'],['Tons handled','tons',10,11,12,13]]:s.append(row)
 b=io.BytesIO();w.save(b);r=analyze({'independent-example.xlsx':b.getvalue()});p=r.package['company_profile']
 assert 'Service contracts' in p['description']
 assert {s['name'] for s in p['segments']}=={'Upland','Lowland'}
 assert all(s['kind']=='geographic territory' and s['citations'] for s in p['segments'])
 assert not r.package['ai_insights'] # descriptive context is never an AI assessment
 assert r.metrics['revenue']==1000 and r.report_period=='FY2028'
 for t in [None,(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes()]:
  data,_,_=export_report(r,t);out=load_workbook(io.BytesIO(data))
  assert 'Investment Summary' in out.sheetnames
  assert any('Upland' in str(row) for row in out['Investment Summary'].values)
 assert fields(r)['INVESTMENT_SUMMARY'] and any(e['Check / use']=='Investment Summary company context' for e in r.evidence)
