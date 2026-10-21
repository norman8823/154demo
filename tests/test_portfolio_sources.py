import io
import re
from copy import deepcopy
from dataclasses import replace
import pytest
from openpyxl import load_workbook
from engine import InputError
from portfolio_data import load_companies
from portfolio_comparison import comparison
from portfolio_excel import export_portfolio, reference_records, SHEETS
from portfolio_sources import snapshot_name, canonical_source


@pytest.fixture(scope='module')
def exported():
    companies=load_companies()
    return companies,load_workbook(io.BytesIO(export_portfolio(companies)))


def target(cell):
    assert cell.hyperlink and cell.hyperlink.location and cell.hyperlink.target is None
    match=re.fullmatch(r"'((?:[^']|'')+)'!([A-Z]+[1-9]\d*)",cell.hyperlink.location)
    assert match,cell.hyperlink.location
    return match[1].replace("''", "'"),match[2]


def test_all_original_cells_and_1755_observed_formula_values_are_retained(exported):
    companies,book=exported
    assert len(book.sheetnames)==33
    assert len({s.casefold() for s in book.sheetnames})==33
    assert all(len(s)<=31 and not re.search(r'[\[\]:*?/\\]',s) for s in book.sheetnames)
    assert not book._external_links
    formulas=0
    for c in companies:
        original=load_workbook(canonical_source(c),data_only=False)
        facts={(f.get('sheet'),f.get('cell')):f for f in c.review.package['facts']}
        for ordinal,sheet in enumerate(original,1):
            snapshot=book[snapshot_name(c.id,ordinal,sheet.title)]
            assert snapshot.sheet_state=='visible' and snapshot.protection.sheet
            assert set(map(str,snapshot.merged_cells.ranges))==set(map(str,sheet.merged_cells.ranges))
            for row in sheet:
                for source in row:
                    dest=snapshot[source.coordinate]
                    assert dest.data_type!='f'
                    if source.data_type=='f':
                        formulas+=1
                        assert dest.value==pytest.approx(facts[(sheet.title,source.coordinate)]['source_value'])
                        assert source.value in dest.comment.text
                        assert 'Original formula' in dest.comment.text
                    else:
                        assert dest.value==source.value
                    if source.value is not None:
                        assert dest.protection.locked
                        assert dest.number_format==source.number_format
    assert formulas==1755


def test_all_three_issue_to_index_to_source_chains_and_recursive_operands(exported):
    companies,book=exported
    records=reference_records(companies)
    indexed={(c.id,e['Fact ID']):row for row,(c,e) in enumerate(records,6)}
    reference=book['Source References']
    def walk(c,sid,visited):
        row=indexed[(c.id,sid)];sheet,coordinate=target(reference.cell(row,4))
        if sheet!='Source References':
            assert book[sheet][coordinate].value is not None
            assert sheet.startswith({'summit':'SHS','tallis':'TBA','harrow':'HRE'}[c.id])
            return {sheet+'!'+coordinate}
        assert sid not in visited
        visited=visited|{sid};start=int(coordinate[1:]);calc=next(f for f in c.review.package['skill_engine']['calculations'] if f['id']==sid)
        assert reference.cell(start,2).value==sid
        assert reference.cell(start,4).value==calc['expression']
        assert reference.cell(start,5).value==pytest.approx(calc['value'])
        leaves=set()
        for index,operand in enumerate(calc['inputs'],start+2):
            assert reference.cell(index,2).value==operand
            assert target(reference.cell(index,8))==target(reference.cell(indexed[(c.id,operand)],4))
            leaves.update(walk(c,operand,visited))
        return leaves
    for c in companies:
        # An original reported revenue formula lands on its actual source coordinate.
        sid=comparison(c)['revenue'].source_ids[0]
        assert len(walk(c,sid,set()))==1
        calc=next(f for f in c.review.package['skill_engine']['calculations'] if (c.id,f['id']) in indexed and len(f['inputs'])>=2)
        assert len(walk(c,calc['id'],set()))>=2
    assert len(walk(companies[0],'calc_51',set()))>=4
    # Every issue index row links first to its correct company-qualified source ID.
    for row in range(6,23):
        sheet,coord=target(book['Investment Issues'].cell(row,12))
        assert sheet=='Source References'
        index_row=int(coord[1:])
        assert reference.cell(index_row,1).value==book['Investment Issues'].cell(row,1).value
        assert reference.cell(index_row,2).value==book['Investment Issues'].cell(row,12).value.split(', ')[0]


def test_unsupported_source_references_and_noncanonical_paths_fail_safely():
    companies=list(load_companies());c=companies[0];r=deepcopy(c.review)
    sid=comparison(c)['revenue'].source_ids[0]
    for e in r.evidence:
        if e['Fact ID']==sid:e['Reference']="'No such source'!A1"
    companies[0]=replace(c,review=r)
    with pytest.raises(InputError,match='Unsupported source worksheet'):export_portfolio(companies)
    companies[0]=replace(c,provenance=dict(c.provenance,workbook='../../outside.xlsx'))
    with pytest.raises(InputError,match='canonical approved manifest'):export_portfolio(companies)


def test_broken_calculation_operands_fail_safely():
    companies=list(load_companies());c=companies[0];r=deepcopy(c.review)
    calc=next(f for f in r.package['skill_engine']['calculations'] if f['id']=='calc_51')
    calc['inputs'].append('unknown_source_operand')
    companies[0]=replace(c,review=r)
    with pytest.raises((InputError,ValueError),match='source reference is unavailable|operand'):export_portfolio(companies)
