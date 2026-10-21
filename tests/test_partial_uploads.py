"""Availability, lineage and export regressions for every supported file subset."""
import csv
import io
from itertools import combinations
import pytest
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest
from engine import analyze, demo_files, InputError, NOT_AVAILABLE, ROOT
from reports import export_report

KINDS=('platform','quickbooks','schedule')
SUBSETS=[subset for n in range(1,4) for subset in combinations(KINDS,n)]


def inputs(kinds):
    files=demo_files()
    names={'platform':next(n for n in files if 'PlatformA' in n),'quickbooks':next(n for n in files if 'QuickBooks' in n),'schedule':next(n for n in files if n.endswith('.csv'))}
    # Filenames carry no business meaning: classification must inspect contents.
    return {f'report_{i}'+('.csv' if kind=='schedule' else '.xlsx'):files[names[kind]] for i,kind in enumerate(kinds)}


def expected_ids(kinds):
    return ({1,2} if 'platform' in kinds else set()) | ({3} if {'quickbooks','schedule'}<=set(kinds) else set())


@pytest.mark.parametrize('kinds',SUBSETS,ids=['+'.join(s) for s in SUBSETS])
def test_supported_subsets_only_produce_available_findings(kinds):
    r=analyze(inputs(kinds)); expected=expected_ids(kinds); full=analyze(demo_files())
    assert {f['id'] for f in r.findings}==expected
    assert len(r.recognized)==len(kinds)
    for row in r.readiness:
        if row['Analysis ID']==0: continue
        if isinstance(row['Analysis ID'],str):
            assert row['Analysis ID'] in {'supporting_schedules','commentary_checks'}
            available=row['Analysis ID']=='supporting_schedules' and 'platform' in kinds
            assert row['Status']==('Available' if available else NOT_AVAILABLE)
            continue
        assert row['Status']==('Available' if row['Analysis ID'] in expected else NOT_AVAILABLE)
        if row['Analysis ID'] not in expected: assert row['Additional data']
    for key,value in r.metrics.items(): assert value==full.metrics[key]
    assert ('billing_total' in r.metrics)==(3 in expected)
    assert ('gm' in r.metrics)==('platform' in kinds)
    assert ('addon_cash_income' in r.metrics)==('quickbooks' in kinds)
    assert len(r.questions)==len(expected)
    if 1 not in expected: assert 'Gross margin 27.8%' not in r.memo
    if 2 not in expected: assert 'Cap-only leverage 4.46x' not in r.memo
    if 3 not in expected: assert '$27,978.75 potential billing exceptions' not in r.memo
    if not expected:
        assert 'not a conclusion that there are no financial exceptions' in r.memo
    for f in r.findings:
        assert all(e['File'] in inputs(kinds) for e in r.evidence)
        assert any(e['Finding']==f['id'] and e['Reference'] for e in r.evidence)


@pytest.mark.parametrize('kinds',SUBSETS,ids=['+'.join(s) for s in SUBSETS])
@pytest.mark.parametrize('custom',[False,True],ids=['standard','custom'])
def test_partial_reports_preserve_available_values_and_omit_unavailable_findings(kinds,custom):
    r=analyze(inputs(kinds)); expected=expected_ids(kinds)
    template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    data,_,warnings=export_report(r,template)
    assert not warnings
    wb=load_workbook(io.BytesIO(data))
    if not custom: assert wb.sheetnames==['Executive Summary','Financial Performance','Findings','Source Evidence','Data Quality','Extracted Detail','Exceptions']
    exceptions=[row for row in wb['Exceptions'].values if isinstance(row[0],int)]
    assert {row[0] for row in exceptions}==expected
    assert all(row[6] for row in exceptions)
    perf={row[0]:row[1] for row in wb['Financial Performance'].values}
    assert ('Platform A revenue' in perf)==('platform' in kinds)
    assert ('Potential billing exceptions' in perf)==(3 in expected)
    if 'platform' in kinds: assert perf['Platform A revenue']==5383752
    if 3 in expected: assert perf['Potential billing exceptions']==27978.75
    detail={row[0]:row[1] for row in wb['Extracted Detail'].values}
    for row in r.readiness: assert detail[row['Analysis']]==row['Status']
    if custom:
        sheet=wb['Partner Review']
        assert sheet['C16'].data_type=='f'
        # Monetary metadata recognition now unlocks an existing receivables insight.
        displayed=min(3,len(expected)+(1 if 'platform' in kinds else 0))
        assert sheet['G16'].value==displayed
        if 'platform' in kinds: assert sheet['C8'].value==5383752
        else:
            assert sheet['C8'].value==NOT_AVAILABLE
            assert sheet['G8'].value==NOT_AVAILABLE
            assert sheet['G10'].value==NOT_AVAILABLE
        for i,cell in [(1,'B19'),(2,'B22'),(3,'B25')]:
            available=(i<=displayed) if r.package.get('view') else (i in expected)
            assert (sheet[cell].value==NOT_AVAILABLE)==(not available)
        original=load_workbook(io.BytesIO(template))
        assert sheet['C16'].value==original['Partner Review']['C16'].value
        assert sheet['C8']._style==original['Partner Review']['C8']._style


@pytest.mark.parametrize('missing,expected',[('Operating KPIs',{2,3}),('Covenant',{1,3}),('P&L Trend',{3}),('Rate Card',{1,2}),('Profit and Loss',{1,2,3})])
def test_missing_sheets_disable_only_dependent_checks(missing,expected):
    files=demo_files()
    name=next(n for n in files if ('QuickBooks' if missing in {'Rate Card','Profit and Loss'} else 'PlatformA') in n)
    wb=load_workbook(io.BytesIO(files[name])); del wb[missing]
    out=io.BytesIO(); wb.save(out); files[name]=out.getvalue()
    r=analyze(files)
    assert {f['id'] for f in r.findings}==expected
    export_report(r)
    if missing=='P&L Trend':
        assert 'revenue' not in r.metrics and 'base_platform' not in r.metrics
        assert not any(e['Finding'] in (1,2) for e in r.evidence)


def test_missing_columns_do_not_disable_independent_checks():
    files=inputs(('platform',))
    name=next(iter(files)); wb=load_workbook(io.BytesIO(files[name]))
    wb['P&L']['C5']=None
    out=io.BytesIO(); wb.save(out); files[name]=out.getvalue()
    r=analyze(files)
    assert {f['id'] for f in r.findings}=={2}
    assert next(row for row in r.readiness if row['Analysis ID']==1)['Status']==NOT_AVAILABLE


@pytest.mark.parametrize('sheet',['Sheet1','P&L'])
def test_unrecognized_workbook_requires_mapping(sheet):
    wb=Workbook(); wb.active.title=sheet; wb.active['A1']='An unrelated financial spreadsheet'
    out=io.BytesIO(); wb.save(out)
    with pytest.raises(InputError,match='No usable dated financial statement'):
        analyze({'unknown.xlsx':out.getvalue()})


def test_empty_and_ambiguous_input_rejected():
    with pytest.raises(InputError,match='at least one'): analyze({})
    files=inputs(('platform',)); files['duplicate.xlsx']=next(iter(files.values()))
    with pytest.raises(InputError,match='Multiple files'): analyze(files)


def test_schedule_recognized_by_columns_in_any_order():
    files=inputs(('quickbooks','schedule')); name=next(n for n in files if n.endswith('.csv'))
    lines=files[name].decode().splitlines(keepends=True)
    reader=csv.DictReader(io.StringIO(''.join(lines[1:])))
    out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=list(reversed(reader.fieldnames)))
    writer.writeheader(); writer.writerows(reader); files[name]=out.getvalue().encode()
    assert analyze(files).metrics['billing_total']==27978.75


@pytest.mark.parametrize('kinds',[('platform',),('quickbooks',),('schedule',),('quickbooks','schedule')])
def test_partial_ui_generation_evidence_and_export(kinds):
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30)
    at.session_state['uploaded_inputs']=inputs(kinds)
    at.run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception and not at.error
    assert {f['id'] for f in at.session_state['review'].findings}==expected_ids(kinds)
    assert 'Data Quality & Coverage' in [e.label for e in at.expander]
    for f in at.session_state['review'].findings:
        if at.radio(key='view').value!='02  Portfolio Review': at.radio(key='view').set_value('02  Portfolio Review').run()
        at.button(key=f'evidence_{f["id"]}').click().run()
        assert not at.exception
        assert at.radio(key='selected').value==f['id']
    at.radio(key='view').set_value('03  Evidence & Excel').run()
    assert not at.exception
    assert 'Downloadable Excel report' in [h.value for h in at.subheader]


def test_schedule_missing_billing_columns_is_unavailable_not_a_financial_exception():
    files=inputs(('quickbooks','schedule')); name=next(n for n in files if n.endswith('.csv'))
    lines=files[name].decode().splitlines(keepends=True)
    reader=csv.DictReader(io.StringIO(''.join(lines[1:])))
    columns=[c for c in reader.fieldnames if c!='Invoice #']
    out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=columns,extrasaction='ignore')
    writer.writeheader(); writer.writerows(reader); files[name]=out.getvalue().encode()
    r=analyze(files)
    assert not r.findings and 'billing_total' not in r.metrics
    status=next(row for row in r.readiness if row['Analysis ID']==3)
    assert status['Status']==NOT_AVAILABLE and 'Invoice #' in status['Additional data']
    export_report(r)


def test_empty_rate_card_does_not_report_zero_dollar_leakage():
    files=inputs(('quickbooks','schedule')); name=next(n for n in files if n.endswith('.xlsx'))
    wb=load_workbook(io.BytesIO(files[name])); wb['Rate Card'].delete_rows(5,100)
    out=io.BytesIO(); wb.save(out); files[name]=out.getvalue()
    r=analyze(files)
    assert not r.findings and 'billing_total' not in r.metrics
    assert next(row for row in r.readiness if row['Analysis ID']==3)['Status']==NOT_AVAILABLE


def test_schedule_only_requires_columns_used_by_existing_billing_check():
    files=inputs(('quickbooks','schedule')); name=next(n for n in files if n.endswith('.csv'))
    lines=files[name].decode().splitlines(keepends=True)
    reader=csv.DictReader(io.StringIO(''.join(lines[1:])))
    columns=['Event ID','Venue','Position','Worked Hours','Bill Rate','Invoice #']
    out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=columns,extrasaction='ignore')
    writer.writeheader(); writer.writerows(reader); files[name]=out.getvalue().encode()
    assert analyze(files).metrics['billing_total']==27978.75
