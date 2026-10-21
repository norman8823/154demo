"""Semantic evidence destinations, beyond mere hyperlink resolution."""
import io
import re
import pytest
from openpyxl import load_workbook
from portfolio_data import load_companies
from portfolio_excel import export_portfolio
from portfolio_comparison import comparison, metric_operand


def destination(book,cell):
    sheet,coordinate=re.fullmatch(r"'((?:[^']|'')+)'!([A-Z]+\d+)",cell.hyperlink.location).groups()
    return book[sheet.replace("''", "'")][coordinate]


def test_budget_links_have_budget_operands_and_margin_has_all_four():
    companies=load_companies();book=load_workbook(io.BytesIO(export_portfolio(companies)))
    matrix=book['Company Comparison'];index=book['Source References']
    for col,c in enumerate(companies,3):
        for row,key in [(8,'revenue'),(11,'ebitda')]:
            proof=destination(book,matrix.cell(row,col));start=proof.row
            assert index.cell(start,5).value==comparison(c)[key].budget
            assert index.cell(start+2,2).value==key+' · budget'
            expected=metric_operand(c,key,'budget')[3][0]
            assert index.cell(start+2,8).value==expected
            source=destination(book,index.cell(start+2,9))
            assert index.cell(source.row,2).value==expected
            original=destination(book,index.cell(source.row,4))
            assert original.value is not None
        proof=destination(book,matrix.cell(16,col))
        assert {index.cell(proof.row+i,2).value for i in range(2,6)}=={
            'ebitda · actual','revenue · actual','ebitda · budget','revenue · budget'}
        actual=comparison(c)['margin'].actual;budget=comparison(c)['margin'].budget
        assert index.cell(proof.row,5).value==pytest.approx((actual-budget)*100)
    # The reported defect: Tallis budgets must end at I12 / I41, never G actuals.
    for row,coordinate in [(8,'I12'),(11,'I41')]:
        proof=destination(book,matrix.cell(row,4))
        source=destination(book,index.cell(proof.row+2,9))
        assert destination(book,index.cell(source.row,4)).coordinate==coordinate
