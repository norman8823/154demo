"""Executive presentation stays short, source-grounded and shared with exports."""
import copy
import io
from pathlib import Path
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from engine import analyze,demo_files,ROOT
from executive import present,briefing,money
from reports import export_report,fields

@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    monkeypatch.setattr('mapping_ui.load_profiles',lambda:{})

@pytest.fixture
def summit():
    p=ROOT/'tests/fixtures/Summit_Home_Services_Sep2026.xlsx'
    review=analyze({p.name:p.read_bytes()})
    review.package={}  # Unit regression for the original financial-only presentation.
    return review


def test_summit_distinct_bounded_cards_and_no_mutation(summit):
    original=copy.deepcopy(summit)
    s=present(summit)
    assert len(s['cards'])==2
    performance,quality=s['cards']
    assert performance['impact']=='−$157K' and performance['kind']=='Business performance'
    assert quality['impact']=='$47K' and quality['kind']=='Data quality' and not quality['adverse']
    assert '$99K' in performance['summary'] and '$83K' in performance['summary']
    assert '5.7 pp' in performance['summary']
    assert len(performance['ids'])==2
    for card in s['cards']:
        assert len(card['title'].split())<=8 and len(card['summary'].split())<=40
        assert all(any(e['Finding']==identity for e in summit.evidence) for identity in card['ids'])
    assert summit==original


def test_briefing_and_questions_are_executive_and_source_grounded(summit):
    s=present(summit);text=briefing(s)
    assert 100<=len(text.split())<=150
    for term in ['financial-statement reader','explicit EBITDA bridge','user-confirmed units','2026-09','gross_profit']:
        assert term not in text
    for value in ['$3.30M','$502K','$157K','$99K','$83K','$47K']:assert text.count(value)==1
    assert 'HVAC equipment' in text and 'beat budget' in text
    assert len(s['questions'])==2
    assert '$99K' in s['questions'][0] and '$83K' in s['questions'][0]
    assert '$47K' in s['questions'][1]
    assert all(q.endswith('?') and len(q.split('?'))==2 for q in s['questions'])
    assert 'Accounting basis is not stated.' in s['notice']


def test_margin_is_not_merged_without_shared_cost_evidence(summit):
    for row in summit.normalized:
        if row['category']=='cogs':row['category']='operating_expenses'
    assert len(present(summit)['cards'])==3


def test_branch_gaps_are_not_summed_or_misrepresented(summit):
    finding=next(f for f in summit.findings if 'Branch/entity' in f['title'])
    finding['summary']=finding['summary'].replace('difference $ 47,000.00','difference $ 60,000.00',1)
    quality=present(summit)['cards'][1]
    assert quality['impact']=='$60K' and 'largest' in quality['summary']
    assert '$154K' not in quality['summary'] and '$141K' not in quality['summary']
    assert 'multiple discrepancies' in briefing(present(summit))


@pytest.mark.parametrize('custom',[False,True])
def test_concise_excel_summary_with_exact_supporting_financials(summit,custom):
    s=present(summit)
    template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    data,_,_=export_report(summit,template);wb=load_workbook(io.BytesIO(data))
    assert fields(summit)['FINDING_COUNT']==2
    assert fields(summit)['PARTNER_BRIEF']==briefing(s)
    assert any(row[1]==3298000 for row in wb['Financial Performance'].values)
    assert any(row[1]==502000 for row in wb['Financial Performance'].values)
    # Supporting checks and full cell lineage are retained, even when grouped in the summary.
    assert len([row for row in wb['Exceptions'].values if isinstance(row[0],int)])==3
    if not custom:
        rows=list(wb['Executive Summary'].values)
        assert len([row for row in rows if str(row[0]).startswith('Finding ')])==2
        assert next(row[1] for row in rows if row[0]=='Investment Assessment')==briefing(s)
        assert not any(str(row[0]).startswith('Data readiness ·') for row in rows)
    else:
        assert any(c.value==briefing(s) for row in wb['Partner Review'] for c in row)
        assert wb['Partner Review']['C16'].data_type=='f'


def test_formatting_never_invents_currency(summit):
    assert money(summit,502000)=='$502K' and money(summit,-157000)=='−$157K'
    summit.currency='unknown'
    assert money(summit,502000)=='502K source units'


def test_original_demo_keeps_three_distinct_supported_issues():
    r=analyze(demo_files());s=present(r)
    assert len(s['cards'])==3 and {c['id'] for c in s['cards']}=={1,2,3}
    assert len(s['questions'])==3
    assert all(len(c['summary'].split())<=40 for c in s['cards'])
    assert 'not a breach forecast' in s['cards'][1]['summary']
    assert 'potential exceptions' in s['cards'][2]['summary']


def test_summit_ui_two_cards_compact_metrics_and_merged_evidence():
    p=ROOT/'tests/fixtures/Summit_Home_Services_Sep2026.xlsx'
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30)
    at.session_state['uploaded_inputs']={p.name:p.read_bytes()};at.run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception
    assert [(m.label,m.value) for m in at.metric][:3]==[('Revenue','$3.30M'),('EBITDA','$502K'),('Gross margin','43.5%')]
    assert at.metric[1].delta=='-$157K vs. budget' and at.metric[2].delta=='-5.7 pp vs. budget'
    assert any('$99K overtime' in m.value and '$83K HVAC equipment' in m.value for m in at.markdown)
    assert len([b for b in at.button if b.label.startswith('View Evidence')])==3
    assert next(e for e in at.expander if e.label=='Data Quality & Coverage').proto.expanded is False
    at.button(key='evidence_101').click().run()
    assert not at.exception
    # Gross margin sources are available in the same one-click evidence path as EBITDA.
    frames=[d.value for d in at.dataframe if 'Fact ID' in d.value.columns]
    assert len(frames)==1 and {101,103}<=set(frames[0]['Finding'])


def test_changed_source_amounts_change_briefing_and_cost_drivers():
    p=ROOT/'tests/fixtures/Summit_Home_Services_Sep2026.xlsx'
    wb=load_workbook(p);wb['Monthly Detail']['K18']=174
    data=io.BytesIO();wb.save(data)
    summary=present(analyze({'changed.xlsx':data.getvalue()}))
    assert summary['cards'][0]['impact']=='−$167K'
    assert '$109K' in summary['cards'][0]['summary']
    assert '$109K' in briefing(summary) and '$99K' not in briefing(summary)


def test_sparse_revenue_question_uses_actual_finding_without_generic_filler():
    r=analyze({'plain.csv':b'Example Company,,\nUSD whole units accrual,,\nAccount,Sep-26 Actual,Aug-26 Actual\nRevenue,100,80\nCost of sales,60,48\nGross profit,40,32\n'})
    s=present(r)
    assert len(s['cards'])==1 and len(s['questions'])==1
    assert '+25.0%' in s['questions'][0] and 'forecast' in s['questions'][0]
    assert 'Discuss the variance' not in s['questions'][0]
    assert 'EBITDA' not in briefing(s)
