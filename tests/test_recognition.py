"""Content-based recognition regressions; source amounts stay local and provider calls are blocked."""
import io
from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from engine import analyze,ROOT
from ingestion import inspect_files,discover_columns,metadata
from reports import export_report

FIXTURES=ROOT/'tests/fixtures'

@pytest.fixture(autouse=True)
def offline(monkeypatch):monkeypatch.delenv('OPENAI_API_KEY',raising=False)

def encoded(w):
    out=io.BytesIO();w.save(out);return out.getvalue()

def tallis():
    p=FIXTURES/'Tallis_Brook_Accounting_Sep2026.xlsx';return {p.name:p.read_bytes()}


def test_tallis_recognizes_all_financial_sources_and_expected_reported_totals():
    intake=inspect_files(tallis())
    assert not intake.unresolved
    assert {t.sheet for t in intake.tables}=={'Dashboard','Practice P&L','Monthly Trend'}
    t=next(t for t in intake.tables if t.sheet=='Practice P&L')
    assert t.currency=='USD' and t.scale==1 and t.basis=='unknown'
    assert {c['column'] for c in t.columns}=={7,8,9} # Entity acquisition dates are not financial periods.
    r=analyze(tallis())
    assert r.company=='TALLIS BROOK ACCOUNTING GROUP' and r.report_period=='September 2026'
    assert r.metrics['revenue']==3475418 and r.metrics['ebitda']==684306
    assert any(e['Reference']=="'Practice P&L'!G12" for e in r.evidence)
    assert any(e['Reference']=="'Practice P&L'!G41" for e in r.evidence)
    assert any(s['sheet']=='Monthly Trend' for s in r.package['stats'])
    assert any(q['sheet']=='MP Letter' for q in r.package['narratives'])
    assert any('accounting basis not stated' in w for w in r.warnings)


@pytest.mark.parametrize('custom',[False,True])
def test_tallis_standard_and_template_excel_export(custom):
    r=analyze(tallis());template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    data,_,_=export_report(r,template);w=load_workbook(io.BytesIO(data))
    assert any(str(row[0]).endswith('actual · revenue') and row[1]==3475418 for row in w['Extracted Detail'].values)
    assert any(str(row[0]).endswith('actual · ebitda') and row[1]==684306 for row in w['Extracted Detail'].values)
    if custom:assert w['Partner Review']['C16'].data_type=='f'


def test_company_and_worksheet_renaming_does_not_change_recognition():
    w=load_workbook(io.BytesIO(next(iter(tallis().values()))));names={s.title:'Schedule '+str(i) for i,s in enumerate(w)}
    for s in w:
        for row in s:
            for c in row:
                if c.data_type=='f':
                    for old,new in names.items():c.value=c.value.replace("'"+old+"'!","'"+new+"'!")
        s['A1']='Example Professional Group';s.title=names[s.title]
    r=analyze({'renamed.xlsx':encoded(w)})
    assert r.company=='Example Professional Group'
    assert r.metrics['revenue']==3475418 and r.metrics['ebitda']==684306


def conventional(units=None):
    w=Workbook();s=w.active;s.title='Economics'
    s.append(['Example Industrial Company']);s.append([units])
    s.append(['Line','July 2028 Actual','August 2028 Actual','September 2028 Actual','August 2028 Budget'])
    s.append(['Revenue',1000,1200,None,1250]);s.append(['Cost of services',600,700,None,750]);s.append(['Gross profit',400,500,None,500]);s.append(['EBITDA (reported)',150,200,None,210])
    return w


def test_latest_complete_period_selected_without_user_configuration():
    r=analyze({'new.xlsx':encoded(conventional('USD whole units'))})
    assert r.report_period=='August 2028' and r.metrics['revenue']==1200 and r.metrics['ebitda']==200
    assert r.metrics['revenue_prior']==1000 and r.metrics['revenue_budget']==1250
    assert any('latest complete actual period' in w for w in r.warnings)


def test_missing_metadata_uses_source_units_without_inventing_basis_currency_or_scale():
    r=analyze({'unstated.xlsx':encoded(conventional())})
    assert r.metrics['revenue']==1200 and r.metrics['ebitda']==200
    assert r.currency=='unknown' and r.scale is None
    assert {row['basis'] for row in r.normalized}=={'unknown'}
    assert 'revenue_variance' not in r.metrics # Unknown multipliers cannot establish comparability.
    assert any('source units' in w for w in r.warnings)
    assert export_report(r)[0]


@pytest.mark.parametrize('text,currency,scale',[
    ('US dollars unless noted','USD',1),('United States dollars in thousands','USD',1000),
    ('U.S. dollars, millions','USD',1000000),('USD whole units','USD',1),
    ('Dollars unless noted','$',1),('USD and EUR, whole units','unknown',1),
])
def test_explicit_denomination_phrases(text,currency,scale):
    w=conventional(text);assert metadata(w.active,3)[2:4]==(currency,scale)


def test_consistent_workbook_units_reused_with_source_reference():
    w=conventional();s=w.create_sheet('Overview');s.append(['Example Industrial Company']);s.append(['US dollars unless noted'])
    i=inspect_files({'workbook.xlsx':encoded(w)});t=next(t for t in i.tables if t.sheet=='Economics')
    assert (t.currency,t.scale,t.basis)==('USD',1,'unknown')
    assert any(ref=="'Overview'!A2" for ref,text in t.metadata_refs)


def test_another_company_metadata_is_not_inherited():
    w=conventional();s=w.create_sheet('Another entity');s.append(['Different Group']);s.append(['USD in millions'])
    t=inspect_files({'different-companies.xlsx':encoded(w)}).tables[0]
    assert t.currency=='unknown' and t.scale is None


def test_conflicting_workbook_units_remain_unknown():
    w=conventional()
    for name,units in [('Overview','USD in thousands'),('Notes','EUR in millions')]:
        s=w.create_sheet(name);s.append(['Example Industrial Company']);s.append([units])
    t=inspect_files({'conflict.xlsx':encoded(w)}).tables[0]
    assert t.currency=='unknown' and t.scale is None


def test_currency_number_format_does_not_invent_a_usd_jurisdiction():
    w=conventional()
    for row in w.active.iter_rows(min_row=4):
        for c in row[1:]:c.number_format='$#,##0.00'
    r=analyze({'formatted.xlsx':encoded(w)})
    assert r.currency=='$' and r.scale==1 and r.metrics['revenue']==1200
    assert any('no USD identity' in w for w in r.warnings)


def test_supported_formula_totals_count_as_table_evidence():
    w=Workbook();s=w.active
    for row in [['Example Wholesale'],['USD whole dollars'],[None,'January 2028 Actual','January 2028 Budget'],['Product sales',700,800],['Service sales',300,350],['Total revenue','=SUM(B4:B5)','=SUM(C4:C5)'],['Total direct costs',400,450],['Reported EBITDA','=B6-B7','=C6-C7']]:s.append(row)
    r=analyze({'formulas.xlsx':encoded(w)})
    assert r.metrics['revenue']==1000 and r.metrics['ebitda']==600


def test_acquisition_dates_are_not_supporting_schedule_months():
    r=analyze(tallis())
    assert all({p['period'] for p in s['points']}<={'2025-09','2026-09'} for s in r.package['stats'] if s['sheet']=='Practice P&L')


def test_automatic_model_layout_uses_verified_source_headline_without_approval(monkeypatch):
    import mapping_ai
    w=Workbook();s=w.active;s.title='Unfamiliar captions'
    for row in [['Example Consulting'],[],['Business captions','August 2028'],['Revenue',500],['Unclassified balance',90]]:s.append(row)
    files={'new-format.xlsx':encoded(w)}
    assert not inspect_files(files).tables
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    monkeypatch.setattr(mapping_ai,'suggest',lambda *a,**k:dict(accounts=[],columns=[],layouts=[dict(file='new-format.xlsx',sheet=s.title,header_row=3,label_column=1,reason='Reported dated revenue')]))
    r=analyze(files)
    assert r.metrics['revenue']==500 and 'ebitda' not in r.metrics
    assert r.currency=='unknown' and r.scale is None
    assert any(not row['accepted'] for row in r.normalized if row['label']=='Unclassified balance')


def test_model_layout_cannot_invent_a_period_or_ordinary_ebitda(monkeypatch):
    import mapping_ai
    from ingestion import MappingRequired
    w=Workbook();s=w.active
    for row in [['Example Services'],[],['Business captions','Current'],['Adjusted EBITDA',900]]:s.append(row)
    files={'undated.xlsx':encoded(w)};monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    monkeypatch.setattr(mapping_ai,'suggest',lambda *a,**k:dict(accounts=[],columns=[],layouts=[dict(file='undated.xlsx',sheet=s.title,header_row=3,label_column=1,reason='Proposed layout')]))
    with pytest.raises(MappingRequired):analyze(files)


def test_populated_but_conflicting_latest_period_is_not_hidden_by_older_data():
    w=conventional('USD whole units');s=w.active
    for row,value in [(4,1400),(5,800),(6,600),(7,250)]:s.cell(row,4,value)
    other=w.create_sheet('Second statement')
    for row in [['Example Industrial Company'],['USD whole units'],['Account','September 2028 Actual'],['Revenue',1500],['EBITDA',250]]:other.append(row)
    r=analyze({'contradiction.xlsx':encoded(w)})
    assert r.report_period=='September 2028' and 'revenue' not in r.metrics
    assert any('conflicting metrics require review' in w for w in r.warnings)


def test_missing_metadata_ui_generates_without_configuration(monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr('mapping_ui.load_profiles',lambda:{})
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30)
    at.session_state['uploaded_inputs']={'missing-units.xlsx':encoded(conventional())}
    at.run();button=next(b for b in at.button if b.label.startswith('Generate Portfolio Review'))
    assert not button.disabled
    button.click().run()
    assert not at.exception
    assert at.session_state['review'].metrics['revenue']==1200
    assert at.session_state['export'][0]


def test_conflicting_source_currency_formats_do_not_silently_override_metadata():
    w=conventional('USD whole units')
    for row in w.active.iter_rows(min_row=4):
        for c in row[1:]:c.number_format='EUR #,##0'
    r=analyze({'format-conflict.xlsx':encoded(w)})
    assert r.currency=='unknown' and r.metrics['revenue']==1200
    assert 'revenue_variance' not in r.metrics


def test_basis_conflict_in_overlapping_source_is_not_accepted_as_unknown():
    w=conventional('USD whole units; accrual and cash basis')
    s=w.create_sheet('Second statement')
    for row in [['Example Industrial Company'],['USD whole units; cash basis'],['Account','August 2028 Actual'],['Revenue',1200],['EBITDA',200]]:s.append(row)
    r=analyze({'basis-conflict.xlsx':encoded(w)})
    assert 'revenue' not in r.metrics
    assert any('Accounting basis conflicts' in w for w in r.warnings)


def test_reported_total_closes_section_before_unfamiliar_accounts():
    w=Workbook();s=w.active
    for row in [['Example Business'],['USD whole units'],['Account','September 2028 Actual'],['Revenue'],['Product sales',1000],['Total revenue',1000],['Unfamiliar operating section'],['Specialist labor',400],['Reported EBITDA',200]]:s.append(row)
    r=analyze({'section-boundaries.xlsx':encoded(w)})
    assert r.metrics['revenue']==1000 and r.metrics['ebitda']==200
    labor=next(row for row in r.normalized if row['label']=='Specialist labor')
    assert not labor['accepted'] and labor['category']=='other'
