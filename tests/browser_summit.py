"""Real upload -> Generate -> evidence -> standard/custom Excel; no mappings or AI calls."""
import io
import json
import os
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts';OUT.mkdir(exist_ok=True)
URL=os.environ.get('COPILOT_TEST_URL','http://127.0.0.1:8501')
with sync_playwright() as p:
    browser=p.chromium.launch(executable_path='/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',headless=True)
    page=browser.new_page(viewport={'width':1440,'height':1100})
    errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    page.goto(URL)
    expect(page.get_by_role('heading',name='Generate a portfolio review',exact=True)).to_be_visible(timeout=30000)
    page.locator('input[type=file]').first.set_input_files(ROOT/'tests/fixtures/Summit_Home_Services_Sep2026.xlsx')
    panel=page.locator('details').filter(has=page.get_by_text('Review Account Mappings',exact=True))
    expect(panel).to_have_count(1)
    assert panel.get_attribute('open') is None
    expect(page.get_by_role('button',name='Generate Portfolio Review')).to_be_enabled()
    panels=[page.locator('.st-key-'+key) for key in ['source_panel','template_panel','deliverable_panel']]
    boxes=[item.bounding_box() for item in panels]
    assert all(boxes) and boxes[0]['y']<boxes[1]['y']<boxes[2]['y']
    upload_boxes=[item.bounding_box() for item in page.locator('[data-testid="stFileUploaderDropzone"]').all()]
    assert len(upload_boxes)==2
    for dimension in ['x','width','height']:
        assert abs(upload_boxes[0][dimension]-upload_boxes[1][dimension])<2,(dimension,upload_boxes)
    expect(page.locator('.deliverable')).to_have_count(4)
    expect(page.locator('.deliverable button,.deliverable a,.deliverable input')).to_have_count(0)
    page.screenshot(path=str(OUT/'summit-upload-optional-mappings.png'),full_page=True)
    page.get_by_role('button',name='Generate Portfolio Review').click()
    expect(page.get_by_role('heading',name='September 2026 Portfolio Review',exact=True)).to_be_visible(timeout=30000)
    expect(page.get_by_text('$3.30M',exact=True)).to_be_visible()
    expect(page.get_by_text('43.5%',exact=True)).to_be_visible()
    expect(page.get_by_text('$502K',exact=True)).to_be_visible()
    expect(page.get_by_text('Reporting totals require reconciliation',exact=False)).to_be_visible()
    expect(page.locator('[data-testid="stException"]')).to_have_count(0)
    expect(page.get_by_role('button',name='View Evidence')).to_have_count(3)
    expect(page.get_by_role('heading',name='Profitability below budget',exact=True)).to_be_visible()
    expect(page.get_by_role('heading',name='Collections risk is concentrated',exact=True)).to_be_visible()
    expect(page.get_by_role('heading',name='Renewal rate deteriorating',exact=True)).to_be_visible()
    expect(page.get_by_text('−$157K',exact=True)).to_be_visible()
    expect(page.get_by_text('$364K',exact=True)).to_be_visible()
    extra=page.locator('details').filter(has=page.get_by_text('Additional insights & commentary',exact=True))
    assert extra.get_attribute('open') is None
    extra.locator('summary').click()
    expect(extra.get_by_text('CONTRADICTED · Management claim',exact=True)).to_have_count(3)
    expect(extra.get_by_text('Branch reporting needs reconciliation · $47K',exact=True)).to_be_visible()
    extra.locator('summary').click()
    expect(page.get_by_text('-$157K vs. budget',exact=True)).to_be_visible()
    expect(page.get_by_text('-5.7 pp vs. budget',exact=True)).to_be_visible()
    expect(page.locator('.st-key-question_panel .katex')).to_have_count(0)
    disclosure=page.locator('details').filter(has=page.get_by_text('Data Quality & Coverage',exact=True))
    assert disclosure.get_attribute('open') is None
    card_boxes=[c.bounding_box() for c in page.locator('.finding').all()]
    assert abs(card_boxes[0]['width']-card_boxes[1]['width'])<2 and abs(card_boxes[0]['y']-card_boxes[1]['y'])<2
    for card in page.locator('.finding').all():
        assert len(card.locator('h3').inner_text().split())<=8
        assert len(card.locator('p').inner_text().split())<=40
    with page.expect_download() as downloaded:page.get_by_role('button',name='Download Excel Report').first.click()
    downloaded.value.save_as(OUT/'summit-standard.xlsx')
    report=load_workbook(OUT/'summit-standard.xlsx')
    assert any(row[1]==3298000 for row in report['Financial Performance'].values)
    assert any(row[1]==502000 for row in report['Financial Performance'].values)
    assert any('47,000' in str(v) for row in report['Exceptions'].values for v in row)
    page.screenshot(path=str(OUT/'summit-one-click-review.png'),full_page=True)
    page.get_by_role('button',name='View Evidence').first.click()
    expect(page.get_by_role('heading',name='See where every conclusion came from',exact=True)).to_be_visible()
    expect(page.locator('[data-testid="stException"]')).to_have_count(0)
    page.get_by_text('01  Upload & Setup',exact=True).click()
    page.locator('input[type=file]').nth(1).set_input_files(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx')
    expect(page.get_by_role('button',name='Remove Sample_Firm_Review_Template.xlsx',exact=True)).to_be_visible()
    expect(page.get_by_role('radio',name='Use existing firm template',exact=True)).to_be_checked()
    page.get_by_role('button',name='Generate Portfolio Review').click()
    expect(page.get_by_role('heading',name='September 2026 Portfolio Review',exact=True)).to_be_visible(timeout=30000)
    with page.expect_download() as downloaded:page.get_by_role('button',name='Download Excel Report').first.click()
    downloaded.value.save_as(OUT/'summit-custom.xlsx')
    custom=load_workbook(OUT/'summit-custom.xlsx')
    assert custom['Partner Review']['C8'].value==3298000
    assert custom['Partner Review']['C16'].data_type=='f'
    page.set_viewport_size({'width':600,'height':1000})
    mobile_boxes=[c.bounding_box() for c in page.locator('.finding').all()]
    assert mobile_boxes[1]['y']>mobile_boxes[0]['y'] and mobile_boxes[2]['y']>mobile_boxes[1]['y']
    page.screenshot(path=str(OUT/'summit-mobile-review.png'),full_page=True)
    assert not errors
    (OUT/'summit-browser-results.json').write_text(json.dumps({'upload_without_account_approval':'passed','optional_mapping_collapsed':'passed','headlines':'passed','material_reconciliation_warning':'passed','evidence_navigation':'passed','standard_excel':'passed','custom_excel':'passed','browser_errors':errors},indent=2))
    browser.close()
    print('PASS: Summit upload -> Generate -> source evidence -> standard/custom Excel, without mapping approval.')
