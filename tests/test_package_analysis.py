"""Source-first package evaluation; no answer key, filenames or model output drives arithmetic."""
import copy
import io
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from engine import analyze,demo_files,ROOT,InputError
from executive import present,briefing
from package_reader import read_package
from package_analysis import build,series_stats,analysis_findings
from analyst_ai import interpret
from reports import export_report

FIXTURES=ROOT/'tests/fixtures'

def source(name='Summit_Home_Services_Sep2026.xlsx'):
    p=FIXTURES/name;return {p.name:p.read_bytes()}

def encoded(w):
    b=io.BytesIO();w.save(b);return b.getvalue()

@pytest.fixture(autouse=True)
def offline(monkeypatch):monkeypatch.delenv('OPENAI_API_KEY',raising=False)


def test_full_package_facts_claims_risks_and_positive_developments():
    r=analyze(source());s=present(r)
    assert r.metrics['revenue']==3298000 and r.metrics['ebitda']==502000
    assert len(s['cards'])==3
    assert {c['theme'] for c in s['cards']}=={'financial','collections','retention'}
    assert any(c['positive'] and c['theme']=='pipeline' for c in s['secondary'])
    assert any('Northern' in c['title'] and c['positive'] for c in s['secondary'])
    assert any('YTD EBITDA' in c['title'] and '+7.1%' in c['impact'] for c in s['secondary'])
    assert any(f.get('roll_forward',{}).get('net')==-4 for f in r.package['findings'])
    assert any(c['quality'] and c['impact']=='$47K' for c in s['secondary'])
    assert 150<=len(briefing(s).split())<=200
    claims=r.package['claims']
    assert any(c['status']=='contradicted' and 'record 13,661' in c['quote'] for c in claims)
    assert any(c['status']=='contradicted' and 'Equipment and material' in c['quote'] for c in claims)
    assert any(c['status']=='supported' and 'majority' in c['quote'] for c in claims)
    assert any(c['status']=='unverifiable' and 'full-year' in c['quote'] for c in claims)
    assert any(f['theme']=='unit' and '$149K' in f['indicator'] for f in r.package['findings'])
    assert r.package['aging'][0]['gl_tie'] is True
    assert any(f['value']==-35000 and 'allowance' in f['label'].lower() for f in r.package['facts'])
    dso=next(s for s in r.package['stats'] if s['label'].startswith('DSO'))
    assert dso['unit']=='days' and dso['latest']['value']==pytest.approx(2484/3298*30)


def test_cross_sheet_profitability_evidence_is_available_in_one_path():
    r=analyze(source());c=next(c for c in present(r)['cards'] if c['theme']=='financial')
    evidence=[e for e in r.evidence if e['Finding'] in c['ids']]
    assert any('Income Statement' in e['Reference'] for e in evidence)
    assert any('Operating Metrics' in e['Reference'] for e in evidence)
    assert all(e['Fact ID'] and e['File'] and e['Reference'] for e in evidence)


def test_synthetic_industry_package_against_recorded_development_expectations():
    expected=json.loads((FIXTURES/'expected_package.json').read_text());r=analyze(source('C_Northgate.xlsx'));s=present(r)
    assert r.metrics['revenue']==expected['revenue'] and r.metrics['ebitda']==expected['ebitda']
    aged=next(f for f in r.package['findings'] if 'aging' in f['key'])['aging']
    assert aged['overdue']==expected['over90'] and aged['ratio']==expected['over90_share']
    renewal=next(st for st in r.package['stats'] if st['theme']=='retention')
    assert renewal['start']['value']==expected['renewal_start'] and renewal['latest']['value']==expected['renewal_end']
    assert len(s['cards'])==3
    assert any(c['status']=='contradicted' and 'record' in c['quote'] for c in r.package['claims'])


def test_renaming_all_tabs_preserves_findings_without_name_rules():
    w=load_workbook(FIXTURES/'Summit_Home_Services_Sep2026.xlsx');names={s.title:'Tab '+str(i) for i,s in enumerate(w,1)}
    for sheet in w:
        for row in sheet:
            for c in row:
                if c.data_type=='f':
                    for old,new in names.items():c.value=c.value.replace("'"+old+"'!","'"+new+"'!")
    for sheet in w:sheet.title=names[sheet.title]
    r=analyze({'renamed.xlsx':encoded(w)})
    assert {c['theme'] for c in present(r)['cards']}=={'financial','collections','retention'}
    assert any(c['status']=='contradicted' and 'record' in c['quote'] for c in r.package['claims'])


def test_aging_gaps_do_not_become_fictitious_zero_or_validated_exposure():
    w=load_workbook(FIXTURES/'C_Northgate.xlsx');w['Customer balances']['F5']=None
    r=analyze({'incomplete.xlsx':encoded(w)})
    assert r.package['aging'][0]['valid'] is False
    assert not any('aging_' in f['key'] for f in r.package['findings'])
    assert any('withheld' in w for w in r.warnings)
    assert r.metrics['revenue']==620000


def test_missing_month_breaks_persistence_and_unknown_metrics_stay_neutral():
    p=read_package(source('C_Northgate.xlsx'));s=next(s for s in p['series'] if s['label']=='Repeat order retention')
    s['points'][-2]['period']='2027-03'
    assert series_stats(s) is None  # duplicate period, cannot assert a sustained trend
    unknown=next(s for s in p['series'] if s['label']=='First-pass yield')
    assert series_stats(unknown)['bad_direction']==0


@pytest.mark.parametrize('custom',[False,True])
def test_package_evidence_and_conclusions_reach_both_excel_modes(custom):
    r=analyze(source());template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    data,_,_=export_report(r,template);w=load_workbook(io.BytesIO(data))
    assert any('Renewal rate' in str(row) for row in w['Extracted Detail'].values)
    assert any('Builder Customer 01' in str(row) for row in w['Extracted Detail'].values)
    assert any('contradicted' in str(row).lower() and 'record 13,661' in str(row) for row in w['Extracted Detail'].values)
    assert any('Operating Metrics' in str(row) for row in w['Source Evidence'].values)
    assert any('AR Aging' in str(row) for row in w['Source Evidence'].values)
    assert any(row[1]==3298000 for row in w['Extracted Detail'].values)
    if custom:assert w['Partner Review']['C16'].data_type=='f'


def response(result):return {'status':'completed','output':[{'content':[{'type':'output_text','text':json.dumps(dict(result,insights=result.get('insights',[]),briefing=result.get('briefing',[])))}]}]}


def test_ai_interprets_unfamiliar_metric_but_local_code_computes_outcome(monkeypatch):
    r=analyze(source('C_Northgate.xlsx'));p=copy.deepcopy(r.package);analysis=build(p,r)
    stat=next(s for s in analysis['stats'] if s['label']=='First-pass yield')
    result={'metrics':[dict(series_id=stat['series_id'],theme='quality',bad_direction=-1,source_label=stat['label'])],'claims':[],'interpretations':[]}
    seen=[]
    def request(req):seen.append(json.loads(req.data));return response(result)
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    semantic=interpret(p,analysis,analysis['findings'],r,request)
    enhanced=build(p,r,semantic)
    f=next(f for f in enhanced['findings'] if f.get('series_id')==stat['series_id'])
    assert f['indicator']=='-22.0 pp' and f['semantic_proposal'] is True
    assert 'AI interpretation' in f['caveat']
    assert seen[0]['store'] is False and seen[0]['text']['format']['strict'] is True
    assert 'answer_key' not in json.dumps(seen).lower()


@pytest.mark.parametrize('bad',['unknown_id','invented_number','causation','fake_quote','forced_status'])
def test_ai_cannot_add_evidence_numbers_causation_or_claim_outcomes(monkeypatch,bad):
    r=analyze(source('C_Northgate.xlsx'));p=r.package;a=build(p,r);c=a['findings'][0]
    item=dict(candidate_id=c['id'],why_it_matters='This may weaken recurring demand.',management_question='Which customer cohorts explain this movement?',fact_ids=c['refs'][:1])
    out=dict(metrics=[],claims=[],interpretations=[item])
    if bad=='unknown_id':item['fact_ids']=['made_up']
    if bad=='invented_number':item['why_it_matters']='This may cost $987M.'
    if bad=='causation':item['why_it_matters']='This proves customer distress and caused lower sales.'
    if bad=='fake_quote':out['claims']=[dict(quote_id='missing',quote='record',series_ids=[],comparison='record_high',theme='retention')]
    if bad=='forced_status':out['claims']=[dict(quote_id=p['narratives'][0]['id'],quote=p['narratives'][0]['text'],series_ids=[],comparison='unsupported',theme='unknown',status='supported')]
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    result=interpret(p,a,a['findings'],r,lambda req:response(out))
    assert not result['insights'] and not result['briefing'] and not result['claims']
    section='claims' if bad in {'fake_quote','forced_status'} else 'interpretations'
    errors=[d for d in result['diagnostics']['items'] if d['section']==section and d['status']=='rejected']
    assert errors and errors[0]['reason']
    if bad not in {'fake_quote','forced_status'}:assert not result['interpretations']


def test_original_demo_keeps_its_guarded_financial_results_and_exports():
    r=analyze(demo_files())
    assert r.metrics['revenue']==5383752 and r.metrics['billing_total']==27978.75
    assert {f['id'] for f in r.findings}=={1,2,3}
    assert export_report(r)[0]


def test_literal_percent_csv_supports_schedule_only_upload_without_invented_financials():
    raw=b'Example Manufacturing,,,,,,,\nMeasure,Unit,Jan-27,Feb-27,Mar-27,Apr-27,May-27,Jun-27\nRepeat customer retention,%,95,93,91,88,85,82\n'
    r=analyze({'operations.csv':raw})
    assert not r.metrics
    st=next(s for s in r.package['stats'] if s['theme']=='retention')
    assert st['latest']['value']==.82
    assert any(c['theme']=='retention' for c in present(r)['cards'])
    assert '85,000' not in briefing(present(r))


def test_ambiguous_fractional_percent_without_excel_format_is_not_guessed():
    p=read_package({'rates.csv':b'Company,,,,,\nMetric,Unit,Jan-27,Feb-27,Mar-27,Apr-27\nRetention,%,0.95,0.9,0.85,0.8\n'})
    assert p['series'][0]['unit']=='unknown'
    assert any('ambiguous' in note for note in p['limitations'])


def test_inventory_aging_is_not_mislabeled_as_customer_credit_risk():
    w=load_workbook(FIXTURES/'C_Northgate.xlsx');w['Customer balances']['A2']='Inventory age as of June 30, 2027'
    assert not read_package({'inventory.xlsx':encoded(w)})['aging']


def test_valid_ai_implication_cites_observed_candidate_and_stays_qualified(monkeypatch):
    r=analyze(source('C_Northgate.xlsx'));p=r.package;a=build(p,r);c=a['findings'][0]
    output=dict(metrics=[],claims=[],interpretations=[dict(candidate_id=c['id'],why_it_matters='This may weaken recurring customer demand and planning visibility.',management_question='Which customer cohorts explain this movement, and what retention actions are planned?',fact_ids=c['refs'][:2])])
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    semantic=interpret(p,a,a['findings'],r,lambda req:response(output))
    assert semantic['interpretations'][c['id']]['why_it_matters'].startswith('This may')


def test_failed_ai_call_preserves_useful_source_analysis(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    r=analyze(source('C_Northgate.xlsx'))
    assert r.metrics['ebitda']==85000 and len(present(r)['cards'])==3
    assert 'unavailable' in r.package['mode']
    assert export_report(r)[0]


def test_summit_model_packet_is_bounded_and_source_referenced(monkeypatch):
    r=analyze(source());a=build(r.package,r);seen=[]
    def request(req):seen.append(json.loads(req.data));return response(dict(metrics=[],claims=[],interpretations=[]))
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    interpret(r.package,a,a['findings'],r,request)
    packet=json.loads(seen[0]['input'][1]['content'])
    assert len(seen[0]['input'][1]['content'])<180000
    assert packet['facts'] and all(f['id'] and f['ref'] for f in packet['facts'])
    assert any(s['label']=='Renewal rate' for s in packet['series'])
    assert any('record 13,661' in q['text'] for q in packet['commentary'])


@pytest.mark.parametrize('company,values,positive',[
    ('Example Software',[72,76,81,87],True),
    ('Example Distribution',[93,87,79,70],False),
    ('Example Services',[80,80,80,80],None),
    ('Example Manufacturing',[80,90,70,85],None),
])
def test_generic_trend_direction_changes_with_values_not_company(company,values,positive):
    raw=(company+',,,,,\nMetric,Unit,Jan-28,Feb-28,Mar-28,Apr-28\nCustomer retention,%, '+','.join(map(str,values))+'\n').encode()
    r=analyze({'arbitrary.csv':raw})
    findings=[f for f in r.package['findings'] if f['theme']=='retention']
    if positive is None:assert findings==[]
    else:
        assert len(findings)==1 and findings[0]['positive'] is positive
        assert findings[0]['indicator']==f'{values[-1]-values[0]:+.1f} pp'


def test_unknown_business_measure_does_not_inherit_a_known_risk():
    raw=b'Another Industry,,,,,\nMetric,Unit,Jan-28,Feb-28,Mar-28,Apr-28\nZeta operating index,count,110,90,70,50\n'
    r=analyze({'unseen-label.csv':raw})
    assert r.package['stats'][0]['theme']=='unknown'
    assert r.package['findings']==[]
    assert 'retention' not in briefing(present(r)).lower()


def test_dimensional_budget_is_not_discarded_by_section_caption():
    from openpyxl import Workbook
    w=Workbook();s=w.active;s.title='Arbitrary schedule'
    for row in [['Example Logistics'],['USD whole units; accrual basis'],['April 2028'],['','East',None,'West',None],['Measure','Actual','Budget','Actual','Budget'],['Operating KPIs'],['Contribution',17000,20000,25000,24000]]:s.append(row)
    p=read_package({'dimensions.xlsx':encoded(w)})
    found=[f for f in p['units'] if f['label']=='Contribution']
    assert {(f['entity'],f['scenario'],f['value']) for f in found}=={('East','actual',17000),('East','budget',20000),('West','actual',25000),('West','budget',24000)}


def test_aging_ledger_mismatch_remains_visible_in_executive_card():
    w=load_workbook(FIXTURES/'C_Northgate.xlsx');w['Customer balances']['G8']=210000
    r=analyze({'ledger-mismatch.xlsx':encoded(w)})
    assert r.package['aging'][0]['gl_tie'] is False
    card=next(c for c in present(r)['cards']+present(r)['secondary'] if 'aging_' in next((f['key'] for f in r.package['findings'] if f['id']==c['id']),'') )
    assert 'does not reconcile' in card['summary'] and 'provisional' in card['summary']


def test_improving_collection_trend_is_not_described_as_increasing():
    w=load_workbook(FIXTURES/'C_Northgate.xlsx');s=w['Customer balances']
    s.append(['Metric','Unit','Jan-27','Feb-27','Mar-27','Apr-27','May-27','Jun-27'])
    s.append(['Days sales outstanding','days',60,55,50,45,40,35])
    r=analyze({'improving-collections.xlsx':encoded(w)})
    aging=next(f for f in r.package['findings'] if f['key'].startswith('aging_'))
    assert 'increased for' not in aging['summary']
    assert any(f['theme']=='collections' and f['positive'] for f in r.package['findings'])
