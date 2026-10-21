"""Portable, zero-provider import and actual dashboard/export reuse of approved results."""
import copy
import io
import json
import pytest
from openpyxl import load_workbook
from engine import ROOT, InputError
from scripts import seed_demo_results as prepared


@pytest.fixture(scope='module')
def approved():
    return prepared.verify_prepared()


def configure(monkeypatch, approved):
    for key, value in prepared.DEMO_SETTINGS.items():
        monkeypatch.setenv(key, value)
    monkeypatch.setenv('OPENAI_MODEL', approved[0][0]['model'])
    monkeypatch.setenv('OPENAI_API_KEY', 'offline-demo-cache-test-scope')
    for key in ('OPENAI_INPUT_USD_PER_M', 'OPENAI_OUTPUT_USD_PER_M', 'COPILOT_AI_MAX_RETRIES', 'COPILOT_SKILL_ROUNDS', 'COPILOT_SKILL_TOOL_CALLS'):
        monkeypatch.delenv(key, raising=False)


def test_verified_snapshot_sources_are_portable_and_complete(approved):
    assert len(approved) == 3
    for entry, review, files in approved:
        assert review.package['skill_engine']['status'] == 'completed'
        assert review.package['inference']['inference_observed']
        assert not review.package['inference'].get('reused')
        assert len(files) == 1
        assert '/Users/' not in json.dumps(prepared.unpack(json.loads(
            __import__('gzip').decompress((ROOT/'demo_results'/entry['snapshot']).read_bytes()))), default=str)


def test_import_requires_separately_configured_credential_scope(monkeypatch, approved):
    configure(monkeypatch, approved)
    monkeypatch.delenv('OPENAI_API_KEY')
    with pytest.raises(InputError, match='local API key'):
        prepared.seed(copy.deepcopy(approved))


def test_import_rejects_changed_model_configuration(monkeypatch, approved):
    configure(monkeypatch, approved)
    monkeypatch.setenv('OPENAI_MODEL', 'different-model')
    with pytest.raises(InputError, match='configuration differs: model'):
        prepared.seed(copy.deepcopy(approved))


@pytest.mark.parametrize('company', range(3))
def test_import_dashboard_evidence_and_both_exports_offline(monkeypatch, tmp_path, approved, company):
    from streamlit.testing.v1 import AppTest
    configure(monkeypatch, approved)
    results = prepared.seed(copy.deepcopy(approved))
    assert all(r['new_model_calls'] == 0 for r in results)
    entry, _, files = approved[company]
    at = AppTest.from_file(str(ROOT/'app.py'), default_timeout=60).run()
    assert not at.exception
    next(x for x in at.selectbox if x.label == 'Company workbook').set_value(next(iter(files))).run()
    next(b for b in at.button if b.label == 'Load company workbook').click().run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception
    review = at.session_state['review']
    receipt = review.package['inference']
    assert review.package['prepared_demo']['run'] == entry['run']
    assert review.package['skill_engine']['status'] == 'completed'
    assert receipt['reused'] and receipt['attempts'] == 0 and receipt['estimated_usd'] == 0
    assert receipt['origin']['inference_observed']
    from executive import present
    dashboard = '\n'.join(str(x.value) for x in at.markdown)
    assert all(c['title'] in dashboard for c in present(review)['cards'])
    outputs = [('standard', at.session_state['export'][0])]
    next(b for b in at.button if b.label.startswith('View Evidence')).click().run()
    assert not at.exception
    at.radio(key='view').set_value('01  Upload & Setup').run()
    at.radio(key='report_format').set_value('Use existing firm template').run()
    next(c for c in at.checkbox if c.label == 'Use supplied sample template').check().run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception
    assert at.session_state['review'].package['inference']['attempts'] == 0
    outputs.append(('custom', at.session_state['export'][0]))
    for name, data in outputs:
        (tmp_path/(entry['run']+'-'+name+'.xlsx')).write_bytes(data)
        book = load_workbook(io.BytesIO(data))
        assert {'Findings', 'Source Evidence', 'Analyst Execution'} <= set(book.sheetnames)
        assert any(c.hyperlink for row in book['Findings'] for c in row)
        assert not any(c.data_type == 'e' for sheet in book for row in sheet for c in row)
        if name == 'custom':
            assert book['Partner Review']['C16'].data_type == 'f'
