import io
import re
from dataclasses import replace
import pytest
from openpyxl import load_workbook
from portfolio_data import load_companies
from portfolio_excel import export_portfolio, SHEETS
from portfolio_comparison import comparison
from portfolio_watchlist import watchlist


@pytest.fixture(scope='module')
def exported():
    companies = load_companies()
    return companies, load_workbook(io.BytesIO(export_portfolio(companies)))


def test_workbook_matches_dashboard_and_has_no_consolidation(exported):
    companies, book = exported
    assert tuple(book.sheetnames) == SHEETS
    overview = book[SHEETS[0]]
    for row, company in enumerate(companies, 6):
        view = comparison(company)
        assert overview.cell(row, 1).value == company.name
        assert overview.cell(row, 2).value == view['period_label']
        assert overview.cell(row, 3).value == view['revenue'].actual
        assert overview.cell(row, 4).value == view['ebitda'].actual
        assert overview.cell(row, 5).value == pytest.approx(view['margin'].actual)
        assert overview.cell(row, 7).value == pytest.approx(view['ebitda'].variance_pct)
    assert 'No combined revenue' in overview['A3'].value
    for sheet in book:
        assert sheet.freeze_panes
        assert sheet.auto_filter.ref
        assert sheet.print_area
        assert sheet['A5'].fill.fgColor.rgb.endswith('18364A')
        assert not sheet.sheet_view.showGridLines
        assert not any(c.data_type in ('e', 'f') for row in sheet for c in row)


def test_original_narratives_and_company_qualified_source_links(exported):
    companies, book = exported
    values = {str(c.value) for row in book['Portfolio Watchlist'] for c in row if c.value is not None}
    for item in watchlist(companies):
        for key in ('title', 'conclusion', 'evidence', 'financial_implication', 'assessment', 'action', 'management_question'):
            assert item.finding[key] in values
        assert item.presentation['caveat'] in values
        assert all(q['quote'] in values for q in item.finding['source_quotes'])
    source_pairs = {(row[0].value, row[1].value) for row in book['Source References'].iter_rows(min_row=6)}
    for item in watchlist(companies):
        assert {(item.company_name, sid) for sid in item.source_ids} <= source_pairs
    assert len([p for p in source_pairs if p[1] == 'calc_2']) == 3
    links = []
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.hyperlink:
                    target = cell.hyperlink.target
                    match = re.fullmatch(r"#'([^']+)'!A(\d+)", target)
                    assert match, target
                    dest, row_num = match.groups()
                    assert book[dest].cell(int(row_num), 1).value is not None
                    links.append(target)
    assert len(links) >= 50


def test_source_text_cannot_become_an_excel_formula():
    companies = list(load_companies())
    companies[0] = replace(companies[0], name='=HYPERLINK("https://invalid", "test")')
    book = load_workbook(io.BytesIO(export_portfolio(companies)))
    assert book['Portfolio Overview']['A6'].data_type == 's'
