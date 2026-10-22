"""Download-route regression for the actual local Streamlit reports (no inference)."""
import io
import json
import os
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/financial-reporting/final-downloads';OUT.mkdir(parents=True,exist_ok=True)
TEMPLATE=ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx'
results=[]

def checked(download,name,custom):
 path=OUT/name;download.save_as(path);book=load_workbook(path);sheet=book['Financial Performance']
 assert 'TALLIS BROOK' in sheet['A3'].value
 assert sheet['B7'].value==3475418 and sheet['C7'].value==3651832
 assert sheet['B12'].value==684306 and sheet['C12'].value==946225
 assert sheet.max_column==4 and book.active.title=='Financial Performance'
 assert sheet['B11'].value=='N/A' and sheet['B17'].value=='N/A'
 assert 'money' not in str(list(sheet.values)) and 'ratio' not in str(list(sheet.values))
 assert 'Statement Evidence' in book and 'Extracted Detail' in book
 if custom: assert book['Partner Review']['C16'].data_type=='f'
 results.append({'route':name,'company':'tallis','revenue':sheet['B7'].value,'reported_ebitda':sheet['B12'].value,'active_sheet':book.active.title,'bytes':path.stat().st_size})

with sync_playwright() as p:
 browser=p.chromium.launch(channel='chrome',headless=True)
 page=browser.new_page(viewport={'width':1440,'height':1100},accept_downloads=True)
 def only_local(route):
  assert urlparse(route.request.url).hostname in ('127.0.0.1','localhost')
  route.continue_()
 page.route('**/*',only_local)
 page.goto('http://127.0.0.1:8502/')
 with page.expect_download() as d:page.get_by_role('button',name='Download Portfolio Excel',exact=True).click(timeout=60000)
 path=OUT/'portfolio-browser.xlsx';d.value.save_as(path);book=load_workbook(path)
 assert book['Portfolio Financial Performance']['E7'].value==3475418
 assert book['Portfolio Financial Performance']['H7'].value==46650000
 assert len(book.sheetnames)==35
 page.locator('.st-key-open_tallis button').click()
 expect(page.get_by_role('heading',name='Tallis Brook Accounting',exact=True)).to_be_visible(timeout=60000)
 page.wait_for_timeout(1000)  # Finish the native desktop-viewport probe rerun.
 expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state','notRunning',timeout=60000)
 for custom in (False,True):
  if custom:
   page.get_by_role('combobox',name='Company Excel format',exact=True).click()
   page.get_by_role('option',name='Use supplied firm template',exact=True).click()
  if custom:expect(page.get_by_text('Optional company Excel template',exact=True)).to_be_visible(timeout=60000)
  expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state','notRunning',timeout=60000)
  with page.expect_download() as d:page.get_by_role('button',name='Download Investment Excel',exact=True).click()
  checked(d.value,'tallis-custom.xlsx' if custom else 'tallis-standard.xlsx',custom)
 page.locator('input[type=file]').set_input_files(str(TEMPLATE))
 expect(page.get_by_role('button',name='Remove '+TEMPLATE.name,exact=True)).to_be_visible(timeout=30000)
 expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state','notRunning',timeout=60000)
 with page.expect_download() as d:page.get_by_role('button',name='Download Investment Excel',exact=True).click()
 checked(d.value,'tallis-uploaded-company.xlsx',True)
 page.get_by_role('button',name='Documents & Reports',exact=True).click()
 page.get_by_role('combobox',name='Company report',exact=True).click()
 page.get_by_role('option',name='Tallis Brook Accounting',exact=True).click()
 page.get_by_role('combobox',name='Company Excel format',exact=True).click()
 page.get_by_role('option',name='Use supplied firm template',exact=True).click()
 expect(page.get_by_text('Optional company Excel template',exact=True)).to_be_visible(timeout=60000)
 page.locator('input[type=file]').set_input_files(str(TEMPLATE))
 expect(page.get_by_role('button',name='Remove '+TEMPLATE.name,exact=True)).to_be_visible(timeout=30000)
 expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state','notRunning',timeout=60000)
 with page.expect_download() as d:page.get_by_role('button',name='Download Investment Excel',exact=True).click()
 checked(d.value,'tallis-uploaded-documents.xlsx',True)
 browser.close()
(OUT/'receipt.json').write_text(json.dumps(results,indent=2)+'\n')
print(json.dumps(results,indent=2))
