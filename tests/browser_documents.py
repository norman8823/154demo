"""Actual Chrome: original downloads, click/keyboard/formula inspection and citations."""
import hashlib
import json
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple,get_column_letter
from playwright.sync_api import sync_playwright,expect

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/interactive-source';OUT.mkdir(parents=True,exist_ok=True)
IMAGES=ROOT/'docs/images/interactive-source';IMAGES.mkdir(parents=True,exist_ok=True)
MANIFEST=json.loads((ROOT/'demo_results/manifest.json').read_text())
TARGETS=[('summit','Summit Home Services','Income Statement','D15'),('tallis','Tallis Brook Accounting','Practice P&L','G12'),('harrow','Harrow Ridge Environmental','P&L Summary','F13')]
receipt={'originals':{},'interaction':[],'citations':[],'mobile':[],'errors':[],'external_requests':[]}


def ready(page):
 expect(page.locator('[data-testid="stApp"]')).to_have_attribute('data-test-script-state','notRunning',timeout=60000)
 expect(page.locator('[data-testid="stException"]')).to_have_count(0)
 page.wait_for_timeout(150)


def choose(page,scope,label,value):
 field=scope.get_by_role('combobox',name=label,exact=True)
 field.evaluate('(e)=>e.scrollIntoView({block:"center"})')
 field.fill(value);field.press('ArrowDown')
 page.get_by_role('option',name=value,exact=True).click()


def close(page):
 page.get_by_role('dialog').get_by_role('button',name='Close',exact=True).last.click()
 expect(page.get_by_role('dialog')).to_have_count(0,timeout=60000)


def assert_cell(page,sheet,cell):
 dialog=page.get_by_role('dialog')
 expect(dialog.locator('.source-target')).to_have_attribute('data-coordinate',cell,timeout=60000)
 expect(dialog.get_by_role('textbox',name='Selected cell',exact=True)).to_have_value(cell)
 original=sheet[cell].value
 raw='' if original is None else ('TRUE' if original is True else ('FALSE' if original is False else str(original)))
 expect(dialog.locator('.formula-bar')).to_have_text(raw)
 expect(dialog.locator('.source-target')).to_have_attribute('aria-selected','true')
 assert dialog.locator('.source-target').count()==1
 assert dialog.locator('.source-grid tbody tr').count()<=12
 assert dialog.locator('.source-grid').evaluate('(e)=>e.scrollWidth<=e.clientWidth+1')
 assert not page.evaluate('document.documentElement.scrollWidth>innerWidth+1')
 return dialog


def jump(page,sheet,cell):
 field=page.get_by_role('dialog').get_by_role('textbox',name='Selected cell',exact=True)
 field.fill(cell);field.press('Enter');return assert_cell(page,sheet,cell)


with sync_playwright() as p:
 browser=p.chromium.launch(channel='chrome',headless=True)
 page=browser.new_page(viewport={'width':1440,'height':1250},accept_downloads=True)
 page.on('pageerror',lambda e:receipt['errors'].append(str(e)))
 def local(route):
  if urlparse(route.request.url).hostname not in ('127.0.0.1','localhost'):
   receipt['external_requests'].append(route.request.url);route.abort()
  else:route.continue_()
 page.route('**/*',local)
 page.goto('http://127.0.0.1:8502/?page=excel')
 expect(page.get_by_role('heading',name='Source Documents',exact=True)).to_be_visible(timeout=60000)
 books={}
 for slug,company,sheet_name,cell in TARGETS:
  card=page.locator('.st-key-source_document_'+slug)
  entry=next(r for r in MANIFEST['runs'] if r['run']==slug+('-17-final' if slug=='harrow' else '-05-final'))
  with page.expect_download() as download:card.get_by_role('button',name='Download Original',exact=True).click()
  path=OUT/(slug+'-original.xlsx');download.value.save_as(path)
  assert download.value.suggested_filename==Path(entry['workbook']).name
  assert hashlib.sha256(path.read_bytes()).hexdigest()==entry['workbook_sha256']
  books[slug]=load_workbook(path,data_only=False,keep_links=False)
  card.get_by_role('button',name='Preview',exact=True).click()
  dialog=page.get_by_role('dialog')
  expect(dialog.get_by_role('heading',name=company,exact=True)).to_be_visible(timeout=60000)
  assert_cell(page,books[slug][dialog.get_by_role('combobox',name='Worksheet',exact=True).input_value()],'A1')
  with page.expect_download() as download:dialog.get_by_role('button',name='Download Original',exact=True).click()
  path=OUT/(slug+'-dialog-original.xlsx');download.value.save_as(path)
  assert hashlib.sha256(path.read_bytes()).hexdigest()==entry['workbook_sha256']
  expect(dialog).to_be_visible() # Download must not dismiss / rerun the viewer.
  close(page)
  receipt['originals'][slug]={'sha256':entry['workbook_sha256'],'library_and_dialog_download':'identical approved bytes'}
 # Exact cited opening, no manual sheet selection / coordinate entry.
 page.goto('http://127.0.0.1:8502/?page=foundation')
 expect(page.locator('.st-key-foundation_source_view button')).to_be_visible(timeout=60000)
 ready(page)
 for slug,company,sheet_name,cell in TARGETS:
  sheet=books[slug][sheet_name]
  choose(page,page,'Company',company)
  expect(page.locator('.foundation-detail-grid')).to_contain_text(sheet_name,timeout=60000)
  ready(page)
  page.locator('.st-key-foundation_source_view button').click()
  dialog=assert_cell(page,sheet,cell)
  expect(dialog.get_by_role('combobox',name='Worksheet',exact=True)).to_have_value(sheet_name)
  dialog.locator('.excel-viewer').evaluate('(e)=>e.dataset.testIdentity="same-mount"')
  r,c=coordinate_to_tuple(cell)
  clicked=get_column_letter(c)+str(r-1)
  dialog.locator('[role="gridcell"][data-coordinate="'+clicked+'"]').click();assert_cell(page,sheet,clicked)
  # Four directions; no backend round trip or formula execution.
  for key,rr,cc in [('ArrowUp',r-2,c),('ArrowRight',r-2,c+1),('ArrowDown',r-1,c+1),('ArrowLeft',r-1,c)]:
   page.keyboard.press(key);assert_cell(page,sheet,get_column_letter(cc)+str(rr))
  assert dialog.locator('.excel-viewer').get_attribute('data-test-identity')=='same-mount'
  old=dialog.get_by_role('textbox',name='Selected cell',exact=True).input_value()
  dialog.locator('.formula-bar').focus();page.keyboard.press('ArrowLeft')
  expect(dialog.get_by_role('textbox',name='Selected cell',exact=True)).to_have_value(old)
  dialog.get_by_role('combobox',name='Worksheet',exact=True).focus();page.keyboard.press('ArrowLeft')
  expect(dialog.get_by_role('textbox',name='Selected cell',exact=True)).to_have_value(old)
  # Navigate beyond each visible edge; range changes while the target remains focused.
  corners=dialog.locator('[role="gridcell"]').evaluate_all('(es)=>[es[0].dataset.coordinate,es[es.length-1].dataset.coordinate]')
  end_row,end_col=coordinate_to_tuple(corners[1]);before=dialog.locator('.grid-range').inner_text()
  dialog.locator('[role="gridcell"][data-coordinate="'+corners[1]+'"]').click()
  page.keyboard.press('ArrowDown');assert_cell(page,sheet,get_column_letter(end_col)+str(end_row+1))
  page.keyboard.press('ArrowRight');assert_cell(page,sheet,get_column_letter(end_col+1)+str(end_row+1))
  assert before!=dialog.locator('.grid-range').inner_text()
  first=dialog.locator('[role="gridcell"]').first.get_attribute('data-coordinate');fr,fc=coordinate_to_tuple(first)
  dialog.locator('[role="gridcell"]').first.click()
  page.keyboard.press('ArrowUp');assert_cell(page,sheet,get_column_letter(fc)+str(fr-1))
  page.keyboard.press('ArrowLeft');assert_cell(page,sheet,get_column_letter(fc-1)+str(fr-1))
  # Worksheet bounds, blank cells and safe invalid coordinate handling.
  jump(page,sheet,'A1');page.keyboard.press('ArrowUp');page.keyboard.press('ArrowLeft');assert_cell(page,sheet,'A1')
  last=get_column_letter(sheet.max_column)+str(sheet.max_row)
  jump(page,sheet,last);page.keyboard.press('ArrowDown');page.keyboard.press('ArrowRight');assert_cell(page,sheet,last)
  field=dialog.get_by_role('textbox',name='Selected cell',exact=True);field.fill('XFD999999');field.press('Enter')
  expect(dialog.locator('.grid-status')).to_contain_text('outside the original worksheet')
  assert_cell(page,sheet,last)
  jump(page,sheet,cell)
  dialog.locator('.source-details summary').click()
  expect(dialog.locator('.source-details')).to_contain_text('Verified standardized value')
  assert sheet[cell].value not in dialog.locator('.source-details').inner_text() # Formula is shown once, above grid.
  page.screenshot(path=str(IMAGES/(slug+'-interactive-desktop.png')),full_page=True)
  dialog.locator('.source-details summary').click()
  # Switch to a one-column narrative source and back; selection always valid.
  other=books[slug].sheetnames[0]
  choose(page,dialog,'Worksheet',other)
  expect(dialog.locator('.grid-range')).to_contain_text(other+' ·',timeout=60000)
  assert_cell(page,books[slug][other],'A1');dialog.locator('.source-target').click();page.keyboard.press('ArrowRight');assert_cell(page,books[slug][other],'A1')
  choose(page,dialog,'Worksheet',sheet_name)
  expect(dialog.locator('.grid-range')).to_contain_text(sheet_name+' ·',timeout=60000)
  assert_cell(page,sheet,cell)
  dialog.locator('.source-target').click();page.keyboard.press('ArrowDown')
  page.keyboard.press('Escape');expect(page.get_by_role('dialog')).to_have_count(0)
  page.locator('.st-key-foundation_source_view button').click();assert_cell(page,sheet,cell);close(page)
  receipt['interaction'].append({'company':slug,'exact_citation':sheet_name+'!'+cell,'click_four_arrows':'pass','immediate_formula_bar':'pass; same component mount','all_four_visible_edges':'pass','worksheet_bounds':'pass','invalid_coordinate':'safe; valid selection unchanged','worksheet_switching':'pass','outside_grid_arrows':'not intercepted','Escape_reopen':'original citation restored'})
  print('interactive PASS',slug,flush=True)
 # Original operand links in both accepted investment workflows.
 for route,key in [('issues','issue_source_summit_'),('company&company=summit&finding=10000','company_source_summit_')]:
  page.goto('http://127.0.0.1:8502/?page='+route)
  button=page.locator('[class*="st-key-'+key+'"][class*="_view"] button').first;expect(button).to_be_visible(timeout=60000)
  selection=page.locator('[class*="st-key-'+key+'"][class*="_operand"]').first.get_by_role('combobox',name='Original source cell',exact=True).input_value()
  button.click();dialog=page.get_by_role('dialog')
  expect(dialog.locator('.source-target')).to_have_count(1,timeout=60000)
  selected=dialog.get_by_role('textbox',name='Selected cell',exact=True).input_value();sheet_name=dialog.get_by_role('combobox',name='Worksheet',exact=True).input_value()
  assert selected in selection and sheet_name in selection
  assert_cell(page,books['summit'][sheet_name],selected);close(page)
  receipt['citations'].append({'route':route,'recorded_operand':sheet_name+'!'+selected})
 # Phone uses a smaller bounded grid, with the same original cells and controls.
 page.set_viewport_size({'width':390,'height':1000})
 page.goto('http://127.0.0.1:8502/?page=foundation')
 expect(page.locator('.st-key-foundation_source_view button')).to_be_visible(timeout=60000)
 sidebar=page.locator('[data-testid="stSidebar"]')
 if sidebar.get_attribute('data-collapsed')=='false':page.get_by_role('button',name='Collapse sidebar',exact=True).click()
 for slug,company,sheet_name,cell in TARGETS:
  choose(page,page,'Company',company)
  expect(page.locator('.foundation-detail-grid')).to_contain_text(sheet_name,timeout=60000)
  ready(page)
  page.locator('.st-key-foundation_source_view button').click();dialog=assert_cell(page,books[slug][sheet_name],cell)
  expect(dialog.locator('.source-grid tbody tr')).to_have_count(8)
  assert dialog.locator('.source-grid tbody tr').first.locator('[role="gridcell"]').count()<=2
  dialog.locator('.source-target').click();page.keyboard.press('ArrowRight')
  r,c=coordinate_to_tuple(cell);assert_cell(page,books[slug][sheet_name],get_column_letter(c+1)+str(r))
  jump(page,books[slug][sheet_name],cell)
  if slug=='summit':page.screenshot(path=str(IMAGES/'summit-interactive-mobile.png'),full_page=True)
  close(page);receipt['mobile'].append({'company':slug,'width':390,'bounded_grid':'8 rows × 2 columns','selection_and_formula_bar':'pass','overflow':False})
 assert not receipt['errors'] and not receipt['external_requests'],receipt
 (OUT/'browser-documents.json').write_text(json.dumps(receipt,indent=2));browser.close();print(json.dumps(receipt,indent=2))
