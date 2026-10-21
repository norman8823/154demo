import io
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from engine import ROOT


def test_company_selector_updates_all_data_and_clears_finding_and_export_state():
    at = AppTest.from_file(str(ROOT/'portfolio_app.py'), default_timeout=60).run()
    at.button(key='open_summit').click().run()
    at.selectbox(key='gp_company_format').set_value('Use supplied firm template').run()
    next(b for b in at.button if b.label.startswith('View Evidence')).click().run()
    at.radio(key='selected').set_value(10001).run()
    for slug, name, revenue in [('tallis','Tallis Brook Accounting',3475418),('harrow','Harrow Ridge Environmental',46650000),('summit','Summit Home Services',3298000)]:
        at.selectbox(key='gp_company_selector').set_value(slug).run()
        assert not at.exception
        assert at.title[0].value == name
        assert at.radio(key='gp_company_tab').value == 'summary'
        assert at.selectbox(key='gp_company_format').value == 'Standard Excel report'
        assert 'finding' not in at.query_params
        book=load_workbook(io.BytesIO(at.session_state['export'][0]))
        assert any(row[1] == revenue for row in book['Financial Performance'].values)
        assert 'Partner Review' not in book.sheetnames
    at.selectbox(key='gp_sidebar_company').set_value('tallis').run()
    assert not at.exception
    assert at.title[0].value == 'Tallis Brook Accounting'


def test_issues_are_grouped_with_original_questions_and_actions():
    at = AppTest.from_file(str(ROOT/'portfolio_app.py'), default_timeout=60).run()
    at.button(key='nav_watchlist').click().run()
    assert not at.exception
    assert at.title[0].value == 'Investment Issues'
    assert [s.value for s in at.subheader][-3:] == ['Summit Home Services','Tallis Brook Accounting','Harrow Ridge Environmental']
    assert len([b for b in at.button if b.label=='View source evidence →']) == 17
    assert len([m for m in at.markdown if m.value=='**Management question**']) == 17
    assert len([m for m in at.markdown if m.value=='**Proposed follow-up**']) == 17
