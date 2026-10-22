import io
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from engine import ROOT


@pytest.mark.parametrize('slug', ['summit', 'tallis', 'harrow'])
def test_all_company_summary_evidence_exports_and_return(slug, monkeypatch):
    def reject(*a, **k):
        raise AssertionError('Company browsing attempted analysis')
    monkeypatch.setattr('engine.analyze', reject)
    at = AppTest.from_file(str(ROOT / 'portfolio_app.py'), default_timeout=60).run()
    at.button(key='open_'+slug).click().run()
    assert not at.exception
    assert any(s.value == 'Investment Summary' for s in at.subheader)
    assert any(s.value == 'Priority findings' for s in at.subheader)
    assert not any(s.value in ('Management Questions','See where every conclusion came from') for s in at.subheader)
    assert not any(r.key=='gp_company_tab' for r in at.radio)
    assert not any(b.label in ('View source evidence','View Evidence  →','Open source evidence  →') for b in at.button)
    assert not any(e.label=='Derived calculations & original operands' for e in at.expander)
    for choice in ['Standard Excel report', 'Use supplied firm template']:
        at.selectbox(key='gp_company_format').set_value(choice).run()
        assert not at.exception
        data = at.session_state['export'][0]
        book = load_workbook(io.BytesIO(data))
        assert 'Source Evidence' in book.sheetnames
        if choice.startswith('Use'):
            assert book['Partner Review']['C16'].data_type == 'f'
    next(b for b in at.button if b.label == '← Portfolio Overview').click().run()
    assert not at.exception
    assert at.title[0].value == 'Portfolio Overview'


def test_watchlist_filter_and_link_to_exact_finding():
    at = AppTest.from_file(str(ROOT / 'portfolio_app.py'), default_timeout=60).run()
    at.button(key='nav_watchlist').click().run()
    next(s for s in at.selectbox if s.label == 'Company').set_value('Harrow Ridge Environmental').run()
    next(s for s in at.selectbox if s.label == 'Original priority').set_value('high').run()
    buttons = [b for b in at.button if b.label == 'View source evidence →']
    assert len(buttons) == 5
    at.button(key='watch_harrow_10003').click().run()
    assert not at.exception
    assert at.session_state['selected'] == 10003
    assert at.query_params['finding']=='10003'
    assert any(s.value=='Priority findings' for s in at.subheader)
    assert not any(r.key=='gp_company_tab' for r in at.radio)
    assert not any(s.value=='See where every conclusion came from' for s in at.subheader)



def test_financial_watchlist_prose_is_literal_currency():
    at = AppTest.from_file(str(ROOT / 'portfolio_app.py'), default_timeout=60).run()
    at.button(key='nav_watchlist').click().run()
    text = '\n'.join(m.value for m in at.markdown)
    assert '&#36;3,916,548 against budget of &#36;5,311,831' in text
    assert '$3,916,548 against budget of $5,311,831' not in text
