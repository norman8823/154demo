"""Synthetic, source-grounded model contract tests. Never submits a paid request."""
import io
import json
import urllib.error
import pytest
from openpyxl import Workbook,load_workbook
from engine import analyze,InputError
from ai_runtime import complete,inference_session
from investment_output import validate,numeric_claim
from executive import present,briefing
from reports import export_report

@pytest.fixture(autouse=True)
def config(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','offline-dummy')
    for key in ['OPENAI_MODEL','OPENAI_INPUT_USD_PER_M','OPENAI_OUTPUT_USD_PER_M','COPILOT_AI_MAX_REVIEW_USD','COPILOT_AI_MAX_RETRIES','COPILOT_AI_TIMEOUT_SECONDS']:monkeypatch.delenv(key,raising=False)


def files():
    w=Workbook();s=w.active;s.title='Economics'
    for row in [['Example Group'],['USD whole units; accrual basis'],['Account','September 2028 Actual','September 2028 Budget','August 2028 Actual'],['Revenue',1000,1200,900],['EBITDA',100,180,140]]:s.append(row)
    s=w.create_sheet('Delivery scorecard')
    for row in [['Example Group'],['September 2028'],['Metric','Company total','North','South'],['Completion rate',.86,.90,.82]]:s.append(row)
    for c in s[4][1:]:c.number_format='0.0%'
    s=w.create_sheet('Management letter')
    for row in [['Example Group'],['September 2028'],['Companywide completion remained above 95% throughout the reporting month, and management considers delivery performance healthy.']]:s.append(row)
    out=io.BytesIO();w.save(out);return {'synthetic.xlsx':out.getvalue()}


def response(result,usage=True):
    out={'status':'completed','model':'gpt-4o-mini','output':[{'content':[{'type':'output_text','text':json.dumps(result)}]}]}
    if usage:out['usage']={'input_tokens':2000,'output_tokens':500}
    return out


def model_result(packet):
    kpi=next(f for f in packet['facts'] if f['label']=='Completion rate' and f['entity']=='Company total')
    revenue=next(f for f in packet['facts'] if f['id'].startswith('financial_') and f['label']=='Revenue' and f['period']=='2028-09')
    quote=next(q for q in packet['commentary'] if q['sheet']=='Management letter')
    item=dict(title='Delivery claim conflicts with operations',theme='quality',kind='relationship',text='Reported completion is {'+kpi['id']+'}, requiring management follow-up.',why_it_matters='This may indicate delivery risk and weaken confidence in the operating explanation.',management_question='Which delivery constraints explain the completion measure, and what corrective actions are being tracked?',fact_ids=[kpi['id']],claim=dict(quote_id=quote['id'],quote='Companywide completion remained above 95%',fact_id=kpi['id'],operator='gt',threshold='95%'))
    brief=[dict(text='The package raises execution risk. Revenue of {'+revenue['id']+'} should be assessed alongside the operating schedule, rather than treating financial results as a complete picture of performance.',fact_ids=[revenue['id']]),dict(text='Reported completion is {'+kpi['id']+'}. The matched numerical management assertion is contradicted by this source measure. Delivery constraints may affect future performance, but the package does not establish a causal earnings impact.',fact_ids=[kpi['id']]),dict(text='Management should reconcile its operating explanation, confirm that the measure covers the same population, and provide a specific recovery plan. The investment team should monitor delivery outcomes before accepting the narrative as evidence of operating stability.',fact_ids=[kpi['id']])]
    return dict(metrics=[],claims=[],interpretations=[],insights=[item],briefing=brief)


def fake_model(monkeypatch,transform=lambda x:x):
    seen=[]
    def invoke(body,phase,request_fn=None):
        def transport(req):
            packet=json.loads(json.loads(req.data)['input'][1]['content']);seen.append(packet)
            return response(transform(model_result(packet)))
        return complete(body,phase,transport)
    monkeypatch.setattr('analyst_ai.complete',invoke)
    return seen


def test_model_selects_cross_sheet_claim_and_python_verifies_the_contradiction(monkeypatch):
    seen=fake_model(monkeypatch);r=analyze(files(),ai_authorized=True)
    assert len(seen)==1
    assert r.metrics['revenue']==1000 and r.metrics['ebitda']==100
    finding=next(f for f in r.package['findings'] if f['key'].startswith('ai_'))
    assert finding['indicator']=='Contradicted' and finding['ai_classification']=='Verified numerical comparison'
    evidence=[e for e in r.evidence if e['Finding']==finding['id']]
    assert any('Delivery scorecard' in e['Reference'] for e in evidence)
    assert any('Management letter' in e['Reference'] for e in evidence)
    assert '86.0%' in briefing(present(r)) and '{grid_' not in briefing(present(r))
    assert r.package['inference']['events'][0]['transport']=='mock'
    assert not r.package['inference']['inference_observed']
    assert r.package['inference']['estimated_usd']==pytest.approx(.0006)
    assert 'answer_key' not in json.dumps(seen).lower()


@pytest.mark.parametrize('custom',[False,True])
def test_ai_briefing_and_source_cells_reach_both_excel_exports(monkeypatch,custom):
    from engine import ROOT
    fake_model(monkeypatch);r=analyze(files(),ai_authorized=True)
    template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    data,_,_=export_report(r,template);w=load_workbook(io.BytesIO(data))
    assert any('Delivery scorecard' in str(row) for row in w['Source Evidence'].values)
    assert any('Management letter' in str(row) for row in w['Source Evidence'].values)
    assert any('86.0%' in str(row) for row in w['Executive Summary' if not custom else 'Partner Review'].values)


@pytest.mark.parametrize('fault',['new_number','wrong_id','unsupported_cause','false_quote','wrong_unit'])
def test_unsupported_model_output_is_rejected_or_remains_unverified(monkeypatch,fault):
    def transform(result):
        item=result['insights'][0]
        if fault=='new_number':item['text']='Completion costs $999,000.'
        if fault=='wrong_id':item['fact_ids']=['invented']
        if fault=='unsupported_cause':item['why_it_matters']='This proves delivery failure caused the earnings miss.'
        if fault=='false_quote':item['claim']['quote']='Management guarantees a recovery.'
        if fault=='wrong_unit':item['claim']['threshold']='95';item['claim']['quote']='Companywide completion remained above 95'
        return result
    fake_model(monkeypatch,transform);r=analyze(files(),ai_authorized=True)
    if fault in {'wrong_unit'}:
        # A truncated or broader clause must not turn a unit mismatch into a verified comparison.
        claim=next(c for c in r.package['claims'] if c['id'].startswith('ai_claim_'))
        if fault=='wrong_unit':assert claim['status']=='unverifiable'
    else:
        assert not r.package['ai_briefing'] and 'Limited' in r.narrative_mode
    assert r.metrics['revenue']==1000


def test_no_external_transmission_without_per_upload_authorization(monkeypatch):
    attempted=[]
    monkeypatch.setattr('urllib.request.urlopen',lambda *a,**k:attempted.append(True))
    r=analyze(files())
    assert not attempted and not r.package['ai_briefing']
    assert r.narrative_mode.startswith('Limited offline') and r.package['inference']['attempts']==0


def test_budget_refuses_before_any_request():
    seen=[]
    with inference_session(True,limit=.000001) as s:
        with pytest.raises(InputError,match='spending limit'):complete({'max_output_tokens':100,'input':[]},'test',lambda req:seen.append(req))
    assert seen==[] and s.events==[]


def test_missing_usage_is_explicitly_unknown_and_reservation_retained():
    with inference_session(True) as s:complete({'max_output_tokens':100,'input':[]},'test',lambda req:response({},usage=False))
    assert s.summary()['estimated_usd'] is None
    assert s.events[0]['input_tokens'] is None and s.reserved>0


def test_retries_are_bounded_and_charged_against_reservation(monkeypatch):
    monkeypatch.setenv('COPILOT_AI_MAX_RETRIES','1');seen=[]
    def fail(req):seen.append(1);raise urllib.error.HTTPError(req.full_url,429,'limited',{},None)
    with inference_session(True) as s:
        with pytest.raises(InputError):complete({'max_output_tokens':100,'input':[]},'test',fail)
    assert len(seen)==2 and len(s.events)==2
    assert s.reserved==pytest.approx(sum(e['reserved_upper_usd'] for e in s.events))


def test_timeout_is_not_automatically_retried_or_logged_with_payload(caplog):
    seen=[]
    def fail(req):seen.append(1);raise TimeoutError('sensitive provider reply')
    with inference_session(True) as s:
        with pytest.raises(InputError):complete({'max_output_tokens':100,'input':'PRIVATE COMPANY FACT'},'test',fail)
    assert len(seen)==1 and s.summary()['estimated_usd'] is None
    assert 'PRIVATE COMPANY FACT' not in caplog.text and 'sensitive provider reply' not in caplog.text


def test_unknown_model_rates_fail_closed(monkeypatch):
    monkeypatch.setenv('OPENAI_MODEL','other-configured-model');seen=[]
    with inference_session(True):
        with pytest.raises(InputError,match='pricing is unknown'):complete({'max_output_tokens':100},'test',lambda req:seen.append(req))
    assert seen==[]


def test_observed_provider_usage_and_model_are_logged(monkeypatch,caplog):
    class Reply:
        def __enter__(self):return io.StringIO(json.dumps(response({})))
        def __exit__(self,*a):pass
    seen=[]
    def send(req,timeout):seen.append(timeout);return Reply()
    monkeypatch.setattr('urllib.request.urlopen',send)
    with inference_session(True) as s:complete({'max_output_tokens':100,'input':[]},'test')
    assert s.summary()['inference_observed'] and seen==[90]
    assert s.events[0]['input_tokens']==2000 and s.events[0]['output_tokens']==500
    assert s.events[0]['model']=='gpt-4o-mini' and 'copilot_inference' in caplog.text


@pytest.mark.parametrize('difference',['period','entity','negation'])
def test_claim_comparison_rejects_mismatched_scope_and_negation(difference):
    f=dict(id='metric',file='book.xlsx',sheet='Metrics',label='Completion rate',unit='ratio',company='Example',entity='Company total',period='2028-09',value=.86)
    q=dict(id='quote',file='book.xlsx',sheet='Letter',company='Example',period='2028-09',text='Companywide completion remained above 95% in the reporting month.')
    proposal=dict(quote_id='quote',quote='Companywide completion remained above 95%',fact_id='metric',operator='gt',threshold='95%')
    if difference=='period':f['period']='2028-08'
    if difference=='entity':f['entity']='Subsidiary'
    if difference=='negation':q['text']='We did not say '+q['text']
    assert numeric_claim(proposal,{'metric':f},{'quote':q})['status']=='unverifiable'


def test_ui_requires_fresh_file_consent_and_reuses_session_analysis(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from engine import ROOT
    seen=fake_model(monkeypatch);monkeypatch.setattr('mapping_ui.load_profiles',lambda:{})
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30)
    source=files();at.session_state['uploaded_inputs']=source;at.run()
    consent=next(c for c in at.checkbox if c.label=='Authorize external analysis of these uploads')
    assert not consent.value
    consent.check().run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception and len(seen)==1
    assert 'Mocked AI' in at.session_state['review'].narrative_mode
    at.radio(key='view').set_value('01  Upload & Setup').run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception and len(seen)==1
    assert at.session_state['review'].package['inference']['reused']
    at.radio(key='view').set_value('01  Upload & Setup').run()
    at.session_state['uploaded_inputs']={'replacement.xlsx':next(iter(source.values()))+b'\x00'}
    at.run()
    assert not next(c for c in at.checkbox if c.label=='Authorize external analysis of these uploads').value


def test_offline_results_and_excel_prominently_disclose_limited_analysis(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from engine import ROOT
    monkeypatch.delenv('OPENAI_API_KEY');monkeypatch.setattr('mapping_ui.load_profiles',lambda:{})
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30);at.session_state['uploaded_inputs']=files();at.run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception and any('AI analysis was unavailable.' in w.value for w in at.warning)
    wb=load_workbook(io.BytesIO(at.session_state['export'][0]))
    assert any('Limited offline analysis' in str(row) for row in wb['Executive Summary'].values)


def test_workbook_and_investment_calls_share_one_review_budget():
    seen=[]
    def send(req):seen.append(1);return response({})
    with inference_session(True,limit=.001) as session:
        complete({'max_output_tokens':100,'input':[]},'workbook_interpretation',send)
        with pytest.raises(InputError,match='spending limit'):complete({'max_output_tokens':100,'input':[]},'investment_analysis',send)
    assert len(seen)==1 and len(session.events)==1


def test_unrecognized_returned_model_does_not_receive_a_fictitious_price():
    def send(req):return dict(response({}),model='unexpected-provider-model')
    with inference_session(True) as session:complete({'max_output_tokens':100,'input':[]},'test',send)
    assert session.summary()['estimated_usd'] is None
    assert 'rates unknown' in session.events[0]['cost_basis']


def test_bad_ai_configuration_does_not_block_the_financial_report(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from engine import ROOT
    monkeypatch.setenv('COPILOT_AI_MAX_REVIEW_USD','invalid')
    monkeypatch.setattr('mapping_ui.load_profiles',lambda:{})
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30);at.session_state['uploaded_inputs']=files();at.run()
    next(c for c in at.checkbox if c.label=='Authorize external analysis of these uploads').check().run()
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception and not at.error
    assert at.session_state['review'].metrics['revenue']==1000
    assert at.session_state['review'].narrative_mode.startswith('Limited')
    assert at.session_state['export'][0]
