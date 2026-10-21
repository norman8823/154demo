"""Independent adaptability expectations were written before ingestion implementation."""
import copy
import io
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook,Workbook
from streamlit.testing.v1 import AppTest
from engine import analyze,InputError,ROOT,NOT_AVAILABLE
from ingestion import inspect_files,apply_decisions,normalized_records,MappingRequired,save_profiles,load_profiles,needs_review
from financial_analysis import analyze_intake
from reports import export_report
from mapping_ai import suggest

FIXTURES=Path(__file__).parent/'fixtures'
EXPECTED=json.loads((FIXTURES/'expected_adaptive.json').read_text())


def files(letter):
    p=next(FIXTURES.glob(letter+'_*.xlsx'));return {p.name:p.read_bytes()}


def bytes_of(wb):
    out=io.BytesIO();wb.save(out);return out.getvalue()


def harbor_intake():
    intake=inspect_files(files('B'))
    ambiguous=next(a for t in intake.tables for a in t.accounts if a.label==EXPECTED['B']['ambiguous_account'])
    return apply_decisions(intake,{ambiguous.id:{'category':EXPECTED['B']['confirmed_category'],'role':'detail','confirmed':True}})


@pytest.mark.parametrize('letter',['A','B'])
def test_independent_expected_financial_results(letter):
    r=analyze(files(letter)) if letter=='A' else analyze_intake(harbor_intake())
    expected=EXPECTED[letter]
    assert r.adaptive and r.company==expected['company']
    for metric in ['revenue','cogs','gross_profit','operating_expenses','ebitda','net_income']:
        assert r.metrics[metric]==expected[metric]
    assert r.metrics['gm']==expected['gross_margin']
    assert r.metrics['revenue_prior']==expected['prior_revenue']
    assert r.metrics['ebitda_prior']==expected['prior_ebitda']
    assert len(r.findings)>0
    assert all(f['id']>=101 for f in r.findings)
    assert all(e['Reference'] for e in r.evidence)
    for f in r.findings:assert any(e['Finding']==f['id'] for e in r.evidence)
    if letter=='A':
        assert r.metrics['gm_budget']==expected['budget_gross_margin']
        assert r.metrics['revenue_prior_year']==expected['prior_year_revenue']
        assert r.metrics['ebitda_budget']==expected['budget_ebitda']
        assert r.metrics['gm_variance_pp']==pytest.approx(-4)
        assert any('Gross margin below budget' in f['title'] for f in r.findings)
        assert r.report_period=='April 2027'
    else:
        assert r.metrics['operating_profit']==108000
        assert r.metrics['ebitda']!=r.metrics['operating_profit']
        assert any('net income + explicit' in row[3] for row in r.financial_rows if row[0].endswith('ebitda'))
        assert r.report_period=='February 2027'


def test_account_structure_avoids_double_counting_totals():
    intake=inspect_files(files('A'));t=intake.tables[0]
    assert [a.role for a in t.accounts if a.label=='Total revenue']==['total']
    r=analyze_intake(intake)
    assert r.metrics['revenue']==1000000 # Not detail + total = 2M.
    assert all(not r.normalized[i]['source_ref'].startswith('Trading') for i in range(len(r.normalized)))
    revenue=next(e for e in r.evidence if e['Reference']=="'Trading results'!C11")
    assert revenue['Original value / formula']==1000 and revenue['Validated value']==1000000
    assert 'scale 1000' in revenue['Check / use']


def test_ambiguous_account_limits_affected_metrics_without_blocking_report():
    intake=inspect_files(files('B'))
    partial=analyze_intake(intake)
    assert partial.metrics['revenue']>0
    assert any('unconfirmed' in w for w in partial.warnings)
    table=intake.tables[0]
    a=next(a for a in table.accounts if a.label=='Mobilization reserve')
    assert 'Review required' in a.status
    apply_decisions(intake,{a.id:{'category':'operating_expenses','role':'detail','confirmed':False}})
    assert needs_review(intake)
    assert not next(x for x in analyze_intake(intake).normalized if x['label']==a.label)['accepted']
    apply_decisions(intake,{a.id:{'category':'operating_expenses','role':'detail','confirmed':True}})
    assert analyze_intake(intake).metrics['ebitda']==118000


def test_confirmed_mapping_reused_but_not_for_another_company(tmp_path):
    intake=harbor_intake();path=tmp_path/'mappings.json';save_profiles(intake,path)
    profiles=load_profiles(path)
    renamed={'next_month_export.xlsx':next(iter(files('B').values()))}
    reused=inspect_files(renamed,profiles)
    a=next(a for t in reused.tables for a in t.accounts if a.label=='Mobilization reserve')
    assert a.confirmed and not needs_review(reused)
    wb=load_workbook(io.BytesIO(next(iter(files('B').values()))));wb['Export Ledger']['B1']='Different Company'
    new=inspect_files({'different.xlsx':bytes_of(wb)},profiles)
    assert needs_review(new)


@pytest.mark.parametrize('custom',[False,True])
@pytest.mark.parametrize('letter',['A','B'])
def test_both_exports_for_independent_layouts(letter,custom):
    r=analyze(files('A')) if letter=='A' else analyze_intake(harbor_intake())
    template=(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes() if custom else None
    data,name,warnings=export_report(r,template);wb=load_workbook(io.BytesIO(data))
    assert not warnings and r.report_period.replace(' ','_') in name
    assert 'September' not in name
    if custom:
        ws=wb['Partner Review'];original=load_workbook(io.BytesIO(template))['Partner Review']
        assert ws['C8'].value==EXPECTED[letter]['revenue'] and ws['G8'].value==EXPECTED[letter]['gross_margin']
        assert ws['B5'].value==EXPECTED[letter]['company']
        assert ws['F5'].value==r.report_period
        assert ws['C16'].value==original['C16'].value and ws['C16'].data_type=='f'
        assert ws['C8']._style==original['C8']._style
        assert ws['G10'].value==NOT_AVAILABLE
        assert ws['B19'].value!=NOT_AVAILABLE
    else:assert [n for n in wb.sheetnames if n!='Investment Summary']==['Executive Summary','Financial Performance','Findings','Source Evidence','Data Quality','Extracted Detail','Exceptions']
    assert any(isinstance(row[1],(int,float)) for row in wb['Financial Performance'].values)
    assert len([row for row in wb['Exceptions'].values if isinstance(row[0],int)])==len(r.findings)


def test_missing_ebitda_components_do_not_promote_operating_profit():
    intake=harbor_intake();t=intake.tables[0]
    # Explicitly ignore D&A, so the required EBITDA bridge is incomplete.
    depreciation=next(a for a in t.accounts if a.category=='depreciation');depreciation.role='ignore'
    r=analyze_intake(intake)
    assert 'net_income' in r.metrics and 'operating_profit' in r.metrics
    assert 'ebitda' not in r.metrics
    assert next(row for row in r.readiness if row['Analysis'].endswith('EBITDA trend'))['Status']==NOT_AVAILABLE


def test_unknown_units_and_basis_are_labeled_and_disable_comparison():
    source=files('A');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));wb['Trading results']['B4']='Management data'
    r=analyze({name:bytes_of(wb)})
    assert r.currency=='unknown' and r.scale is None
    assert r.metrics['revenue']==1000 # No invented thousands scaling.
    assert r.metrics['gm']==.5
    assert 'revenue_variance' not in r.metrics and 'revenue_prior' not in r.metrics
    assert any('not fully stated' in w for w in r.warnings)
    assert not any('budget' in f['title'] for f in r.findings)
    assert all(row[2].startswith('source units') or row[2]=='ratio' for row in r.financial_rows)


def test_reconciliation_discrepancy_is_a_finding_not_silently_repaired():
    source=files('A');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));wb['Trading results']['C11']=1100
    r=analyze({name:bytes_of(wb)})
    assert r.metrics['revenue']==1100000
    assert 'gm' not in r.metrics
    assert any('reconciliation' in f['title'] for f in r.findings)
    assert any('differs' in f['summary'] for f in r.findings)


def test_negative_expense_presentation_requires_and_respects_mapping_choice():
    source=files('A');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));ws=wb['Trading results']
    for row in [13,14,15,18,19,20,21]:
        for col in range(3,7):ws.cell(row,col,-ws.cell(row,col).value)
    modified={name:bytes_of(wb)}
    intake=inspect_files(modified,layouts={name+'|Trading results':{'header_row':7,'label_column':2,'expense_sign':'negative'}})
    r=analyze_intake(intake)
    assert r.metrics['cogs']==500000 and r.metrics['operating_expenses']==180000 and r.metrics['gm']==.5


def test_manual_table_and_header_mapping_without_code_changes():
    wb=Workbook();ws=wb.active;ws.title='Unfamiliar'
    ws['B1']='Maple Manufacturing';ws['B2']='USD whole dollars; Accrual basis'
    ws.append([])
    ws['B5']='Business captions';ws['D5']='CURRENT';ws['E5']='TARGET'
    for row,label,actual,budget in [(6,'Sales',500,600),(7,'Cost of goods sold',200,210),(8,'Gross profit',300,390),(9,'EBITDA',100,150)]:
        ws.cell(row,2,label);ws.cell(row,4,actual);ws.cell(row,5,budget)
    source={'new.xlsx':bytes_of(wb)}
    with pytest.raises(MappingRequired):analyze(source)
    layout={'new.xlsx':{'sheet':'Unfamiliar','header_row':5,'label_column':2,'columns':[{'column':4,'header':'CURRENT','period':'2028-03','scenario':'actual'},{'column':5,'header':'TARGET','period':'2028-03','scenario':'budget'}]}}
    r=analyze(source,layouts=layout)
    assert r.metrics['revenue']==500 and r.metrics['revenue_variance']==pytest.approx(-1/6)
    assert r.report_period=='March 2028'


def test_period_year_not_guessed_from_machine_date():
    wb=Workbook();ws=wb.active;ws['A1']='No dates';ws['A3']='Account';ws['B3']='Actual';ws['A4']='Revenue';ws['B4']=100
    with pytest.raises(MappingRequired):analyze({'yearless.xlsx':bytes_of(wb)})


def test_multiple_monthly_workbooks_form_trends_but_overlapping_statements_require_review():
    source=files('A');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));ws=wb['Trading results']
    # Split the independently authored report into one actual month per workbook.
    first=load_workbook(io.BytesIO(source[name]));first['Trading results'].delete_cols(4,3)
    second=load_workbook(io.BytesIO(source[name]));second['Trading results'].delete_cols(3,2);second['Trading results'].delete_cols(4,1)
    r=analyze({'april.xlsx':bytes_of(first),'march.xlsx':bytes_of(second)})
    assert r.metrics['revenue']==1000000 and r.metrics['revenue_prior']==900000
    duplicate=analyze({'one.xlsx':bytes_of(first),'two.xlsx':bytes_of(first)})
    assert duplicate.metrics['revenue']==1000000
    assert any('not added together' in row['Recognized as'] for row in duplicate.recognized)


def test_multiple_companies_are_never_consolidated():
    intake=inspect_files({**files('A'),**files('B')})
    a=next(a for t in intake.tables for a in t.accounts if a.label=='Mobilization reserve')
    apply_decisions(intake,{a.id:{'category':'operating_expenses','role':'detail','confirmed':True}})
    r=analyze_intake(intake)
    assert not r.metrics and 'Northstar' in r.company and 'Harbor' in r.company
    assert len(r.financial_rows)>0
    data,_,_=export_report(r,(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes())
    assert load_workbook(io.BytesIO(data))['Partner Review']['C8'].value==NOT_AVAILABLE


def test_ai_suggestions_are_structured_bounded_and_do_not_send_amounts(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','test-key')
    intake=inspect_files(files('B'));a=next(a for t in intake.tables for a in t.accounts if a.label=='Mobilization reserve')
    def fake(req):
        payload=json.loads(req.data)
        text=payload['input'][1]['content']
        assert '270000' not in text and '12000' not in text and '400000' not in text
        assert payload['text']['format']['strict'] and payload['store'] is False
        return {'output':[{'content':[{'type':'output_text','text':json.dumps({'accounts':[{'id':a.id,'category':'operating_expenses','role':'detail','reason':'Reserve requires controller confirmation'}],'columns':[],'layouts':[]})}]}]}
    result=suggest(intake,fake)
    apply_decisions(intake,{a.id:dict(result['accounts'][0],confirmed=False)})
    assert needs_review(intake)


def test_ai_failure_and_invented_ids_do_not_accept_mappings(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','test-key');intake=inspect_files(files('B'))
    def fake(req):return {'output':[{'content':[{'type':'output_text','text':'{"accounts":[{"id":"invented","category":"ebitda","role":"total"}],"columns":[],"layouts":[]}'}]}]}
    with pytest.raises(InputError,match='invalid suggestion'):suggest(intake,fake)
    assert needs_review(intake)
    monkeypatch.delenv('OPENAI_API_KEY')
    with pytest.raises(InputError,match='OPENAI_API_KEY'):suggest(intake)


def test_mapping_review_ui_can_resolve_one_ambiguous_account_and_generate(monkeypatch):
    monkeypatch.setattr('mapping_ui.load_profiles',lambda: {})
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30);at.session_state['uploaded_inputs']=files('B');at.run()
    assert not at.exception
    assert not next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).disabled
    next(s for s in at.selectbox if s.label=='Classification for Mobilization reserve').set_value('operating_expenses')
    next(c for c in at.checkbox if c.label=='Confirm classification for Mobilization reserve').check()
    next(b for b in at.button if b.label=='Apply reviewed classifications').click().run()
    assert not at.exception
    next(b for b in at.button if b.label.startswith('Generate Portfolio Review')).click().run()
    assert not at.exception and 'February 2027 Portfolio Review' in [t.value for t in at.title]
    assert at.session_state['review'].metrics['ebitda']==118000
    at.button(key='evidence_101').click().run()
    assert not at.exception


def test_generic_financials_and_known_billing_module_can_coexist_without_consolidation():
    from engine import demo_files
    known={name:data for name,data in demo_files().items() if 'QuickBooks' in name or name.endswith('.csv')}
    r=analyze({**files('A'),**known})
    assert any(f['id']==3 and '$27,978.75' in f['summary'] for f in r.findings)
    assert any(f['id']>=101 for f in r.findings)
    assert not r.metrics
    export_report(r)


def test_non_usd_and_unknown_units_do_not_mislabel_dollar_template_fields():
    source=files('A');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));wb['Trading results']['B4']='GBP thousands | accrual basis'
    r=analyze({name:bytes_of(wb)})
    data,_,warnings=export_report(r,(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes())
    result=load_workbook(io.BytesIO(data))['Partner Review']
    assert result['C8'].value==1000000 and 'GBP' in result['C8'].number_format and '$' not in result['C8'].number_format
    assert warnings and 'no currency conversion' in warnings[0]
    wb['Trading results']['B4']='Unit information absent'
    r=analyze({name:bytes_of(wb)})
    data,_,_=export_report(r,(ROOT/'sample_templates/Sample_Firm_Review_Template.xlsx').read_bytes())
    result=load_workbook(io.BytesIO(data))['Partner Review']
    assert 'currency/scale not confirmed' in result['C8'].value
    assert result['G8'].value==.5


def test_saved_manual_layout_reuses_structure_but_never_stale_periods(tmp_path):
    wb=Workbook();ws=wb.active;ws.title='Custom ledger'
    ws['B1']='Maple Manufacturing';ws['B2']='USD whole dollars; accrual basis';ws['B3']='Month: March 2028'
    ws['B5']='Business captions';ws['D5']='CURRENT';ws['E5']='TARGET'
    for row,label,actual,budget in [(6,'Sales',500,600),(7,'Cost of goods sold',200,210),(8,'Gross profit',300,390)]:
        ws.cell(row,2,label);ws.cell(row,4,actual);ws.cell(row,5,budget)
    manual={'first.xlsx':{'sheet':'Custom ledger','header_row':5,'label_column':2,'columns':[{'column':4,'header':'CURRENT','period':'2028-03','scenario':'actual'},{'column':5,'header':'TARGET','period':'2028-03','scenario':'budget'}]}}
    intake=inspect_files({'first.xlsx':bytes_of(wb)},layouts=manual)
    path=tmp_path/'profiles.json';save_profiles(intake,path)
    ws['B3']='Month: April 2028'
    r=analyze({'second.xlsx':bytes_of(wb)},profiles=load_profiles(path))
    assert r.report_period=='April 2028' and r.metrics['revenue']==500
    ws['B3']=None
    with pytest.raises(MappingRequired):analyze({'third.xlsx':bytes_of(wb)},profiles=load_profiles(path))


def test_missing_account_values_and_dashes_are_not_imputed_as_zero():
    source=files('A');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));ws=wb['Trading results']
    ws.delete_rows(11,1) # Remove reported total revenue, so revenue would need all detail accounts.
    ws['C10']='—' # Service retainers actual is missing; budget/prior cells still present.
    r=analyze({name:bytes_of(wb)})
    assert 'revenue' not in r.metrics and 'gm' not in r.metrics
    assert any('no zero values' in w for w in r.warnings)
    export_report(r)


def test_quickbooks_blank_account_header_is_detected_from_structure():
    source=files('B');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));wb['Export Ledger']['B9']=None
    intake=inspect_files({name:bytes_of(wb)})
    assert len(intake.tables)==1 and not intake.unresolved
    table=intake.tables[0]
    assert table.label_column==2 and table.header_row==9
    a=next(a for a in table.accounts if a.label=='Mobilization reserve')
    apply_decisions(intake,{a.id:{'category':'operating_expenses','role':'detail','confirmed':True}})
    assert analyze_intake(intake).metrics['ebitda']==118000


def test_other_company_with_demo_like_quickbooks_layout_uses_adaptive_path():
    from engine import demo_files,is_known_package
    source=demo_files();name=next(n for n in source if 'QuickBooks' in n)
    wb=load_workbook(io.BytesIO(source[name]));wb['Profit and Loss']['A2']='Larch Services Ltd'
    # A Rate Card is an optional specialized sheet, not evidence that the company is Add-on B.
    general={'larch.xlsx':bytes_of(wb)}
    assert not is_known_package(general)
    result=analyze(general)
    assert result.company=='Larch Services Ltd' and result.adaptive


def test_conflicting_metadata_dates_do_not_assign_yearless_actual_headers():
    wb=Workbook();ws=wb.active;ws['A1']='Larch Services Ltd';ws['A2']='Report month: March 2028';ws['A3']='Prepared April 15, 2028'
    ws['A5']='Account';ws['B5']='Actual';ws['A6']='Revenue';ws['B6']=100
    with pytest.raises(MappingRequired):analyze({'ambiguous_dates.xlsx':bytes_of(wb)})


def test_ai_or_manual_layout_without_known_headers_exposes_columns_for_review():
    wb=Workbook();ws=wb.active;ws.title='Unfamiliar';ws['A1']='Maple Manufacturing';ws['A3']='Descriptions';ws['C3']='CURRENT';ws['D3']='TARGET'
    ws['A4']='Sales';ws['C4']=100;ws['D4']=200
    intake=inspect_files({'new.xlsx':bytes_of(wb)},layouts={'new.xlsx':{'sheet':'Unfamiliar','header_row':3,'label_column':1,'ai_pending':True}})
    assert len(intake.tables)==1 and len(intake.tables[0].columns)==2
    assert {c['header'] for c in intake.tables[0].columns}=={'CURRENT','TARGET'}
    assert needs_review(intake)


def test_conflicting_column_currency_is_not_silently_compared_or_converted():
    source=files('A');name=next(iter(source));wb=load_workbook(io.BytesIO(source[name]));wb['Trading results']['D7']='Apr 2027 Plan GBP'
    result=analyze({name:bytes_of(wb)})
    assert result.metrics['revenue']==1000000 and result.currency=='USD'
    assert 'revenue_variance' not in result.metrics
    assert any('currency conflicts' in w for w in result.warnings)


def test_numeric_text_does_not_guess_decimal_comma_conventions():
    from ingestion import amount
    assert amount('1,234.56')==1234.56
    assert amount('($1,234.56)')==-1234.56
    assert amount('1234,56') is None
    assert amount('1.234,56') is None
